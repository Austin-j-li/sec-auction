#!/usr/bin/env python3
"""Mechanical checks on a deal-ledger workbook produced under SEC_Deal_Ledger_Extraction_Instruction.md
(full or short form). Usage:  python3 tools/check_ledger.py path/to/<deal>_ledger.xlsx [more.xlsx ...]

It checks what a script can check; it does not judge whether the extraction is right.
FAIL = a rule of the instruction is broken.  WARN = worth a look (heuristic).  Exit code 1 if any FAIL.
"""
import sys, re, datetime
import openpyxl

VISIBLE = ['#','When','Who','What happened','Process','Round','Type','Terms or outcome','Formality',
           'Conditions level','Why and evidence','Source','Review','Reviewer note']
EXPAND = ['Row id','Include','Count','Date from','Date to','Working date','Date basis','Date method',
          'Price low','Price high','Price kind','Price origin','Currency','All cash','Cash at closing',
          'Conditions detail','Due date','Deadline treatment','Round finality','Decided by','Exit reason',
          'Outcome basis','Page','Related rows','Deal']
LABELS = {'Target interest','Bidder interest','Target sale decision','Activist pressure','Activist involvement',
 'Adviser engaged','Adviser service observed','Adviser ended','Contact','NDA signed','Round opened','Deadline set',
 'Deadline revised','Deadline','Target decision','Information access changed','Material process update',
 'Exclusivity changed','Bid','Bid reaffirmed','Offer update','Other-scope bid','Valuation statement',
 'Bidding group changed','Joined group','Dropped by target','Withdrew','Did not submit','Participation paused',
 'Re-entered','Not selected at signing','Sale process announced','Bid announced','Merger announced',
 'Merger agreement signed','Go-shop changed','Process terminated','Process restarted','Agreement terminated','Closed'}
CONTROLLED = {
 'What happened': LABELS,
 'Type': {'Strategic','Financial','Mixed','Unknown'},
 'Formality': {'Formal','Informal','Insufficient evidence','Varies'},
 'Conditions level': {'None','Light','Heavy','Insufficient evidence','Varies'},
 'Include': {'Yes','No'},
 'Date basis': {'Reported day','Inferred day','Reported interval','Approximate window','Relative only','Undated'},
 'Date method': {'Reported','Inferred','Assigned: deadline','Assigned: decision day','Assigned: midpoint',
                 'Assigned: bound','Assigned: sequence'},
 'Price kind': {'Point','Bidder range','Group envelope','Bound only','Undisclosed'},
 'Price origin': {'Stated','Carried forward','Inferred'},
 'All cash': {'Yes','No','Not stated'},
 'Deadline treatment': {'Enforced','Extended','Late bids accepted','Passed without action','Unclear',
                        'Extended; Late bids accepted'},
 'Round finality': {'Announced as final','Inferred final','Not final'},
 'Decided by': {'Bidder','Target','Both','Unknown'},
 'Exit reason': {'Value below market price','Value at or below market price','Value below earlier offer',
                 'Value at earlier offer','Would not improve earlier offer','Lower offer than rivals',
                 'Terms or process','Other stated reason','Not stated'},
 'Outcome basis': {'Stated','Inferred: residual','Inferred: exclusivity','Inferred: silent','Inferred: identity'},
}
OUTCOMES = {'Dropped by target','Withdrew','Did not submit','Participation paused','Not selected at signing'}
EXITS = {'Dropped by target','Withdrew','Did not submit','Not selected at signing','Joined group'}
COND_RE = re.compile(r'^Fin: (committed|no contingency|not needed|represented|not committed|not stated); '
                     r'DD required: (none|confirmatory|substantive \(\d+ ?wks?\)|substantive \(N wks\)|not stated); '
                     r'DD open: (yes|no|not disclosed); '
                     r'Excl: (requested \(\d+ ?wks?\)|requested \(N wks\)|granted|none|not stated)$')
DATECOLS = ['Date from','Date to','Working date','Due date']

def norm(s):
    s = str(s).replace('“','"').replace('”','"').replace('‘',"'").replace('’',"'").replace('\xa0',' ')
    return re.sub(r'\s+',' ',s).strip()

class Report:
    def __init__(self): self.items=[]
    def add(self, level, check, msg): self.items.append((level,check,msg))
    def show(self, path):
        print('='*100); print(path)
        order={'FAIL':0,'WARN':1,'PASS':2,'INFO':3}
        for lvl,chk,msg in sorted(self.items,key=lambda x:order[x[0]]): print(f'  [{lvl}] {chk}: {msg}')
        n=sum(1 for i in self.items if i[0]=='FAIL'); w=sum(1 for i in self.items if i[0]=='WARN')
        print(f'  -> {n} FAIL, {w} WARN'); return n

def rows_of(ws):
    rows=list(ws.iter_rows(values_only=True))
    hi=next((i for i,r in enumerate(rows) if r and r[0]=='#'),None)
    if hi is None: return None,None,None
    hdr=[h for h in rows[hi]]
    data=[(hi+2+k,r) for k,r in enumerate(rows[hi+1:]) if any(c is not None for c in r)]
    return hdr,data,hi+1

def check(path):
    R=Report()
    wb=openpyxl.load_workbook(path); wbv=openpyxl.load_workbook(path,data_only=True)
    names=wb.sheetnames
    for s in ['Deal ledger','Summary','Questions','Source text']:
        if s not in names: R.add('FAIL','sheets',f'missing sheet "{s}"')
        elif wb[s].sheet_state!='visible': R.add('FAIL','sheets',f'"{s}" is not visible')
    vis=[s for s in names if wb[s].sheet_state=='visible']
    if len(vis)!=4: R.add('FAIL','sheets',f'{len(vis)} visible sheets {vis}; the instruction asks for four')
    if 'Lists' not in names: R.add('WARN','sheets','no Lists sheet (drop-downs)')
    if 'Deal ledger' not in names: R.show(path); return 1
    hdr,data,hrow=rows_of(wbv['Deal ledger'])
    if hdr is None: R.add('FAIL','ledger','no header row starting with "#"'); R.show(path); return 1
    hdr_clean=[h for h in hdr if h is not None]
    want=VISIBLE+EXPAND
    missing=[c for c in want if c not in hdr_clean]; extra=[c for c in hdr_clean if c not in want]
    if missing: R.add('FAIL','columns',f'missing columns: {missing}')
    if extra: R.add('WARN','columns',f'unexpected columns: {extra}')
    if not missing and [c for c in hdr_clean if c in want]!=want: R.add('FAIL','columns','columns are not in the prescribed left-to-right order')
    ix={h:i for i,h in enumerate(hdr) if h is not None}
    def g(r,c): return r[ix[c]] if c in ix else None
    rid=lambda n,r: f'{g(r,"Row id") or "?"} (Excel row {n})'
    R.add('INFO','ledger',f'{len(data)} ledger rows')

    # Row ids, Include
    ids=[g(r,'Row id') for n,r in data]
    dup={i for i in ids if i is not None and ids.count(i)>1}
    if None in ids: R.add('FAIL','row ids',f'{ids.count(None)} rows without a Row id')
    if dup: R.add('FAIL','row ids',f'duplicate Row ids: {sorted(dup)}')
    if not dup and None not in ids: R.add('PASS','row ids','unique and filled')

    # controlled values
    bad=[]
    for col,allowed in CONTROLLED.items():
        if col not in ix: continue
        for n,r in data:
            v=g(r,col)
            if v is None or str(v).strip()=='' : continue
            if str(v) not in allowed: bad.append(f'{rid(n,r)} {col}="{v}"')
    R.add('FAIL' if bad else 'PASS','controlled values', (f'{len(bad)} values not in the allowed lists: '+'; '.join(bad[:12])+(' …' if len(bad)>12 else '')) if bad else 'all within the allowed lists')

    # dates
    if 'Working date' in ix:
        notdate=[];blank=[];outside=[];back=[];prev=None;prevr=None
        for n,r in data:
            for c in DATECOLS:
                v=g(r,c)
                if v is not None and not isinstance(v,(datetime.datetime,datetime.date)): notdate.append(f'{rid(n,r)} {c}="{v}"')
            wd=g(r,'Working date')
            if wd is None: blank.append(rid(n,r)); continue
            if not isinstance(wd,(datetime.datetime,datetime.date)): continue
            df,dt=g(r,'Date from'),g(r,'Date to')
            if isinstance(df,datetime.datetime) and wd<df: outside.append(f'{rid(n,r)} before Date from')
            if isinstance(dt,datetime.datetime) and wd>dt: outside.append(f'{rid(n,r)} after Date to')
            if prev is not None and wd<prev: back.append(f'{rid(n,r)} {wd:%m/%d/%Y} after {prevr} {prev:%m/%d/%Y}')
            prev,prevr=wd,rid(n,r)
        R.add('FAIL' if notdate else 'PASS','native dates','; '.join(notdate[:10]) if notdate else 'date cells are real dates')
        R.add('FAIL' if blank else 'PASS','working date filled',f'{len(blank)} rows blank: '+', '.join(blank[:15]) if blank else 'every row has a Working date')
        R.add('FAIL' if outside else 'PASS','working date inside window','; '.join(outside[:10]) if outside else 'ok')
        R.add('FAIL' if back else 'PASS','working dates never decrease',f'{len(back)} reversals: '+'; '.join(back[:10]) if back else 'sorted order reproduces #')
    if 'Date method' in ix and 'Date basis' in ix:
        bad=[]
        for n,r in data:
            b,m=g(r,'Date basis'),g(r,'Date method')
            if m is None: bad.append(f'{rid(n,r)} no Date method'); continue
            if (m=='Reported')!=(b=='Reported day') or (m=='Inferred')!=(b=='Inferred day'): bad.append(f'{rid(n,r)} basis="{b}" method="{m}"')
        R.add('FAIL' if bad else 'PASS','date basis/method pairs',f'{len(bad)}: '+'; '.join(bad[:10]) if bad else 'permitted pairs only')

    # outcome rows
    if 'Outcome basis' in ix:
        bad=[]
        for n,r in data:
            lab=g(r,'What happened')
            if lab in OUTCOMES:
                for c in ['Decided by','Exit reason','Outcome basis']:
                    if g(r,c) in (None,''): bad.append(f'{rid(n,r)} no {c}')
                ob=str(g(r,'Outcome basis') or ''); terms=norm(g(r,'Terms or outcome') or '')
                if ob.startswith('Inferred') and not terms.startswith(ob): bad.append(f'{rid(n,r)} Terms does not open with "{ob}"')
                if ob=='Stated' and terms.startswith('Stated'): bad.append(f'{rid(n,r)} Stated row has an opener')
                if ob=='Inferred: identity' and g(r,'Count') not in (0,'0'): bad.append(f'{rid(n,r)} identity row Count is not 0')
        R.add('FAIL' if bad else 'PASS','outcome rows',f'{len(bad)}: '+'; '.join(bad[:12]) if bad else 'basis, agency, reason and openers consistent')

    # rounds and deadlines
    seen={};bad=[]
    for n,r in data:
        lab=g(r,'What happened')
        if lab=='Round opened':
            key=(g(r,'Process'),str(g(r,'Round')))
            if key in seen: bad.append(f'{rid(n,r)} second Round opened for process/round {key}')
            seen[key]=1
            if 'Round finality' in ix and g(r,'Round finality') in (None,''): bad.append(f'{rid(n,r)} no Round finality')
            if str(g(r,'Round'))=='1' and 'R1 anchor:' not in str(g(r,'Terms or outcome') or ''): bad.append(f'{rid(n,r)} round-1 row has no "R1 anchor:" tag')
        if lab=='Deadline' and 'Deadline treatment' in ix and g(r,'Deadline treatment') in (None,''): bad.append(f'{rid(n,r)} Deadline row without a Deadline treatment')
    rounds_used={(g(r,'Process'),str(g(r,'Round'))) for n,r in data if str(g(r,'Round')) not in ('0','post','None','')}
    for k in sorted(rounds_used-set(seen),key=str): bad.append(f'round {k} has rows but no Round opened row')
    R.add('FAIL' if bad else 'PASS','rounds and deadlines','; '.join(bad[:12]) if bad else 'one opening per round; finality, anchor and treatments present')

    # bids
    bad=[]
    for n,r in data:
        lab=g(r,'What happened')
        if lab in ('Bid','Bid reaffirmed'):
            for c in ['Formality','Conditions level','Price kind','Price origin']:
                if c in ix and g(r,c) in (None,''): bad.append(f'{rid(n,r)} no {c}')
            if 'Conditions detail' in ix:
                cd=g(r,'Conditions detail')
                if cd in (None,''): bad.append(f'{rid(n,r)} no Conditions detail')
                elif not COND_RE.match(norm(cd)): bad.append(f'{rid(n,r)} Conditions detail not in the fixed format: "{cd}"')
            lo,hi,kind=g(r,'Price low'),g(r,'Price high'),g(r,'Price kind')
            terms=str(g(r,'Terms or outcome') or '')
            if kind=='Point' and lo!=hi: bad.append(f'{rid(n,r)} Point but Price low != Price high')
            if kind=='Bound only':
                if (lo is None)==(hi is None): bad.append(f'{rid(n,r)} Bound only needs exactly one price cell')
                if 'Bound: ≥' not in terms and 'Bound: ≤' not in terms: bad.append(f'{rid(n,r)} Bound only without a "Bound:" tag')
                if g(r,'Cash at closing') is not None: bad.append(f'{rid(n,r)} Bound only with Cash at closing filled')
            if kind=='Bidder range' and lo is not None and hi is not None and lo>=hi: bad.append(f'{rid(n,r)} range endpoints not increasing')
            if lab=='Bid reaffirmed' and g(r,'Price origin')=='Stated' and 'Reaffirmed by:' in terms: bad.append(f'{rid(n,r)} late reconfirmation with Price origin Stated')
        if lab in ('Merger agreement signed','Offer update') and (g(r,'Price low') is not None or g(r,'Price high') is not None):
            bad.append(f'{rid(n,r)} {lab} row carries numeric price cells')
    R.add('FAIL' if bad else 'PASS','bid rows',f'{len(bad)}: '+'; '.join(bad[:12]) if bad else 'classifications, detail strings and price cells consistent')

    # sources, pages, quotations
    if 'Source text' in names:
        st=list(wbv['Source text'].iter_rows(values_only=True))
        sh=next((i for i,r in enumerate(st) if r and r[0] and str(r[0]).strip().lower()=='paragraph'),0)
        paras={str(r[0]).strip():norm(r[3] if len(r)>3 and r[3] else '') for r in st[sh+1:] if r and r[0]}
        alltext=' '.join(paras.values())
        bad_src=[];bad_page=[];bad_q=[];nq=0
        for n,r in data:
            src=g(r,'Source')
            if src:
                for pid in re.findall(r'[PX]-?\d+',str(src)):
                    if pid not in paras: bad_src.append(f'{rid(n,r)} cites {pid}')
                if 'Page' in ix and g(r,'Page') in (None,''): bad_page.append(rid(n,r))
            for q in re.findall(r'"([^"]{25,})"', norm(g(r,'Why and evidence') or '')):
                nq+=1; q2=q.rstrip('.,;: ')
                pieces=[p.strip() for p in re.split(r'\.\.\.|…|\[[^\]]*\]',q2) if len(p.strip())>=15]
                if pieces and not all(p in alltext for p in pieces): bad_q.append(f'{rid(n,r)} "{q[:60]}…"')
        R.add('FAIL' if bad_src else 'PASS','source ids',f'{len(bad_src)} unknown paragraph ids: '+'; '.join(bad_src[:10]) if bad_src else 'every cited paragraph exists in Source text')
        if 'Page' in ix: R.add('FAIL' if bad_page else 'PASS','page filled',f'{len(bad_page)} rows with Source but no Page: '+', '.join(bad_page[:12]) if bad_page else 'Page filled wherever Source is')
        R.add('FAIL' if bad_q else 'PASS','quotations',f'{len(bad_q)} of {nq} quotations not found in Source text: '+'; '.join(bad_q[:8]) if bad_q else f'all {nq} quotations found in Source text')

    # per-stage balance (heuristic: cohorts and pre-NDA bidders make an exact automatic count impossible)
    if 'Count' in ix:
        def num(v):
            try: return float(v)
            except (TypeError,ValueError): return 0.0
        procs=sorted({g(r,'Process') for n,r in data if g(r,'Process') is not None},key=str)
        for p in procs:
            rowsp=[(n,r) for n,r in data if g(r,'Process')==p and g(r,'Include')!='No']
            nda=sum(num(g(r,'Count')) for n,r in rowsp if g(r,'What happened')=='NDA signed')
            nda_who={norm(g(r,'Who') or '') for n,r in rowsp if g(r,'What happened')=='NDA signed'}
            bid_no_nda={norm(g(r,'Who') or '') for n,r in rowsp if g(r,'What happened')=='Bid'} - nda_who
            exits=sum(num(g(r,'Count')) for n,r in rowsp if g(r,'What happened') in EXITS)
            term=sum(num(g(r,'Count')) for n,r in rowsp if g(r,'What happened')=='Process terminated')
            reent=sum(max(num(g(r,'Count')),1) for n,r in rowsp if g(r,'What happened')=='Re-entered')
            units=sum(int(m) for n,r in rowsp if g(r,'What happened')=='Bidding group changed' for m in re.findall(r'units [−-](\d+)',str(g(r,'Terms or outcome') or '')))
            signed=any(g(r,'What happened')=='Merger agreement signed' for n,r in rowsp)
            stock=nda-exits-term+reent-units
            msg=f'process {p}: NDA entrants {nda:g} − exits {exits:g} − closed at termination {term:g} + re-entries {reent:g} − group units {units} = {stock:g}; bidders with a Bid but no NDA row of their own (normally inside an NDA cohort; they add entrants only if they never signed): {sorted(bid_no_nda) or "none"}'
            target=1 if signed else 0
            R.add('PASS' if stock==target else 'WARN','per-stage balance (heuristic)',msg+f'; expected {target}')

    # summary formulas
    if 'Summary' in names:
        bad=[];nf=0
        for row in wb['Summary'].iter_rows():
            for c in row:
                if isinstance(c.value,str) and c.value.startswith('='):
                    nf+=1; f=c.value
                    if f.count('(')!=f.count(')') or re.search(r'\)\s*[A-Za-z→>-]{2,}',f) or '->' in f: bad.append(f'Summary!{c.coordinate}')
                    cv=wbv['Summary'][c.coordinate].value
                    if isinstance(cv,str) and cv.startswith('#'): bad.append(f'Summary!{c.coordinate} = {cv}')
        R.add('FAIL' if bad else 'PASS','summary formulas',f'{nf} formulas; suspect: '+', '.join(bad) if bad else f'{nf} formulas, none malformed')
    return R.show(path)

if __name__=='__main__':
    if len(sys.argv)<2: print(__doc__); sys.exit(2)
    total=sum(check(p) for p in sys.argv[1:])
    sys.exit(1 if total else 0)
