#!/usr/bin/env python3
"""Preserve originals; make randomized neutral grading copies, dumps and checks."""
import argparse,datetime,hashlib,json,pathlib,re,secrets,shutil,subprocess,sys,zipfile,xml.etree.ElementTree as ET
import openpyxl
BASE=pathlib.Path(__file__).resolve().parents[1]
MODELS=['sol','opus','deepseek']

def original(model,deal):
 r=BASE/'runs'/f'{model}_{deal}'
 return r/('work/extraction' if model=='deepseek' else 'extraction')/f'{deal}.xlsx'

def dump(src,dst):
 w=openpyxl.load_workbook(src,data_only=False)
 lines=[]
 for s in w:
  lines.append(f'SHEET {s.title!r}; state={s.sheet_state}; rows={s.max_row}; columns={s.max_column}')
  headers=[str(c.value) if c.value is not None else f'column {c.column_letter}' for c in s[1]]
  for row in s:
   if not any(c.value is not None for c in row):continue
   lines.append(f'EXCEL ROW {row[0].row}')
   for c in row:
    if c.value is None:continue
    v=c.value.isoformat() if isinstance(c.value,(datetime.date,datetime.datetime)) else str(c.value)
    lines.append(f'  {c.coordinate} [{headers[c.column-1]}]: {v}')
    if c.comment:lines.append(f'  COMMENT {c.coordinate}: {c.comment.text}')
 dst.write_text('\n'.join(lines)+'\n')

def neutral_copy(src,dst):
 with zipfile.ZipFile(src) as zin,zipfile.ZipFile(dst,'w',zipfile.ZIP_DEFLATED) as zout:
  for item in zin.infolist():
   data=zin.read(item.filename)
   if item.filename=='docProps/core.xml':
    root=ET.fromstring(data)
    for child in list(root):
     if child.tag.split('}')[-1] in ['creator','lastModifiedBy','created','modified','title','description','subject','keywords','revision']:root.remove(child)
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   elif item.filename=='docProps/app.xml':
    root=ET.fromstring(data)
    for child in list(root):
     if child.tag.split('}')[-1] in ['Application','AppVersion','Company','Manager']:root.remove(child)
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   elif 'comments' in item.filename and item.filename.endswith('.xml'):
    root=ET.fromstring(data)
    for child in root.iter():
     if child.tag.split('}')[-1]=='author':child.text='Reviewer'
    data=ET.tostring(root,encoding='utf-8',xml_declaration=True)
   zi=zipfile.ZipInfo(item.filename,date_time=(2000,1,1,0,0,0));zi.compress_type=zipfile.ZIP_DEFLATED
   zout.writestr(zi,data)

def main():
 p=argparse.ArgumentParser();p.add_argument('deal');a=p.parse_args();deal=a.deal
 root=BASE/'grading'/deal;work=root/'work';ref=work/'output/reference/reference.json'
 status=json.loads((root/'reference_status.json').read_text())
 if status.get('state')!='finished' or status.get('exit_code')!=0 or not ref.exists():raise SystemExit('Reference not complete')
 tests=json.loads(ref.read_text())['tests'];total=sum(t['weight'] for t in tests)
 if abs(total-100)>1e-6:raise SystemExit(f'Reference weights total {total}')
 mapfile=BASE/'admin'/f'blind_map_{deal}.json'
 if mapfile.exists():raise SystemExit('Already prepared; refusing to reshuffle')
 for model in MODELS:
  if not original(model,deal).exists():raise SystemExit(f'Missing candidate {model}/{deal}')
 mapping={};models=MODELS.copy();secrets.SystemRandom().shuffle(models)
 (work/'candidates').mkdir(exist_ok=True);(work/'checks').mkdir(exist_ok=True)
 refhash=hashlib.sha256(ref.read_bytes()).hexdigest()
 (root/'reference_lock.json').write_text(json.dumps({'sha256':refhash,'tests':len(tests),'total_weight':total,'locked_utc':datetime.datetime.now(datetime.timezone.utc).isoformat()},indent=2))
 for i,model in enumerate(models):
  label=f'Candidate-{chr(65+i)}';src=original(model,deal);dest=work/'candidates'/f'{label}.xlsx'
  neutral_copy(src,dest);dump(dest,dest.with_suffix('.txt'))
  text=dest.with_suffix('.txt').read_text()
  identity_hits=re.findall(r'(?i)\b(?:deepseek|claude|anthropic|openai|gpt[- ]?\d|gpt[- ]?sol)\b',text)
  if identity_hits:raise SystemExit(f'Identity clue in {label}; requires documented neutralization: {identity_hits}')
  mapping[label]={'model':model,'original':str(src),'original_sha256':hashlib.sha256(src.read_bytes()).hexdigest(),'neutral_sha256':hashlib.sha256(dest.read_bytes()).hexdigest()}
  checker=BASE/'admin/check_lean.py'
  if checker.exists():
   filing=next((work/'raw_filing').glob('*'))
   cp=subprocess.run([sys.executable,str(checker),'--workbook',str(dest),'--filing',str(filing),'--output',str(work/'checks'/f'{label}.json')],capture_output=True,text=True)
   (BASE/'admin'/f'check_{deal}_{label}.log').write_text(cp.stdout+cp.stderr)
 (work/'alex').mkdir(exist_ok=True)
 previous=BASE.parent/'clean_run_comparison_2026-09-19/score/deals'/deal
 for name in ['alex_hand_rows.txt','alex_voice_notes.txt']:
  if (previous/name).exists():shutil.copyfile(previous/name,work/'alex'/name)
 summary=previous.parent/'alex_closing_summary.txt'
 if summary.exists():
  raw=summary.read_text();marker="[] Alex's summary"
  if marker not in raw:raise SystemExit('Cannot isolate Alex-authored summary')
  (work/'alex'/summary.name).write_text(raw[raw.index(marker):])
 (work/'alex/PROVENANCE.txt').write_text('Primary-source transcripts previously prepared from ref/deal_details_Alex_2026.xlsx and ref/alex_voice_notes_2026-08.docx. These are researcher-source transcripts, not old candidate scores or old grading references. The closing-summary transcript begins at the explicit Alex\'s summary heading; the preceding Claude-authored interpretation has been excluded. These sources supplement only the separate compatibility assessment; the current instruction governs primary grading.\n')
 mapfile.write_text(json.dumps(mapping,indent=2)+'\n')
 print(json.dumps({'deal':deal,'labels':list(mapping),'reference_sha256':refhash}))
if __name__=='__main__':main()
