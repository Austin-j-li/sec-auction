#!/usr/bin/env python3
"""Small synthetic-fixture tests for check_lean.py (no candidate workbooks)."""

from __future__ import annotations

import datetime as dt
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

import check_lean


QUOTES = [
    "On January 2, the target requested proposals.",
    "On January 3, Alpha submitted an offer of $10 per share.",
    "On January 4, the target signed the merger agreement.",
]


def finish_sheet(ws, width: int) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, width).column_letter}{max(ws.max_row, 1)}"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=width):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(wrap_text=True)


def build_valid_fixture(directory: Path) -> tuple[Path, Path]:
    filing = directory / "filing.htm"
    filing.write_text(
        "<html><body>"
        f"<p>{QUOTES[0]}</p>"
        "<p>On January 3, Alpha submitted an offer of $<b>10</b> per share.</p>"
        f"<p>{QUOTES[2]}</p>"
        "</body></html>",
        encoding="utf-8",
    )

    workbook = directory / "valid.xlsx"
    wb = Workbook()
    ledger = wb.active
    ledger.title = "Deal ledger"
    ledger.append(check_lean.LEDGER_COLUMNS)
    rows = [
        [
            1,
            "01/02/2020",
            "Target",
            None,
            "Round opened",
            1,
            1,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "Requested proposals from Alpha.",
            f"\u201c{QUOTES[0]}\u201d (p. 10)",
            "Q1",
            None,
            dt.date(2020, 1, 2),
            dt.date(2020, 1, 2),
            dt.date(2020, 1, 2),
        ],
        [
            2,
            "01/03/2020",
            "Alpha",
            "Strategic",
            "Bid",
            1,
            1,
            10,
            10,
            "Yes",
            "Informal",
            "Heavy",
            1,
            None,
            None,
            "Preliminary bid before substantive diligence.",
            f"\u201c{QUOTES[1]}\u201d (p. 10)",
            "Q1",
            None,
            dt.date(2020, 1, 3),
            dt.date(2020, 1, 3),
            dt.date(2020, 1, 3),
        ],
        [
            3,
            "01/04/2020",
            "Target",
            None,
            "Merger agreement signed",
            1,
            1,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            None,
            "$10 per share, all cash.",
            f"\u201c{QUOTES[2]}\u201d (p. 10)",
            "Q1",
            None,
            dt.date(2020, 1, 4),
            dt.date(2020, 1, 4),
            dt.date(2020, 1, 4),
        ],
    ]
    for row in rows:
        ledger.append(row)
    for row in range(2, 5):
        for column in (20, 21, 22):
            ledger.cell(row, column).number_format = "MM/DD/YYYY"
    finish_sheet(ledger, len(check_lean.LEDGER_COLUMNS))

    rounds = wb.create_sheet("Rounds")
    rounds.append(check_lean.ROUND_COLUMNS)
    rounds.append(
        [
            1,
            1,
            dt.date(2020, 1, 2),
            "Target requested proposals.",
            "1 strategic: Alpha",
            "none stated",
            "No deadline stated",
            "Not final",
            "1: Alpha",
            "Signing",
        ]
    )
    rounds["C2"].number_format = "MM/DD/YYYY"
    finish_sheet(rounds, len(check_lean.ROUND_COLUMNS))

    questions = wb.create_sheet("Questions")
    questions.append(check_lean.QUESTION_COLUMNS)
    questions.append(
        [
            "Q1",
            "Is one process and one round the best map?",
            "Yes; retain one process and one round.",
            "The target requested proposals once before signing (p. 10).",
            "Rows 1-3",
            "A different answer would move all three rows.",
            None,
        ]
    )
    finish_sheet(questions, len(check_lean.QUESTION_COLUMNS))

    facts = wb.create_sheet("Deal facts")
    facts.append(check_lean.FACT_COLUMNS)
    fact_values = [
        "Target",
        "Alpha",
        "Strategic",
        "$10 per share in cash",
        "01/04/2020",
        "01/04/2020",
        "DEFM14A filed 01/10/2020",
        "10",
        "target-led",
        1,
        "Not met: 0 qualifying NDAs",
        "Yes",
        "US dollars per share",
        "Bank",
        "Law Firm",
        "The target opened one round. Alpha submitted one bid. The offer was all cash. The target selected Alpha. The parties signed on January 4.",
    ]
    for field, value in zip(check_lean.FACT_FIELDS, fact_values, strict=True):
        facts.append([field, value])
    finish_sheet(facts, len(check_lean.FACT_COLUMNS))

    wb.save(workbook)
    return workbook, filing


class LeanCheckerTests(unittest.TestCase):
    def test_valid_minimal_fixture_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])
            self.assertEqual(report["summary"]["errors"], 0)

    def test_important_failures_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            ledger = wb["Deal ledger"]
            ledger["E3"] = "Invented event"
            ledger["T3"] = "01/03/2020"
            ledger["T4"] = dt.date(2020, 1, 1)
            ledger["Q4"] = "\u201cOn January 4 ... merger agreement.\u201d (p. 10)"
            wb["Questions"]["E2"] = "Rows 1-99"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            codes = {issue["code"] for issue in report["issues"]}
            self.assertEqual(report["status"], "fail")
            self.assertIn("controlled.event", codes)
            self.assertIn("date.not_excel_date", codes)
            self.assertIn("date.sort_outside_window", codes)
            self.assertIn("date.sort_decreases", codes)
            self.assertIn("quote.not_contiguous_in_filing", codes)
            self.assertIn("questions.unknown_row", codes)

    def test_malformed_required_sheet_fails_gracefully(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal ledger"]["A1"] = "Wrong header"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertIn("schema.columns", {issue["code"] for issue in report["issues"]})

    def test_account_and_quote_wrapper_variants_pass(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            ledger = wb["Deal ledger"]
            ledger["Q2"] = f"{QUOTES[0]} (p. 10)"
            ledger["Q3"] = f"\u201c{QUOTES[1]} (pp. 10\u201311)\u201d"
            ledger["Q4"] = f'"{QUOTES[2]}" p. 10'
            wb["Deal facts"]["A17"] = "Account"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_deal_fact_guidance_labels_and_explanations_pass(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            facts = wb["Deal facts"]
            facts["B4"] = "Strategic - operating-company acquirer"
            facts["A10"] = "Initiation (target-led, bidder-led, activist-influenced, mixed or unclear)"
            facts["B10"] = "Target-led - the board initiated outreach"
            facts["B11"] = 1
            facts["B11"].number_format = "MM/DD/YYYY"
            facts["A12"] = "Auction screen (C1)"
            facts["A13"] = "Whole-company bids (Yes, or No with what was bid for)"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_multiple_flags_and_link_mismatches_are_review_leads(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal ledger"]["R3"] = "Q1, Q2"
            questions = wb["Questions"]
            questions.append(
                [
                    "Q2",
                    "Is Alpha's classification supported?",
                    "Yes; retain the classification.",
                    "The filing describes Alpha as an operating company (p. 10).",
                    "Rows 1",
                    "A different answer would change Alpha's Type.",
                    None,
                ]
            )
            finish_sheet(questions, len(check_lean.QUESTION_COLUMNS))
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            issues = {issue["code"]: issue["severity"] for issue in report["issues"]}
            self.assertEqual(report["status"], "pass_with_warnings", report["issues"])
            self.assertEqual(issues["questions.flag_mismatch"], "warning")
            self.assertEqual(issues["questions.reverse_link"], "warning")
            self.assertNotIn("controlled.flag", issues)

            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal ledger"]["R3"] = "Q1, Q9"
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            errors = {
                issue["code"] for issue in report["issues"] if issue["severity"] == "error"
            }
            self.assertIn("questions.unknown_flag", errors)

    def test_missing_sheets_fail_gracefully(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            filing = directory / "filing.htm"
            filing.write_text("<p>Source text.</p>", encoding="utf-8")
            workbook = directory / "malformed.xlsx"
            Workbook().save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertEqual(report["issues"][0]["code"], "schema.sheets")

    def test_cli_writes_json_for_unreadable_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            workbook = directory / "broken.xlsx"
            workbook.write_text("not an xlsx", encoding="utf-8")
            filing = directory / "filing.htm"
            filing.write_text("<p>Source text.</p>", encoding="utf-8")
            output = directory / "report.json"

            status = check_lean.main(
                ["--workbook", str(workbook), "--filing", str(filing), "--output", str(output)]
            )
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(status, 2)
            self.assertEqual(report["status"], "error")
            self.assertEqual(report["issues"][0]["code"], "input.workbook_unreadable")


if __name__ == "__main__":
    unittest.main()
