"""Grader's mechanical re-check of the blinded workbooks (26 Sep 2026).
Run from this folder: python3 mechanical_check.py
Checks: quotation containment (punctuation-insensitive), <=30-word quotes with a page,
non-decreasing Sort date, required cells on bid rows, Round opened = Rounds lines, Flag<->Question.
Status, 26 Sep 2026: v1.14 only. It has no Note cap and is not a v1.14.1 grader; use checker 1.8
(_dev/tools/check_lean.py) and the retest plan's acceptance script instead."""
import sys, re, json, unicodedata, datetime, os
import openpyxl
from bs4 import BeautifulSoup
HERE=os.path.dirname(os.path.abspath(__file__))
P=os.path.join(HERE,'..','blinded')
RAW=os.path.join(HERE,'..','inputs','raw_filing')
FIL={'mac-gray':'mac-gray_2013-12-04_DEFM14A.htm','providence-worcester':'providence-worcester_2016-09-20_DEFM14A.htm','stec':'stec_2013-08-08_DEFM14A.htm'}
def html_text(path):
    soup=BeautifulSoup(open(path,encoding='utf-8',errors='replace').read(),'lxml')
    for t in soup(['script','style']): t.decompose()
    return soup.get_text(' ')
def norm(s):
    s=unicodedata.normalize('NFKC',s)
    s=s.replace('’',"'").replace('‘',"'").replace('“','"').replace('”','"').replace('–','-').replace('—','-').replace('\xa0',' ')
    s=re.sub(r'[^a-z0-9]+',' ',s.lower())
    return s.strip()
ftext={d:norm(html_text(os.path.join(RAW,f))) for d,f in FIL.items()}
BIDROWS={'Bid','Bid reaffirmed','Other-scope bid'}
REQ=['Stock %','Formality','Conditions','Due diligence','Financing','Regulatory','Exclusivity']
out={}
for deal in FIL:
    for L in 'ABCDE':
        wb=openpyxl.load_workbook(f'{P}/{deal}/{L}.xlsx',data_only=True)
        res={'sheets':wb.sheetnames}
        ws=wb['Deal ledger']
        hdr=[c.value for c in ws[1]]
        idx={h:i for i,h in enumerate(hdr)}
        rows=[r for r in ws.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
        res['n_rows']=len(rows)
        # quote check
        qmiss=[];qlong=0;qnopage=0
        for r in rows:
            q=r[idx['Quote and page']]
            if not q: qmiss.append((r[0],'EMPTY'));continue
            m=re.match(r'^(.*)\((pp?\.\s*[^)]*)\)\s*$',str(q).strip(),re.S)
            if not m: qnopage+=1; body=str(q)
            else: body=m.group(1)
            body=body.strip().strip('“”"').strip()
            if len(body.split())>30: qlong+=1
            nb=norm(body)
            # allow ellipsis splits? instruction says no splicing; check full containment
            if nb and nb not in ftext[deal]:
                qmiss.append((r[0], body[:80]))
        res['quote_mismatch']=qmiss; res['quote_over30w']=qlong; res['quote_nopage']=qnopage
        # sort date monotonic
        sd=[r[idx['Sort date']] for r in rows]
        bad=[rows[i][0] for i in range(1,len(sd)) if sd[i] is None or sd[i-1] is None or sd[i]<sd[i-1]]
        res['sortdate_violations']=bad
        res['sortdate_missing']=[r[0] for r in rows if r[idx['Sort date']] is None]
        # bid rows required fields
        missing=[]
        for r in rows:
            if r[idx['Event']] in BIDROWS:
                for f in REQ:
                    if r[idx[f]] in (None,''):
                        missing.append((r[0],f))
                if r[idx['Event']]=='Other-scope bid' and (r[idx['Price low']] is not None or r[idx['Price high']] is not None):
                    missing.append((r[0],'price on other-scope'))
        res['bidrow_missing_fields']=missing
        # counts
        ev={}
        for r in rows: ev[r[idx['Event']]]=ev.get(r[idx['Event']],0)+1
        res['events']=ev
        # rounds sheet consistency
        rs=wb['Rounds']; rrows=[r for r in rs.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
        res['rounds_lines']=len(rrows)
        res['round_opened_rows']=ev.get('Round opened',0)
        # flags vs questions
        qs=wb['Questions']; qrows=[r for r in qs.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]
        qids={str(r[0]).strip() for r in qrows if r[0]}
        res['n_questions']=len(qids)
        flagged=set()
        for r in rows:
            f=r[idx['Flag']]
            if f:
                for t in re.split(r'[;,]\s*',str(f)): 
                    if t.strip(): flagged.add(t.strip())
        res['flags_without_question']=sorted(flagged-qids)
        res['questions_without_flag']=sorted(qids-flagged)
        # inferred rows
        res['inferred_Y']=sum(1 for r in rows if r[idx['Inferred']]=='Y')
        # note lengths
        nl=[len(str(r[idx['Note']]).split()) for r in rows if r[idx['Note']]]
        res['note_words_mean']=round(sum(nl)/len(nl),1) if nl else None
        res['note_over60']=sum(1 for n in nl if n>60)
        # deal facts fields
        df=wb['Deal facts']; dfrows=[r for r in df.iter_rows(min_row=2,values_only=True) if r[0]]
        res['dealfacts_fields']=len(dfrows)
        out[f'{deal}/{L}']=res
json.dump(out,open(os.path.join(HERE,'mechanical_check_results.json'),'w'),indent=1,default=str)
for k,v in out.items():
    print(f"\n=== {k}: rows={v['n_rows']} events={sum(v['events'].values())} rounds_lines={v['rounds_lines']} round_opened={v['round_opened_rows']} Q={v['n_questions']} inferredY={v['inferred_Y']} note_mean={v['note_words_mean']} note>60={v['note_over60']}")
    print(f"  quote mismatches={len(v['quote_mismatch'])} over30w={v['quote_over30w']} nopage={v['quote_nopage']}; sortdate viol={v['sortdate_violations']} missing={v['sortdate_missing']}")
    print(f"  bidrow missing={v['bidrow_missing_fields']}")
    print(f"  flags w/o Q={v['flags_without_question']} Q w/o flag={v['questions_without_flag']} sheets={v['sheets']}")
    for m in v['quote_mismatch'][:12]: print('    MISMATCH', m)
