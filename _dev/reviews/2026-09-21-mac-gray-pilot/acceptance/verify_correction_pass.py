"""Independent read-only verification of the acceptance-correction pass.

The expected mutations below are transcribed from the lead's frozen packet,
not inferred from the provider's output or revision notes.
"""
from pathlib import Path
from copy import copy
from datetime import datetime
import hashlib
import importlib.util
import json
import re
import zipfile
import openpyxl
from openpyxl.xml.functions import tostring

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[3]
OUT = HERE / "verification"
OUT.mkdir(exist_ok=True)
BEFORE = HERE.parent / "revision/extraction/mac-gray.xlsx"
AFTER = HERE / "extraction/mac-gray.xlsx"
a = openpyxl.load_workbook(BEFORE)
b = openpyxl.load_workbook(AFTER)

def jsonval(v):
    return v.isoformat() if isinstance(v, datetime) else v

def write(name, obj):
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, indent=2, default=str) + "\n")

def cellval(c):
    return [c.data_type, jsonval(c.value)]

def mapped(n):
    return n if n < 54 else n + 3

# Key: sheet, old Excel row, column; value: required replacement and ruling.
expected = {}
def set_expected(sheet, row, col, value, rule):
    expected[(sheet, row, col)] = (value, rule)

def note(n, value, rule):
    set_expected("Deal ledger", n+1, 16, value, rule)

def replace_note(n, old, new, rule):
    text = a["Deal ledger"].cell(n+1, 16).value
    assert text.count(old) == 1
    note(n, text.replace(old, new), rule)

set_expected("Deal ledger",50,2,"by 09/25/2013","A02")
set_expected("Deal ledger",50,21,None,"A02")
note(49,"Buyer outside legal counsel. Received Goodwin's first merger draft 09/25; also negotiated family voting agreements during 09/24–09/27 (pp. 38–39), so earlier involvement is possible. Engagement day is not stated. Later negotiated Moab's voting agreement.","A02")
set_expected("Deal ledger",4,17,"“representatives of BofA Merrill Lynch, as instructed by the Board, telephoned a representative of Party A to discuss generally a possible business combination” (p. 27)","A03")
replace_note(8,"Unsolicited written proposal","Unsolicited proposal","A04")
replace_note(23,"Written preliminary indication","Preliminary indication","A04")
note(29,"BofA letter to all four bidders (authorized 08/15): revised written proposals due 09/09/2013 to continue to in-depth diligence. Follow-up diligence: Party B telephone call 08/27; Party C in-person session 08/29; Party A telephone call 09/03 (p. 35).","A04")
replace_note(19,"Communicated to CSC/Pamplona (by BofA) and Rothenberg (by Goodwin).","Instructed BofA and Goodwin to convey the decision to CSC/Pamplona and Rothenberg, respectively.","A05")
note(7,"Moab won a proxy contest: its nominees Rothenberg (Moab general partner) and Hyman elected. Moab owned about 9%. Possible exclusion for rollover conflicts was discussed 06/12; Rothenberg was excluded from the sale review on 06/24 (pp. 29–30).","A06")
replace_note(46,"It then got full data-room and management access for confirmatory diligence.","During 09/25–10/07 it got full data-room and management access for confirmatory diligence.","A07")
replace_note(47,"Its $18.00–$19.00 best and final was below CSC/Pamplona's $20.75; not mentioned again.","Its $18.00–$19.00 best and final was below CSC/Pamplona's $20.75 on 09/18 and $21.25 on 09/21; not mentioned again.","A08")
for row,value in {
    5:"No supported alternative deadline outcome under E9: the filing reports late bids that the target considered. Party C's later 07/25 written revision is assigned to round 1 because it answers that solicitation; its actual arrival follows the round-2 opening.",
    7:"No supported alternative under E9: the 09/10 submissions arrived after the 09/09 due date and were considered. Ignoring a one-day delay would change the fixed convention.",
    9:"No supported alternative under E9: the target acted on the bids in hand on 09/19. Subsequent price negotiation with the preferred bidder does not undo that action."
}.items():
    set_expected("Questions",row,6,value,"A09")
set_expected("Questions",6,3,"Unresolved timing. Account for the cohort once by an inferred Dropped by target by 09/11/2013 (#37), a conservative final-stage bound. Leave individual signing dates, 07/23 eligibility and actual exit dates unknown; do not treat the bound as observed timing or a revealed exit reason.","A10")
v = a["Deal facts"]["B18"].value
assert v.count("and Party C dropped out.") == 1
set_expected("Deal facts",18,2,v.replace("and Party C dropped out.","and Party C did not submit or reiterate its earlier offer."),"A11")

assert a.sheetnames == b.sheetnames == ["Deal ledger","Rounds","Questions","Deal facts"]
diff=[]; errors=[]; styles=[]; checked=0; matched_expected=set()
for sa in a:
    sb=b[sa.title]
    assert sa.max_column==sb.max_column
    assert sb.max_row == sa.max_row + (3 if sa.title=="Deal ledger" else 0)
    for row in sa:
        for ca in row:
            rb=ca.row if sa.title!="Deal ledger" or ca.row==1 else mapped(ca.row-1)+1
            cb=sb.cell(rb,ca.column)
            checked+=1
            key=(sa.title,ca.row,ca.column)
            ev,rule=expected.get(key,(ca.value,"unchanged"))
            if key in expected:
                matched_expected.add(key)
            if sa.title=="Deal ledger" and ca.column==1 and ca.row>1:
                ev=mapped(ca.value)
                if ev!=ca.value:rule="dependent_event_number"
            if isinstance(ev,str):
                revised=re.sub(r"#(\d+)",lambda m:"#"+str(mapped(int(m.group(1)))),ev)
                if revised!=ev:
                    ev=revised
                    if rule=="unchanged":rule="dependent_event_reference"
            if cb.value!=ev or (ev is not None and type(cb.value)!=type(ev)):
                errors.append({'sheet':sa.title,'before_cell':ca.coordinate,'after_cell':cb.coordinate,'expected':jsonval(ev),'actual':jsonval(cb.value),'rule':rule})
            if cellval(ca)!=cellval(cb):
                diff.append({'sheet':sa.title,'before_cell':ca.coordinate,'after_cell':cb.coordinate,'before':cellval(ca),'after':cellval(cb),'disposition':rule,'verified':cb.value==ev})
            style_keys=['font','fill','border','alignment','protection','number_format']
            mismatch=[k for k in style_keys if copy(getattr(ca,k))!=copy(getattr(cb,k))]
            if mismatch:styles.append({'sheet':sa.title,'before_cell':ca.coordinate,'after_cell':cb.coordinate,'mismatch':mismatch})
            for k in ['hyperlink','comment']:
                if getattr(ca,k)!=getattr(cb,k):errors.append({'cell':cb.coordinate,'sheet':sa.title,'unexpected_changed_attribute':k})
    for k in ['freeze_panes','sheet_format','sheet_view','sheet_properties','sheet_state','merged_cells','data_validations','print_options','page_margins','page_setup','print_area','print_title_rows','print_title_cols']:
        if getattr(sa,k)!=getattr(sb,k):errors.append({'sheet':sa.title,'unexpected_changed_control':k})
    # ConditionalFormattingList has identity equality; compare serialized rules.
    assert [tostring(x.to_tree()) for x in sa.conditional_formatting] == [tostring(x.to_tree()) for x in sb.conditional_formatting]
    assert {k:str(v) for k,v in sa.column_dimensions.items()}=={k:str(v) for k,v in sb.column_dimensions.items()}
    expected_filter='A1:V59' if sa.title=='Deal ledger' else sa.auto_filter.ref
    assert sb.auto_filter.ref==expected_filter
assert matched_expected==set(expected)

quote="“Except for Morgan Stanley & Co. LLC, Deutsche Bank Securities Inc. and Evercore Group L.L.C., there is no investment banker, broker, finder or other agent or intermediary” (p. A-30)"
adviser_note="Named buyer-side intermediary in Annex A §5.11's transaction-fee clause, acting for Parent or its affiliates. Specific role and engagement day are not disclosed; relationship is established by the 10/14/2013 signed agreement."
added=[]
for n,who in zip([54,55,56],["Morgan Stanley & Co. LLC","Deutsche Bank Securities Inc.","Evercore Group L.L.C."]):
    vals=[n,"by 10/14/2013",who,None,"Adviser",1,3,None,None,None,None,None,None,None,None,adviser_note,quote,None,None,datetime(2013,10,14),None,datetime(2013,10,14)]
    for c,ev in enumerate(vals,1):
        cb=b['Deal ledger'].cell(n+1,c)
        if cb.value!=ev:errors.append({'new_event':n,'column':c,'expected':jsonval(ev),'actual':jsonval(cb.value)})
        template=a['Deal ledger'].cell(50,c)
        for k in ['font','fill','border','alignment','protection','number_format']:
            if copy(getattr(template,k))!=copy(getattr(cb,k)):styles.append({'new_event':n,'column':c,'style':k})
    added.append({'event':n,'cells':{b['Deal ledger'].cell(n+1,c).coordinate:jsonval(v) for c,v in enumerate(vals,1)},'disposition':'A01 source checked Annex A5.11, date bound from October14 signed agreement; no bidder entry','verified':True})

spec=importlib.util.spec_from_file_location('cockpit_data',ROOT/'_dev/tools/cockpit/data.py')
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
f=m.Filing((ROOT/'raw_filing/mac-gray_2013-12-04_DEFM14A.htm').read_bytes())
quotes=[]
for row in b['Deal ledger'].iter_rows(min_row=2):
    q,p=m.check_lean.parse_quote_and_page(row[16].value)
    located=f.locate(q,p)
    rec={'event':row[0].value,'word_count':len(q.split()),'cited_page':p,'found_page':located.get('found_page'),'located':located.get('located')}
    quotes.append(rec)
    assert rec['word_count']<=30 and rec['located'] and rec['cited_page']==rec['found_page'],rec

references=[]
for s in b:
    for row in s:
        for c in row:
            if isinstance(c.value,str):
                for n in re.findall(r'#(\d+)',c.value):
                    assert 1<=int(n)<=58
                    references.append({'sheet':s.title,'cell':c.coordinate,'event':int(n)})
flags={}
for r in b['Deal ledger'].iter_rows(min_row=2,values_only=True):
    for q in re.findall(r'Q\d+',r[17] or ''):flags.setdefault(q,set()).add(r[0])
for r in b['Questions'].iter_rows(min_row=2,values_only=True):
    assert set(map(int,re.findall(r'#(\d+)',r[4])))==flags.get(r[0],set())

bidcols=[2,3,4,5,6,7,8,9,10,11,12,13,14,15,20,21,22]
oldbids=[r for r in a['Deal ledger'].iter_rows(min_row=2) if r[4].value=='Bid']
newbids=[r for r in b['Deal ledger'].iter_rows(min_row=2) if r[4].value=='Bid']
assert len(oldbids)==len(newbids)==13
assert [[cellval(r[c-1]) for c in bidcols] for r in oldbids]==[[cellval(r[c-1]) for c in bidcols] for r in newbids]

with zipfile.ZipFile(BEFORE) as za,zipfile.ZipFile(AFTER) as zb:
    assert za.namelist()==zb.namelist()
    changed_parts=[n for n in za.namelist() if za.read(n)!=zb.read(n)]
    assert set(changed_parts)<= {'docProps/core.xml','xl/worksheets/sheet1.xml','xl/worksheets/sheet3.xml','xl/worksheets/sheet4.xml','xl/workbook.xml'},changed_parts

write('cell-diff.json',diff);write('added-rows.json',added);write('quote-pages.json',quotes);write('event-references.json',references)
write('comparison.json',{'before_sha256':hashlib.sha256(BEFORE.read_bytes()).hexdigest(),'after_sha256':hashlib.sha256(AFTER.read_bytes()).hexdigest(),'existing_cells_checked':checked,'changed_existing_cells':len(diff),'added_rows':3,'added_cells_checked':66,'unexpected_content_or_control_changes':errors,'resolved_style_changes':styles,'changed_package_parts':changed_parts,'bid_rows_unchanged':13,'all_quotes_found_on_cited_page':len(quotes),'internal_reference_occurrences_valid':len(references),'flags_equal_question_row_lists':True,'R01':'not decided; no economic-term bids added'})
lines=['# Complete acceptance-correction diff','', '| Sheet | Before cell | After cell | Ruling | Before | After |','| --- | --- | --- | --- | --- | --- |']
def esc(v):return str(v).replace('|','\\|').replace('\n','<br>')
for d in diff:lines.append('| '+' | '.join(esc(d[k]) for k in ['sheet','before_cell','after_cell','disposition','before','after'])+' |')
lines+=['','Three new rows are fully itemized in `added-rows.json`; all 66 fields checked against A01.','']
(OUT/'CELL_DIFF.md').write_text('\n'.join(lines))
print(json.dumps({'existing_cells_checked':checked,'changed_cells':len(diff),'new_rows':3,'errors':errors,'style_changes':styles,'quotes':len(quotes),'references':len(references)},indent=2))
assert not errors and not styles
