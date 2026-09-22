#!/usr/bin/env python3
"""Blind the sweep's completed workbooks for grading, and score graded bundles.

blind:  every completed cell's workbook becomes <out>/bundles/<label>/workbook.json, a dump of
        its four sheets with Excel row numbers, under a random label. Two workbooks are dumped a
        second time under fresh labels to measure grader noise. The grading directory is
        self-contained: copies of the references, filings and instruction sit beside the bundles,
        and each grader gets a private work directory, so graders never read the repository (where
        run receipts name the arms) or each other's files. The label key is written outside the
        grading directory and never shown to graders.
score:  turns per-grader grade files into per-bundle scores (pass 1, partial 0.5, fail 0, times
        each test's weight), then unblinds with the key into the packet's grades.json.
audit:  lists every file path each grader's tool calls touched (Astra's Codex event logs and the
        Claude workflow transcripts) and flags any that reach outside the grader's own bundle,
        the shared references, filings and instruction, and its own work directory.
"""

import argparse
import csv
import datetime as dt
import hashlib
import json
from pathlib import Path
import random
import re
import secrets
import shutil
import statistics

from openpyxl import load_workbook

REPO = Path("/home/uctpiaj/work/Projects/sec-extraction")
SHEETS = ["Deal ledger", "Rounds", "Questions", "Deal facts"]
GRADERS = ("claude1", "claude2", "astra")
CREDIT = {"pass": 1.0, "partial": 0.5, "fail": 0.0}


def cell(value):
    if isinstance(value, (dt.datetime, dt.date)):
        return value.strftime("%m/%d/%Y")
    return value


def dump(workbook: Path) -> dict:
    book = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheets = {}
        for name in SHEETS:
            rows = list(book[name].iter_rows(values_only=True)) if name in book.sheetnames else []
            header = [cell(v) for v in rows[0]] if rows else []
            body = []
            for number, row in enumerate(rows[1:], start=2):
                if any(v not in (None, "") for v in row):
                    body.append({"excel_row": number, **{str(h): cell(v) for h, v in zip(header, row) if v not in (None, "")}})
            sheets[name] = {"header": header, "rows": body}
        return {"sheets": sheets, "missing_sheets": [n for n in SHEETS if n not in book.sheetnames]}
    finally:
        book.close()


def blind(args):
    packet, out = Path(args.packet), Path(args.out)
    bundles = out / "bundles"
    bundles.mkdir(parents=True, exist_ok=False)
    key = {}
    cells = []
    for receipt_path in sorted((packet / "runs").glob("*/receipt.json")):
        receipt = json.loads(receipt_path.read_text())
        if receipt["state"] != "completed" or receipt["mismatches"]:
            continue
        cells.append((receipt["cell"], receipt_path.parent / f"{receipt['cell']['deal']}.xlsx"))
    rng = random.Random(secrets.randbits(64))
    duplicates = []
    for provider in ("opus", "sol"):
        candidates = [c for c in cells if c[0]["provider"] == provider]
        if candidates:
            duplicates.append(rng.choice(candidates))
    for cell_info, workbook in cells + duplicates:
        label = secrets.token_hex(4)
        while label in key:
            label = secrets.token_hex(4)
        (bundles / label).mkdir()
        (bundles / label / "workbook.json").write_text(json.dumps(dump(workbook), indent=1, ensure_ascii=False))
        key[label] = {"cell": cell_info["id"], "deal": cell_info["deal"], "arm": cell_info["arm"],
                      "duplicate": (cell_info, workbook) in duplicates and any(v["cell"] == cell_info["id"] for v in key.values())}
    order = list(key)
    rng.shuffle(order)
    with (REPO / "raw_filing" / "MANIFEST.csv").open(newline="") as stream:
        filings = {row["deal"]: row["file"] for row in csv.DictReader(stream)}
    for folder in ("reference", "filings", "work"):
        (out / folder).mkdir()
    copied = {}
    for deal in sorted({v["deal"] for v in key.values()}):
        sources = {out / "reference" / f"{deal}.json": packet / "reference" / f"{deal}.json",
                   out / "filings" / f"{deal}.htm": REPO / "raw_filing" / filings[deal]}
        for target, source in sources.items():
            shutil.copyfile(source, target)
            copied[str(target.relative_to(out))] = hashlib.sha256(source.read_bytes()).hexdigest()
    shutil.copyfile(REPO / "SEC_Deal_Ledger_Extraction_Instruction.md", out / "instruction.md")
    copied["instruction.md"] = hashlib.sha256((out / "instruction.md").read_bytes()).hexdigest()
    for label in key:
        for grader in GRADERS:
            (out / "work" / f"{label}.{grader}").mkdir()
    Path(args.key).parent.mkdir(parents=True, exist_ok=True)
    Path(args.key).write_text(json.dumps(key, indent=1))
    (out / "inputs.json").write_text(json.dumps(copied, indent=1, sort_keys=True))
    (out / "queue.json").write_text(json.dumps([{"label": label, "deal": key[label]["deal"]} for label in order], indent=1))
    print(f"{len(key)} bundles ({len(duplicates)} duplicate controls); key at {args.key}")


def score_one(grades: list, tests: list) -> tuple[float, list]:
    by_id = {g["id"]: g for g in grades}
    missing = [t["id"] for t in tests if t["id"] not in by_id]
    total = sum(t["weight"] * CREDIT.get(by_id.get(t["id"], {}).get("grade"), 0.0) for t in tests)
    return round(total, 2), missing


def score(args):
    packet, out = Path(args.packet), Path(args.out)
    key = json.loads(Path(args.key).read_text())
    references = {deal: json.loads((packet / "reference" / f"{deal}.json").read_text())["tests"]
                  for deal in {v["deal"] for v in key.values()}}
    per_bundle = {}
    for label, info in key.items():
        tests = references[info["deal"]]
        graders = {}
        for grade_file in sorted((out / "grades").glob(f"{label}.*.json")):
            grader = grade_file.name.split(".")[1]
            total, missing = score_one(json.loads(grade_file.read_text())["grades"], tests)
            graders[grader] = {"score": total, "missing": missing}
        claude = [g["score"] for name, g in graders.items() if name.startswith("claude")]
        astra = [g["score"] for name, g in graders.items() if name.startswith("astra")]
        balanced = None
        if claude and astra:
            balanced = round((statistics.mean(claude) + statistics.mean(astra)) / 2, 2)
        per_bundle[label] = {**info, "graders": graders, "claude_mean": round(statistics.mean(claude), 2) if claude else None,
                             "astra": astra[0] if astra else None, "score": balanced}
    (out / "bundle-scores.json").write_text(json.dumps(per_bundle, indent=1))
    grades = {}
    for label, entry in per_bundle.items():
        if not entry["duplicate"]:
            grades[entry["cell"]] = {k: entry[k] for k in ("score", "claude_mean", "astra", "graders")} | {"label": label}
    (packet / "grades.json").write_text(json.dumps(grades, indent=1, sort_keys=True) + "\n")
    print(f"scored {len(per_bundle)} bundles; wrote {packet / 'grades.json'}")


def tool_calls(args) -> dict:
    """(label, grader) -> list of tool-call texts, from Astra event logs and Claude transcripts."""
    out = Path(args.out).resolve()
    calls = {}
    for raw in sorted((out / "grades" / "raw").glob("*.astra.attempt*.jsonl")):
        texts = calls.setdefault((raw.name.split(".")[0], "astra"), [])
        for line in raw.read_text().splitlines():
            try:
                event = json.loads(line)
            except ValueError:
                continue
            item = event.get("item") or {}
            if event.get("type") == "item.completed" and item.get("type") != "agent_message":
                texts.append(json.dumps({k: v for k, v in item.items() if k not in ("aggregated_output", "text")}))
    for folder in args.transcripts:
        for meta in sorted(Path(folder).glob("agent-*.meta.json")):
            match = re.fullmatch(r"grade:([0-9a-f]{8})#(\d)", json.loads(meta.read_text()).get("description", ""))
            if not match:
                continue
            texts = calls.setdefault((match[1], f"claude{match[2]}"), [])
            for line in meta.with_name(meta.name.replace(".meta.json", ".jsonl")).read_text().splitlines():
                message = json.loads(line).get("message")
                if isinstance(message, dict) and isinstance(message.get("content"), list):
                    texts += [f"{c['name']}: {json.dumps(c['input'])}" for c in message["content"]
                              if c.get("type") == "tool_use" and c["name"] != "StructuredOutput"]
    return calls


def audit(args):
    out = Path(args.out).resolve()
    labels = set(json.loads(Path(args.key).read_text()))
    shared = [str(out / "reference"), str(out / "filings"), str(out / "instruction.md")]
    report, flagged = {}, 0
    for (label, grader), texts in sorted(tool_calls(args).items()):
        allowed = shared + [str(out / "bundles" / label), str(out / "work" / f"{label}.{grader}")]
        paths = sorted({m for text in texts for m in re.findall(r"/home/[^\s\"'\\;|&)<>]+", text)})
        # The bare grading root is allowed (graders cd there); a relative path from it to another
        # bundle or work directory is caught by the label check below.
        outside = [p for p in paths if p.rstrip("/") != str(out) and not any(p.startswith(a) for a in allowed)
                   and "/tool-results/" not in p]  # Claude Code's store of long tool outputs
        others = sorted({m for text in texts for m in re.findall(r"\b[0-9a-f]{8}\b", text)} & (labels - {label}))
        relative = sorted({m for text in texts for m in re.findall(r"(?:\.\./|_dev/|receipt|key\.json|/runs/)\S*", text)})
        report[f"{label}.{grader}"] = {"calls": len(texts), "outside": outside, "other_labels": others, "suspect_relative": relative}
        if outside or others or relative:
            flagged += 1
            print(f"{label}.{grader}: outside={outside} other_labels={others} relative={relative}")
    (out / "access-audit.json").write_text(json.dumps(report, indent=1, sort_keys=True))
    print(f"audited {len(report)} grader sessions; {flagged} flagged; wrote {out / 'access-audit.json'}")


if __name__ == "__main__":
    root = argparse.ArgumentParser()
    sub = root.add_subparsers(dest="command", required=True)
    for name, func in (("blind", blind), ("score", score), ("audit", audit)):
        p = sub.add_parser(name)
        p.add_argument("--packet", required=name != "audit")
        p.add_argument("--out", required=True)
        p.add_argument("--key", required=True)
        if name == "audit":
            p.add_argument("--transcripts", nargs="*", default=[], help="Claude workflow transcript directories")
        p.set_defaults(func=func)
    parsed = root.parse_args()
    parsed.func(parsed)
