# -*- coding: utf-8 -*-
"""Dump workbook sheets for review."""
from openpyxl import load_workbook
WB = '/home/uctpiaj/work/extraction/petsmart.xlsx'
wb = load_workbook(WB)
mode = __import__('sys').argv[1] if len(__import__('sys').argv) > 1 else 'ledger'
if mode == 'ledger':
    ws = wb['Deal ledger']
    H = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(H, r))
        print(f"--- #{d['#']} {d['Row id']} | {d['When']} | {d['Who']} | {d['What happened']} | P{d['Process']}/R{d['Round']} | type={d['Type']} | fmt={d['Formality']} cond={d['Conditions level']} | count={d['Count']} | dates={d['Date from']}..{d['Date to']} wd={d['Working date']} ({d['Date basis']}/{d['Date method']})")
        print(f"    terms: {d['Terms or outcome']}")
        print(f"    why: {d['Why and evidence']}")
        print(f"    src: {d['Source']} | page={d['Page']} | rel: {d['Related rows']} | review={d['Review']}")
        extra = {k: d[k] for k in ('Price low','Price high','Price kind','Price origin','Currency','All cash','Cash at closing','Conditions detail','Due date','Deadline treatment','Round finality','Decided by','Exit reason','Outcome basis') if d[k] not in (None, '')}
        if extra: print(f"    extra: {extra}")
elif mode == 'summary':
    ws = wb['Summary']
    for r in ws.iter_rows(values_only=True):
        a, b = (list(r) + [None, None])[:2]
        c = r[2] if len(r) > 2 else None
        if a is None and b is None: print()
        else:
            print(f"[{a}] {b}")
            if c: print(f"    formula: {c}")
elif mode == 'questions':
    ws = wb['Questions']
    for r in ws.iter_rows(values_only=True):
        cells = [str(x) for x in r if x is not None]
        print(' || '.join(cells))
        print()
