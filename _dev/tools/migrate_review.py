#!/usr/bin/env python3
"""Move reviewed v1.13.2 work onto v1.14.1 runs, deal by deal: register, aligner, triage and port batch.

    python3 _dev/tools/migrate_review.py register [deal ...]
    python3 _dev/tools/migrate_review.py triage <deal> <copy of the run's workbook> --run-id <version id> [--rules v1.14]

`register` writes, per deal, the latest working-copy snapshot's Deal ledger rows with their row
marks, and its reviewed facts in Part A's order (participation and exits; the round map and
deadline outcomes; each bid's price, Formality and terms; the order of events). Each fact carries
its filing key (page, cockpit block, quote), the reviewed value, its basis (reported, inference or
convention), its status (supported, reverted or unresolved) and its impact under the target instruction,
v1.14.1 (unchanged, re-judge under named decisions, or a new column with no reviewed value). The
re-judge tags are V114_SPEC's D1–D27 and v1.14.1's changes: R1–R6, V1141_SPEC D1–D6 (written
"v1.14.1 D1" …) and the structural cuts (IMPACT_TAGS_V1141).

`triage` aligns a v1.14.1 run's rows (or a v1.14 run's, with --rules v1.14) to the register (event family, Who, overlapping date windows
and price with or without the CVR; unique matches only), sorts the facts into four buckets and
prepares a port batch in the cockpit edit API's format for the accepted items.

Reads only: the cockpit database through a read-only (mode=ro) URI, a copy of the run's workbook,
the filings in raw_filing/ (for block ids) and the independent audit's files in lesson/. Writes only
under --out/<deal>/. Nothing here applies an edit: Austin applies the port batch, or has it applied,
after the rebase.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import io
import json
import random
import re
import sqlite3
import sys
import unicodedata
from pathlib import Path
from typing import Any
from urllib.parse import quote

import openpyxl

import check_lean
import diff_workbooks
from cockpit import data

PROJECT = Path(__file__).resolve().parents[2]
STATE_DB = PROJECT / "_dev/cockpit/state/workspace.sqlite3"
AUDIT_DIR = PROJECT / "lesson/independent-audit-2026-09-23"
FILINGS = PROJECT / "raw_filing"
OUT = PROJECT / "_dev/maintenance/2026-09-24-bid-terms-taxonomy/migration"
LEDGER, ROUNDS, QUESTIONS, FACTS = data.LEDGER_SHEET, data.ROUNDS_SHEET, data.QUESTIONS_SHEET, data.FACTS_SHEET
MAX_OPERATIONS = 100  # one cockpit save takes at most 100 operations (workspace.edit)

ENTRY_EVENTS = {"Target interest", "Bidder interest", "Contact", "NDA signed", "Re-entered", "Bidding group changed"}
FAMILY = {
    **{event: "bid" for event in check_lean.BID_EVENTS},
    **{event: "exit" for event in check_lean.EXIT_EVENTS},
    **{event: "entry" for event in ENTRY_EVENTS},
    **{event: "round" for event in ("Round opened", "Deadline set", "Deadline revised", "Deadline")},
    **{event: "adviser" for event in ("Adviser", "Adviser ended", "Activist")},
    **{event: "public" for event in ("Sale process announced", "Bid announced", "Merger announced", "Merger agreement signed")},
    **{event: "process" for event in ("Target sale decision", "Process terminated", "Process restarted", "Go-shop changed")},
}  # anything else (Other material event, Exclusivity changed) is "other"

# The fields each fact kind compares. "terms" holds the 29-column ledger's columns that have no reviewed value.
FACT_FIELDS = {
    "participation": ("Event", "Type", "Count", "Exit reason"),
    "deal fact": ("Value",),
    "round": ("Process", "Round", "Opened", "Who was in", "Finality"),
    "deadline": ("Due dates", "Deadline outcome"),
    "assignment": ("Process", "Round"),
    "price": ("Price low", "Price high"),
    "stock": ("All cash",),
    "formality": ("Formality",),
    "conditions": ("Conditions",),
    "terms": (),
    "order": ("When", "Sort date", "Date from", "Date to"),
}
SECTIONS = ("participation and exits", "round map and deadline outcomes", "bids", "order of events")
NEW_TERM_COLUMNS = ("CVR/earnout", "CVR/earnout value", "Due diligence", "Financing", "Regulatory", "Antitrust", "Exclusivity")
DEAL_FACT_SECTIONS = {"Initiation": 0, "Auction screen": 0, "Whole-company bids": 0, "Number of processes": 1}
ROUND_KINDS = {"round", "deadline", "assignment"}
TEXT_FIELDS = {"Note", "Quote and page", "Flag", "Reviewer note", "Event", "Who", "Inferred", ""}
LEDGER_FIELD_KINDS = {
    "Type": {"participation"}, "Count": {"participation"}, "Exit reason": {"participation"},
    "Process": {"assignment"}, "Round": {"assignment"},
    "Price low": {"price"}, "Price high": {"price"}, "All cash": {"stock"},
    "Formality": {"formality"}, "Conditions": {"conditions"},
    "When": {"order"}, "Sort date": {"order"}, "Date from": {"order"}, "Date to": {"order"},
}
# The open convention questions of the audit's README §2 and the facts each one bears on.
CONVENTION_KINDS = {
    "A1": ROUND_KINDS, "A2": ROUND_KINDS, "A3": ROUND_KINDS, "A4": ROUND_KINDS,
    "A5": {"formality"}, "A6": {"formality", "conditions"}, "A7": {"conditions"}, "A8": {"conditions"}, "A9": {"conditions"},
    "A10": {"participation"}, "A11": {"participation"}, "A12": {"participation"}, "A13": {"participation"},
    "A14": {"participation"}, "A16": {"deadline", "order"}, "A17": {"price"},
}  # any other id (A15, Q9, SYN-…) bears on every fact of the rows it names
AUDIT_DEALS = {"mac-gray": "mac-gray", "providence": "providence-worcester", "p&w": "providence-worcester",
               "petsmart": "petsmart", "stec": "stec", "penford": "penford", "synacor": "synacor",
               "kraton": "kraton", "meredith": "meredith"}

BUCKETS = ("agrees", "differs, v1.14.1 changed the rule", "differs, rule unchanged", "omitted or inserted")
BUCKET_ACTIONS = ("accept its accept list after the seeded spot check; its review list stays a review item", "check the source under v1.14.1",
                  "check the source: a regression or an earlier review error", "check completeness")
# Aligner passes, strictest first: (what must match: the Event itself or its family; date slack in days;
# price must agree where both rows have one). A pass accepts only rows with a single candidate each way.
PASSES = (("event", 0, True), ("family", 0, True), ("family", 3, True), ("family", 0, False))
# The agrees bucket's two lists. Only a supported fact whose v1.14.1 impact is unchanged may be accepted on agreement;
# an agreement never settles an unresolved or reverted fact, a Needs decision row, a rule v1.14.1 changed or a new column.
AGREES_LISTS = ("accept", "review")
# The v1.14.1 changes (V1141_SPEC §3–§4, CHANGELOG_v1.14.1_candidate.md) and the facts each tag marks for re-judging.
# V1141_SPEC's D1–D6 carry the "v1.14.1 " prefix so they never read as V114_SPEC's D1–D27. Two change no fact the
# register holds: v1.14.1 D4 (five Questions; Questions are not facts) and v1.14.1 D5 (the nine exit reasons stay).
IMPACT_TAGS_V1141 = {
    "R1": "Same offer copies the earlier row, 'Same as #n'; Bid reaffirmed only after definitive negotiation: price, "
          "Formality and Conditions of Bid reaffirmed rows and same-price revisions",
    "R2": "the evidence window up to the bidder's next row, forecasts counting: every Conditions value",
    "R3": "unnamed cohort closure, Count = total minus named rows, no ranges: participation and dates of cohort rows",
    "R4": "H2 only for a period tied to diligence alone: Conditions whose Note names diligence or a period",
    "R5": "not invited into a stage: Dropped by target at the opening; Re-entered; Would not improve earlier offer: "
          "Dropped by target, Did not submit, Re-entered and inferred exit rows, and their dates",
    "R6": "Due diligence Not begun only before an NDA: Conditions whose Note names diligence or an NDA",
    "v1.14.1 D1": "a revision, price-only included, is Formal only if it meets a route itself: Formality of revisions",
    "v1.14.1 D2": "a condition on proceeding is a Bid, H3, price blank: Conditions naming one, and Other material "
                  "event rows that may now be such a Bid",
    "v1.14.1 D3": "the second Light route is gone (committed financing with silent diligence is Unclear): Conditions Light",
    "v1.14.1 D6": "E6 trigger (d), 30 days or an ended exclusivity: every round line, assignment and Round opened date",
    "not invited: Dropped by target": "Who was in lists invited bidders only; the rest get exit rows: every round line",
    "Formality routes": "E11's narrower routes, Unclear only on a cohort whose members differ: Formality Formal or "
                        "Unclear, and every Bid reaffirmed",
    "Sort-date ladder": "E8's new ladder: the dates of rows without a single reported day, and of inferred rows",
    "bidders' advisers to Notes": "D2 Adviser rows are the target's only: Adviser rows that name a bidder",
    "Contact vs interest": "Target and Bidder interest before round 1 opens, Contact after: those three events",
    "Other material event list": "D2's closed list; anything else goes in a Note: every Other material event row",
    "Initiation from the first row": "D5 Initiation derived from the earliest ledger rows; mixed and unclear dropped",
    "express incorporation removed": "a revision's cells come from its own communication; only a Same-offer row "
                                     "copies (E10, Part B): Formality and Conditions of revisions and Bid reaffirmed",
    "E5 process test": "E5's three-part test with 90 days: Number of processes, Process terminated and restarted rows",
}
ADVISER_EVENTS = {"Adviser", "Adviser ended"}
BIDDER_WORDS = re.compile(r"\b(bidder|buyer|acquir\w*|sponsor|purchaser|parent|lender)s?\b")

SUFFIXES = {"inc", "llc", "lp", "llp", "corp", "corporation", "co", "ltd", "plc", "the"}
NOT_ALIASES = {"incl", "including", "approx", "approximately", "other", "unnamed", "later", "from", "via", "and", "e", "g"}


# -- reading ------------------------------------------------------------------------------------

def connect(path: Path | str) -> sqlite3.Connection:
    """The cockpit database, read-only: the one place this module opens SQLite."""
    conn = sqlite3.connect(f"file:{quote(str(Path(path).resolve()))}?mode=ro", uri=True)
    conn.row_factory = sqlite3.Row
    return conn


def _decode(value: Any) -> Any:
    """A snapshot cell ({"$type": "datetime", …} for dates) as a plain value; dates become dt.date."""
    if isinstance(value, dict) and value.get("$type") in ("datetime", "date"):
        value = dt.datetime.fromisoformat(value["value"])
    elif isinstance(value, dict) and value.get("$type") == "formula":
        value = value["value"]
    return value.date() if isinstance(value, dt.datetime) else value


def show(value: Any) -> str:
    return data.display_value(value).strip()


def working_copy(conn: sqlite3.Connection, slug: str) -> dict[str, Any]:
    row = conn.execute("SELECT * FROM revisions WHERE slug=? ORDER BY revision DESC LIMIT 1", (slug,)).fetchone()
    if row is None:
        raise SystemExit(f"{slug}: no working-copy revision")
    state = json.loads(row["snapshot"])
    for part in state["sheets"].values():
        for record in part["rows"]:
            record["values"] = {key: _decode(value) for key, value in record["values"].items()}
    meta = {key: row[key] for key in ("revision", "base_id", "base_sha256", "at", "actor", "reason", "summary")}
    return {"meta": meta, "state": state}


def threads(conn: sqlite3.Connection, slug: str) -> list[dict[str, Any]]:
    try:
        rows = conn.execute("SELECT id,target_kind,target_sheet,target_uid,target_label,resolved_at FROM threads WHERE slug=? ORDER BY created_at", (slug,)).fetchall()
    except sqlite3.OperationalError:
        return []
    return [{**dict(row), "resolved": row["resolved_at"] is not None} for row in rows]


def load_run(path: Path, slug: str, rules: str | None = None) -> dict[str, Any]:
    """A run workbook's sheets, with the row uids the cockpit gives them once it is a deal's base. A 29-column run is
    read under the rules asked for (v1.14 or v1.14.1), else check_lean's default, v1.14.1."""
    raw = Path(path).read_bytes()
    sha = hashlib.sha256(raw).hexdigest()
    wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    sheets = {}
    for sheet in (LEDGER, ROUNDS, QUESTIONS, FACTS):
        ws = wb[sheet]
        columns = [data.display_value(cell.value) for cell in ws[1]]
        while columns and not columns[-1]: columns.pop()
        rows = []
        for excel_row in range(2, ws.max_row + 1):  # as workspace._base_state reads a base
            cells = [ws.cell(excel_row, column + 1) for column in range(len(columns))]
            if all(check_lean.is_blank(cell.value) for cell in cells): continue
            rows.append({"uid": hashlib.sha256(f"{slug}|{sha}|{sheet}|{excel_row}".encode()).hexdigest()[:24], "excel_row": excel_row,
                         "values": {col: _decode(cell.value) for col, cell in zip(columns, cells) if col}})
        sheets[sheet] = {"columns": columns, "rows": rows}
    wb.close()
    return {"sha256": sha, "schema": check_lean.schema_for_header(sheets[LEDGER]["columns"], rules), "sheets": sheets}


def filing_for(slug: str, filings: Path) -> data.Filing | None:
    manifest = filings / "MANIFEST.csv"
    if not manifest.is_file():
        return None
    with manifest.open(newline="", encoding="utf-8") as stream:
        name = next((row["file"] for row in csv.DictReader(stream) if row.get("deal") == slug), None)
    return data.Filing((filings / name).read_bytes()) if name and (filings / name).is_file() else None


def audit_changes(audit_dir: Path, slug: str) -> list[dict[str, str]]:
    """The audit's per-change inventory for one deal, each change with its post-challenge verdict."""
    inventory, verdicts = audit_dir / "inventory" / f"{slug}.csv", audit_dir / "deals" / f"{slug}-verdicts.csv"
    if not inventory.is_file():
        return []
    judged = {}
    if verdicts.is_file():
        with verdicts.open(newline="", encoding="utf-8") as stream:
            judged = {row["change_id"]: row for row in csv.DictReader(stream)}
    with inventory.open(newline="", encoding="utf-8") as stream:
        return [{**row, **{key: judged.get(row["change_id"], {}).get(key, "") for key in ("verdict", "basis", "filing_cite", "rule_or_decision")}}
                for row in csv.DictReader(stream)]


def audit_conventions(audit_dir: Path) -> dict[str, list[dict[str, Any]]]:
    """The open convention questions (README §2), per deal: the rows, Rounds lines and Deal facts each names."""
    readme = audit_dir / "README.md"
    if not readme.is_file():
        return {}
    found: dict[str, list[dict[str, Any]]] = {}
    inside = False
    for line in readme.read_text(encoding="utf-8").splitlines():
        if line.startswith("## "):
            inside = line.startswith("## 2")
            continue
        cells = [cell.strip() for cell in line.strip().strip("|").split("|")] if inside and line.startswith("|") else []
        if len(cells) < 4 or cells[0] in ("Id", "") or set(cells[0]) <= set("-: "):
            continue
        ident, deal = cells[0].replace(" (optional)", ""), None
        for segment in cells[3].split(";"):
            segment = segment.strip()
            named = next((slug for name, slug in AUDIT_DEALS.items() if segment.lower().startswith(name)), None)
            deal = named or deal
            if deal is None: continue
            ref = {"id": ident, "rows": set(), "rounds": set(), "questions": set(), "facts": False, "unnamed": [], "text": segment}
            for first, last in re.findall(r"#(\d+)\s*[–-]\s*#?(\d+)", segment):
                ref["rows"].update(range(int(first), int(last) + 1))
            ref["rows"].update(int(n) for n in re.findall(r"#(\d+)(?!\d|\s*[–-]\s*#?\d)", segment))
            rounds = re.search(r"Rounds((?:\s*/?R\d+)*)", segment)
            if rounds:
                ref["rounds"] = {int(n) for n in re.findall(r"R(\d+)", rounds.group(1))} or {"all"}
            ref["questions"] = set(re.findall(r"\b(Q\d+)'s\b", segment))
            ref["unnamed"] = [part.strip() for part in segment.split(",")  # "Meredith 26 LMG rows": rows named without numbers
                              if re.search(r"\brows\b", part) and "#" not in part and not re.search(r"\bQ\d+'s\b", part)]
            ref["facts"] = "Deal facts" in segment
            found.setdefault(deal, []).append(ref)
    return found


# -- parties, dates and prices ------------------------------------------------------------------

def _words(text: str) -> tuple[str, ...]:
    text = re.sub(r"\bl\.\s?p\.", "lp", unicodedata.normalize("NFKC", text).lower().replace("&", " and "))
    return tuple(word for word in re.findall(r"[a-z0-9]+", text) if word not in SUFFIXES)


def aliases(who: Any) -> set[tuple[str, ...]]:
    """Normalized names for a Who cell: the whole, the part outside parentheses and a short parenthetical name."""
    text = show(who)
    found = {_words(text), _words(re.sub(r"\([^)]*\)", " ", text))}
    for inner in re.findall(r"\(([^)]*)\)", text):
        words = _words(re.split(r"[;,]", inner)[0])
        if words == ("target",):
            found.add(("<target>",))
        elif words and len(words) <= 4 and words[0] not in NOT_ALIASES and not words[0].isdigit():
            found.add(words)
    return {name for name in found if name}


def same_party(left: Any, right: Any) -> bool:
    for a in aliases(left):
        for b in aliases(right):
            short, long = sorted((a, b), key=len)
            if short == long[:len(short)] and (len(short) == len(long) or not short[-1].isdigit()):
                return True
            if len(short) >= 2 and not short[0].isdigit() and _tokens(short) <= set(long) | _tokens(long):
                return True  # "Moab / CSC/Pamplona" and "Moab Partners, L.P. / CSC/Pamplona"
    return False


def _tokens(words: tuple[str, ...]) -> set[str]:
    """Words for the subset test, a single letter joined to the word before it, so "Party A" is not in "Party C (a buyer)"."""
    return {f"{words[i - 1]} {word}" if len(word) == 1 and word.isalpha() and i else word for i, word in enumerate(words)}


def window(values: dict[str, Any]) -> tuple[dt.date, dt.date] | None:
    days = [values.get(field) for field in check_lean.DATE_COLUMNS]
    days = [day for day in days if isinstance(day, dt.date)]
    return (min(days), max(days)) if days else None


def overlaps(a: tuple[dt.date, dt.date] | None, b: tuple[dt.date, dt.date] | None, slack: int) -> bool:
    if a is None or b is None:
        return False
    pad = dt.timedelta(days=slack)
    return a[0] - pad <= b[1] and b[0] - pad <= a[1]


def _number(value: Any) -> float | None:
    if isinstance(value, bool): return None
    if isinstance(value, (int, float)): return float(value)
    try: return float(str(value).replace("$", "").replace(",", "").strip())
    except ValueError: return None


def price_match(reviewed: dict[str, Any], run: dict[str, Any]) -> str | None:
    """"equal", "with CVR" (the run's price plus its CVR/earnout value), "n/a" where either row has no price, else None."""
    old = [_number(reviewed.get(field)) for field in ("Price low", "Price high")]
    new = [_number(run.get(field)) for field in ("Price low", "Price high")]
    if all(value is None for value in old) or all(value is None for value in new):
        return "n/a"
    if [round(v, 4) if v is not None else None for v in old] == [round(v, 4) if v is not None else None for v in new]:
        return "equal"
    cvr = _number(run.get("CVR/earnout value"))
    if cvr is not None and [round(v, 4) if v is not None else None for v in old] == [round(v + cvr, 4) if v is not None else None for v in new]:
        return "with CVR"
    return None


def stock_from_all_cash(all_cash: Any) -> Any:
    """The Stock % value All cash determines under S7's crosswalk (Yes → 0, Not stated → Not stated); None where it
    does not (No stands for any figure above 0, a range or Part stock)."""
    return next((stock for stock in (0, "Not stated") if diff_workbooks.all_cash_for_stock(stock) == show(all_cash)), None)


def stock_agrees(all_cash: Any, stock: Any) -> bool | None:
    """S7's crosswalk (diff_workbooks.crosswalk): Yes ↔ 0; No ↔ a figure above 0, a range or Part stock; Not stated ↔
    Not stated. None where the run has Varies, which has no v1.13.2 equivalent."""
    if show(stock) == "Varies":
        return None
    return diff_workbooks.crosswalk("Stock %", all_cash, stock, {})[0] == "same"


def members(text: Any, vocabulary: set[tuple[str, ...]]) -> set[tuple[str, ...]]:
    """The party names of a vocabulary that a Rounds cell mentions; a name inside a longer found name is dropped."""
    padded = f" {' '.join(_words(show(text)))} "
    found = {name for name in vocabulary if f" {' '.join(name)} " in padded}
    return {name for name in found if not any(other != name and other[:len(name)] == name for other in found)}


def due_dates(text: Any) -> set[str]:
    """The due dates a Rounds cell names, leaving out parenthetical remarks such as "(set 08/27/2013)"."""
    text = re.sub(r"\([^)]*\)", " ", show(text))
    return set(re.findall(r"\d{1,2}/\d{1,2}/\d{4}", text)) or {_normalized_text(text)}


# -- register -----------------------------------------------------------------------------------

def _mark(state: dict[str, Any], uid: str) -> dict[str, Any]:
    return state["row_review"].get(uid) or {"status": "unreviewed", "note": "", "actor": None, "at": None}


def _filing_key(values: dict[str, Any], filing: data.Filing | None) -> dict[str, Any]:
    parsed = check_lean.parse_quote_and_page(values.get("Quote and page"))
    if parsed is None:
        return {"page": None, "block": None, "quote": show(values.get("Quote and page")) or None}
    passage, page = parsed
    located = filing.locate(passage, page) if filing is not None else {"located": False}
    return {"page": page, "block": f"b{located['start']['block']}" if located["located"] else None, "quote": passage}


def _cohort(values: dict[str, Any]) -> bool:
    count = values.get("Count")
    return not check_lean.is_blank(count) and (_number(count) is None or _number(count) > 1)


def _impact(kind: str, values: dict[str, Any], context: dict[str, Any]) -> dict[str, Any]:
    """The decisions under which a fact must be judged again for v1.14.1, or the new column it becomes: V114_SPEC's
    D1–D27 and v1.14.1's changes (IMPACT_TAGS_V1141)."""
    event, note = show(values.get("Event")), show(values.get("Note")).lower()
    decisions, new, notes = [], [], []
    revision, same_price = context.get("revision", False), context.get("same_price", False)
    reaffirmed, inferred = event == "Bid reaffirmed", show(values.get("Inferred")).upper() == "Y"
    exit_row = event in check_lean.EXIT_EVENTS
    not_invited = event in ("Dropped by target", "Did not submit", "Re-entered") or (exit_row and inferred)
    if kind == "participation":
        if event == "Did not submit": decisions.append("D10")
        if context.get("partial") or re.search(r"\b(partial|segment|scope)\b", note): decisions.append("D7")
        if _cohort(values): decisions.append("R3")
        if not_invited: decisions.append("R5")
        if event in ("Target interest", "Bidder interest", "Contact"): decisions.append("Contact vs interest")
    elif kind == "deal fact":
        field = show(values.get("Field"))
        if field.startswith(("Auction screen", "Whole-company bids")): decisions.append("D7")
        if field.startswith("Initiation"): decisions.append("Initiation from the first row")
        if field.startswith("Number of processes"): decisions.append("E5 process test")
    elif kind in ("round", "assignment"):
        decisions += ["D8", "v1.14.1 D6"]
        if kind == "round": decisions.append("not invited: Dropped by target")
    elif kind == "deadline":
        decisions.append("D11")
        if show(values.get("Deadline outcome")) == "Late bids accepted":
            notes.append("v1.14 and v1.14.1 record an accepted late bid as Extended (late bid accepted): shown, not equated")
    elif kind == "price":
        if event == "Other-scope bid": decisions.append("D18")
        if reaffirmed or same_price: decisions += ["D13", "R1"]
        if re.search(r"\b(cvr|contingent value|earn-?out|contingent)\b", note): new.append("CVR/earnout")
    elif kind == "stock":
        new.append("Stock %")
        if _cohort(values): decisions.append("D16")
        notes.append(f"All cash {show(values.get('All cash')) or '(blank)'} crosswalks to Stock % "
                     + {"Yes": "0", "No": "above 0, a range or Part stock", "Not stated": "Not stated"}.get(show(values.get("All cash")), "(no value)"))
    elif kind == "formality":
        if revision or reaffirmed: decisions += ["D9", "v1.14.1 D1", "express incorporation removed"]
        if reaffirmed or same_price: decisions.append("R1")
        if reaffirmed or show(values.get("Formality")) in ("Formal", "Unclear"): decisions.append("Formality routes")
    elif kind == "conditions":
        decisions += ["H1–H3", "R2"]
        if re.search(r"financ|commitment letter|highly confident|\bdebt\b|equity commitment", note): decisions.append("D5")
        if "exclusiv" in note: decisions.append("D15")
        if _cohort(values): decisions.append("D16")
        if revision or reaffirmed: decisions.append("express incorporation removed")
        if same_price: decisions.append("D13")
        if reaffirmed or same_price: decisions.append("R1")
        if re.search(r"dilig|\bh2\b|\b(day|week|month)s?\b", note): decisions.append("R4")
        if re.search(r"dilig|\bnda\b|confidentiality", note): decisions.append("R6")
        if re.search(r"proceed|unless|standstill|waiver|repric", note): decisions.append("v1.14.1 D2")
        if show(values.get("Conditions")) == "Light": decisions.append("v1.14.1 D3")
    elif kind == "terms":
        new += list(NEW_TERM_COLUMNS)
    elif kind == "order":
        if event == "Exclusivity changed": decisions.append("D15")
        if event == "Did not submit": decisions.append("D10")
        if event == "Other-scope bid": decisions.append("D7")
        if event == "Round opened": decisions += ["D8", "v1.14.1 D6"]  # E6 may move where a round opens
        if event in ("Deadline set", "Deadline revised", "Deadline"): decisions.append("D11")
        if event == "Other material event":
            if re.search(r"reverse termination|commitment|guarantee|financing condition|damages cap", note): decisions.append("D13")
            if re.search(r"proceed|unless|standstill|waiver|repric", note): decisions.append("v1.14.1 D2")
            decisions.append("Other material event list")
        if event in ("Process terminated", "Process restarted"): decisions.append("E5 process test")
        if event in ADVISER_EVENTS and context.get("names_bidder"): decisions.append("bidders' advisers to Notes")
        if not_invited: decisions.append("R5")
        if event == "Did not submit" and _cohort(values): decisions.append("R3")
        days = [values.get(field) for field in ("Date from", "Date to")]
        if inferred or not all(isinstance(day, dt.date) for day in days) or days[0] != days[1]:
            decisions.append("Sort-date ladder")
    cls = "re-judge" if decisions else "new column" if new else "unchanged"
    return {"class": cls, "decisions": decisions + [f"new column: {name}" for name in new], "note": "; ".join(notes)}


def _normalized_text(value: Any) -> str:
    """Text for comparing a cell with the audit's recorded value: references renumbered, whitespace collapsed."""
    text = re.sub(r"#\s*\d+", "#n", show(value))  # a restored Note may carry renumbered references
    return re.sub(r"\s+", " ", text).strip()


def _status(mark: str | None, hits: list[dict[str, str]], values: dict[str, Any], conventions: list[str]) -> tuple[str, list[str]]:
    reasons, reverted = [], False
    for ident in conventions:
        reasons.append(f"rests on open convention question {ident} (audit README §2)")
    for hit in hits:
        verdict = hit["verdict"] or "not judged"
        if verdict == "incorrect":
            current = values.get(hit["field"]) if hit["field"] else None
            if hit["change_type"] == "update" and _normalized_text(current) == _normalized_text(hit["original_value"]):
                reverted = True
                continue
            reasons.append(f"{hit['change_id']} judged incorrect and not reverted")
        elif verdict != "supported":
            reasons.append(f"{hit['change_id']} judged {verdict}")
    if mark is not None and mark != "reviewed":
        reasons.append(f"row marked {mark.replace('_', ' ')}")
    if reasons:
        return "unresolved", reasons
    return ("reverted", ["an incorrect change was reverted to the original value (revert of 24 September)"]) if reverted else ("supported", [])


def build_register(slug: str, conn: sqlite3.Connection, audit_dir: Path = AUDIT_DIR, filings: Path = FILINGS) -> dict[str, Any]:
    working = working_copy(conn, slug)
    state = working["state"]
    ledger = state["sheets"][LEDGER]["rows"]
    rounds = state["sheets"][ROUNDS]["rows"]
    filing = filing_for(slug, filings)
    changes = audit_changes(audit_dir, slug)
    by_uid: dict[str, list[dict[str, str]]] = {}
    for change in changes:
        by_uid.setdefault(change["uid"], []).append(change)
    number_of = {record["values"].get("#"): record for record in ledger}
    current = {record["uid"]: str(record["values"].get("#")) for record in ledger}
    mismatched = [change["change_id"] for change in changes if change["uid"] in current and current[change["uid"]] != change["final_ref"]]
    fact_uids = {record["uid"] for record in state["sheets"][FACTS]["rows"]}
    # Convention questions by row uid, Rounds uid and Deal facts uid.
    conventions: dict[str, list[tuple[str, set[str] | None]]] = {}
    unparsed = []
    questions = {show(r["values"].get("Q")): r["values"] for r in state["sheets"][QUESTIONS]["rows"]}
    for ref in audit_conventions(audit_dir).get(slug, []):
        kinds = CONVENTION_KINDS.get(ref["id"])
        targets = [number_of[n]["uid"] for n in ref["rows"] if n in number_of]
        for q in ref["questions"]:
            refs, _ = check_lean.parse_affected_rows(questions.get(q, {}).get("Rows affected"))
            targets += [number_of[n]["uid"] for n in refs if n in number_of]
        targets += [r["uid"] for r in rounds if "all" in ref["rounds"] or r["values"].get("Round") in ref["rounds"]]
        if ref["facts"]:
            targets += [r["uid"] for r in state["sheets"][FACTS]["rows"]]
        if not targets:
            unparsed.append(f"{ref['id']}: {ref['text']}")
        elif ref["unnamed"]:
            unparsed.append(f"{ref['id']}: {'; '.join(ref['unnamed'])} (the rest of \"{ref['text']}\" is tied to rows)")
        for uid in targets:
            conventions.setdefault(uid, []).append((ref["id"], None if uid in fact_uids else kinds))
    # Revisions and Other-scope parties, for the impact rules.
    partial = [record["values"].get("Who") for record in ledger if record["values"].get("Event") == "Other-scope bid"]
    bidders = {name for record in ledger if FAMILY.get(record["values"].get("Event")) in ("bid", "entry", "exit")
               for name in aliases(record["values"].get("Who"))}
    earlier: list[dict[str, Any]] = []
    context: dict[str, dict[str, Any]] = {}
    for record in ledger:
        values = record["values"]
        before = [prior for prior in earlier if same_party(prior.get("Who"), values.get("Who"))]
        context[record["uid"]] = {"partial": any(same_party(values.get("Who"), who) for who in partial),
                                  "revision": bool(before) and values.get("Event") in check_lean.BID_EVENTS,
                                  "same_price": bool(before) and values.get("Event") == "Bid" and price_match(before[-1], values) == "equal",
                                  "names_bidder": values.get("Event") in ADVISER_EVENTS and bool(
                                      members(f"{show(values.get('Who'))} {show(values.get('Note'))}", bidders)
                                      or BIDDER_WORDS.search(f"{show(values.get('Who'))} {show(values.get('Note'))}".lower()))}
        if values.get("Event") in check_lean.BID_EVENTS:
            earlier.append(values)
    thread_list = threads(conn, slug)
    rows = []
    for record in ledger:
        values, mark = record["values"], _mark(state, record["uid"])
        rows.append({"#": values.get("#"), "uid": record["uid"], "When": show(values.get("When")), "Who": show(values.get("Who")),
                     "Event": show(values.get("Event")), "mark": mark["status"], "mark_note": mark.get("note", ""),
                     "mark_actor": mark.get("actor"), "mark_at": mark.get("at"),
                     "threads": sum(1 for t in thread_list if t["target_uid"] == record["uid"])})

    facts: list[dict[str, Any]] = []

    def add(section: int, kind: str, sheet: str, record: dict[str, Any], key: dict[str, Any], mark: str | None, crosswalk: Any = None) -> None:
        values = record["values"]
        hits = [hit for hit in by_uid.get(record["uid"], [])
                if sheet != LEDGER or hit["field"] in TEXT_FIELDS or hit["change_type"] != "update" or kind in LEDGER_FIELD_KINDS.get(hit["field"], set())]
        if sheet == ROUNDS:
            hits = [hit for hit in hits if (hit["field"] in ("Due dates", "Deadline outcome")) == (kind == "deadline") or not hit["field"]]
        named = [ident for ident, kinds in conventions.get(record["uid"], []) if kinds is None or kind in kinds]
        text = " ".join(show(values.get(field)) for field in FACT_FIELDS[kind])
        convention_basis = [hit["change_id"] for hit in hits if re.search(r"convention|R01|Alex-stated rule", hit["basis"] + hit["rule_or_decision"])]
        if named or convention_basis:
            basis = "convention"
        elif show(values.get("Inferred")).upper() == "Y" or re.match(r"\s*Inferred\b", show(values.get("Note"))) or re.search(r"\b(inferred|assum)", text, re.I):
            basis = "inference"
        else:
            basis = "reported"
        status, reasons = _status(mark, hits, values, named)
        fact = {"id": f"{kind.replace(' ', '-')}:{record['uid']}", "section": SECTIONS[section], "kind": kind, "sheet": sheet,
                "uid": record["uid"], "row": values.get("#") if sheet == LEDGER else None,
                "label": (f"#{values.get('#')} {show(values.get('Event'))} · {show(values.get('Who'))}" if sheet == LEDGER
                          else f"Process {show(values.get('Process'))} / Round {show(values.get('Round'))}" if sheet == ROUNDS
                          else show(values.get("Field"))),
                "filing_key": key, "value": {field: show(values.get(field)) for field in FACT_FIELDS[kind]},
                "basis": basis, "basis_refs": named + convention_basis, "status": status, "status_reasons": reasons,
                "audit": [hit["change_id"] for hit in hits], "impact": _impact(kind, values, context.get(record["uid"], {}))}
        if crosswalk is not None:
            fact["crosswalk"] = crosswalk
        facts.append(fact)

    keys = {record["uid"]: _filing_key(record["values"], filing) for record in ledger}
    marks = {record["uid"]: _mark(state, record["uid"])["status"] for record in ledger}
    for record in state["sheets"][FACTS]["rows"]:
        field = next((name for name in DEAL_FACT_SECTIONS if show(record["values"].get("Field")).startswith(name)), None)
        if field and DEAL_FACT_SECTIONS[field] == 0:
            add(0, "deal fact", FACTS, record, {"page": None, "block": None, "quote": None}, None)
    for record in ledger:
        if FAMILY.get(record["values"].get("Event")) in ("entry", "exit"):
            add(0, "participation", LEDGER, record, keys[record["uid"]], marks[record["uid"]])
    for record in state["sheets"][FACTS]["rows"]:
        if show(record["values"].get("Field")).startswith("Number of processes"):
            add(1, "deal fact", FACTS, record, {"page": None, "block": None, "quote": None}, None)
    for record in rounds:
        values = record["values"]
        def anchor(events: tuple[str, ...]) -> dict[str, Any]:
            found = next((r for event in events for r in ledger if r["values"].get("Event") == event
                          and (r["values"].get("Process"), r["values"].get("Round")) == (values.get("Process"), values.get("Round"))), None)
            return keys[found["uid"]] if found else {"page": None, "block": None, "quote": None}
        add(1, "round", ROUNDS, record, anchor(("Round opened",)), None)
        add(1, "deadline", ROUNDS, record, anchor(("Deadline", "Deadline revised", "Deadline set")), None)
    for record in ledger:
        if FAMILY.get(record["values"].get("Event")) in ("bid", "exit"):
            add(1, "assignment", LEDGER, record, keys[record["uid"]], marks[record["uid"]])
    for record in ledger:
        if record["values"].get("Event") in check_lean.BID_EVENTS:
            for kind in ("price", "stock", "formality", "conditions", "terms"):
                add(2, kind, LEDGER, record, keys[record["uid"]], marks[record["uid"]],
                    {"All cash": show(record["values"].get("All cash")), "Stock %": stock_from_all_cash(record["values"].get("All cash"))} if kind == "stock" else None)
    for record in ledger:
        add(3, "order", LEDGER, record, keys[record["uid"]], marks[record["uid"]])

    counts = {"reviewed": 0, "needs_decision": 0, "unreviewed": 0}
    for row in rows:
        counts[row["mark"]] = counts.get(row["mark"], 0) + 1
    return {"deal": slug, "tool": "migrate_review.py",
            "source": {**working["meta"], "filing_located": filing is not None},
            "marks": counts, "findings": state.get("findings", {}), "threads": thread_list,
            "rows": rows, "facts": facts,
            "audit_checks": {"inventory_changes": len(changes), "changes_on_current_rows": sum(len(by_uid.get(r["uid"], [])) for r in ledger + rounds + state["sheets"][FACTS]["rows"]),
                             "final_ref_differs_from_current_number": mismatched, "convention_refs_not_resolved": unparsed},
            "_state": state}


def _flat(value: dict[str, Any]) -> str:
    return "; ".join(f"{key}={text}" for key, text in value.items() if text)


def write_register(register: dict[str, Any], out: Path) -> Path:
    folder = out / register["deal"]
    folder.mkdir(parents=True, exist_ok=True)
    (folder / "register.json").write_text(json.dumps({k: v for k, v in register.items() if k != "_state"}, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8")
    with (folder / "rows.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.DictWriter(stream, fieldnames=list(register["rows"][0]) if register["rows"] else ["#"])
        writer.writeheader()
        writer.writerows(register["rows"])
    with (folder / "facts.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["id", "section", "kind", "label", "reviewed value", "basis", "basis refs", "status", "status reasons",
                         "v1.14.1 impact", "decisions", "impact note", "page", "block", "quote", "audit changes"])
        for fact in register["facts"]:
            key = fact["filing_key"]
            writer.writerow([fact["id"], fact["section"], fact["kind"], fact["label"], _flat(fact["value"]), fact["basis"],
                             ", ".join(fact["basis_refs"]), fact["status"], "; ".join(fact["status_reasons"]), fact["impact"]["class"],
                             ", ".join(fact["impact"]["decisions"]), fact["impact"]["note"], key["page"] or "", key["block"] or "",
                             key["quote"] or "", ", ".join(fact["audit"])])
    return folder


# -- aligner ------------------------------------------------------------------------------------

def _mutual(left: list[dict[str, Any]], right: list[dict[str, Any]], fits) -> list[tuple[dict[str, Any], dict[str, Any]]]:
    """Pairs whose members fit only each other among the rows still unmatched; repeated until none is added."""
    pairs, left, right = [], list(left), list(right)
    while True:
        options = {a["uid"]: [b for b in right if fits(a, b)] for a in left}
        back = {b["uid"]: [a for a in left if b in options[a["uid"]]] for b in right}
        found = [(a, options[a["uid"]][0]) for a in left if len(options[a["uid"]]) == 1 and len(back[options[a["uid"]][0]["uid"]]) == 1]
        if not found:
            return pairs
        pairs += found
        left = [a for a in left if all(a is not x for x, _ in found)]
        right = [b for b in right if all(b is not y for _, y in found)]


def align_ledger(reviewed: list[dict[str, Any]], run: list[dict[str, Any]]) -> dict[str, Any]:
    matches: dict[str, dict[str, Any]] = {}
    for number, (level, slack, priced) in enumerate(PASSES, 1):
        left = [r for r in reviewed if r["uid"] not in matches]
        right = [r for r in run if r["uid"] not in {m["uid"] for m in matches.values()}]

        def fits(a: dict[str, Any], b: dict[str, Any]) -> bool:
            va, vb = a["values"], b["values"]
            if level == "event" and va.get("Event") != vb.get("Event"): return False
            if level == "family" and FAMILY.get(va.get("Event"), "other") != FAMILY.get(vb.get("Event"), "other"): return False
            if not same_party(va.get("Who"), vb.get("Who")) or not overlaps(window(va), window(vb), slack): return False
            return not priced or price_match(va, vb) is not None
        for a, b in _mutual(left, right, fits):
            matches[a["uid"]] = {"uid": b["uid"], "pass": number, "price": price_match(a["values"], b["values"])}
    # Order: matched pairs outside the longest run-order-preserving chain are out of order.
    pairs = [(i, next(j for j, r in enumerate(run) if r["uid"] == matches[a["uid"]]["uid"]), a["uid"]) for i, a in enumerate(reviewed) if a["uid"] in matches]
    tails, back, chain = [], [None] * len(pairs), set()
    for k, (_, j, _) in enumerate(pairs):
        lo, hi = 0, len(tails)
        while lo < hi:
            mid = (lo + hi) // 2
            if pairs[tails[mid]][1] < j: lo = mid + 1
            else: hi = mid
        back[k] = tails[lo - 1] if lo else None
        tails[lo:lo + 1] = [k]
    k = tails[-1] if tails else None
    while k is not None:
        chain.add(pairs[k][2])
        k = back[k]
    # An inversion between rows of one Sort date is a same-day reordering, not a change in the order of events.
    day = {a["uid"]: a["values"].get("Sort date") for a in reviewed}
    for i, j, uid in pairs:
        matches[uid]["in_order"] = uid in chain or not any(
            (i - k) * (j - m) < 0 and day[uid] != day[other] for k, m, other in pairs)
    matched_run = {m["uid"] for m in matches.values()}
    return {"matches": matches, "omitted": [r["uid"] for r in reviewed if r["uid"] not in matches],
            "inserted": [r["uid"] for r in run if r["uid"] not in matched_run]}


def align_rounds(reviewed: list[dict[str, Any]], run: list[dict[str, Any]], vocabulary: set[tuple[str, ...]]) -> dict[str, Any]:
    def opened(record: dict[str, Any]) -> dt.date | None:
        value = record["values"].get("Opened")
        return value if isinstance(value, dt.date) else None

    def similar(a: dict[str, Any], b: dict[str, Any], named: bool = False) -> bool:
        ma, mb = members(a["values"].get("Who was in"), vocabulary), members(b["values"].get("Who was in"), vocabulary)
        return (not named and not ma and not mb) or bool(ma | mb) and len(ma & mb) * 2 >= len(ma | mb)
    matches: dict[str, dict[str, Any]] = {}
    for number, fits in enumerate((
            lambda a, b: opened(a) is not None and opened(b) is not None and abs((opened(a) - opened(b)).days) <= 3 and similar(a, b),
            lambda a, b: opened(a) is not None and opened(a) == opened(b),
            lambda a, b: similar(a, b, named=True)), 1):  # an opening moved by more than three days, the same members
        left = [r for r in reviewed if r["uid"] not in matches]
        right = [r for r in run if r["uid"] not in {m["uid"] for m in matches.values()}]
        for a, b in _mutual(left, right, fits):
            matches[a["uid"]] = {"uid": b["uid"], "pass": number}
    matched_run = {m["uid"] for m in matches.values()}
    return {"matches": matches, "omitted": [r["uid"] for r in reviewed if r["uid"] not in matches],
            "inserted": [r["uid"] for r in run if r["uid"] not in matched_run]}


# -- triage and port batch ----------------------------------------------------------------------

def _compare(fact: dict[str, Any], old: dict[str, Any], new: dict[str, Any], match: dict[str, Any], vocabulary: set[tuple[str, ...]]) -> tuple[bool, dict[str, str], str]:
    """(agrees, the run's values, a note on a difference that the schema explains)."""
    kind = fact["kind"]
    shown = {field: show(new.get(field)) for field in FACT_FIELDS[kind]}
    if kind == "stock":
        if "All cash" in new:
            return show(old.get("All cash")) == show(new.get("All cash")), shown, ""
        shown = {"Stock %": show(new.get("Stock %"))}
        verdict, label = diff_workbooks.crosswalk("Stock %", old.get("All cash"), new.get("Stock %"), new)
        return verdict == "same", shown, label if stock_agrees(old.get("All cash"), new.get("Stock %")) is None else ""
    if kind == "price":
        shown["CVR/earnout value"] = show(new.get("CVR/earnout value"))
        found = price_match(old, new)
        if found == "with CVR":
            return False, shown, "the run's price plus its CVR/earnout value equals the reviewed price: v1.14.1 shows the CVR apart, not equated"
        suppressed = [diff_workbooks.crosswalk(f, old.get(f), new.get(f), new) for f in ("Price low", "Price high")]
        if found == "n/a" and all(check_lean.is_blank(new.get(f)) for f in ("Price low", "Price high")) \
                and any(cross and cross[0] == "suppressed" for cross in suppressed):
            return False, shown, next(cross[1] for cross in suppressed if cross and cross[0] == "suppressed")
        return found == "equal" or (found == "n/a" and all(show(old.get(f)) == show(new.get(f)) for f in ("Price low", "Price high"))), shown, ""
    if kind == "round":
        same = all(show(old.get(f)) == show(new.get(f)) for f in ("Process", "Round", "Opened", "Finality"))
        return same and members(old.get("Who was in"), vocabulary) == members(new.get("Who was in"), vocabulary), shown, ""
    if kind == "deadline":
        cross = diff_workbooks.crosswalk("Deadline outcome", old.get("Deadline outcome"), new.get("Deadline outcome"), new)
        note = cross[1] if cross and show(old.get("Deadline outcome")) != show(new.get("Deadline outcome")) else ""
        return due_dates(old.get("Due dates")) == due_dates(new.get("Due dates")) and show(old.get("Deadline outcome")) == show(new.get("Deadline outcome")), shown, note
    if kind == "deal fact":
        lead = lambda value: _normalized_text(re.sub(r"\([^)]*\)", " ", show(value))).lower()
        return lead(old.get("Value")) == lead(new.get("Value")), shown, ""
    if kind == "order":  # When is free text: shown and ported with the dates, not compared
        agrees = all(show(old.get(f)) == show(new.get(f)) for f in FACT_FIELDS[kind] if f != "When") and match.get("in_order", True)
        return agrees, {**shown, "in order": "yes" if match.get("in_order", True) else "no"}, ""
    return all(show(old.get(f)) == show(new.get(f)) for f in FACT_FIELDS[kind]), shown, ""


def _holds(fact: dict[str, Any]) -> list[str]:
    """Why an agreeing fact stays a review item: its status is not supported, or v1.14.1 changes its rule or adds its column."""
    reasons = []
    if fact["status"] != "supported":
        reasons.append(f"status {fact['status']}" + (f" ({'; '.join(fact['status_reasons'])})" if fact["status_reasons"] else ""))
    if fact["impact"]["class"] != "unchanged":
        reasons.append(f"v1.14.1 impact {fact['impact']['class']} ({', '.join(fact['impact']['decisions'])})")
    return reasons


def triage(register: dict[str, Any], run: dict[str, Any], run_id: str, seed: int = 1, sample: int = 5) -> dict[str, Any]:
    state = register["_state"]
    old_ledger, new_ledger = state["sheets"][LEDGER]["rows"], run["sheets"][LEDGER]["rows"]
    old_rounds, new_rounds = state["sheets"][ROUNDS]["rows"], run["sheets"][ROUNDS]["rows"]
    vocabulary = {name for record in old_ledger + new_ledger if FAMILY.get(record["values"].get("Event")) in ("bid", "entry", "exit")
                  for name in aliases(record["values"].get("Who"))}
    ledger = align_ledger(old_ledger, new_ledger)
    rounds = align_rounds(old_rounds, new_rounds, vocabulary)
    field = lambda record: re.sub(r"\s*\(.*$", "", show(record["values"].get("Field")))  # "Auction screen (E1)" is "Auction screen"
    new_facts = {field(r): r for r in run["sheets"][FACTS]["rows"]}
    old_by_uid = {r["uid"]: r for r in old_ledger + old_rounds + state["sheets"][FACTS]["rows"]}
    new_by_uid = {r["uid"]: r for r in new_ledger + new_rounds + run["sheets"][FACTS]["rows"]}
    matches = {**ledger["matches"], **rounds["matches"]}
    for record in state["sheets"][FACTS]["rows"]:
        found = new_facts.get(field(record))
        if found: matches[record["uid"]] = {"uid": found["uid"], "pass": 1}
    buckets: dict[str, list[dict[str, Any]]] = {bucket: [] for bucket in BUCKETS}
    new_columns = []
    for fact in register["facts"]:
        match = matches.get(fact["uid"])
        item = {"fact": fact["id"], "kind": fact["kind"], "label": fact["label"], "reviewed": fact["value"], "status": fact["status"],
                "impact": fact["impact"]["class"], "decisions": fact["impact"]["decisions"], "filing_key": fact["filing_key"]}
        if match is None:
            if fact["kind"] != "terms":
                buckets[BUCKETS[3]].append({**item, "side": "omitted from the run"})
            continue
        new = new_by_uid[match["uid"]]["values"]
        item.update(run_uid=match["uid"], run_row=new.get("#") if fact["sheet"] == LEDGER else None, match_pass=match["pass"])
        if fact["kind"] == "terms":
            new_columns.append({**item, "run": {field: show(new.get(field)) for field in NEW_TERM_COLUMNS if field in new}})
            continue
        agrees, shown, note = _compare(fact, old_by_uid[fact["uid"]]["values"], new, match, vocabulary)
        item.update(run=shown, note=note)
        bucket = 0 if agrees else 1 if fact["impact"]["class"] != "unchanged" or note else 2
        if agrees:
            item["hold"] = _holds(fact)
            item["list"] = AGREES_LISTS[1] if item["hold"] else AGREES_LISTS[0]
        buckets[BUCKETS[bucket]].append(item)
    for uid in ledger["inserted"] + rounds["inserted"]:
        values = new_by_uid[uid]["values"]
        label = f"#{values.get('#')} {show(values.get('Event'))} · {show(values.get('Who'))}" if "#" in values else f"Process {show(values.get('Process'))} / Round {show(values.get('Round'))}"
        buckets[BUCKETS[3]].append({"side": "inserted in the run", "run_uid": uid, "label": label, "run_row": values.get("#")})
    chooser = random.Random(f"{seed}|{register['deal']}|{run_id}")
    accept = [item for item in buckets[BUCKETS[0]] if item["list"] == AGREES_LISTS[0]]
    spot = chooser.sample(accept, min(sample, len(accept)))

    def count(items: list[dict[str, Any]]) -> dict[str, int]:
        return {"facts": len(items), "rows": len({i.get("run_uid") if i.get("side") == "inserted in the run" else i.get("fact", "").split(":", 1)[-1] for i in items})}
    summary = {bucket: count(items) for bucket, items in buckets.items()}
    summary[BUCKETS[0]].update({name: count([i for i in buckets[BUCKETS[0]] if i["list"] == name]) for name in AGREES_LISTS})
    return {"deal": register["deal"], "run_id": run_id, "run_sha256": run["sha256"], "run_schema": run["schema"],
            "register_revision": register["source"]["revision"], "seed": seed, "summary": summary,
            "alignment": {"ledger_matched": len(ledger["matches"]), "ledger_omitted": len(ledger["omitted"]), "ledger_inserted": len(ledger["inserted"]),
                          "rounds_matched": len(rounds["matches"]), "rounds_omitted": len(rounds["omitted"]), "rounds_inserted": len(rounds["inserted"]),
                          "by_pass": {str(n): sum(1 for m in ledger["matches"].values() if m["pass"] == n) for n in range(1, len(PASSES) + 1)},
                          "pairs": [{"reviewed_uid": uid, "reviewed_row": old_by_uid[uid]["values"].get("#"), "run_uid": m["uid"],
                                     "run_row": new_by_uid[m["uid"]]["values"].get("#"), "pass": m["pass"], "price": m.get("price"), "in_order": m.get("in_order")}
                                    for uid, m in ledger["matches"].items()]},
            "spot_check": [item["fact"] for item in spot], "buckets": buckets, "new_columns": new_columns}


def _api_value(value: Any) -> Any:
    if isinstance(value, dt.date): return value.isoformat()
    return None if check_lean.is_blank(value) else value


def port_batch(register: dict[str, Any], result: dict[str, Any], accepted: set[str] = frozenset()) -> dict[str, Any]:
    """Edit requests for the cockpit's edit API: the reviewed values of accepted differing facts, then row marks.

    A row gets a mark only when every one of its facts agrees or was accepted and ported. The mark is "reviewed" only
    where the old mark was reviewed and every fact is supported with v1.14.1 impact unchanged; a row with a fact to
    re-judge or a new column with no reviewed value (every bid row) gets "needs_decision", so an agreement on the agrees
    bucket's review list carries no acceptance. Never applied here."""
    state = register["_state"]
    old_by_uid = {r["uid"]: r for part in state["sheets"].values() for r in part["rows"]}
    facts_by_row: dict[str, list[dict[str, Any]]] = {}
    for fact in register["facts"]:
        facts_by_row.setdefault(fact["uid"], []).append(fact)
    placed = {item["fact"]: item for bucket in BUCKETS[:3] for item in result["buckets"][bucket]}
    placed.update({item["fact"]: item for item in result["new_columns"]})
    agreed = {item["fact"] for item in result["buckets"][BUCKETS[0]]}
    choices = check_lean.choice_lists(result["run_schema"])
    operations, skipped = [], []
    for fact in register["facts"]:
        if fact["id"] not in accepted or fact["id"] in agreed:
            continue
        if fact["id"] not in placed:
            skipped.append({"fact": fact["id"], "reason": "the run has no matching row: an insert must be prepared by hand"})
            continue
        item, values = placed[fact["id"]], old_by_uid[fact["uid"]]["values"]
        if item.get("note"):
            skipped.append({"fact": fact["id"], "reason": f"the difference comes from the {result['run_schema']} schema ({item['note']}); not ported"})
            continue
        update = {field: _api_value(values.get(field)) for field in FACT_FIELDS[fact["kind"]]
                  if field in item["run"] and show(values.get(field)) != item["run"][field]}
        if fact["kind"] == "order" and set(update) <= {"When"}:
            update = {}  # a difference in When's wording alone is not ported
        invalid = [field for field, value in update.items() if field in choices and value is not None and value not in choices[field]]
        if invalid:
            skipped.append({"fact": fact["id"], "reason": f"the reviewed value of {', '.join(invalid)} is not a {result['run_schema']} choice"})
            continue
        if fact["kind"] == "stock" and "Stock %" in item["run"]:
            if fact["crosswalk"]["Stock %"] is None:
                skipped.append({"fact": fact["id"], "reason": f"All cash {fact['crosswalk']['All cash'] or '(blank)'} does not determine a Stock % figure"})
                continue
            update = {"Stock %": fact["crosswalk"]["Stock %"]}
        if not update:
            skipped.append({"fact": fact["id"], "reason": "no field of this fact differs from the run's, or none exists in the run"})
            continue
        operations.append({"type": "update", "sheet": fact["sheet"], "uid": item["run_uid"], "values": update})
    ported = set(accepted) - {item["fact"] for item in skipped}
    pairs = {p["reviewed_uid"]: p for p in result["alignment"]["pairs"]}
    for row in register["rows"]:
        pair = pairs.get(row["uid"])
        facts = facts_by_row.get(row["uid"], [])
        if pair is None or row["mark"] == "unreviewed":
            continue
        if not all(f["id"] in agreed or f["id"] in ported or f["kind"] == "terms" for f in facts):
            continue
        settled = all(f["status"] == "supported" and f["impact"]["class"] == "unchanged" for f in facts)
        status = "reviewed" if row["mark"] == "reviewed" and settled else "needs_decision"
        operations.append({"type": "review", "uid": pair["run_uid"], "status": status,
                           "note": f"Ported from the v1.13.2 working copy, revision {register['source']['revision']}, row #{row['#']} ({row['uid'][:8]})."})
    reason = (f"Port reviewed work from the v1.13.2 working copy (revision {register['source']['revision']}) onto {result['run_id']}; "
              f"prepared by migrate_review.py from the triage (seed {result['seed']})")
    requests = [{"revision": None, "base_sha256": result["run_sha256"], "reason": reason + (f", part {i // MAX_OPERATIONS + 1}" if len(operations) > MAX_OPERATIONS else ""),
                 "operations": operations[i:i + MAX_OPERATIONS]} for i in range(0, len(operations), MAX_OPERATIONS)]
    return {"deal": register["deal"], "target_version": result["run_id"], "base_sha256": result["run_sha256"],
            "applied": False,
            "note": ("Prepared, never applied. After the deal is rebased onto target_version, set each request's revision to the working "
                     "copy's revision and POST it to /api/deal/<deal>/edit, as Austin or at his command (V114_SPEC §13, gate 11). "
                     "More than one request means more than 100 operations: one save takes at most 100."),
            "accepted_facts": sorted(accepted), "skipped": skipped, "requests": requests}


def _cell(item: dict[str, Any], key: str) -> str:
    return _flat(item.get(key) or {}).replace("|", "/")


def write_triage(result: dict[str, Any], batch: dict[str, Any], out: Path) -> Path:
    folder = out / result["deal"]
    folder.mkdir(parents=True, exist_ok=True)
    stem = result["run_id"]
    (folder / f"triage-{stem}.json").write_text(json.dumps(result, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8")
    (folder / f"port-batch-{stem}.json").write_text(json.dumps(batch, ensure_ascii=False, indent=1, default=str) + "\n", encoding="utf-8")
    align = result["alignment"]
    lines = [f"# Triage: {result['deal']}, run `{result['run_id']}` against the working copy (revision {result['register_revision']})", "",
             f"Run SHA-256 `{result['run_sha256']}`, schema {result['run_schema']}. Built by `migrate_review.py triage`; spot-check seed {result['seed']}.", "",
             f"Alignment: {align['ledger_matched']} ledger rows matched (by pass: {align['by_pass']}), {align['ledger_omitted']} omitted, "
             f"{align['ledger_inserted']} inserted; Rounds {align['rounds_matched']} matched, {align['rounds_omitted']} omitted, {align['rounds_inserted']} inserted.", "",
             "| Bucket | Action | Facts | Rows |", "|---|---|---|---|"]
    lines += [f"| {bucket} | {action} | {result['summary'][bucket]['facts']} | {result['summary'][bucket]['rows']} |" for bucket, action in zip(BUCKETS, BUCKET_ACTIONS)]
    agrees = result["summary"][BUCKETS[0]]
    lines += ["", f"Of the {agrees['facts']} agreeing facts, {agrees['accept']['facts']} ({agrees['accept']['rows']} rows) may be accepted: supported, "
              f"with v1.14.1 impact unchanged. The other {agrees['review']['facts']} ({agrees['review']['rows']} rows) stay review items: an agreement "
              "does not settle an unresolved or reverted fact, a Needs decision row, a rule v1.14.1 changed or a new column."]
    lines += ["", f"New 29-column terms with no reviewed value, for review under {result['run_schema']}: {len(result['new_columns'])} bid rows.",
              f"Port batch: {sum(len(r['operations']) for r in batch['requests'])} operations in {len(batch['requests'])} request(s), not applied.", ""]
    def table(items: list[dict[str, Any]]) -> list[str]:
        if not items:
            return ["None.", ""]
        rows = ["| Fact | Row | Run row | Reviewed | Run | Status | Impact | Note |", "|---|---|---|---|---|---|---|---|"]
        for item in items:
            rows.append("| {} | {} | {} | {} | {} | {} | {} | {} |".format(
                item.get("fact") or item.get("side"), item.get("label", "").replace("|", "/"), item.get("run_row") or "",
                _cell(item, "reviewed"), _cell(item, "run"), item.get("status", ""),
                ", ".join(item.get("decisions") or []) or item.get("impact", ""),
                (item.get("note") or "; ".join(item.get("hold") or []) or item.get("side") or "").replace("|", "/")))
        return rows + [""]
    for bucket in BUCKETS:
        items = result["buckets"][bucket]
        lines += [f"## {bucket} ({len(items)})", ""]
        if bucket != BUCKETS[0]:
            lines += table(items)
            continue
        accept, review = ([i for i in items if i["list"] == name] for name in AGREES_LISTS)
        lines += [f"### accept, after the seeded spot check ({len(accept)})", "",
                  f"Spot-check sample (seed {result['seed']}): " + (", ".join(f"`{f}`" for f in result["spot_check"]) or "none"), ""]
        lines += table(accept)
        lines += [f"### review: agrees, but not settled ({len(review)})", "",
                  "The run repeats the reviewed value, but the fact is not supported or v1.14.1 changes its rule or adds its column; "
                  "the Note column says why. Check it as a review item; it inherits no acceptance.", ""]
        lines += table(review)
    (folder / f"triage-{stem}.md").write_text("\n".join(lines), encoding="utf-8")
    return folder


# -- command line -------------------------------------------------------------------------------

def register_command(args: argparse.Namespace) -> None:
    conn = connect(args.state_db)
    try:
        deals = args.deals or [row[0] for row in conn.execute("SELECT DISTINCT slug FROM revisions ORDER BY slug")]
        for slug in deals:
            register = build_register(slug, conn, Path(args.audit_dir), Path(args.filings))
            folder = write_register(register, Path(args.out))
            print(f"{slug}: revision {register['source']['revision']}, {len(register['rows'])} rows {register['marks']}, "
                  f"{len(register['facts'])} facts -> {folder}")
    finally:
        conn.close()


def verify_command(args: argparse.Namespace) -> None:
    """Re-read the database and check that each written row list carries the database's row marks."""
    conn = connect(args.state_db)
    failures, totals = 0, {}
    try:
        for slug in args.deals or [row[0] for row in conn.execute("SELECT DISTINCT slug FROM revisions ORDER BY slug")]:
            state = working_copy(conn, slug)["state"]
            expected = {record["uid"]: _mark(state, record["uid"])["status"] for record in state["sheets"][LEDGER]["rows"]}
            path = Path(args.out) / slug / "rows.csv"
            with path.open(newline="", encoding="utf-8") as stream:
                written = {row["uid"]: row["mark"] for row in csv.DictReader(stream)}
            counts = {status: list(expected.values()).count(status) for status in sorted(set(expected.values()))}
            for status, n in counts.items():
                totals[status] = totals.get(status, 0) + n
            failures += written != expected
            print(f"{slug}: marks {'match' if written == expected else 'DIFFER'} the database's {counts}")
    finally:
        conn.close()
    print(f"total: {totals}")
    if failures:
        raise SystemExit(f"{failures} row list(s) differ from the database")


def triage_command(args: argparse.Namespace) -> None:
    conn = connect(args.state_db)
    try:
        register = build_register(args.deal, conn, Path(args.audit_dir), Path(args.filings))
    finally:
        conn.close()
    run = load_run(Path(args.run), args.deal, args.rules)
    result = triage(register, run, args.run_id, args.seed, args.sample)
    accepted = set()
    if args.accept:
        accepted = {line.strip() for line in Path(args.accept).read_text(encoding="utf-8").splitlines() if line.strip() and not line.startswith("#")}
    folder = write_triage(result, port_batch(register, result, accepted), Path(args.out))
    print(f"{args.deal} vs {args.run_id}: " + ", ".join(f"{bucket} {n['facts']} facts/{n['rows']} rows" for bucket, n in result["summary"].items()) + f" -> {folder}")


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    common = argparse.ArgumentParser(add_help=False)
    common.add_argument("--state-db", default=str(STATE_DB), help="the cockpit database, opened read-only")
    common.add_argument("--audit-dir", default=str(AUDIT_DIR), help="lesson/independent-audit-2026-09-23")
    common.add_argument("--filings", default=str(FILINGS), help="the folder holding MANIFEST.csv and the filings")
    common.add_argument("--out", default=str(OUT), help="output root; files go to <out>/<deal>/")
    sub = root.add_subparsers(dest="command", required=True)
    p = sub.add_parser("register", parents=[common], help="write each deal's reviewed-facts register")
    p.add_argument("deals", nargs="*", help="deal slugs (default: every deal with a working-copy revision)")
    p.set_defaults(func=register_command)
    p = sub.add_parser("verify", parents=[common], help="check the written row lists against the database's row marks")
    p.add_argument("deals", nargs="*")
    p.set_defaults(func=verify_command)
    p = sub.add_parser("triage", parents=[common], help="align a v1.14.1 (or v1.14) run to the register, triage and prepare the port batch")
    p.add_argument("deal")
    p.add_argument("run", help="a copy of the run's workbook (never the file in the state folder)")
    p.add_argument("--run-id", required=True, help="the run's cockpit version id")
    p.add_argument("--rules", choices=check_lean.RULES_29_COLUMN,
                   help="the rules of the run's instruction (check_lean.rules_for_instruction of its SHA-256); default v1.14.1")
    p.add_argument("--accept", help="a file of fact ids the reviewer accepts (one per line) for the port batch")
    p.add_argument("--seed", type=int, default=1)
    p.add_argument("--sample", type=int, default=5, help="size of the agreed items' spot-check sample")
    p.set_defaults(func=triage_command)
    return root


if __name__ == "__main__":
    arguments = parser().parse_args()
    arguments.func(arguments)
