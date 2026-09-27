#!/usr/bin/env python3
"""Small synthetic-fixture tests for check_lean.py (no candidate workbooks)."""

from __future__ import annotations

import datetime as dt
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment

import check_lean


QUOTES = [
    "On January 2, the target requested proposals.",
    "On January 3, Alpha submitted an offer of $10 per share.",
    "On January 4, the target signed the merger agreement.",
]


BID_TERMS = {
    "Stock %": 0,
    "Due diligence": "Not begun",  # the bid came before an NDA (E12)
    "Financing": "Not stated",
    "Regulatory": "Not stated",
    "Exclusivity": "Not stated",
}


def finish_sheet(ws, width: int) -> None:
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, width).column_letter}{max(ws.max_row, 1)}"
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=width):
        for cell in row:
            if cell.value is not None:
                cell.alignment = Alignment(wrap_text=True)


def build_valid_fixture(directory: Path) -> tuple[Path, Path]:
    filing = directory / "filing.htm"
    filing.write_text(
        "<html><body>"
        f"<p>{QUOTES[0]}</p>"
        "<p>On January 3, Alpha submitted an offer of $<b>10</b> per share.</p>"
        f"<p>{QUOTES[2]}</p>"
        "</body></html>",
        encoding="utf-8",
    )

    workbook = directory / "valid.xlsx"
    wb = Workbook()
    ledger = wb.active
    ledger.title = "Deal ledger"
    ledger.append(check_lean.LEDGER_COLUMNS)
    day = dt.date
    rows = [
        {"#": 1, "When": "01/02/2020", "Who": "Target", "Event": "Round opened", "Round": 1, "day": day(2020, 1, 2),
         "Note": "Requested proposals from Alpha.", "q": 0},
        {"#": 2, "When": "01/03/2020", "Who": "Alpha", "Type": "Strategic", "Event": "Bid", "Round": 1, "day": day(2020, 1, 3),
         "Price low": 10, "Price high": 10, "Formality": "Informal", "Conditions": "Unclear", "Count": 1, **BID_TERMS,
         "Note": "Preliminary bid before an NDA.", "q": 1},
        {"#": 3, "When": "01/04/2020", "Who": "Target", "Event": "Merger agreement signed", "Round": 1, "day": day(2020, 1, 4),
         "Note": "$10 per share, all cash.", "q": 2},
    ]
    for values in rows:
        record = dict(values)
        when_day, quote = record.pop("day"), QUOTES[record.pop("q")]
        record.update({"Process": 1, "Quote and page": f"\u201c{quote}\u201d (p. 10)", "Flag": "Q1",
                       **{name: when_day for name in check_lean.DATE_COLUMNS}})
        ledger.append([record.get(column) for column in check_lean.LEDGER_COLUMNS])
    dates = [check_lean.LEDGER_COLUMNS.index(name) + 1 for name in check_lean.DATE_COLUMNS]
    for row in range(2, ledger.max_row + 1):
        for column in dates:
            ledger.cell(row, column).number_format = "MM/DD/YYYY"
    finish_sheet(ledger, len(check_lean.LEDGER_COLUMNS))

    rounds = wb.create_sheet("Rounds")
    rounds.append(check_lean.ROUND_COLUMNS)
    rounds.append(
        [
            1,
            1,
            dt.date(2020, 1, 2),
            "Target requested proposals.",
            "1 strategic: Alpha",
            "none stated",
            "No deadline stated",
            "Not final",
            "1: Alpha",
            "Signing",
        ]
    )
    rounds["C2"].number_format = "MM/DD/YYYY"
    finish_sheet(rounds, len(check_lean.ROUND_COLUMNS))

    questions = wb.create_sheet("Questions")
    questions.append(check_lean.QUESTION_COLUMNS)
    questions.append(
        [
            "Q1",
            "Is one process and one round the best map?",
            "Yes; retain one process and one round.",
            "The target requested proposals once before signing (p. 10).",
            "Rows 1-3",
            "A different answer would move all three rows.",
            None,
        ]
    )
    finish_sheet(questions, len(check_lean.QUESTION_COLUMNS))

    facts = wb.create_sheet("Deal facts")
    facts.append(check_lean.FACT_COLUMNS)
    fact_values = [
        "Target",
        "Alpha",
        "Strategic",
        "$10 per share in cash",
        "01/04/2020",
        "01/04/2020",
        "DEFM14A filed 01/10/2020",
        "10",
        "target-led",
        1,
        "None reported",
        "Not met: 0 qualifying NDAs",
        "Yes",
        "US dollars per share",
        "Bank",
        "Law Firm",
        "The target opened one round. Alpha submitted one bid. The offer was all cash. The target selected Alpha. The parties signed on January 4.",
    ]
    for field, value in zip(check_lean.FACT_FIELDS, fact_values, strict=True):
        facts.append([field, value])
    finish_sheet(facts, len(check_lean.FACT_COLUMNS))

    wb.save(workbook)
    return workbook, filing


def set_ledger(workbook: Path, row: int, values: dict) -> None:
    wb = check_lean.openpyxl.load_workbook(workbook)
    ws = wb["Deal ledger"]
    header = [cell.value for cell in ws[1]]
    for column, value in values.items():
        cell = ws.cell(row, header.index(column) + 1)
        cell.value = value
        cell.alignment = Alignment(wrap_text=True)
    wb.save(workbook)


def append_ledger(workbook: Path, records: list[dict]) -> None:
    """Append ledger rows (quoting a passage the fixture filing contains) with dated cells formatted."""
    wb = check_lean.openpyxl.load_workbook(workbook)
    ws = wb["Deal ledger"]
    header = [cell.value for cell in ws[1]]
    for values in records:
        day = values.pop("day")
        row = {"#": ws.max_row, "When": f"{day:%m/%d/%Y}", "Process": 1, "Round": 1, "Note": "Synthetic row.",
               "Quote and page": f"\u201c{QUOTES[2]}\u201d (p. 10)", **{name: day for name in check_lean.DATE_COLUMNS}, **values}
        ws.append([row.get(name) for name in header])
        for name in check_lean.DATE_COLUMNS:
            ws.cell(ws.max_row, header.index(name) + 1).number_format = "MM/DD/YYYY"
    finish_sheet(ws, len(header))
    wb.save(workbook)


def set_rounds(workbook: Path, values: dict) -> None:
    wb = check_lean.openpyxl.load_workbook(workbook)
    ws = wb["Rounds"]
    for column, value in values.items():
        cell = ws.cell(2, check_lean.ROUND_COLUMNS.index(column) + 1)
        cell.value = value
        cell.alignment = Alignment(wrap_text=True)
    wb.save(workbook)


def add_questions(workbook: Path, questions: list[list]) -> None:
    wb = check_lean.openpyxl.load_workbook(workbook)
    ws = wb["Questions"]
    for values in questions:
        ws.append(values)
    finish_sheet(ws, len(check_lean.QUESTION_COLUMNS))
    wb.save(workbook)


def question(number: int, text: str = "Is Alpha's Type Strategic?", rows: str = "Rows 2") -> list:
    return [f"Q{number}", text, "Yes.", "The filing calls Alpha an operating company (p. 10).", rows, "Row 2's Type changes.", None]


def issue_set(workbook: Path, filing: Path) -> set[tuple[str, str]]:
    return {(issue["code"], issue["severity"]) for issue in check_lean.LeanChecker(workbook, filing).run()["issues"]}


class LeanCheckerTests(unittest.TestCase):
    def test_valid_minimal_fixture_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])
            self.assertEqual(report["summary"]["errors"], 0)

    def test_future_deadline_has_no_outcome_but_reached_deadline_requires_one(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Rounds"]["F2"] = "01/20/2020 (future at filing)"
            wb["Rounds"]["G2"] = None
            wb["Rounds"]["J2"] = "Bidding remains open at filing."
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["summary"]["errors"], 0, report["issues"])

            wb["Deal ledger"]["E4"] = "Deadline"
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertIn("rounds.deadline_count", {issue["code"] for issue in report["issues"]})

    def test_no_deadline_pair(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_rounds(workbook, {"Deadline outcome": None})
            self.assertIn(("rounds.no_deadline_pair", "error"), issue_set(workbook, filing))

    def test_inferred_exit_may_carry_a_reported_reason(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            day = dt.date(2020, 1, 4)
            set_ledger(workbook, 4, {
                "Who": "Alpha", "Event": "Withdrew", "Inferred": "Y", "Exit reason": "Terms or process",
                "Date from": day - dt.timedelta(days=1),
            })
            found = issue_set(workbook, filing)
            self.assertIn(("date.exact_day_mismatch", "error"), found)
            self.assertNotIn("exit.inferred_reason", {code for code, _ in found})

    def test_important_failures_are_reported(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Event": "Invented event", "Sort date": "01/03/2020"})
            set_ledger(workbook, 4, {"Sort date": dt.date(2020, 1, 1),
                                     "Quote and page": "“On January 4 ... merger agreement.” (p. 10)"})
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Questions"]["E2"] = "Rows 1-99"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            codes = {issue["code"] for issue in report["issues"]}
            self.assertEqual(report["status"], "fail")
            self.assertIn("controlled.event", codes)
            self.assertIn("date.not_excel_date", codes)
            self.assertIn("date.sort_outside_window", codes)
            self.assertIn("date.sort_decreases", codes)
            self.assertIn("quote.not_contiguous_in_filing", codes)
            self.assertIn("questions.unknown_row", codes)

    def test_malformed_required_sheet_fails_gracefully(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal ledger"]["A1"] = "Wrong header"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertIn("schema.columns", {issue["code"] for issue in report["issues"]})

    def test_ledger_without_the_v0_header_is_an_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            # Drop the ten bid-term columns, Stock % to Exclusivity: an older ledger shape.
            wb["Deal ledger"].delete_cols(check_lean.LEDGER_COLUMNS.index("Stock %") + 1, 10)
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual((report["status"], report["ledger_schema"]), ("fail", "v0"))
            schema = [issue for issue in report["issues"] if issue["code"] == "schema.columns"]
            self.assertEqual(len(schema), 1, report["issues"])
            self.assertIn("reads only the v0 ledger", schema[0]["message"])

    def test_account_and_quote_wrapper_variants_pass(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 2, {"Quote and page": f"{QUOTES[0]} (p. 10)"})
            set_ledger(workbook, 3, {"Quote and page": f"“{QUOTES[1]} (pp. 10–11)”"})
            set_ledger(workbook, 4, {"Quote and page": f'"{QUOTES[2]}" p. 10'})
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal facts"]["A18"] = "Account (five or six plain sentences)"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_deal_fact_guidance_labels_and_explanations_pass(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            facts = wb["Deal facts"]
            facts["B4"] = "Strategic - operating-company acquirer"
            facts["B10"] = "Target-led - the board initiated outreach"
            facts["B11"] = 1
            facts["A12"] = "Earlier approaches (E5; “None reported” if none)"
            facts["A13"] = "Auction screen (E1)"
            facts["A14"] = "Whole-company bids (Yes, or No with what was bid for)"
            wb.save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_superseded_deal_fact_labels_are_rejected(self) -> None:
        for cell, label in (("A13", "Auction screen (C1)"), ("A12", "Earlier approaches (C7)")):
            with self.subTest(label=label), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_valid_fixture(Path(tmp))
                wb = check_lean.openpyxl.load_workbook(workbook)
                wb["Deal facts"][cell] = label
                wb.save(workbook)
                report = check_lean.LeanChecker(workbook, filing).run()
                self.assertEqual(report["status"], "fail")
                self.assertIn("facts.fields", {issue["code"] for issue in report["issues"]})

    def test_missing_earlier_approaches_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal facts"].delete_rows(12)
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertIn("facts.fields", {issue["code"] for issue in report["issues"]})

    def test_date_formatted_process_count_is_not_silently_accepted(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal facts"]["B11"] = 1
            wb["Deal facts"]["B11"].number_format = "MM/DD/YYYY"
            wb.save(workbook)
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertIn("facts.process_count", {issue["code"] for issue in report["issues"]})

    def test_multiple_flags_and_link_mismatches_are_review_leads(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Flag": "Q1, Q2"})
            add_questions(workbook, [[
                "Q2",
                "Is Alpha's classification supported?",
                "Yes; retain the classification.",
                "The filing describes Alpha as an operating company (p. 10).",
                "Rows 1",
                "A different answer would change Alpha's Type.",
                None,
            ]])

            report = check_lean.LeanChecker(workbook, filing).run()
            issues = {issue["code"]: issue["severity"] for issue in report["issues"]}
            self.assertEqual(report["status"], "pass_with_warnings", report["issues"])
            self.assertEqual(issues["questions.flag_mismatch"], "warning")
            self.assertEqual(issues["questions.reverse_link"], "warning")
            self.assertNotIn("controlled.flag", issues)

            set_ledger(workbook, 3, {"Flag": "Q1, Q9"})
            report = check_lean.LeanChecker(workbook, filing).run()
            errors = {
                issue["code"] for issue in report["issues"] if issue["severity"] == "error"
            }
            self.assertIn("questions.unknown_flag", errors)

    def test_missing_sheets_fail_gracefully(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            filing = directory / "filing.htm"
            filing.write_text("<p>Source text.</p>", encoding="utf-8")
            workbook = directory / "malformed.xlsx"
            Workbook().save(workbook)

            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "fail")
            self.assertEqual(report["issues"][0]["code"], "schema.sheets")

    def test_value_lists_and_consistency_rules(self) -> None:
        cases = [
            ({"Stock %": "abc"}, "controlled.stock_pct", "error"),
            ({"Stock %": "40"}, "controlled.stock_pct", "error"),
            ({"Stock %": 120}, "controlled.stock_pct", "error"),
            ({"Stock %": "75-50"}, "controlled.stock_pct", "error"),
            ({"Stock %": None}, "controlled.stock_pct", "error"),
            ({"Due diligence": None}, "controlled.due_diligence", "error"),
            ({"Financing": "Highly confident"}, "controlled.financing", "error"),
            ({"CVR/earnout": "Yes"}, "controlled.cvr_earnout", "error"),
            ({"CVR/earnout value": 1.25}, "bid.cvr_value_marker", "error"),
            ({"CVR/earnout": "Varies", "CVR/earnout value": 1.25}, "bid.cvr_value_marker", "error"),
            ({"Financing": "Contingent", "Conditions": "Light"}, "conditions.financing_heavy", "error"),
            ({"Conditions": "None", "Due diligence": "Incomplete", "Financing": "Committed"}, "conditions.none_support", "error"),
            ({"Exclusivity": "Varies"}, "bid.varies_single", "error"),
            ({"Antitrust": "Y"}, "conditions.antitrust_regulatory", "error"),
            ({"Regulatory": "Concern", "Conditions": "None", "Due diligence": "Complete", "Financing": "Committed"}, "conditions.none_support", "error"),
            ({"Event": "Other-scope bid"}, "bid.other_scope_per_share", "error"),
            ({"Event": "Other-scope bid", "Price low": None, "Price high": None, "CVR/earnout": "Y", "CVR/earnout value": 1.25}, "bid.other_scope_per_share", "error"),
            ({"Event": "Other-scope bid", "Price low": None, "Price high": None, "Note": None}, "bid.other_scope_note", "warning"),
        ]
        for values, code, severity in cases:
            with self.subTest(values=values):
                with tempfile.TemporaryDirectory() as tmp:
                    workbook, filing = build_valid_fixture(Path(tmp))
                    set_ledger(workbook, 3, values)
                    report = check_lean.LeanChecker(workbook, filing).run()
                    found = {(issue["code"], issue["severity"]) for issue in report["issues"]}
                    self.assertIn((code, severity), found, report["issues"])

    def test_accepts_codes_and_contingent_payments(self) -> None:
        accepted = [
            {"Stock %": 33.3},
            {"Stock %": "Part stock"},
            {"CVR/earnout": "Y", "CVR/earnout value": 1.25},
            {"Regulatory": "Concern", "Antitrust": "Y"},
            {"Financing": "Contingent", "Conditions": "Heavy", "Note": "H1: financing not committed."},
            {"Conditions": "None", "Due diligence": "Complete", "Financing": "Not needed", "Regulatory": "No concern"},
        ]
        for values in accepted:
            with self.subTest(values=values):
                with tempfile.TemporaryDirectory() as tmp:
                    workbook, filing = build_valid_fixture(Path(tmp))
                    set_ledger(workbook, 3, values)
                    report = check_lean.LeanChecker(workbook, filing).run()
                    self.assertEqual(report["status"], "pass", report["issues"])

    def test_terms_must_be_blank_on_other_rows(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 4, {"Financing": "Committed"})
            report = check_lean.LeanChecker(workbook, filing).run()
            codes = {(issue["code"], issue["column"]) for issue in report["issues"]}
            self.assertIn(("bid.fields_on_nonbid", "Financing"), codes, report["issues"])

    def test_markers_and_exclusivity_add_no_issue_at_any_level(self) -> None:
        levels = {
            "None": {"Conditions": "None", "Due diligence": "Complete", "Financing": "Committed", "Regulatory": "No concern"},
            "Light": {"Conditions": "Light", "Due diligence": "Incomplete"},
            "Heavy": {"Conditions": "Heavy", "Note": "H3: may reprice after diligence."},
            "Unclear": {"Conditions": "Unclear"},
        }
        extras = [
            {"CVR/earnout": "Y", "CVR/earnout value": 1.25},
            *({"Exclusivity": value} for value in ("Required", "Requested", "Not stated")),
        ]
        for level, base in levels.items():
            for extra in extras:
                with self.subTest(level=level, extra=extra), tempfile.TemporaryDirectory() as tmp:
                    workbook, filing = build_valid_fixture(Path(tmp))
                    set_ledger(workbook, 3, base)
                    before = check_lean.LeanChecker(workbook, filing).run()["issues"]
                    set_ledger(workbook, 3, extra)
                    after = check_lean.LeanChecker(workbook, filing).run()["issues"]
                    self.assertEqual(before, after)
                    self.assertEqual(after, [], after)

    def test_concern_and_cohort_financing_pass_where_allowed(self) -> None:
        accepted = [
            {"Regulatory": "Concern", "Conditions": "Light", "Due diligence": "Incomplete"},
            {"Regulatory": "Concern", "Conditions": "Heavy", "Note": "H3: may reprice after diligence."},
            {"Regulatory": "Concern", "Conditions": "Unclear"},
            {"Who": "Two financial bidders", "Count": 2, "Financing": "Varies", "Conditions": "Unclear"},
        ]
        for values in accepted:
            with self.subTest(values=values), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_valid_fixture(Path(tmp))
                set_ledger(workbook, 3, values)
                report = check_lean.LeanChecker(workbook, filing).run()
                self.assertEqual(report["status"], "pass", report["issues"])

    def test_other_scope_bid_without_per_share_cells_passes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Event": "Other-scope bid", "Price low": None, "Price high": None})
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_other_scope_price_high_alone_is_an_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Event": "Other-scope bid", "Price low": None, "Price high": 10, "Note": "Ceiling of $10 for the segment."})
            issues = check_lean.LeanChecker(workbook, filing).run()["issues"]
            found = [(issue["code"], issue["severity"], issue["column"]) for issue in issues if issue["code"] == "bid.other_scope_per_share"]
            self.assertEqual(found, [("bid.other_scope_per_share", "error", "Price high")], issues)

    def test_deadline_outcome_values(self) -> None:
        cases = [
            ("Extended (late bid accepted)", set(), {"controlled.deadline_outcome"}),
            ("Enforced", set(), {"controlled.deadline_outcome"}),
            ("Late bids accepted", {("controlled.deadline_outcome", "error")}, set()),
        ]
        for outcome, present, absent in cases:
            with self.subTest(outcome=outcome), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_valid_fixture(Path(tmp))
                set_ledger(workbook, 4, {"Event": "Deadline"})
                set_rounds(workbook, {"Due dates": "01/04/2020", "Deadline outcome": outcome})
                found = issue_set(workbook, filing)
                self.assertTrue(present <= found, found)
                self.assertFalse(absent & {code for code, _ in found}, found)

    def test_review_warning_extended_without_new_date(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 4, {"Event": "Deadline", "Who": "Target"})
            set_rounds(workbook, {"Due dates": "01/04/2020", "Deadline outcome": "Extended"})
            self.assertIn(("rounds.extended_without_new_date", "warning"), issue_set(workbook, filing))
            append_ledger(workbook, [{"day": dt.date(2020, 1, 4), "Who": "Target", "Event": "Deadline revised"}])
            self.assertNotIn("rounds.extended_without_new_date", {code for code, _ in issue_set(workbook, filing)})

    def test_review_warning_exclusivity_repeating_a_bid_request(self) -> None:
        day = dt.date(2020, 1, 3)
        repeat = {"Who": "Alpha", "Event": "Exclusivity changed", "When": "01/03/2020", "Sort date": day, "Date from": day, "Date to": day}
        cases = [
            ({"Exclusivity": "Requested"}, repeat, True),
            ({"Exclusivity": "Required", "Event": "Bid reaffirmed"}, repeat, True),
            ({"Exclusivity": "Requested", "Event": "Other-scope bid", "Price low": None, "Price high": None}, repeat, True),
            ({"Exclusivity": "Not stated"}, repeat, False),
            ({"Exclusivity": "Requested"}, {**repeat, "Who": "Beta"}, False),
            ({"Exclusivity": "Requested"}, {"Who": "Alpha", "Event": "Exclusivity changed"}, False),
        ]
        for bid, row, expected in cases:
            with self.subTest(bid=bid, row=row), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_valid_fixture(Path(tmp))
                set_ledger(workbook, 3, bid)
                set_ledger(workbook, 4, row)
                codes = {code for code, _ in issue_set(workbook, filing)}
                self.assertEqual("ledger.exclusivity_duplicate" in codes, expected, codes)

    def test_review_warning_activity_after_exit(self) -> None:
        day = dt.date(2020, 1, 4)
        bid = {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "Bid", "Count": 1, "Formality": "Informal",
               "Conditions": "Heavy", **BID_TERMS}
        exit_row = {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "Withdrew", "Count": 1, "Exit reason": "Not stated"}
        cases = [
            ([exit_row, bid], True),
            ([exit_row, {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "NDA signed", "Count": 1}], True),
            ([exit_row, {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "Re-entered", "Count": 1}, bid], False),
            ([exit_row, {**bid, "Event": "Other-scope bid"}], False),
            ([exit_row, {**bid, "Who": "Beta"}], False),
        ]
        for rows, expected in cases:
            with self.subTest(rows=[row["Event"] for row in rows]), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_valid_fixture(Path(tmp))
                append_ledger(workbook, [dict(row) for row in rows])
                found = issue_set(workbook, filing)
                self.assertEqual(("exit.activity_without_reentry", "warning") in found, expected, found)
                self.assertNotIn("error", {severity for _, severity in found}, found)

    def test_messages_name_the_rule(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_valid_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Financing": "Contingent", "Conditions": "Light"})
            messages = {issue["code"]: issue["message"] for issue in check_lean.LeanChecker(workbook, filing).run()["issues"]}
            self.assertIn("(E12, H1)", messages["conditions.financing_heavy"])
            set_ledger(workbook, 3, {"Stock %": "abc"})
            messages = {issue["code"]: issue["message"] for issue in check_lean.LeanChecker(workbook, filing).run()["issues"]}
            self.assertIn("Part stock, Not stated or Varies (E13)", messages["controlled.stock_pct"])

    def test_cli_writes_json_for_unreadable_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            directory = Path(tmp)
            workbook = directory / "broken.xlsx"
            workbook.write_text("not an xlsx", encoding="utf-8")
            filing = directory / "filing.htm"
            filing.write_text("<p>Source text.</p>", encoding="utf-8")
            output = directory / "report.json"

            status = check_lean.main(
                ["--workbook", str(workbook), "--filing", str(filing), "--output", str(output)]
            )
            report = json.loads(output.read_text(encoding="utf-8"))
            self.assertEqual(status, 2)
            self.assertEqual(report["status"], "error")
            self.assertEqual(report["issues"][0]["code"], "input.workbook_unreadable")


# ---- rules checked one at a time ---------------------------------------------------------------


def build_no_question_fixture(directory: Path) -> tuple[Path, Path]:
    """The valid fixture with no Question, since one process needs none."""
    workbook, filing = build_valid_fixture(directory)
    wb = check_lean.openpyxl.load_workbook(workbook)
    ledger = wb["Deal ledger"]
    flag = check_lean.LEDGER_COLUMNS.index("Flag") + 1
    for row in range(2, ledger.max_row + 1):
        ledger.cell(row, flag).value = None
    wb["Questions"].delete_rows(2, wb["Questions"].max_row)
    wb.save(workbook)
    return workbook, filing


def check(workbook: Path, filing: Path) -> dict:
    return check_lean.LeanChecker(workbook, filing).run()


def codes_of(report: dict) -> set[tuple[str, str]]:
    return {(issue["code"], issue["severity"]) for issue in report["issues"]}


EXAMPLE_PASSAGES = [
    "On January 10, 2021, the board decided to explore a sale of the Company.",
    "Over the following six weeks, 14 parties, including Party A and Party B, signed confidentiality agreements and were asked to submit indications by March 3, 2021.",
    "On March 3, 2021, Party A submitted an indication of $25.00 per share in cash.",
    "On March 3, 2021, Party B submitted an indication of $24.00 per share in cash.",
    "On March 3, 2021, Party C submitted an indication of $30.00 per share in cash, with financing not committed.",
    "On March 3, 2021, Party D submitted the lowest indication, $20.00 per share in cash.",
    "On March 3, 2021, Party E submitted an indication of $28.00 per share in cash.",
    "On March 3, 2021, Party G submitted an indication of $27.00 per share in cash.",
    "The Company informed Party D its proposal was insufficient but that it could submit a revised proposal.",
    "On March 15, 2021, final-round letters asking for best and final offers by March 29, 2021 went to Parties C, E and G.",
    "Party C confirmed its prior proposal remained its best and final offer.",
    "Party G, which had signed a confidentiality agreement, made a proposal subject to a 45-day exclusivity period to complete due diligence and negotiate the merger agreement.",
    "On March 29, 2021, Party E submitted a best and final offer of $29.00 per share in cash with a markup of the merger agreement and committed financing.",
    "On April 2, 2021, Parent's counsel proposed that the sponsor's liability be capped at $40 million.",
    "On April 9, 2021, the Company and Party E signed the merger agreement.",
]


def build_examples_fixture(directory: Path) -> tuple[Path, Path]:
    """The instruction's five synthetic Examples, entered as one small deal.

    Example 1 is row 9 (the cohort closure), Example 2 row 14 (Same as #5), Example 3 row 10
    (Party D not invited), Example 4 row 15 (Party G's exclusivity period) and Example 5 row 18
    (the sponsor liability cap with no price).
    """
    directory.mkdir(parents=True, exist_ok=True)
    filing = directory / "examples.htm"
    filing.write_text("<html><body>" + "".join(f"<p>{text}</p>" for text in EXAMPLE_PASSAGES) + "</body></html>", encoding="utf-8")
    day = dt.date
    cash = {"Stock %": 0, "Due diligence": "Incomplete", "Financing": "Not stated", "Regulatory": "Not stated", "Exclusivity": "Not stated"}

    def bid(who, when, price, quote, **extra):
        return {"Who": who, "Type": "Financial", "Event": "Bid", "Round": 1, "Price low": price, "Price high": price, "Formality": "Informal",
                "Conditions": "Unclear", "Count": 1, "day": when, "q": quote, **cash, **extra}

    def exit_(who, event, when, quote, round_=1, **extra):
        return {"Who": who, "Type": "Financial", "Event": event, "Round": round_, "Count": 1, "Exit reason": "Not stated", "Inferred": "Y",
                "When": f"by {when:%m/%d/%Y}", "day": when, "from": None, "q": quote, **extra}

    rows = [
        {"Who": "Target", "Event": "Target sale decision", "Round": 0, "day": day(2021, 1, 10), "q": 0, "Note": "Board to explore a sale."},
        {"Who": "Target", "Event": "Round opened", "Round": 1, "day": day(2021, 1, 20), "q": 1, "When": "late January 2021",
         "from": day(2021, 1, 20), "to": day(2021, 1, 31), "Note": "Outreach; indications due 03/03/2021."},
        bid("Party A", day(2021, 3, 3), 25, 2),
        bid("Party B", day(2021, 3, 3), 24, 3),
        bid("Party C", day(2021, 3, 3), 30, 4, Financing="Contingent", Conditions="Heavy", Note="H1: financing not committed."),
        bid("Party D", day(2021, 3, 3), 20, 5),
        bid("Party E", day(2021, 3, 3), 28, 6),
        bid("Party G", day(2021, 3, 3), 27, 7),
        # Example 1: unnamed members of a reported total with no reported offer.
        {"Who": "7 other NDA signers", "Type": "Unknown", "Event": "Did not submit", "Round": 1, "Count": 7, "Exit reason": "Not stated",
         "Inferred": "Y", "When": "by 03/03/2021", "day": day(2021, 3, 3), "from": None, "q": 1, "Note": "Count: 14 signers less Parties A to G."},
        # Example 3: not invited into the stage when it opens.
        exit_("Party D", "Dropped by target", day(2021, 3, 15), 8),
        exit_("Party A", "Dropped by target", day(2021, 3, 15), 9),
        exit_("Party B", "Dropped by target", day(2021, 3, 15), 9),
        {"Who": "Target", "Event": "Round opened", "Round": 2, "day": day(2021, 3, 15), "q": 9, "Note": "Best and final offers due 03/29/2021."},
        # Example 2: the same offer, restated in answer to the final-round letter.
        bid("Party C", day(2021, 3, 29), 30, 10, Round=2, When="by 03/29/2021", Formality="Formal", Financing="Contingent", Conditions="Heavy",
            **{"Due diligence": "Incomplete"}, Note="Same as #5. H1: financing not committed.", **{"from": None}),
        # Example 4: a period that also covers negotiation is not H2.
        bid("Party G", day(2021, 3, 29), None, 11, Round=2, When="by 03/29/2021", Formality="Formal", Exclusivity="Required",
            **{"Stock %": "Not stated"}, Note="45 days' exclusivity.", **{"from": None}),
        bid("Party E", day(2021, 3, 29), 29, 12, Round=2, Formality="Formal", Financing="Committed", Conditions="Unclear"),
        {"Who": "Target", "Event": "Deadline", "Round": 2, "day": day(2021, 3, 29), "q": 9, "Note": "Final bids due."},
        # Example 5: a commitment-only revision; no price, not a price observation.
        bid("Party E", day(2021, 4, 2), None, 13, Round=2, **{"Stock %": "Not stated", "Due diligence": "Not stated"}, Note="Sponsor liability cap $40m proposed."),
        {"Who": "Target", "Event": "Merger agreement signed", "Round": 2, "day": day(2021, 4, 9), "q": 14, "Note": "$29.00 per share in cash with Party E."},
        exit_("Party C", "Not selected at signing", day(2021, 4, 9), 14, round_=2),
        exit_("Party G", "Not selected at signing", day(2021, 4, 9), 14, round_=2),
    ]
    workbook = directory / "examples.xlsx"
    wb = Workbook()
    ledger = wb.active
    ledger.title = "Deal ledger"
    ledger.append(check_lean.LEDGER_COLUMNS)
    for number, values in enumerate(rows, start=1):
        values = dict(values)
        when_day, quote = values.pop("day"), EXAMPLE_PASSAGES[values.pop("q")]
        start, end = values.pop("from", when_day), values.pop("to", when_day)
        # A quotation holds at most 30 words (B).
        quote = " ".join(quote.split()[:30])
        record = {"#": number, "When": f"{when_day:%m/%d/%Y}", "Process": 1, "Quote and page": f"“{quote}” (p. 20)",
                  "Sort date": when_day, "Date from": start, "Date to": end, **values}
        ledger.append([record.get(column) for column in check_lean.LEDGER_COLUMNS])
    for row in range(2, ledger.max_row + 1):
        for name in check_lean.DATE_COLUMNS:
            ledger.cell(row, check_lean.LEDGER_COLUMNS.index(name) + 1).number_format = "MM/DD/YYYY"
    finish_sheet(ledger, len(check_lean.LEDGER_COLUMNS))

    rounds = wb.create_sheet("Rounds")
    rounds.append(check_lean.ROUND_COLUMNS)
    rounds.append([1, 1, day(2021, 1, 20), "Outreach to prospective buyers", "14 NDA signers, including Parties A to G", "03/03/2021", None,
                   "Not final", "6: Parties A, B, C, D, E and G", "Three advanced; eleven out"])
    rounds.append([1, 2, day(2021, 3, 15), "Final-round letters", "3: Parties C, E and G", "03/29/2021", "Enforced",
                   "Announced as final", "3: Parties C, E and G", "Signing with Party E"])
    for row in (2, 3):
        rounds.cell(row, 3).number_format = "MM/DD/YYYY"
    finish_sheet(rounds, len(check_lean.ROUND_COLUMNS))

    questions = wb.create_sheet("Questions")
    questions.append(check_lean.QUESTION_COLUMNS)
    finish_sheet(questions, len(check_lean.QUESTION_COLUMNS))

    facts = wb.create_sheet("Deal facts")
    facts.append(check_lean.FACT_COLUMNS)
    values = ["Example Co", "Party E", "Financial", "$29.00 per share in cash", "04/09/2021", "not reported", "Synthetic, 2021", "20",
              "target-led", 1, "None reported", "Met (process 1): 14 parties", "Yes", "US dollars per share", "Bank", "Law Firm",
              "The board decided to sell in January. Fourteen parties signed agreements. Six bid in March. Three were invited to a final round. "
              "Party E won and signed in April."]
    for field, value in zip(check_lean.FACT_FIELDS, values, strict=True):
        facts.append([field, value])
    finish_sheet(facts, len(check_lean.FACT_COLUMNS))
    wb.save(workbook)
    return workbook, filing


class LeanCheckerRuleTests(unittest.TestCase):
    """Each rule on the fixture with no Question."""

    def check_rule(self, values: dict, code: str, severity: str | None, row: int = 3) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_ledger(workbook, row, values)
            found = {sev for c, sev in issue_set(workbook, filing) if c == code}
            self.assertEqual(found, {severity} if severity else set(), (values, issue_set(workbook, filing)))

    def test_fixture_passes_and_reports_v0(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            report = check_lean.LeanChecker(workbook, filing).run()
            self.assertEqual((report["status"], report["ledger_schema"], report["checker_version"]), ("pass", "v0", "v0"), report["issues"])

    def test_five_examples_pass_with_no_errors(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_examples_fixture(Path(tmp))
            report = check(workbook, filing)
            self.assertEqual(report["summary"]["errors"], 0, report["issues"])
            self.assertEqual(report["status"], "pass", report["issues"])

    def test_cli_exit_status_and_schema(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            output = Path(tmp) / "report.json"
            argv = ["--workbook", str(workbook), "--filing", str(filing), "--output", str(output)]
            self.assertEqual(check_lean.main(argv), 0)
            set_ledger(workbook, 3, {"Stock %": "40-60"})
            self.assertEqual(check_lean.main(argv), 1)
            self.assertEqual(json.loads(output.read_text())["ledger_schema"], "v0")

    def test_note_over_40_words_is_an_error(self) -> None:
        self.check_rule({"Note": " ".join(["word"] * 41)}, "ledger.note_length", "error")
        self.check_rule({"Note": " ".join(["word"] * 40)}, "ledger.note_length", None)

    def test_inferred_only_on_exit_round_opened_and_process_restarted(self) -> None:
        self.check_rule({"Inferred": "Y"}, "ledger.inferred_event", "error")
        self.check_rule({"Inferred": "Y"}, "ledger.inferred_event", None, row=2)  # Round opened
        self.check_rule({"Event": "Withdrew", "Inferred": "Y", "Exit reason": "Not stated", "Who": "Alpha", "Note": None,
                         "Price low": None, "Price high": None, "Formality": None, "Conditions": None, "Stock %": None,
                         **{c: None for c in check_lean.CONDITION_COLUMNS}}, "ledger.inferred_event", None)

    def test_formality_unclear_only_on_cohort_rows(self) -> None:
        self.check_rule({"Formality": "Unclear"}, "bid.formality_unclear", "error")
        self.check_rule({"Formality": "Unclear", "Count": 2}, "bid.formality_unclear", None)
        self.check_rule({"Formality": "Unclear", "Who": "Two financial bidders"}, "bid.formality_unclear", None)

    def test_antitrust_needs_regulatory_concern(self) -> None:
        self.check_rule({"Antitrust": "Y", "Regulatory": "No concern"}, "conditions.antitrust_regulatory", "error")
        self.check_rule({"Antitrust": "Y", "Regulatory": "Concern"}, "conditions.antitrust_regulatory", None)
        self.check_rule({"Antitrust": "Y"}, "conditions.antitrust_regulatory", "error")

    def test_stock_range_is_an_error(self) -> None:
        self.check_rule({"Stock %": "40-60"}, "controlled.stock_pct", "error")
        self.check_rule({"Stock %": "50–75"}, "controlled.stock_pct", "error")
        self.check_rule({"Stock %": "Part stock", "Note": "Stock 40-60% of value."}, "controlled.stock_pct", None)
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Stock %": "40-60"})
            message = next(i["message"] for i in check(workbook, filing)["issues"] if i["code"] == "controlled.stock_pct")
            self.assertIn("Part stock", message)

    def test_markers_take_y_or_blank_only(self) -> None:
        cohort = {"Who": "Two financial bidders", "Count": 2}
        self.check_rule({**cohort, "CVR/earnout": "Varies"}, "controlled.cvr_earnout", "error")
        self.check_rule({**cohort, "Regulatory": "Varies", "Antitrust": "Varies"}, "controlled.antitrust", "error")

    def test_initiation_takes_three_values(self) -> None:
        for value, codes in (("mixed", {"controlled.initiation"}), ("unclear", {"controlled.initiation"}), ("bidder-led", set())):
            with self.subTest(value=value), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                wb = check_lean.openpyxl.load_workbook(workbook)
                wb["Deal facts"]["B10"] = value
                wb.save(workbook)
                self.assertEqual({c for c, _ in issue_set(workbook, filing)} & {"controlled.initiation"}, codes)
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal facts"]["A10"] = "Initiation (target-led, bidder-led, activist-influenced, mixed or unclear)"
            wb.save(workbook)
            self.assertIn(("facts.fields", "error"), issue_set(workbook, filing))
    def test_auction_screen_is_met_or_not_met_with_a_number(self) -> None:
        cases = [("Uncertain: count unknown", True), ("Met: count unknown", True), ("Met (process 1): 3 parties", False),
                 ("Not met (process 1): 1 party; Met (process 2): 4 parties", False)]
        for value, error in cases:
            with self.subTest(value=value), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                wb = check_lean.openpyxl.load_workbook(workbook)
                wb["Deal facts"]["B13"] = value
                wb.save(workbook)
                self.assertEqual(("facts.auction_screen", "error") in issue_set(workbook, filing), error)
    def test_question_count_excludes_the_process_question(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            add_questions(workbook, [question(n) for n in range(1, 6)])
            set_ledger(workbook, 3, {"Flag": "Q1; Q2; Q3; Q4; Q5"})
            self.assertNotIn("questions.count", {c for c, _ in issue_set(workbook, filing)})
            add_questions(workbook, [question(6, "Process: is this one process with one round?", "Rows 1")])
            set_ledger(workbook, 2, {"Flag": "Q6"})
            self.assertNotIn("questions.count", {c for c, _ in issue_set(workbook, filing)})
            add_questions(workbook, [question(7)])
            set_ledger(workbook, 3, {"Flag": "Q1; Q2; Q3; Q4; Q5; Q7"})
            self.assertIn(("questions.count", "warning"), issue_set(workbook, filing))

    def test_process_question_is_recognised_and_expected_for_two_processes(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            day = dt.date(2020, 1, 4)
            append_ledger(workbook, [{"day": day, "Who": "Target", "Event": "Process restarted", "Process": 2, "Round": 0, "Count": 1,
                                      "Inferred": "Y", "Note": "Last reported contact 01/03/2020."}])
            wb = check_lean.openpyxl.load_workbook(workbook)
            wb["Deal facts"]["B11"] = 2
            wb.save(workbook)
            self.assertIn(("questions.process_missing", "warning"), issue_set(workbook, filing))
            add_questions(workbook, [question(n) for n in range(1, 6)] + [question(6, "Is the break a new process?", "#4")])
            set_ledger(workbook, 3, {"Flag": "Q1; Q2; Q3; Q4; Q5"})
            set_ledger(workbook, 5, {"Flag": "Q6"})
            found = {c for c, _ in issue_set(workbook, filing)}
            self.assertNotIn("questions.process_missing", found)
            self.assertNotIn("questions.count", found)

    def test_no_mandatory_map_or_deadline_question(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_ledger(workbook, 4, {"Event": "Deadline"})
            set_rounds(workbook, {"Due dates": "01/04/2020", "Deadline outcome": "Enforced"})
            self.assertEqual(check(workbook, filing)["status"], "pass")
    def test_question_entry_is_at_most_60_words(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            long = question(1, "Is Alpha " + " ".join(["really"] * 50) + " Strategic?")
            add_questions(workbook, [long])
            set_ledger(workbook, 3, {"Flag": "Q1"})
            self.assertIn(("questions.length", "warning"), issue_set(workbook, filing))

    def test_rounds_due_dates_compare_by_prefix(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_rounds(workbook, {"Due dates": "none stated (the letter set no date)"})
            self.assertNotIn("rounds.no_deadline_pair", {c for c, _ in issue_set(workbook, filing)})

    def test_inference_note_is_a_warning_for_restarts_and_cohort_closures(self) -> None:
        day = dt.date(2020, 1, 4)
        cases = [
            ({"Who": "Target", "Event": "Process restarted", "Process": 2, "Round": 0, "Count": 1}, "warning"),
            ({"Who": "5 other NDA signers", "Type": "Unknown", "Event": "Did not submit", "Count": 5, "Exit reason": "Not stated"}, "warning"),
            ({"Who": "Beta", "Type": "Unknown", "Event": "Did not submit", "Count": 1, "Exit reason": "Not stated"}, None),
        ]
        for values, severity in cases:
            with self.subTest(event=values["Event"], who=values["Who"]), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                append_ledger(workbook, [{"day": day, "Inferred": "Y", "Note": None, **values}])
                found = {sev for c, sev in issue_set(workbook, filing) if c == "ledger.inference_note"}
                self.assertEqual(found, {severity} if severity else set())

    def test_count_on_process_markers_and_group_changes(self) -> None:
        day = dt.date(2020, 1, 4)
        cases = [
            ({"Who": "Alpha and Beta", "Event": "Bidding group changed", "Count": None}, ("ledger.count_bidder", "error")),
            ({"Who": "Alpha and Beta", "Event": "Bidding group changed", "Count": 1}, None),
            ({"Who": "Target", "Event": "Process terminated", "Count": None}, ("ledger.count_marker", "warning")),
            ({"Who": "Target", "Event": "Process terminated", "Count": 1}, None),
        ]
        for values, expected in cases:
            with self.subTest(values=values), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                append_ledger(workbook, [{"day": day, **values}])
                found = {(c, sev) for c, sev in issue_set(workbook, filing) if c in {"ledger.count_bidder", "ledger.count_marker"}}
                self.assertEqual(found, {expected} if expected else set())

    def test_count_note_accepts_a_qualifier_and_warns_on_a_range(self) -> None:
        cases = [("Count: more than ten.", None), ("Count: approximately 20.", None), ("Count: 11–14.", ("ledger.count_range", "warning")),
                 ("Count: unknown.", ("ledger.count_range", "warning")), ("", ("ledger.count_bidder", "error"))]
        for note, expected in cases:
            with self.subTest(note=note), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                set_ledger(workbook, 3, {"Who": "Financial bidders", "Count": None, "Note": note})
                found = {(c, s) for c, s in issue_set(workbook, filing) if c in {"ledger.count_range", "ledger.count_bidder", "ledger.count_uncertain"}}
                self.assertEqual(found, {expected} if expected else set())

    def test_round_zero_after_round_one_opens_is_an_error(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_ledger(workbook, 4, {"Round": 0})
            self.assertIn(("round.zero_after_opening", "error"), issue_set(workbook, filing))

    def test_same_as_points_to_an_earlier_bid_of_the_same_bidder(self) -> None:
        day = dt.date(2020, 1, 4)
        again = {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "Bid", "Count": 1, "Price low": 10, "Price high": 10,
                 "Formality": "Informal", "Conditions": "Unclear", **BID_TERMS}
        cases = [("Same as #2.", False), ("Same as #1.", True), ("Same as #9.", True), ("Same as #5.", True)]
        for note, error in cases:
            with self.subTest(note=note), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                append_ledger(workbook, [{**again, "Note": note}])
                self.assertEqual(("ledger.same_as", "error") in issue_set(workbook, filing), error)
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            append_ledger(workbook, [{**again, "Who": "Beta", "Note": "Same as #2."}])
            self.assertIn(("ledger.same_as", "error"), issue_set(workbook, filing))

    def test_heavy_note_begins_with_its_trigger(self) -> None:
        self.check_rule({"Conditions": "Heavy", "Note": "Financing not committed."}, "conditions.heavy_trigger", "warning")
        self.check_rule({"Conditions": "Heavy", "Note": "H3: may reprice after diligence."}, "conditions.heavy_trigger", None)
        self.check_rule({"Conditions": "Heavy", "Note": "Same as #2. H1: financing not committed."}, "conditions.heavy_trigger", None)

    def test_bid_reaffirmed_is_formal_and_a_same_offer_row(self) -> None:
        day = dt.date(2020, 1, 4)
        base = {"day": day, "Who": "Alpha", "Type": "Strategic", "Event": "Bid reaffirmed", "Count": 1, "Price low": 10, "Price high": 10,
                "Conditions": "Unclear", **BID_TERMS}
        cases = [({"Formality": "Formal", "Note": "Same as #2."}, set()),
                 ({"Formality": "Informal", "Note": "Same as #2."}, {"bid.reaffirmed_formal"}),
                 ({"Formality": "Formal", "Note": "Reaffirmed."}, {"bid.reaffirmed_same_as"})]
        for values, expected in cases:
            with self.subTest(values=values), tempfile.TemporaryDirectory() as tmp:
                workbook, filing = build_no_question_fixture(Path(tmp))
                append_ledger(workbook, [{**base, **values}])
                found = {c for c, _ in issue_set(workbook, filing)} & {"bid.reaffirmed_formal", "bid.reaffirmed_same_as"}
                self.assertEqual(found, expected)

    def test_light_needs_diligence_complete_or_incomplete(self) -> None:
        self.check_rule({"Conditions": "Light", "Due diligence": "Not stated", "Financing": "Committed"}, "conditions.light_support", "warning")
        self.check_rule({"Conditions": "Light", "Due diligence": "Incomplete"}, "conditions.light_support", None)

    def test_other_scope_message_cites_the_current_part(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            workbook, filing = build_no_question_fixture(Path(tmp))
            set_ledger(workbook, 3, {"Event": "Other-scope bid"})
            message = next(i["message"] for i in check(workbook, filing)["issues"] if i["code"] == "bid.other_scope_per_share")
            self.assertIn("(D2)", message)
            self.assertNotIn("F.4", message)


if __name__ == "__main__":
    unittest.main()
