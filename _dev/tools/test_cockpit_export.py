"""The export-to-repository script, on a fixture repository with a real workspace."""
from __future__ import annotations

import contextlib
import csv
import hashlib
import io
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from cockpit import data, export_repo
from test_cockpit import tree_digest
from test_cockpit_workspace import fixture

SLUG = "beta-holdings"
FILE = f"{SLUG}_2021-03-04_DEFM14A.htm"
FILING = b"<h2>Background of the Merger</h2><p>Beta met Party A.</p>"


def sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


class ExportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit, self.original = fixture(self.root)
        self.ws, self.instructions = self.cockpit.workspace, self.cockpit.instructions
        self.base = self.instructions.list()["items"][0]  # imports the repository instruction

    def tearDown(self): self.temp.cleanup()

    def export(self, *args):
        out, err = io.StringIO(), io.StringIO()
        with contextlib.redirect_stdout(out), contextlib.redirect_stderr(err):
            try:
                code = export_repo.main(["--repo-root", str(self.root), *args])
            except SystemExit as exc:  # an argument error
                code = exc.code
        return code, out.getvalue(), err.getvalue()

    def write(self, *args):
        """An export with --write into the fixture checkout itself (--write requires --out-root)."""
        return self.export("--out-root", str(self.root), *args, "--write")

    def published(self, text="# Synthetic instruction Version 1\n\nRead the filing twice.\n", name="Version 1"):
        draft = self.instructions.request("alex", {"action": "draft", "from": self.base["id"]})
        self.instructions.request("alex", {"action": "save", "id": draft["item"]["id"], "text": text, "base_sha256": draft["item"]["sha256"]})
        if name:
            self.instructions.request("austin", {"action": "publish", "id": draft["item"]["id"], "name": name, "note": "Test"})
        return self.instructions.detail(draft["item"]["id"])["item"]

    def add_deal(self, edit=False):
        """An added deal with one imported version, as the cockpit leaves it after its first run."""
        folder = self.root / "_dev/cockpit/state/filings" / SLUG
        folder.mkdir(parents=True)
        (folder / FILE).write_bytes(FILING)
        version = self.root / "_dev/cockpit/state/versions" / SLUG / "run1"
        version.mkdir(parents=True)
        wb = load_workbook(self.root / "extraction/alpha-deal.xlsx")
        wb["Deal ledger"]["A2"] = 1
        wb.save(version / f"{SLUG}.xlsx")
        workbook = (version / f"{SLUG}.xlsx").read_bytes()
        conn = self.ws._connect(write=True)
        conn.execute("INSERT INTO added_deals VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (SLUG, "Beta Holdings", "DEFM14A", "2021-03-04", FILE, "link", None, "https://www.sec.gov/x-index.htm",
                      "https://www.sec.gov/Archives/edgar/data/42/0000000042-21-000007.txt", "beta-proxy.htm",
                      "2026-09-23T10:00:00+00:00", len(FILING), sha(FILING), "alex", "2026-09-23T10:00:05+00:00"))
        conn.execute("INSERT INTO versions (slug, id, label, path, sha256, kind, engine, model, effort, instruction_version, instruction_sha256, filing_sha256, started_by, started_at, finished_at, receipts, checker) VALUES (?,?,?,?,?,'raw','Opus 5.5','claude-opus-5-5','medium','Version 0','i','f','alex','2026-09-23T10:00:00+00:00','2026-09-23T10:10:00+00:00',?,'{}')",
                     (SLUG, "run1", "Opus 5.5 · medium", str((version / f"{SLUG}.xlsx").relative_to(self.root)), sha(workbook), str(version.relative_to(self.root))))
        conn.commit(); conn.close()
        if edit:
            row = self.cockpit.deal(SLUG)["ledger"]["rows"][0]
            self.ws.edit(SLUG, {"revision": 0, "base_sha256": sha(workbook), "reason": "Fix price", "operations": [
                {"type": "update", "sheet": "Deal ledger", "uid": row["uid"], "values": {"Price low": "99"}}]}, "austin")
        return workbook

    # ---- instructions -----------------------------------------------------------------

    def test_dry_run_changes_nothing(self):
        item = self.published()
        self.add_deal(edit=True)
        before = tree_digest(self.root)
        code, out, _ = self.export("instruction", "Version 1")
        self.assertEqual(code, 0)
        target = self.root / "SEC_Deal_Ledger_Extraction_Instruction.md"
        self.assertIn(f"SEC_Deal_Ledger_Extraction_Instruction.md  {sha(target.read_bytes())} -> {item['sha256']}", out)
        self.assertIn("dry run", out)
        code, out, _ = self.export("deal", SLUG, "--version", "working")
        self.assertEqual(code, 0)
        self.assertIn(f"extraction/{SLUG}.xlsx  new -> ", out)
        self.assertIn(f"raw_filing/{FILE}  new -> {sha(FILING)}", out)
        self.assertIn("raw_filing/MANIFEST.csv  ", out)
        self.assertEqual(tree_digest(self.root), before)

    def test_write_published_instruction(self):
        text = "# Synthetic instruction Version 1\n\nRead the filing twice. — “quoted”\n"
        item = self.published(text)
        code, out, _ = self.write("instruction", item["id"])
        self.assertEqual(code, 0, out)
        self.assertEqual((self.root / "SEC_Deal_Ledger_Extraction_Instruction.md").read_bytes(), text.encode("utf-8"))
        self.assertIn(f"-> {item['sha256']}", out)
        self.assertIn("written", out)
        code, out, _ = self.export("instruction", "VERSION 1")  # names match without case; now unchanged
        self.assertIn("(unchanged)", out)

    def test_draft_is_refused(self):
        draft = self.published(name=None)
        before = tree_digest(self.root)
        code, _, err = self.write("instruction", draft["id"])
        self.assertEqual(code, 1)
        self.assertIn("draft", err)
        self.assertEqual(tree_digest(self.root), before)

    def test_hash_mismatch_is_refused(self):
        item = self.published()
        self.instructions.path(item["sha256"]).write_text("tampered\n", encoding="utf-8")
        before = tree_digest(self.root)
        code, _, err = self.write("instruction", "Version 1")
        self.assertEqual(code, 1)
        self.assertIn("does not match its content address", err)
        self.assertEqual(tree_digest(self.root), before)

    # ---- deals -------------------------------------------------------------------------

    def test_added_deal_writes_workbook_filing_and_manifest_row(self):
        workbook = self.add_deal(edit=True)
        manifest = self.root / "raw_filing/MANIFEST.csv"
        before = manifest.read_text(encoding="utf-8")
        code, out, _ = self.write("deal", SLUG, "--version", "working")
        self.assertEqual(code, 0, out)
        written = (self.root / f"extraction/{SLUG}.xlsx").read_bytes()
        self.assertEqual(written, self.ws.export(SLUG))  # the four-sheet working copy; the cockpit's download adds a Source sheet unless ?source=0
        self.assertEqual(load_workbook(io.BytesIO(written)).sheetnames, ["Deal ledger", "Rounds", "Questions", "Deal facts"])
        self.assertEqual(load_workbook(io.BytesIO(written))["Deal ledger"]["H2"].value, 99)
        self.assertEqual((self.root / "raw_filing" / FILE).read_bytes(), FILING)
        text = manifest.read_text(encoding="utf-8")
        self.assertEqual(text.splitlines()[0], before.splitlines()[0])
        self.assertTrue(set(before.splitlines()) <= set(text.splitlines()))  # other rows untouched
        rows = list(csv.DictReader(io.StringIO(text)))
        self.assertEqual([r["file"] for r in rows], sorted(r["file"] for r in rows))
        row = next(r for r in rows if r["deal"] == SLUG)
        self.assertEqual(row, {"file": FILE, "deal": SLUG, "form_type": "DEFM14A", "date_filed": "2021-03-04",
                               "source_url": "https://www.sec.gov/Archives/edgar/data/42/0000000042-21-000007.txt", "document": "beta-proxy.htm",
                               "fetched_utc": "2026-09-23T10:00:00Z", "bytes": str(len(FILING)), "sha256": sha(FILING)})
        # A version export replaces the workbook and updates the row in place, not a second one.
        code, out, _ = self.write("deal", SLUG, "--version", "run1")
        self.assertEqual(code, 0, out)
        self.assertEqual((self.root / f"extraction/{SLUG}.xlsx").read_bytes(), workbook)
        self.assertEqual(manifest.read_text(encoding="utf-8"), text)
        self.assertIn("raw_filing/MANIFEST.csv  " + sha(text.encode()) + " -> " + sha(text.encode()) + "  (unchanged)", out)
        self.assertIn(SLUG, data.Cockpit(self.root).manifest())

    def test_catalog_referenced_target_is_refused(self):
        before = tree_digest(self.root)
        for version in ("working", "base-raw"):
            for export in (self.export, self.write):
                code, _, err = export("deal", "alpha-deal", "--version", version)
                self.assertEqual(code, 1)
                self.assertIn("extraction/alpha-deal.xlsx is an immutable original referenced by _dev/cockpit/catalog.json", err)
                self.assertIn("separate, requested edit of the catalog", err)
        self.assertEqual(tree_digest(self.root), before)
        self.assertEqual((self.root / "extraction/alpha-deal.xlsx").read_bytes(), self.original)

    # ---- a separate output checkout (after the switch-over, the development clone) --------

    def test_out_root_receives_the_export_and_the_state_checkout_is_untouched(self):
        workbook = self.add_deal()
        item = self.published()
        with tempfile.TemporaryDirectory() as folder:
            out = Path(folder)
            (out / "raw_filing").mkdir()
            (out / "raw_filing/MANIFEST.csv").write_text("file,deal,form_type,date_filed,source_url,document,fetched_utc,bytes,sha256\nzz.htm,zeta,DEFM14A,2020-01-01,u,d,t,1,x\n", encoding="utf-8")
            before = tree_digest(self.root)
            for args in (("deal", SLUG, "--version", "run1", "--write"), ("instruction", "Version 1", "--write")):
                code, out_text, err = self.export("--out-root", str(out), *args)
                self.assertEqual(code, 0, err)
            self.assertEqual(tree_digest(self.root), before)
            self.assertEqual((out / f"extraction/{SLUG}.xlsx").read_bytes(), workbook)
            self.assertEqual((out / "raw_filing" / FILE).read_bytes(), FILING)
            self.assertEqual((out / "SEC_Deal_Ledger_Extraction_Instruction.md").read_bytes(), self.instructions.path(item["sha256"]).read_bytes())
            rows = list(csv.DictReader(io.StringIO((out / "raw_filing/MANIFEST.csv").read_text(encoding="utf-8"))))
            self.assertEqual([r["deal"] for r in rows], [SLUG, "zeta"])  # the output checkout's manifest, updated

    def test_write_requires_out_root(self):
        self.add_deal()
        self.published()
        before = tree_digest(self.root)
        for args in (("deal", SLUG, "--version", "run1"), ("instruction", "Version 1")):
            code, out_text, _ = self.export(*args)
            self.assertEqual(code, 0)
            self.assertIn("dry run", out_text)
            code, _, err = self.export(*args, "--write")
            self.assertEqual(code, 2)
            self.assertIn("--write requires --out-root", err)
        self.assertEqual(tree_digest(self.root), before)
        self.assertEqual(self.export("--out-root", str(self.root / "missing"), "deal", SLUG, "--version", "run1")[0], 1)

    def test_unknown_deal_and_version_are_refused(self):
        self.add_deal()
        self.assertEqual(self.export("deal", "nope", "--version", "working")[0], 1)
        self.assertEqual(self.export("deal", SLUG, "--version", "missing")[0], 1)


if __name__ == "__main__":
    unittest.main()
