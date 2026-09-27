"""End-to-end HTTP contract checks against a disposable synthetic repository.

Run from the repository root: python3 -m pytest -q _dev/tools/cockpit/acceptance/test_http.py
No real filing, workbook, catalog, or working-state database is opened by this suite.
"""

from __future__ import annotations

import csv
import datetime as dt
import hashlib
import io
import json
import sys
import tempfile
import threading
import time
from pathlib import Path

import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import PatternFill
import pytest
import requests

TOOLS = Path(__file__).resolve().parents[2]
if str(TOOLS) not in sys.path:
    sys.path.insert(0, str(TOOLS))

from cockpit import data, server  # noqa: E402
import check_lean  # noqa: E402


def synthetic_repo(root: Path) -> tuple[Path, str]:
    """Make a small, traceable four-sheet deal with no research content."""
    extraction = root / "extraction"
    filings = root / "raw_filing"
    catalog_dir = root / "_dev/cockpit"
    docs = root / "_dev/reviews/synthetic"
    for directory in (extraction, filings, catalog_dir, docs):
        directory.mkdir(parents=True, exist_ok=True)
    (root / "SEC_Deal_Ledger_Extraction_Instruction.md").write_text("# Version 1\n\nRead the filing.\nSave the workbook.\n", encoding="utf-8")

    filing = ("""<html><body><h1>Background of the Merger</h1>
    <p>1</p><p>Acme invited bids from three parties.</p>
    <p>Acme received an offer from Party A.</p>
    <p>Acme signed the merger agreement.</p>
    """ + "".join(f"<p>Synthetic page one filler paragraph {n}.</p>" for n in range(60)) + """
    <p>2</p><p>Later supplemental discussion is searchable.</p>
    </body></html>""").encode("utf-8")
    filing_path = filings / "synthetic.htm"
    filing_path.write_bytes(filing)
    with (filings / "MANIFEST.csv").open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["file", "deal", "form_type", "date_filed", "source_url", "document", "fetched_utc", "bytes", "sha256"])
        writer.writeheader()
        writer.writerow({"file": filing_path.name, "deal": "synthetic", "form_type": "DEFM14A", "date_filed": "2026-01-10", "source_url": "https://example.invalid/synthetic", "document": filing_path.name, "fetched_utc": "2026-01-10T00:00:00Z", "bytes": len(filing), "sha256": hashlib.sha256(filing).hexdigest()})

    wb = openpyxl.Workbook()
    ledger = wb.active
    ledger.title = "Deal ledger"
    ledger.append(check_lean.LEDGER_COLUMNS)
    rows = [
        {"#": 1, "When": "2026-01-01", "Who": "Acme", "Type": "Contact", "Event": "Invitation", "Process": 1, "Round": 1, "Count": 3, "Note": "Three parties contacted", "Quote and page": "Acme invited bids from three parties. (p. 1)", "Sort date": dt.datetime(2026, 1, 1), "Date from": dt.datetime(2026, 1, 1)},
        {"#": 2, "When": "Date uncertain", "Who": "Party A", "Type": "Bid", "Event": "Offer", "Process": 1, "Round": 1, "Price low": 12.5, "Price high": 12.5, "Count": None, "Note": "Follow-up to #1", "Quote and page": "Acme received an offer from Party A. (p. 1)", "Flag": "Q1 R1", "Sort date": None, "Date from": None},
        {"#": 3, "When": "2026-01-03", "Who": "Acme", "Type": "Agreement", "Event": "Signing", "Process": 1, "Round": 1, "Note": "After #2", "Quote and page": "Acme signed the merger agreement. (p. 1)", "Flag": "Q1", "Sort date": dt.datetime(2026, 1, 3), "Date from": dt.datetime(2026, 1, 3)},
    ]
    for values in rows:
        ledger.append([values.get(field) for field in check_lean.LEDGER_COLUMNS])
    for field in ("Sort date", "Date from", "Date to"):
        column = check_lean.LEDGER_COLUMNS.index(field) + 1
        for excel_row in range(2, 5):
            ledger.cell(excel_row, column).number_format = "mm/dd/yyyy"
    ledger.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(check_lean.LEDGER_COLUMNS))}4"
    ledger.freeze_panes = "A2"
    ledger.cell(2, check_lean.LEDGER_COLUMNS.index("Who") + 1).fill = PatternFill("solid", fgColor="FFCCDDEE")
    ledger.cell(2, check_lean.LEDGER_COLUMNS.index("Reviewer note") + 1).value = "=1+1"
    ledger.cell(4, check_lean.LEDGER_COLUMNS.index("Event") + 1).hyperlink = "https://example.invalid/record"
    ledger.cell(4, check_lean.LEDGER_COLUMNS.index("Note") + 1).comment = Comment("Preserve synthetic annotation", "Fixture")

    rounds = wb.create_sheet("Rounds")
    rounds.append(check_lean.ROUND_COLUMNS)
    rounds.append([{"Process": 1, "Round": 1, "Opened": dt.datetime(2026, 1, 1), "How opened": "Invitation #1", "Who was in": "Party A", "Bids received": 1}.get(field) for field in check_lean.ROUND_COLUMNS])
    rounds.cell(2, check_lean.ROUND_COLUMNS.index("Opened") + 1).number_format = "mm/dd/yyyy"
    rounds.auto_filter.ref = f"A1:J2"

    questions = wb.create_sheet("Questions")
    questions.append(check_lean.QUESTION_COLUMNS)
    questions.append([{"Q": "Q1", "Question": "When did Party A bid?", "Recommended answer": "Date unknown", "Rows affected": "#2-#3"}.get(field) for field in check_lean.QUESTION_COLUMNS])
    questions.append([{"Q": "R1", "Question": "Review the omitted source event", "Recommended answer": "Omit the event", "Why, with page": "Synthetic source p. 1", "Rows affected": "Omitted source event"}.get(field) for field in check_lean.QUESTION_COLUMNS])
    questions.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(check_lean.QUESTION_COLUMNS))}3"

    facts = wb.create_sheet("Deal facts")
    facts.append(check_lean.FACT_COLUMNS)
    facts.append(["Target", "Synthetic Acme"])
    facts.append(["Number of processes", 1])
    facts.auto_filter.ref = "A1:B3"
    workbook = extraction / "synthetic.xlsx"
    wb.save(workbook)
    wb.close()
    original_hash = hashlib.sha256(workbook.read_bytes()).hexdigest()

    (docs / "finding.md").write_text("# Synthetic finding\nCheck uncertain bid date.\n", encoding="utf-8")
    catalog = {"schema_version": 1, "deals": {"synthetic": {
        "name": "Synthetic Acme", "default_base": "version-1-raw",
        "versions": [{"id": "version-1-raw", "label": "Synthetic raw", "path": "extraction/synthetic.xlsx", "sha256": original_hash, "instruction_version": "Version 1", "kind": "raw", "review_status": "unreviewed"}],
        "findings": [{"id": "F1", "title": "Uncertain date", "detail": "Date remains unknown", "rule": "E8", "source_version": "version-1-raw", "source_label": "Synthetic raw", "source_rows": [2], "evidence": [{"quote": "Acme received an offer from Party A.", "page": "1"}], "proposed_change": "Review date", "needs_recheck": True, "judgment": "unreviewed", "implementation": "not_applied", "verification": "unchecked", "note": "", "actor": None, "at": None}],
        "documents": [{"id": "finding", "label": "Synthetic finding", "source_version": "version-1-raw", "kind": "report", "path": "_dev/reviews/synthetic/finding.md"}],
    }}}
    (catalog_dir / "catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    return workbook, original_hash


def add_pending_catalog_deal(root: Path) -> None:
    """Add a second synthetic filing with no workbook or version."""
    source = root / "raw_filing/synthetic.htm"
    pending = root / "raw_filing/pending.htm"
    pending.write_bytes(source.read_bytes())
    filing = {"file": pending.name, "form_type": "DEFM14A", "date_filed": "2026-01-10", "source_url": "https://example.invalid/pending", "document": pending.name, "sha256": hashlib.sha256(pending.read_bytes()).hexdigest()}
    with (root / "raw_filing/MANIFEST.csv").open("a", newline="", encoding="utf-8") as handle:
        csv.DictWriter(handle, fieldnames=["file", "deal", "form_type", "date_filed", "source_url", "document", "fetched_utc", "bytes", "sha256"]).writerow({**filing, "deal": "pending", "fetched_utc": "2026-01-10T00:00:00Z", "bytes": pending.stat().st_size})
    catalog_path = root / "_dev/cockpit/catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    catalog["deals"]["pending"] = {"name": "Pending Acme", "filing": filing, "versions": [], "findings": [], "documents": []}
    catalog_path.write_text(json.dumps(catalog), encoding="utf-8")


class HttpFixture:
    def __init__(self, root: Path):
        self.root = root
        self.cockpit = data.Cockpit(root)
        self.httpd = server.make_server(0, self.cockpit, quiet=True)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()
        self.base = f"http://127.0.0.1:{self.httpd.server_port}"
        self.client = requests.Session()

    def close(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join(timeout=5)
        self.client.close()

    def get(self, path: str, **kwargs):
        return self.client.get(self.base + path, timeout=30, **kwargs)

    def session(self):
        response = self.get("/api/session")
        assert response.status_code == 200, response.text
        return response.json()

    def post(self, path: str, payload, *, csrf=None, origin=None, content_type="application/json", headers=None):
        csrf = self.session()["csrf_token"] if csrf is None else csrf
        hdrs = {"X-Cockpit-CSRF": csrf, "Content-Type": content_type}
        if origin != "omit":
            hdrs["Origin"] = origin or self.base
        if headers:
            hdrs.update(headers)
        return self.client.post(self.base + path, data=json.dumps(payload) if content_type == "application/json" else payload, headers=hdrs, timeout=30)

    def edit(self, payload, **kwargs):
        return self.post("/api/deal/synthetic/edit", payload, **kwargs)


@pytest.fixture
def env():
    with tempfile.TemporaryDirectory(prefix="cockpit-acceptance-") as folder:
        root = Path(folder)
        workbook, original_hash = synthetic_repo(root)
        http = HttpFixture(root)
        try:
            yield http, workbook, original_hash
        finally:
            http.close()


def deal(http, version="working"):
    response = http.get(f"/api/deal/synthetic?version={version}")
    assert response.status_code == 200, response.text
    return response.json()


def save(http, snapshot, operations, reason="Synthetic acceptance change"):
    payload = {"revision": snapshot["workspace"]["revision"], "base_sha256": snapshot["workspace"]["base_sha256"], "reason": reason, "operations": operations}
    return http.edit(payload)


def row(snapshot, sheet, index=0):
    rows = snapshot["facts"] if sheet == "Deal facts" else snapshot[{"Deal ledger": "ledger", "Rounds": "rounds", "Questions": "questions"}[sheet]]["rows"]
    return rows[index]


def xlsx(response):
    assert response.status_code == 200, response.text
    return openpyxl.load_workbook(io.BytesIO(response.content), data_only=False)


def test_baseline_catalog_source_and_raw_are_read_only(env):
    http, workbook, original_hash = env
    initial = deal(http)
    assert initial["workspace"]["revision"] == 0
    assert initial["workspace"]["editable"] is True
    assert {v["id"] for v in initial["versions"]} == {"working", "version-1-raw"}
    assert initial["findings"][0]["judgment"] == "unreviewed"
    assert initial["ledger"]["rows"][1]["cells"]["Count"] == ""
    assert initial["ledger"]["rows"][1]["cells"]["Date from"] == ""
    raw_payload = deal(http, "version-1-raw")
    original_report = check_lean.LeanChecker(workbook, http.root / "raw_filing/synthetic.htm").run()
    assert raw_payload["check"]["summary"] == original_report["summary"]
    attached_sheets = {"Deal ledger", "Rounds", "Questions"}
    expected_other = [issue for issue in original_report["issues"] if issue.get("sheet") not in attached_sheets or not isinstance(issue.get("row"), int) or issue.get("row") == 1]
    assert [(issue["code"], issue.get("sheet"), issue.get("row")) for issue in raw_payload["check"]["other_issues"]] == [(issue["code"], issue.get("sheet"), issue.get("row")) for issue in expected_other]
    assert initial["check"]["summary"] == raw_payload["check"]["summary"]
    assert http.get("/api/document/synthetic/finding").json()["text"].startswith("# Synthetic finding")
    assert "Acme invited bids" in json.dumps(http.get("/api/filing/synthetic").json())
    assert hashlib.sha256(http.get("/api/deal/synthetic/export?version=version-1-raw").content).hexdigest() == original_hash
    with (http.root / "raw_filing/MANIFEST.csv").open(newline="", encoding="utf-8") as handle:
        source_hash = next(csv.DictReader(handle))["sha256"]
    assert hashlib.sha256((http.root / "raw_filing/synthetic.htm").read_bytes()).hexdigest() == source_hash
    raw = deal(http, "version-1-raw")
    assert raw["workspace"]["editable"] is False
    assert workbook.read_bytes() == http.get("/api/deal/synthetic/export?version=version-1-raw").content


def test_catalog_deal_without_a_version_is_pending():
    with tempfile.TemporaryDirectory(prefix="cockpit-pending-catalog-") as folder:
        root = Path(folder)
        synthetic_repo(root)
        add_pending_catalog_deal(root)
        http = HttpFixture(root)
        try:
            listed = next(item for item in http.get("/api/deals").json() if item["slug"] == "pending")
            assert listed["pending"] is True and listed["name"] == "Pending Acme"
            payload = http.get("/api/deal/pending")
            assert payload.status_code == 200, payload.text
            deal = payload.json()
            assert deal["pending"] is True and deal["versions"] == []
            assert deal["added"] is None
            assert deal["workspace"]["base_sha256"] is None and deal["workspace"]["editable"] is False
            assert deal["ledger"]["rows"] == [] and deal["rounds"]["rows"] == [] and deal["questions"]["rows"] == []
            assert deal["filing"]["file"] == "pending.htm"
        finally:
            http.close()


def test_four_sheet_edits_typed_export_decision_and_restart(env):
    http, workbook, original_hash = env
    initial = deal(http)
    ops = [
        {"type": "update", "sheet": "Deal ledger", "uid": row(initial, "Deal ledger", 1)["uid"], "values": {"Price low": 13.25, "Price high": 13.25, "Count": "", "Date from": "", "Sort date": "2026-01-02"}},
        {"type": "update", "sheet": "Rounds", "uid": row(initial, "Rounds")["uid"], "values": {"Opened": "2026-01-02", "Bids received": 2}},
        {"type": "update", "sheet": "Questions", "uid": row(initial, "Questions")["uid"], "values": {"Recommended answer": "Still uncertain"}},
        {"type": "update", "sheet": "Deal facts", "uid": row(initial, "Deal facts", 0)["uid"], "values": {"Value": "Synthetic Acme Revised"}},
        {"type": "finding", "id": "F1", "judgment": "supported", "implementation": "not_applied", "verification": "unchecked", "note": "Synthetic-only judgment"},
        {"type": "review", "uid": row(initial, "Deal ledger", 1)["uid"], "status": "needs_decision", "note": "Date unknown"},
    ]
    saved = save(http, initial, ops, "Edit four sheets and record judgment")
    assert saved.status_code == 200, saved.text
    payload = saved.json()
    assert payload["workspace"]["revision"] == 1
    assert payload["ledger"]["rows"][1]["cells"]["Count"] == ""
    assert payload["ledger"]["rows"][1]["cells"]["Date from"] == ""
    assert payload["findings"][0]["judgment"] == "supported"
    assert payload["findings"][0]["actor"] == http.session()["user"]
    assert payload["row_review"][row(initial, "Deal ledger", 1)["uid"]]["status"] == "needs_decision"
    assert len(http.get("/api/deal/synthetic/changes").json()["changes"]) >= 6
    history = http.get("/api/deal/synthetic/history").json()["history"]
    assert history[0]["reason"] == "Edit four sheets and record judgment"
    assert history[0]["actor"] == http.session()["user"]
    assert any(change["field"] == "Value" and change["after"] == "Synthetic Acme Revised" for change in history[0]["changes"])

    exported = xlsx(http.get("/api/deal/synthetic/export?version=working"))
    ledger = exported["Deal ledger"]
    def cell(sheet, excel_row, field):
        columns = [c.value for c in exported[sheet][1]]
        return exported[sheet].cell(excel_row, columns.index(field) + 1)
    assert cell("Deal ledger", 3, "Price low").value == 13.25
    assert cell("Deal ledger", 3, "Price low").data_type == "n"
    assert cell("Deal ledger", 3, "Count").value is None
    assert cell("Deal ledger", 3, "Date from").value is None
    assert isinstance(cell("Deal ledger", 3, "Sort date").value, dt.datetime)
    assert isinstance(cell("Rounds", 2, "Opened").value, dt.datetime)
    assert cell("Rounds", 2, "Bids received").data_type == "n"
    assert cell("Deal facts", 2, "Value").value == "Synthetic Acme Revised"
    assert cell("Deal ledger", 2, "Who").fill.fgColor.rgb == "FFCCDDEE"
    assert cell("Deal ledger", 2, "Reviewer note").value == "=1+1"
    assert cell("Deal ledger", 4, "Event").hyperlink.target == "https://example.invalid/record"
    assert cell("Deal ledger", 4, "Note").comment.text == "Preserve synthetic annotation"
    assert list(exported.sheetnames) == ["Deal ledger", "Rounds", "Questions", "Deal facts", "Source"]
    assert ledger.auto_filter.ref.endswith("4")
    exported.close()
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == original_hash

    http.close()
    restarted = HttpFixture(http.root)
    try:
        reloaded = deal(restarted)
        assert reloaded["workspace"]["revision"] == 1
        assert reloaded["findings"][0]["note"] == "Synthetic-only judgment"
        assert reloaded["ledger"]["rows"][1]["cells"]["Price low"] == "13.25"
    finally:
        restarted.close()


def test_insert_move_references_delete_restore_and_history(env):
    http, workbook, original_hash = env
    initial = deal(http)
    uid1, uid2, uid3 = [r["uid"] for r in initial["ledger"]["rows"]]
    inserted = save(http, initial, [{"type": "insert", "sheet": "Deal ledger", "after_uid": uid1, "values": {"When": "Date unknown", "Who": "Party B", "Event": "Synthetic insertion", "Count": "", "Date from": "", "Note": "New event"}}], "Insert uncertain event")
    assert inserted.status_code == 200, inserted.text
    state = inserted.json()
    assert [r["id"] for r in state["ledger"]["rows"]] == [1, 2, 3, 4]
    new_uid = state["ledger"]["rows"][1]["uid"]
    assert state["ledger"]["rows"][2]["uid"] == uid2
    assert state["ledger"]["rows"][2]["cells"]["Note"] == "Follow-up to #1"
    assert state["ledger"]["rows"][3]["cells"]["Note"] == "After #3"
    assert "#3" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert "#4" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert state["ledger"]["rows"][1]["cells"]["Count"] == ""
    assert state["ledger"]["rows"][1]["cells"]["Date from"] == ""

    moved = save(http, state, [{"type": "move", "sheet": "Deal ledger", "uid": uid3, "after_uid": uid1}], "Move signing before bid")
    assert moved.status_code == 200, moved.text
    state = moved.json()
    assert [r["uid"] for r in state["ledger"]["rows"]] == [uid1, uid3, new_uid, uid2]
    assert state["ledger"]["rows"][1]["cells"]["Note"] == "After #4"
    assert "#4" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert "#2" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert "#3" not in state["questions"]["rows"][0]["cells"]["Rows affected"]

    rejection = save(http, state, [{"type": "delete", "sheet": "Deal ledger", "uid": uid2}], "Reject dangling reference")
    assert rejection.status_code == 400, rejection.text
    assert deal(http)["workspace"]["revision"] == 2
    invalid = save(http, state, [{"type": "delete", "sheet": "Deal ledger", "uid": uid2, "replacement_uid": "not-live"}], "Invalid replacement")
    assert invalid.status_code == 400, invalid.text
    replacement = save(http, state, [{"type": "delete", "sheet": "Deal ledger", "uid": uid2, "replacement_uid": uid3}], "Replace deleted reference")
    assert replacement.status_code == 200, replacement.text
    state = replacement.json()
    assert [r["uid"] for r in state["ledger"]["rows"]] == [uid1, uid3, new_uid]
    assert state["ledger"]["rows"][1]["cells"]["Note"] == "After #2"
    assert "#2" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert workbook.exists() and hashlib.sha256(workbook.read_bytes()).hexdigest() == original_hash

    restored = save(http, state, [{"type": "restore", "target_revision": 0}], "Restore original snapshot")
    assert restored.status_code == 200, restored.text
    state = restored.json()
    assert state["workspace"]["revision"] == 4
    assert [r["uid"] for r in state["ledger"]["rows"]] == [uid1, uid2, uid3]
    assert state["ledger"]["rows"][2]["cells"]["Note"] == "After #2"
    history = http.get("/api/deal/synthetic/history").json()["history"]
    assert [item["revision"] for item in history] == [4, 3, 2, 1, 0]
    assert http.get("/api/deal/synthetic/changes").json()["changes"] == []


def test_same_batch_insert_references_client_uid_and_unknown_refs(env):
    http, _, _ = env
    initial = deal(http)
    uid1, uid2, uid3 = [r["uid"] for r in initial["ledger"]["rows"]]
    operations = [
        {"type": "insert", "sheet": "Deal ledger", "after_uid": uid1, "client_uid": "new-temporary-a", "values": {"Event": "Discarded synthetic A", "Note": "temporary"}},
        {"type": "insert", "sheet": "Deal ledger", "after_uid": "new-temporary-a", "client_uid": "new-temporary-b", "values": {"Event": "Kept synthetic B", "Note": "temporary"}},
        {"type": "move", "sheet": "Deal ledger", "uid": "new-temporary-a", "after_uid": "new-temporary-b"},
        {"type": "delete", "sheet": "Deal ledger", "uid": "new-temporary-a"},
        {"type": "update", "sheet": "Deal ledger", "uid": uid3, "values": {"Note": "see #2; unknown #999 remains"}},
    ]
    response = save(http, initial, operations, "Resolve client UIDs and original references")
    assert response.status_code == 200, response.text
    state = response.json()
    assert [r["cells"]["Event"] for r in state["ledger"]["rows"]] == ["Invitation", "Kept synthetic B", "Offer", "Signing"]
    assert state["ledger"]["rows"][2]["uid"] == uid2
    assert state["ledger"]["rows"][3]["uid"] == uid3
    assert state["ledger"]["rows"][3]["cells"]["Note"] == "see #3; unknown #999 remains"
    assert state["ledger"]["rows"][1]["uid"] not in ("new-temporary-a", "new-temporary-b")
    assert "#3" in state["questions"]["rows"][0]["cells"]["Rows affected"]
    assert "#4" in state["questions"]["rows"][0]["cells"]["Rows affected"]


def test_question_id_rename_updates_ledger_flags_and_duplicate_blocks(env):
    http, _, _ = env
    initial = deal(http)
    renamed = save(http, initial, [{"type": "update", "sheet": "Questions", "uid": row(initial, "Questions")["uid"], "values": {"Q": "Q2"}}], "Rename synthetic question")
    assert renamed.status_code == 200, renamed.text
    state = renamed.json()
    assert state["questions"]["rows"][0]["cells"]["Q"] == "Q2"
    assert [item["cells"]["Flag"] for item in state["ledger"]["rows"]] == ["", "Q2 R1", "Q2"]
    duplicate = save(http, state, [{"type": "insert", "sheet": "Questions", "after_uid": row(state, "Questions")["uid"], "values": {"Q": "Q2", "Question": "Duplicate synthetic ID"}}], "Reject duplicate question")
    assert duplicate.status_code == 400, duplicate.text
    assert deal(http)["workspace"]["revision"] == 1


def test_review_item_r_id_save_rename_delete_and_flag_refs(env):
    http, _, _ = env
    initial = deal(http)
    assert [item["cells"]["Q"] for item in initial["questions"]["rows"]] == ["Q1", "R1"]
    review = row(initial, "Questions", 1)
    renamed = save(http, initial, [{"type": "update", "sheet": "Questions", "uid": review["uid"], "values": {"Q": "R2"}}], "Rename review item")
    assert renamed.status_code == 200, renamed.text
    state = renamed.json()
    assert row(state, "Questions", 1)["cells"]["Q"] == "R2"
    assert row(state, "Deal ledger", 1)["cells"]["Flag"] == "Q1 R2"
    duplicate = save(http, state, [{"type": "insert", "sheet": "Questions", "after_uid": review["uid"], "values": {"Q": "R2", "Question": "Duplicate review item"}}])
    assert duplicate.status_code == 400
    blocked = save(http, state, [{"type": "delete", "sheet": "Questions", "uid": review["uid"]}], "Referenced review item")
    assert blocked.status_code == 400
    cleared = save(http, state, [{"type": "update", "sheet": "Deal ledger", "uid": row(state, "Deal ledger", 1)["uid"], "values": {"Flag": "Q1"}}], "Clear review item flag")
    assert cleared.status_code == 200, cleared.text
    deleted = save(http, cleared.json(), [{"type": "delete", "sheet": "Questions", "uid": review["uid"]}], "Delete review item")
    assert deleted.status_code == 200, deleted.text
    assert row(deleted.json(), "Deal ledger", 1)["cells"]["Flag"] == "Q1"
    assert [item["cells"]["Q"] for item in deleted.json()["questions"]["rows"]] == ["Q1"]


def test_unassessed_finding_implementation_is_attributed(env):
    http, _, _ = env
    initial = deal(http)
    decision = {"type": "finding", "id": "F1", "judgment": "deferred", "implementation": "unassessed", "verification": "unchecked", "note": "Synthetic source review pending"}
    response = save(http, initial, [decision], "Record provisional implementation state")
    assert response.status_code == 200, response.text
    finding = response.json()["findings"][0]
    assert finding["implementation"] == "unassessed"
    assert finding["judgment"] == "deferred"
    assert finding["actor"] == "local" and finding["at"]
    assert deal(http)["findings"][0]["implementation"] == "unassessed"


def test_post_round_can_be_edited_cloned_and_exported_as_text(env):
    http, _, _ = env
    initial = deal(http)
    event = row(initial, "Deal ledger", 2)
    post = save(http, initial, [{"type": "update", "sheet": "Deal ledger", "uid": event["uid"], "values": {"Round": "post", "Note": "Synthetic post-signing event"}}], "Mark synthetic post round")
    assert post.status_code == 200, post.text
    state = post.json()
    assert state["ledger"]["rows"][2]["cells"]["Round"] == "post"
    cloned = save(http, state, [{"type": "insert", "sheet": "Deal ledger", "after_uid": event["uid"], "client_uid": "new-post-clone", "values": {"When": "Date unknown", "Event": "Synthetic post clone", "Round": "post", "Count": "", "Date from": ""}}], "Clone post round event")
    assert cloned.status_code == 200, cloned.text
    state = cloned.json()
    assert state["ledger"]["rows"][3]["cells"]["Round"] == "post"
    exported = xlsx(http.get("/api/deal/synthetic/export?version=working"))
    column = check_lean.LEDGER_COLUMNS.index("Round") + 1
    assert exported["Deal ledger"].cell(4, column).value == "post"
    assert exported["Deal ledger"].cell(4, column).data_type == "s"
    assert exported["Deal ledger"].cell(5, column).value == "post"
    assert exported["Deal ledger"].cell(5, column).data_type == "s"
    exported.close()


def test_atomic_stale_formula_and_question_id_guards(env):
    http, _, _ = env
    initial = deal(http)
    uid = row(initial, "Deal ledger", 0)["uid"]
    base = {"revision": 0, "base_sha256": initial["workspace"]["base_sha256"], "reason": "Rejected batch"}
    invalid_batches = [
        [{"type": "update", "sheet": "Deal facts", "uid": row(initial, "Deal facts", 0)["uid"], "values": {"Value": "Would change"}}, {"type": "update", "sheet": "Questions", "uid": row(initial, "Questions")["uid"], "values": {"Q": "Q0"}}],
        [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Nonexistent column": "x"}}],
    ]
    for ops in invalid_batches:
        response = http.edit({**base, "operations": ops})
        assert response.status_code == 400, response.text
        assert deal(http)["workspace"]["revision"] == 0
        assert row(deal(http), "Deal ledger")["cells"]["Note"] == "Three parties contacted"
    good = save(http, initial, [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "Saved once"}}])
    assert good.status_code == 200, good.text
    stale = save(http, initial, [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "Stale overwrite"}}])
    assert stale.status_code == 409, stale.text
    wrong_hash = http.edit({"revision": 1, "base_sha256": "0" * 64, "reason": "Wrong base", "operations": [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "Wrong base"}}]})
    assert wrong_hash.status_code == 409, wrong_hash.text
    assert row(deal(http), "Deal ledger")["cells"]["Note"] == "Saved once"

    literal = '=HYPERLINK("https://example.invalid", "not a link")'
    current = deal(http)
    edited = save(http, current, [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": literal}}], "Literal formula-like source text")
    assert edited.status_code == 200, edited.text
    exported = xlsx(http.get("/api/deal/synthetic/export?version=working"))
    note = exported["Deal ledger"].cell(2, check_lean.LEDGER_COLUMNS.index("Note") + 1)
    original_formula = exported["Deal ledger"].cell(2, check_lean.LEDGER_COLUMNS.index("Reviewer note") + 1)
    assert note.value == literal and note.data_type == "s"
    assert original_formula.value == "=1+1" and original_formula.data_type == "f"
    exported.close()


def test_huge_range_is_bounded_without_state_change(env):
    http, _, _ = env
    initial = deal(http)
    payload = {"revision": 0, "base_sha256": initial["workspace"]["base_sha256"], "reason": "Reject huge range", "operations": [{"type": "update", "sheet": "Questions", "uid": row(initial, "Questions")["uid"], "values": {"Rows affected": "#1-#1000000000"}}]}
    started = time.monotonic()
    response = http.edit(payload)
    elapsed = time.monotonic() - started
    assert response.status_code == 400, response.text
    assert elapsed < 5, f"unbounded range validation took {elapsed:.1f}s"
    assert deal(http)["workspace"]["revision"] == 0


def test_session_security_origin_csrf_and_path_boundaries(env, monkeypatch):
    http, _, _ = env
    initial = deal(http)
    uid = row(initial, "Deal ledger")["uid"]
    payload = {"revision": 0, "base_sha256": initial["workspace"]["base_sha256"], "reason": "Security probe", "operations": [{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "Must not save"}}]}
    assert http.session()["can_edit"] is True  # explicitly local test server
    denied = [
        http.edit(payload, csrf="wrong"),
        http.edit(payload, origin="https://evil.example"),
        http.edit(payload, origin="http://127.0.0.1:9999"),
        http.edit(payload, origin=http.base.replace("http://", "https://")),
        http.edit(payload, origin="https://evil.example", headers={"X-Forwarded-Host": "evil.example"}),
        http.edit(payload, origin="omit"),
        http.edit(payload, content_type="text/plain"),
        http.edit(payload, headers={"Host": "lines.example.invalid", "Origin": "https://lines.example.invalid"}),
        http.post("/api/deal/../synthetic/edit", payload),
        http.post("/api/deal/%2e%2e%2fsynthetic/edit", payload),
    ]
    assert all(response.status_code in (400, 403, 404, 405, 415) for response in denied), [(i, r.status_code, r.text[:120]) for i, r in enumerate(denied)]
    assert http.get("/api/document/synthetic/%2e%2e%2f%2e%2e%2fref%2fanswers").status_code == 404
    assert http.get("/assets/%2e%2e/%2e%2e/ref/answers").status_code == 404
    assert http.get("/api/deal/%2e%2e%2fsynthetic").status_code == 404
    oversized = http.post("/api/deal/synthetic/edit", {"padding": "x" * (server.MAX_JSON + 1)})
    assert oversized.status_code == 413, oversized.text[:200]
    assert deal(http)["workspace"]["revision"] == 0
    monkeypatch.setenv("COCKPIT_PUBLIC_ORIGIN", "https://lines.example.invalid")
    public_headers = {"Host": "lines.example.invalid", "Origin": "https://lines.example.invalid"}
    unknown = http.edit(payload, headers={**public_headers, "Cf-Access-Authenticated-User-Email": "intruder@example.invalid"})
    missing = http.edit(payload, headers=public_headers)
    assert unknown.status_code == 403 and missing.status_code == 403
    assert deal(http)["workspace"]["revision"] == 0
    # Synthetic request headers alone must not be reported as a real Cloudflare login test.


def test_catalog_allowlist_rejects_ref_and_symlink_escape(env):
    http, _, _ = env
    secret = "SYNTHETIC-SECRET-MUST-NOT-LEAK"
    ref = http.root / "ref"
    ref.mkdir()
    (ref / "secret.txt").write_text(secret, encoding="utf-8")
    catalog_path = http.root / "_dev/cockpit/catalog.json"
    catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
    doc = catalog["deals"]["synthetic"]["documents"][0]
    doc["path"] = "ref/secret.txt"
    catalog_path.write_text(json.dumps(catalog), encoding="utf-8")
    response = http.get("/api/document/synthetic/finding")
    assert response.status_code in (400, 404), response.text
    assert secret not in response.text

    outside = http.root.parent / (http.root.name + "-outside.txt")
    outside.write_text(secret, encoding="utf-8")
    try:
        link = http.root / "_dev/reviews/synthetic/escape.txt"
        link.symlink_to(outside)
        doc["path"] = "_dev/reviews/synthetic/escape.txt"
        catalog_path.write_text(json.dumps(catalog), encoding="utf-8")
        response = http.get("/api/document/synthetic/finding")
        assert response.status_code in (400, 404), response.text
        assert secret not in response.text
    finally:
        outside.unlink(missing_ok=True)


def test_instruction_versions_and_engines_over_http(env, monkeypatch):
    http, _, _ = env
    listing = http.get("/api/instructions").json()
    [base] = listing["items"]
    assert (base["name"], base["is_default"]) == ("Version 1", True)
    assert http.get("/instructions").status_code == 200
    assert http.get("/api/instructions/zzz").status_code == 404
    assert http.get(f"/api/instructions/{base['id']}?seq=x").status_code == 400
    draft = http.post("/api/instructions", {"action": "draft", "from": base["id"]})
    assert draft.status_code == 200, draft.text
    item = draft.json()["item"]
    saved = http.post("/api/instructions", {"action": "save", "id": item["id"], "text": "New text\n", "base_sha256": item["sha256"]})
    stale = http.post("/api/instructions", {"action": "save", "id": item["id"], "text": "Other\n", "base_sha256": item["sha256"]})
    assert (saved.status_code, stale.status_code) == (200, 409)
    assert http.get(f"/api/instructions/{item['id']}?seq=1").json()["text"] == draft.json()["text"]
    assert http.post("/api/instructions", {"action": "save", "id": item["id"], "text": "x"}, csrf="wrong").status_code == 403
    monkeypatch.setenv("COCKPIT_PUBLIC_ORIGIN", "https://lines.example.invalid")
    public = {"Host": "lines.example.invalid", "Origin": "https://lines.example.invalid", "Cf-Access-Authenticated-User-Email": "intruder@example.invalid"}
    assert http.post("/api/instructions", {"action": "default", "id": base["id"]}, headers=public).status_code == 403
    assert http.post("/api/account/chatgpt", {"action": "connect"}, headers=public).status_code == 403
    account = http.get("/api/account").json()
    assert [engine["id"] for engine in account["engines"]] == ["opus55", "fable51", "sol6", "astra6"]
    assert account["chatgpt"]["connected"] is False
    refused = http.post("/api/deal/synthetic/jobs", {"action": "extract", "engine": "sol6"})
    assert refused.status_code == 409 and "ChatGPT" in refused.json()["error"]


def test_hide_and_unhide_a_deal_over_http(env, monkeypatch):
    http, _, _ = env
    route = "/api/deals/synthetic/visibility"
    listed = lambda: next(d for d in http.get("/api/deals").json() if d["slug"] == "synthetic")
    assert (listed()["hidden"], listed()["hidden_by"], deal(http)["hidden"]) == (False, None, False)
    assert http.get(route).status_code == 405
    assert http.post(route, {"action": "hide"}, csrf="wrong").status_code == 403
    assert http.post(route, {"action": "hide"}, origin="https://evil.example").status_code == 403
    assert http.post(route, {"action": "hide"}, content_type="text/plain").status_code == 415
    assert http.post("/api/deals/nosuch/visibility", {"action": "hide"}).status_code == 404
    assert http.post("/api/deals/%2e%2e%2fsynthetic/visibility", {"action": "hide"}).status_code == 404
    assert http.post(route, {"action": "delete"}).status_code == 400
    monkeypatch.setenv("COCKPIT_PUBLIC_ORIGIN", "https://lines.example.invalid")
    public = {"Host": "lines.example.invalid", "Origin": "https://lines.example.invalid", "Cf-Access-Authenticated-User-Email": "intruder@example.invalid"}
    assert http.post(route, {"action": "hide"}, headers=public).status_code == 403
    assert listed()["hidden"] is False

    hidden = http.post(route, {"action": "hide"})
    assert hidden.status_code == 200, hidden.text
    assert (hidden.json()["hidden"], hidden.json()["hidden_by"]) == (True, "local")
    assert http.post(route, {"action": "hide"}).status_code == 409
    assert listed()["hidden"] is True and listed()["hidden_at"]
    opened = deal(http)  # a hidden deal still opens by URL, its working copy untouched
    assert (opened["hidden"], opened["hidden_by"], opened["workspace"]["revision"]) == (True, "local", 0)
    refused = http.post("/api/deal/synthetic/jobs", {"action": "extract"})
    assert refused.status_code == 409 and "unhide the deal first" in refused.json()["error"]
    shown = http.post(route, {"action": "unhide"})
    assert shown.status_code == 200 and shown.json()["hidden"] is False
    assert http.post(route, {"action": "unhide"}).status_code == 409
    kinds = [item["kind"] for item in http.get("/api/activity?slug=synthetic").json()["items"]]
    assert kinds[:2] == ["unhide_deal", "hide_deal"]


def test_deal_review_status_is_recorded_at_a_revision(env):
    http, _, _ = env
    route = "/api/deal/synthetic/review"
    listed = lambda: next(d for d in http.get("/api/deals").json() if d["slug"] == "synthetic")["deal_review"]
    initial = deal(http)["deal_review"]
    assert (initial["status"], initial["actor"], initial["edited_since"]) == ("unreviewed", None, False)
    assert http.post(route, {"status": "reviewed", "revision": 0}, csrf="wrong").status_code == 403
    assert http.post(route, {"status": "done", "revision": 0}).status_code == 400
    assert http.post(route, {"status": "reviewed"}).status_code == 400
    assert http.post(route, {"status": "reviewed", "revision": 1}).status_code == 409

    snapshot = deal(http)
    saved = save(http, snapshot, [{"type": "review", "uid": row(snapshot, "Deal ledger")["uid"], "status": "reviewed"}])
    assert saved.status_code == 200, saved.text
    assert saved.json()["deal_review"]["status"] == "in_review"  # edited, nothing recorded yet
    assert http.post(route, {"status": "reviewed", "revision": 0}).status_code == 409  # stale view

    marked = http.post(route, {"status": "reviewed", "revision": 1})
    assert marked.status_code == 200, marked.text
    review = marked.json()["deal_review"]
    assert (review["status"], review["revision"], review["actor"], review["edited_since"]) == ("reviewed", 1, "local", False)
    assert http.post(route, {"status": "reviewed", "revision": 1}).status_code == 400  # no change
    assert listed()["status"] == "reviewed"

    snapshot = deal(http)
    later = save(http, snapshot, [{"type": "review", "uid": row(snapshot, "Deal ledger")["uid"], "status": "needs_decision"}])
    assert later.status_code == 200, later.text
    stale = later.json()["deal_review"]
    assert (stale["status"], stale["revision"], stale["current_revision"], stale["edited_since"]) == ("reviewed", 1, 2, True)
    assert listed()["edited_since"] is True
    again = http.post(route, {"status": "reviewed", "revision": 2}).json()["deal_review"]
    assert (again["revision"], again["edited_since"]) == (2, False)
    items = http.get("/api/activity?slug=synthetic").json()["items"]
    assert [(i["kind"], i["summary"]) for i in items[:2]] == [("deal_review", "Marked reviewed at revision 2"), ("revision", "Synthetic acceptance change")]


def test_bulk_process_round_edit_is_one_attributed_revision(env):
    http, workbook, original_hash = env
    snapshot = deal(http)
    marked = save(http, snapshot, [{"type": "review", "uid": row(snapshot, "Deal ledger", 1)["uid"], "status": "needs_decision"}])
    assert marked.status_code == 200, marked.text
    snapshot = marked.json()
    uids = [row(snapshot, "Deal ledger", index)["uid"] for index in (1, 2)]
    for values in ({"Note": "x"}, {"Process": 0}, {"Round": ""}):
        rejected = save(http, snapshot, [{"type": "bulk_update", "sheet": "Deal ledger", "uids": uids, "values": values}])
        assert rejected.status_code == 400, rejected.text
    assert deal(http)["workspace"]["revision"] == 1
    saved = save(http, snapshot, [{"type": "bulk_update", "sheet": "Deal ledger", "uids": uids, "values": {"Process": 1, "Round": "2"}}], "Second round starts at #2")
    assert saved.status_code == 200, saved.text
    payload = saved.json()
    assert payload["workspace"]["revision"] == 2
    assert [r["cells"]["Round"] for r in payload["ledger"]["rows"]] == ["1", "2", "2"]
    assert payload["row_review"][uids[0]]["status"] == "needs_decision"
    user = http.session()["user"]
    assert {payload["field_authors"][uid]["Round"]["actor"] for uid in uids} == {user}
    assert "Process" not in payload["field_authors"][uids[0]]  # unchanged values are not recorded
    history = http.get("/api/deal/synthetic/history").json()["history"][0]
    assert (history["actor"], history["reason"], history["summary"]) == (user, "Second round starts at #2", "2 changes")
    assert sorted((c["uid"], c["field"], c["before"], c["after"], c["type"]) for c in history["changes"]) == sorted((uid, "Round", "1", "2", "update") for uid in uids)
    exported = xlsx(http.get("/api/deal/synthetic/export?version=working&source=0"))
    column = check_lean.LEDGER_COLUMNS.index("Round") + 1
    assert [(exported["Deal ledger"].cell(r, column).value, exported["Deal ledger"].cell(r, column).data_type) for r in (2, 3, 4)] == [(1, "n"), (2, "n"), (2, "n")]
    exported.close()
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == original_hash

def source_values(workbook):
    return {field: value for field, value in workbook["Source"].iter_rows(min_row=2, values_only=True)}


def sheet_values(workbook):
    """Every sheet's cell values; a rendered workbook's bytes also carry the time it was saved."""
    return {ws.title: list(ws.iter_rows(values_only=True)) for ws in workbook.worksheets}


def _saved(root, content):
    path = root / "download-check.xlsx"
    path.write_bytes(content)
    return path


def test_download_adds_a_source_sheet_to_the_working_copy_only(env):
    http, workbook, original_hash = env
    submission = "https://www.sec.gov/Archives/edgar/data/77/0000000077-26-000001.txt"
    index = "https://www.sec.gov/Archives/edgar/data/77/0000000077-26-000001-index.htm"
    manifest = http.root / "raw_filing/MANIFEST.csv"
    with manifest.open(newline="", encoding="utf-8") as handle:
        entries = list(csv.DictReader(handle))
    filing_hash = entries[0]["sha256"]
    manifest.write_text(manifest.read_text(encoding="utf-8").replace("https://example.invalid/synthetic", submission), encoding="utf-8")
    (http.root / "ref").mkdir()
    (http.root / "ref/seed.csv").write_text(f"deal,target_name,deal_number,form_type,date_filed,index_url,status\nsynthetic,SYNTHETIC ACME,1,DEFM14A,2026-01-10,{index},ok\n", encoding="utf-8")
    route = "/api/deal/synthetic/export"

    response = http.get(route)  # an unedited working copy: revision 0, no review status
    assert response.headers["Content-Disposition"] == 'attachment; filename="synthetic-working-r0.xlsx"'
    exported = xlsx(response)
    assert exported.sheetnames == ["Deal ledger", "Rounds", "Questions", "Deal facts", "Source"]
    values = source_values(exported)
    assert values["EDGAR filing index"] == index and exported["Source"]["B2"].hyperlink.target == index
    assert values["Complete submission (.txt)"] == submission and exported["Source"]["B3"].hyperlink.target == submission
    assert values["Background pages"] == "not recorded"  # the synthetic Deal facts have no such row
    assert values["Filing SHA-256"] == filing_hash
    assert (values["Instruction"], values["Instruction SHA-256"]) == ("Version 1", "not recorded")  # no published instruction in this state
    assert (values["Version ID (working-copy base)"], values["Raw workbook SHA-256"]) == ("version-1-raw", original_hash)
    assert (values["Working revision"], values["Review status"]) == ("0", "not set")
    assert dt.datetime.fromisoformat(values["Exported at"]).tzinfo is not None
    exported.close()

    initial = deal(http)
    pages = {"type": "insert", "sheet": "Deal facts", "after_uid": row(initial, "Deal facts", 1)["uid"], "values": {"Field": "Background pages", "Value": "pp. 1-2"}}
    assert save(http, initial, [pages], "Record background pages").status_code == 200
    assert http.post("/api/deal/synthetic/review", {"status": "reviewed", "revision": 1}).status_code == 200
    current = deal(http)
    assert save(http, current, [{"type": "update", "sheet": "Deal facts", "uid": row(current, "Deal facts", 0)["uid"], "values": {"Value": "Synthetic Acme Later"}}], "Edit after review").status_code == 200
    response = http.get(route + "?version=working")
    assert response.headers["Content-Disposition"] == 'attachment; filename="synthetic-working-r2.xlsx"'
    exported = xlsx(response)
    values = source_values(exported)
    assert (values["Background pages"], values["Working revision"]) == ("pp. 1-2", "2")
    assert values["Review status"] == "Reviewed at revision 1; edited since (working revision 2)"
    assert exported["Deal facts"]["B2"].value == "Synthetic Acme Later"
    exported.close()

    four = http.get(route + "?source=0")
    assert four.headers["Content-Disposition"] == 'attachment; filename="synthetic-working-r2.xlsx"'
    assert four.content == http.cockpit.workspace.export("synthetic")  # the export_repo.py bytes: four sheets, no Source
    exported = xlsx(four)
    assert exported.sheetnames == ["Deal ledger", "Rounds", "Questions", "Deal facts"]
    exported.close()
    report = check_lean.LeanChecker(_saved(http.root, four.content), http.root / "raw_filing/synthetic.htm").run()
    assert not any(issue["code"] == "schema.sheets" for issue in report["issues"])
    five = check_lean.LeanChecker(_saved(http.root, response.content), http.root / "raw_filing/synthetic.htm").run()
    assert any(issue["code"] == "schema.sheets" for issue in five["issues"])  # the Source sheet is not checker-valid, by design

    raw = http.get(route + "?version=version-1-raw")
    assert raw.headers["Content-Disposition"] == 'attachment; filename="synthetic-version-1-raw.xlsx"'
    assert raw.content == workbook.read_bytes()
    with_source = http.get(route + "?version=version-1-raw&source=1")
    assert with_source.headers["Content-Disposition"] == 'attachment; filename="synthetic-version-1-raw-with-source.xlsx"'
    exported = xlsx(with_source)
    assert exported.sheetnames == ["Deal ledger", "Rounds", "Questions", "Deal facts", "Source"]
    assert exported["Deal facts"]["B2"].value == "Synthetic Acme"
    values = source_values(exported)
    assert (values["Version ID"], values["Raw workbook SHA-256"], values["Background pages"]) == ("version-1-raw", original_hash, "not recorded")
    assert values["Working revision"] == values["Review status"] == "not applicable: a raw version, not the working copy"
    exported.close()
    assert http.get(route + "?version=version-1-raw&source=0").content == workbook.read_bytes()
    assert http.get(route + "?source=2").status_code == 400
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == original_hash


def test_past_revision_downloads_like_the_working_copy(env):
    http, workbook, original_hash = env
    route = "/api/deal/synthetic/export"
    initial = deal(http)
    pages = {"type": "insert", "sheet": "Deal facts", "after_uid": row(initial, "Deal facts", 1)["uid"], "values": {"Field": "Background pages", "Value": "pp. 1-2"}}
    assert save(http, initial, [pages], "Record background pages").status_code == 200
    # Revision 1 is marked Reviewed, then In review; only the latest marking is kept once revision 2 is marked.
    for status in ("reviewed", "in_review"):
        assert http.post("/api/deal/synthetic/review", {"status": status, "revision": 1}).status_code == 200
    current = deal(http)
    facts_uid = row(current, "Deal facts", 0)["uid"]
    assert save(http, current, [{"type": "update", "sheet": "Deal facts", "uid": facts_uid, "values": {"Value": "Synthetic Acme Two"}}], "Second edit").status_code == 200
    assert http.post("/api/deal/synthetic/review", {"status": "reviewed", "revision": 2}).status_code == 200
    current = deal(http)
    assert save(http, current, [{"type": "update", "sheet": "Deal facts", "uid": facts_uid, "values": {"Value": "Synthetic Acme Three"}}], "Edit after review").status_code == 200

    # Each past revision: its own name, a Source sheet by default, and the review status related to it.
    expected = {
        0: ("0 (a past revision; the working copy is at revision 3)", "Reviewed at revision 2 (set after this revision; the status at revision 0 is not recorded here)", "Synthetic Acme"),
        1: ("1 (a past revision; the working copy is at revision 3)", "Reviewed at revision 2 (set after this revision; the status at revision 1 is not recorded here)", "Synthetic Acme"),
        2: ("2 (a past revision; the working copy is at revision 3)", "Reviewed at revision 2", "Synthetic Acme Two"),
        3: ("3", "Reviewed at revision 2; edited since (this file is revision 3)", "Synthetic Acme Three"),
    }
    for number, (revision_text, status, name) in expected.items():
        response = http.get(f"{route}?version=rev:{number}")
        assert response.headers["Content-Disposition"] == f'attachment; filename="synthetic-working-r{number}.xlsx"'
        exported = xlsx(response)
        assert exported.sheetnames == ["Deal ledger", "Rounds", "Questions", "Deal facts", "Source"]
        assert exported["Deal facts"]["B2"].value == name
        values = source_values(exported)
        assert (values["Working revision"], values["Review status"]) == (revision_text, status)
        assert (values[f"Version ID (base of revision {number})"], values["Raw workbook SHA-256"]) == ("version-1-raw", original_hash)
        assert values["Background pages"] == ("not recorded" if number == 0 else "pp. 1-2")
        exported.close()
        four = http.get(f"{route}?version=rev:{number}&source=0")
        assert four.headers["Content-Disposition"] == f'attachment; filename="synthetic-working-r{number}.xlsx"'
        assert sheet_values(xlsx(four)) == sheet_values(openpyxl.load_workbook(io.BytesIO(http.cockpit.workspace.export("synthetic", f"rev:{number}"))))
        assert xlsx(four).sheetnames == ["Deal ledger", "Rounds", "Questions", "Deal facts"]  # four sheets, no Source
    assert http.get(f"{route}?version=rev:0&source=0").content == workbook.read_bytes()
    # The latest revision's file is the working copy's, apart from how the Source sheet describes it.
    working = source_values(xlsx(http.get(route)))
    assert (working["Working revision"], working["Review status"]) == ("3", "Reviewed at revision 2; edited since (working revision 3)")
    assert sheet_values(xlsx(http.get(f"{route}?version=rev:3&source=0"))) == sheet_values(xlsx(http.get(route + "?source=0")))
    assert http.get(f"{route}?version=rev:4").status_code == 404
    assert [item["revision"] for item in http.get("/api/deal/synthetic/history").json()["history"]] == [3, 2, 1, 0]  # nothing saved
    # The earlier markings of revision 1 remain in the deal's activity, which the Source sheet points to.
    marks = [i["summary"] for i in http.get("/api/activity?slug=synthetic").json()["items"] if i["kind"] == "deal_review"]
    assert marks == ["Marked reviewed at revision 2", "Marked in review at revision 1", "Marked reviewed at revision 1"]


def test_download_never_shows_an_unrecorded_or_disagreeing_link(env):
    http, _, _ = env
    values = source_values(xlsx(http.get("/api/deal/synthetic/export")))  # the fixture's source is not an EDGAR link
    assert (values["EDGAR filing index"], values["Complete submission (.txt)"]) == ("not recorded", "not recorded")
    manifest = http.root / "raw_filing/MANIFEST.csv"
    manifest.write_text(manifest.read_text(encoding="utf-8").replace("https://example.invalid/synthetic", "https://www.sec.gov/Archives/edgar/data/77/0000000077-26-000001.txt"), encoding="utf-8")
    (http.root / "ref").mkdir()
    (http.root / "ref/seed.csv").write_text("deal,target_name,deal_number,form_type,date_filed,index_url,status\nsynthetic,SYNTHETIC ACME,1,DEFM14A,2026-01-10,https://www.sec.gov/Archives/edgar/data/77/0000000077-26-000002-index.htm,ok\n", encoding="utf-8")
    values = source_values(xlsx(http.get("/api/deal/synthetic/export")))
    assert values["EDGAR filing index"] == "not recorded"  # the seed disagrees with the filing's own link
    assert values["Complete submission (.txt)"] == "https://www.sec.gov/Archives/edgar/data/77/0000000077-26-000001.txt"


def add_current_version(http, ident="opus55-second-run"):
    """Import a second current-format synthetic workbook as a run version."""
    source = openpyxl.load_workbook(http.root / "extraction/synthetic.xlsx")
    ledger = source["Deal ledger"]
    ledger.cell(3, check_lean.LEDGER_COLUMNS.index("Who") + 1, "Party A (new run)")
    ledger.cell(3, check_lean.LEDGER_COLUMNS.index("Stock %") + 1, "50–75")
    ledger.cell(3, check_lean.LEDGER_COLUMNS.index("Financing") + 1, "Contingent")
    folder = http.root / "_dev/cockpit/state/versions/synthetic" / ident
    folder.mkdir(parents=True)
    path = folder / "synthetic.xlsx"
    source.save(path)
    report = check_lean.LeanChecker(path, http.root / "raw_filing/synthetic.htm").run()
    (folder / "check.json").write_text(json.dumps(report), encoding="utf-8")
    conn = http.cockpit.workspace._connect(write=True)
    conn.execute("INSERT INTO versions (slug, id, label, path, sha256, kind, engine, model, effort, instruction_version, instruction_sha256, filing_sha256, started_by, started_at, finished_at, receipts, checker) VALUES (?,?,?,?,?,'raw','Opus 5.5','claude-opus-5-5','medium',NULL,?,'f','alex','2026-09-24T22:41:00+00:00','2026-09-24T22:52:00+00:00',?,?)",
                 ("synthetic", ident, "Opus 5.5 · medium · draft — Alex", str(path.relative_to(http.root)), hashlib.sha256(path.read_bytes()).hexdigest(), "d" * 64,
                  str(folder.relative_to(http.root)), json.dumps({"errors": report["summary"]["errors"], "warnings": report["summary"]["warnings"]})))
    conn.commit(); conn.close()
    return ident, hashlib.sha256(path.read_bytes()).hexdigest(), report


def test_current_choices_compare_and_the_checker_shown(env):
    http, _, _ = env
    ident, digest, report = add_current_version(http)
    working = deal(http)
    assert (working["ledger_schema"], working["workspace"]["base_ledger_schema"], working["workspace"]["base_instruction_version"]) == ("Version 1", "Version 1", "Version 1")
    assert "mixed" in working["choices"]["Initiation"] and "Financing" in working["choices"]
    shown = deal(http, ident)
    assert shown["ledger_schema"] == "Version 1" and shown["choices"] == working["choices"]
    assert "No deadline stated" in shown["choices"]["Deadline outcome"] and "No deadline stated" in working["choices"]["Deadline outcome"]
    assert shown["check"]["checker_version"] == check_lean.CHECKER_VERSION
    at_import = next(v for v in shown["versions"] if v["id"] == ident)["checker"]
    assert at_import == {"checker_version": check_lean.CHECKER_VERSION, "ledger_schema": "Version 1", "errors": report["summary"]["errors"], "warnings": report["summary"]["warnings"]}
    assert next(v for v in shown["versions"] if v["id"] == "version-1-raw")["checker"] is None  # no receipt: "At import: not recorded"
    listed = next(d for d in http.get("/api/deals").json() if d["slug"] == "synthetic")["check"]
    assert (listed["checker_version"], listed["ledger_schema"]) == (check_lean.CHECKER_VERSION, "Version 1")

    # The server keeps not enforcing the lists: an unlisted value and a two-part deadline outcome save.
    saved = save(http, working, [{"type": "update", "sheet": "Rounds", "uid": row(working, "Rounds")["uid"], "values": {"Deadline outcome": "Extended; Enforced", "Finality": "Something new"}},
                                 {"type": "update", "sheet": "Deal ledger", "uid": row(working, "Deal ledger", 1)["uid"], "values": {"Stock %": "Not stated"}}])
    assert saved.status_code == 200, saved.text
    assert (row(saved.json(), "Rounds")["cells"]["Deadline outcome"], row(saved.json(), "Rounds")["cells"]["Finality"]) == ("Extended; Enforced", "Something new")

    # Compare walks the current columns in either direction.
    for left, right in (("working", ident), (ident, "working")):
        compared = http.get(f"/api/deal/synthetic/compare?from={left}&to={right}").json()["changes"]
        fields = {change["field"] for change in compared if change["type"] == "update"}
        assert {"Who", "Stock %", "Financing"} <= fields, fields
    assert hashlib.sha256(http.get(f"/api/deal/synthetic/export?version={ident}").content).hexdigest() == digest


def test_mig_c_rebase_preview_carried_judgments_past_revisions_and_orphaned_threads(env):
    http, workbook, original_hash = env
    ident, _, _ = add_current_version(http)
    initial = deal(http)
    first = row(initial, "Deal ledger")["uid"]
    edited = save(http, initial, [{"type": "update", "sheet": "Deal ledger", "uid": first, "values": {"Note": "Reviewed wording"}},
                                  {"type": "review", "uid": first, "status": "reviewed"},
                                  {"type": "review", "uid": row(initial, "Deal ledger", 1)["uid"], "status": "needs_decision"},
                                  {"type": "finding", "id": "F1", "judgment": "supported", "implementation": "applied", "verification": "verified", "note": ""}], "Review pass")
    assert edited.status_code == 200, edited.text
    assert http.post("/api/deal/synthetic/comments", {"action": "create", "target": {"kind": "row", "sheet": "Deal ledger", "uid": first}, "body": "Check the page"}).status_code == 200

    # rev:N is a read-only compare source and a download; neither saves a revision.
    compared = http.get("/api/deal/synthetic/compare?from=rev:0&to=rev:1").json()
    assert compared["from_label"].startswith("Revision 0") and compared["to_label"] == "Revision 1"
    assert [(c["field"], c["after"]) for c in compared["changes"]] == [("Note", "Reviewed wording")]
    past = xlsx(http.get("/api/deal/synthetic/export?version=rev:1"))
    assert past["Deal ledger"].cell(2, check_lean.LEDGER_COLUMNS.index("Note") + 1).value == "Reviewed wording"
    past.close()
    assert hashlib.sha256(http.get("/api/deal/synthetic/export?version=rev:0&source=0").content).hexdigest() == original_hash
    assert http.get("/api/deal/synthetic/export?version=rev:7").status_code == 404
    assert deal(http)["workspace"]["revision"] == 1
    assert [item["revision"] for item in http.get("/api/deal/synthetic/history").json()["history"]] == [1, 0]

    # The rebase dialog's payload: what stops applying, with counts. A GET saves nothing; POST is not a route.
    preview = http.get(f"/api/deal/synthetic/rebase?to={ident}")
    assert preview.status_code == 200, preview.text
    stops = preview.json()["stops_applying"]
    assert stops["revisions"] == 1
    assert stops["row_marks"] == {"total": 2, "reviewed": 1, "needs_decision": 1, "unreviewed": 0}
    assert stops["finding_decisions"] == {"total": 1, "judgments_kept": 1, "reset": 1}
    assert stops["row_threads"] == {"total": 1, "open": 1, "resolved": 0}
    assert (preview.json()["current"]["ledger_schema"], preview.json()["target"]["ledger_schema"]) == ("Version 1", "Version 1")
    assert http.get("/api/deal/synthetic/rebase?to=version-1-raw").status_code == 400
    assert http.get("/api/deal/synthetic/rebase?to=nosuch").status_code == 404
    assert http.post("/api/deal/synthetic/rebase", {"to": ident}).status_code == 405
    assert deal(http)["workspace"]["revision"] == 1

    current = deal(http, ident)  # the dialog saves from the original's view, which reports the working copy's revision
    assert (current["workspace"]["revision"], current["workspace"]["editable"]) == (1, False)
    rebased = save(http, current, [{"type": "rebase", "target_version": ident}], "Rebase onto the second run")
    assert rebased.status_code == 200, rebased.text
    finding = rebased.json()["findings"][0]
    assert (finding["judgment"], finding["implementation"], finding["verification"]) == ("supported", "unassessed", "unchecked")
    assert finding["carried_over"]["revision"] == 2 and finding["carried_over"]["from_base"] == "version-1-raw"
    assert rebased.json()["row_review"] == {} and rebased.json()["ledger_schema"] == "Version 1"
    [thread] = http.get("/api/deal/synthetic/comments").json()["threads"]
    assert (thread["target_missing"], thread["target_context"]) == (True, "on an earlier base (revision 1)")
    assert hashlib.sha256(workbook.read_bytes()).hexdigest() == original_hash
