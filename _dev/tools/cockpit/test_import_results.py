"""Synthetic evidence tests for the cockpit catalog importer."""

from __future__ import annotations

import json
from pathlib import Path
import tempfile
import unittest

import openpyxl

from cockpit import import_results as imp


def write_json(path: Path, value: dict | list) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value), encoding="utf-8")


def make_book(path: Path) -> str:
    path.parent.mkdir(parents=True, exist_ok=True)
    book = openpyxl.Workbook()
    book.active.title = imp.SHEETS[0]
    for name in imp.SHEETS[1:]:
        book.create_sheet(name)
    book.save(path)
    return imp.sha256(path)


class ImportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)

    def tearDown(self) -> None:
        self.temp.cleanup()

    def test_forbidden_document_path(self) -> None:
        safe = self.root / "_dev/review.md"
        safe.parent.mkdir()
        safe.write_text("safe")
        (self.root / "ref").mkdir()
        (self.root / "ref/answer.md").write_text("answer")
        self.assertEqual(imp.checked_path(self.root, "_dev/review.md"), safe)
        for path in ("../outside", "/etc/passwd", "ref/answer.md", "_dev/missing.md"):
            with self.subTest(path=path), self.assertRaises(imp.ImportErrorEvidence):
                imp.checked_path(self.root, path)

    def test_workbook_hash_and_schema(self) -> None:
        path = self.root / "raw.xlsx"
        digest = make_book(path)
        self.assertEqual(imp.workbook(self.root, "raw.xlsx", digest), digest)
        with self.assertRaises(imp.ImportErrorEvidence):
            imp.workbook(self.root, "raw.xlsx", "0" * 64)
        bad = openpyxl.Workbook()
        bad.save(self.root / "bad.xlsx")
        with self.assertRaises(imp.ImportErrorEvidence):
            imp.workbook(self.root, "bad.xlsx")

    def synthetic_repo(self) -> dict:
        """A nine-deal checkout with filings, re-extraction receipts and decision documents."""
        root = self.root
        rows = ["file,deal,sha256"]
        manifest = {}
        deals = {}
        for deal in imp.DEALS:
            filing = root / "raw_filing" / f"{deal}.htm"
            filing.parent.mkdir(parents=True, exist_ok=True)
            filing.write_text(f"synthetic filing for {deal}")
            manifest[deal] = {"file": filing.name, "sha256": imp.sha256(filing)}
            rows.append(f"{filing.name},{deal},{manifest[deal]['sha256']}")
            folder = root / imp.REEXTRACT / "receipts" / deal
            digest = make_book(root / "extraction" / f"{deal}.xlsx")
            summary = {"errors": 0, "warnings": 2}
            deals[deal] = {"new_sha256": digest, "old_sha256": "0" * 64, "check_summary": summary}
            write_json(folder / "metadata.json", dict(mode="extract", provider="opus",
                model="claude-opus-5-5", effort="medium", instruction_sha256=imp.INSTRUCTION_HASH,
                filing_name=filing.name, filing_sha256=manifest[deal]["sha256"]))
            write_json(folder / "status.json", dict(state="completed", exit_code=0, continuations=0,
                provider={"served_models": ["claude-opus-5-5"]}, usage={"cost_usd": 1.5}))
            write_json(folder / "command.json", {})
            write_json(folder / "check.json", {"summary": summary})
            write_json(folder / "provider-results.json", [dict(subtype="success", is_error=False,
                modelUsage={"claude-opus-5-5": {}})])
        (root / "raw_filing/MANIFEST.csv").write_text("\n".join(rows) + "\n")
        write_json(root / imp.REEXTRACT / "reextraction.json", dict(model="claude-opus-5-5",
            effort="medium", instruction_sha256=imp.INSTRUCTION_HASH, deals=deals))
        for relative in (f"{imp.DATA}/ADJUDICATION.md", f"{imp.MAC_ACCEPTANCE}/RESEARCH_DECISION.md"):
            (root / relative).parent.mkdir(parents=True, exist_ok=True)
            (root / relative).write_text("recorded decision\n")
        return manifest

    def test_catalog_has_only_the_current_extraction(self) -> None:
        self.synthetic_repo()
        catalog, provenance = imp.build_catalog(self.root)
        self.assertEqual(list(catalog["deals"]), list(imp.DEALS))
        for deal, item in catalog["deals"].items():
            self.assertEqual(item["default_base"], "opus55-medium")
            self.assertEqual(len(item["versions"]), 1)
            only = item["versions"][0]
            self.assertEqual((only["id"], only["path"], only["sha256"]),
                             ("opus55-medium", f"extraction/{deal}.xlsx",
                              imp.sha256(self.root / "extraction" / f"{deal}.xlsx")))
            # Only case-level decisions survive, and none is tied to a removed version.
            for record in item["findings"] + item["documents"]:
                self.assertIsNone(record["source_version"])
            self.assertEqual(provenance["opus55_medium"][deal]["workbook_sha256"], only["sha256"])
        self.assertEqual([f["id"] for f in catalog["deals"]["datalink"]["findings"]], ["datalink-f9"])
        f9 = catalog["deals"]["datalink"]["findings"][0]
        self.assertEqual(f9["recorded_decision"]["actor"], "Austin Li")
        self.assertEqual((f9["judgment"], f9["implementation"]), ("unreviewed", "unassessed"))
        self.assertNotIn("recorded_correction", f9)
        mac = catalog["deals"]["mac-gray"]
        self.assertEqual([f["id"] for f in mac["findings"]], ["mac-gray-r01"])
        self.assertEqual([d["id"] for d in mac["documents"]], ["mac-gray-r01"])
        for deal in set(imp.DEALS) - {"datalink", "mac-gray"}:
            self.assertEqual((catalog["deals"][deal]["findings"], catalog["deals"][deal]["documents"]), ([], []))
        text = json.dumps(catalog)
        for retired in ("v1132-raw", "v113-baseline", "datalink-verified", "mac-gray-verified",
                        "mac-gray-candidate", "previous/"):
            self.assertNotIn(retired, text)

    def test_main_writes_only_catalog_and_current_verification(self) -> None:
        self.synthetic_repo()
        before = {p for p in self.root.rglob("*") if p.is_file()}
        self.assertEqual(imp.main(["--root", str(self.root)]), 0)
        written = {p.relative_to(self.root).as_posix()
                   for p in self.root.rglob("*") if p.is_file()} - {
                       p.relative_to(self.root).as_posix() for p in before}
        self.assertEqual(written, {"_dev/cockpit/catalog.json",
                                   f"{imp.REEXTRACT}/import-verification.json"})

    def test_filing_must_match_manifest(self) -> None:
        self.synthetic_repo()
        (self.root / "raw_filing/stec.htm").write_text("altered filing")
        with self.assertRaisesRegex(imp.ImportErrorEvidence, "Filing hash differs from manifest: stec"):
            imp.build_catalog(self.root)

    def test_reextraction_refuses_wrong_model_or_workbook(self) -> None:
        manifest = self.synthetic_repo()
        result = imp.verify_reextraction(self.root, manifest)
        self.assertEqual(result["kraton"]["workbook_sha256"],
                         imp.sha256(self.root / "extraction/kraton.xlsx"))
        self.assertEqual(result["kraton"]["reported_cost_usd"], 1.5)
        meta = self.root / imp.REEXTRACT / "receipts/stec/metadata.json"
        good = json.loads(meta.read_text())
        write_json(meta, {**good, "effort": "high"})
        with self.assertRaisesRegex(imp.ImportErrorEvidence, "prepared inputs differ: stec"):
            imp.verify_reextraction(self.root, manifest)
        write_json(meta, good)
        other = openpyxl.load_workbook(self.root / "extraction/penford.xlsx")
        other[imp.SHEETS[0]]["A1"] = "not the run's output"
        other.save(self.root / "extraction/penford.xlsx")
        with self.assertRaisesRegex(imp.ImportErrorEvidence, "Hash mismatch"):
            imp.verify_reextraction(self.root, manifest)


if __name__ == "__main__":
    unittest.main()
