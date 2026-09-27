"""The cockpit's Excel download, with the filing's EDGAR links and the run's provenance on a Source sheet.

The sheet is added here, on download, and never by the model or by `Workspace.export`, which
`export_repo.py` and `verify_catalog.py` share and which keeps producing the four-sheet workbook the
checker accepts. A workbook with a Source sheet fails the checker's sheet rule by design; the
four-sheet file is one option away. Links come from the filing's records and hashes from receipts
(the manifest, the catalog, the versions table); a value that is not recorded says so.
"""
from __future__ import annotations

import io
from typing import Any

import openpyxl
from openpyxl.styles import Font

import fetch_filing
from cockpit import data
from cockpit.trace import DEAL_STATUSES, _deal_review
from cockpit.workspace import REV_ID, Conflict, _decode, _now

SHEET = "Source"
NOT_RECORDED = "not recorded"
RAW_ONLY = "not applicable: a raw version, not the working copy"


def download(cockpit: data.Cockpit, slug: str, version: str = "working", source: bool | None = None) -> tuple[bytes, str]:
    """(workbook bytes, file name) for GET /api/deal/<slug>/export.

    The working copy and a past revision ("rev:N") get the Source sheet unless `source` is False,
    and are named by their revision; a version is its raw bytes, unchanged, unless `source` is True.
    """
    workspace = cockpit.workspace
    past = REV_ID.fullmatch(version or "")
    if source is None:
        source = version == "working" or bool(past)
    if version != "working" and not past:
        content = workspace.export(slug, version)
        if not source:
            return content, f"{slug}-{version}.xlsx"
        item = workspace.item(slug)
        base = workspace.version(item, version)
        state = workspace._base_state(slug, item, base)
        return add_sheet(content, rows(cockpit, slug, item, base, state, None, None)), f"{slug}-{version}-with-source.xlsx"
    item = workspace.item(slug)
    conn = workspace._connect()
    try:
        if conn: conn.execute("BEGIN")  # one read snapshot, so the bytes, the revision and the review status agree
        if past:
            # As Workspace.export: revision 0 is the starting base's raw bytes, revision N its snapshot on its own base.
            revision = int(past.group(1))
            state, base, _ = workspace._past(slug, item, conn, version)
            content = workspace._path(base["path"]).read_bytes() if revision == 0 else workspace._render_xlsx(base, state)
        else:
            state, row, base = workspace._state(slug, item, conn)
            content = workspace._render_xlsx(base, state) if row else workspace._path(base["path"]).read_bytes()
            revision = row["revision"] if row else 0
        review = _deal_review(conn, slug)
    finally:
        if conn: conn.close()
    name = f"{slug}-working-r{revision}.xlsx"
    if not source:
        return content, name
    return add_sheet(content, rows(cockpit, slug, item, base, state, revision, review, past=bool(past))), name


def links(cockpit: data.Cockpit, slug: str, item: dict[str, Any], entry: dict[str, str]) -> tuple[str | None, str | None]:
    """(EDGAR filing index, complete submission .txt) as recorded for the deal, or None where not recorded.

    The .txt link is the filing's recorded source. The index link is the added deal's recorded
    index_url, or the seed's index_url for a catalog deal, and is shown only when it equals the one
    derived from the .txt link.
    """
    submission = (entry.get("source_url") or "").strip()
    try:
        derived = fetch_filing.index_link(submission)
    except fetch_filing.FetchError:
        return None, None
    if item.get("added"):
        stated = item["added"].get("index_url")
    else:
        stated = next((seed["index_url"] for seed in cockpit.deals.seed() if seed["deal"] == slug), None)
    return (stated if stated == derived else None), submission


def review_text(review: dict[str, Any] | None, revision: int, past: bool = False) -> str:
    """The deal's review status as it relates to the working revision shown.

    Only the deal's latest status is kept, with the revision it was set at (always the working copy's
    revision then, so later markings never carry an earlier revision). A revision at or after that one
    therefore had this status, "edited since" where it is later. For a past revision older than the
    status, the status it had then is not kept here (the deal's activity lists each marking), so none is claimed.
    """
    if review is None or review.get("revision") is None:
        return "not set"
    label = DEAL_STATUSES.get(review["status"], review["status"])
    if revision < review["revision"]:
        return f"{label} at revision {review['revision']} (set after this revision; the status at revision {revision} is not recorded here)"
    status = f"{label} at revision {review['revision']}"
    if revision > review["revision"]:
        status += f"; edited since ({'this file is revision' if past else 'working revision'} {revision})"
    return status


def rows(cockpit: data.Cockpit, slug: str, item: dict[str, Any], base: dict[str, Any], state: dict[str, Any],
         revision: int | None, review: dict[str, Any] | None, past: bool = False) -> list[tuple[str, str, bool]]:
    """(field, value, is a link) for the Source sheet; revision and review are None for a raw version, and
    `past` marks a past revision ("rev:N") rather than the working copy's latest."""
    _, _, entry = cockpit.resolve(slug)
    index, submission = links(cockpit, slug, item, entry)
    facts = [record["values"] for record in state["sheets"][data.FACTS_SHEET]["rows"]]
    pages = next((data.display_value(_decode(values.get("Value"))) for values in facts if data.display_value(_decode(values.get("Field"))) == "Background pages"), "")
    instruction = base.get("instruction_version") or base.get("instruction_id")
    if base.get("instruction_version") and base.get("instruction_id"):
        instruction = f"{base['instruction_version']} (cockpit instruction {base['instruction_id']})"
    if revision is None:
        revision_text = status = RAW_ONLY
    else:
        revision_text = str(revision)
        current = (review or {}).get("current_revision")
        if past and current is not None and revision < current:
            revision_text += f" (a past revision; the working copy is at revision {current})"
        status = review_text(review, revision, past)
    return [
        ("EDGAR filing index", index or NOT_RECORDED, bool(index)),
        ("Complete submission (.txt)", submission or NOT_RECORDED, bool(submission)),
        ("Background pages", pages or NOT_RECORDED, False),
        ("Filing SHA-256", entry.get("sha256") or NOT_RECORDED, False),
        ("Instruction", instruction or NOT_RECORDED, False),
        ("Instruction SHA-256", cockpit.workspace._instruction_hash(base) or NOT_RECORDED, False),
        ("Version ID" if revision is None else f"Version ID (base of revision {revision})" if past else "Version ID (working-copy base)", base.get("id") or NOT_RECORDED, False),
        ("Raw workbook SHA-256", base.get("sha256") or NOT_RECORDED, False),
        ("Working revision", revision_text, False),
        ("Review status", status, False),
        ("Exported at", _now(), False),
    ]


def add_sheet(content: bytes, values: list[tuple[str, str, bool]]) -> bytes:
    """The workbook with a Source sheet appended after its own sheets."""
    wb = openpyxl.load_workbook(io.BytesIO(content))
    if SHEET in wb.sheetnames:
        wb.close()
        raise Conflict(f"the workbook already has a sheet named {SHEET}")
    ws = wb.create_sheet(SHEET)
    ws.append(["Field", "Value"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for field, value, link in values:
        ws.append([field, value])
        cell = ws.cell(ws.max_row, 2)
        if value.startswith("="):
            cell.data_type = "s"  # text copied from the workbook, never a formula
        if link:
            cell.hyperlink = value
            cell.style = "Hyperlink"
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()
