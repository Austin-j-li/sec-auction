"""Catalog-backed, transactional editable copies of SEC ledgers."""
from __future__ import annotations

import copy
import datetime as dt
import hashlib
import io
import json
import re
import sqlite3
import tempfile
import uuid
from pathlib import Path
from typing import Any

import openpyxl

from cockpit import data
import check_lean

SHEETS = (data.LEDGER_SHEET, data.ROUNDS_SHEET, data.QUESTIONS_SHEET, data.FACTS_SHEET)
DATE_FIELDS = {"Sort date", "Date from", "Date to", "Opened"}
NUMBER_FIELDS = {"#", "Process", "Round", "Price low", "Price high", "Count"}
QID = re.compile(r"Q[1-9][0-9]*\Z")
REF_EXPR = re.compile(r"(?<![\w])#\s*(\d+)(?:\s*[-–—]\s*#?\s*(\d+))?\b")
QREF = re.compile(r"(?<![\w])Q[1-9][0-9]*\b")
SAFE_ID = re.compile(r"[A-Za-z0-9][A-Za-z0-9_.-]*\Z")

class WorkspaceError(ValueError):
    status = 400
class Conflict(WorkspaceError):
    status = 409
class Missing(WorkspaceError):
    status = 404


def _encode(value: Any) -> Any:
    if isinstance(value, dt.datetime):
        return {"$type": "datetime", "value": value.isoformat()}
    if isinstance(value, dt.date):
        return {"$type": "date", "value": value.isoformat()}
    return value


def _decode(value: Any) -> Any:
    if isinstance(value, dict) and value.get("$type") == "datetime":
        return dt.datetime.fromisoformat(value["value"])
    if isinstance(value, dict) and value.get("$type") == "date":
        return dt.date.fromisoformat(value["value"])
    if isinstance(value, dict) and value.get("$type") == "formula":
        return value["value"]
    return value


def _affected(value: Any) -> tuple[set[int], bool]:
    """Call the checker parser only for bounded ranges."""
    if isinstance(value, str):
        for start, end in re.findall(r"(?:#\s*)?(\d+)\s*-\s*(?:#\s*)?(\d+)", value):
            if abs(int(end) - int(start)) > 10000:
                raise WorkspaceError("Rows affected range is too large")
    return check_lean.parse_affected_rows(value)


def _dump(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def _now() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds")


def _hash(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def record_label(sheet: str, previous: dict[str, Any] | None, current: dict[str, Any] | None) -> str:
    """A readable label for a sheet record, as used in change lists and comment targets."""
    record = current or previous
    if record is None: return sheet
    values = {key: data.display_value(_decode(value)) for key, value in record["values"].items()}
    if sheet == data.LEDGER_SHEET:
        earlier = data.display_value(_decode(previous["values"].get("#"))) if previous else ""
        later = data.display_value(_decode(current["values"].get("#"))) if current else ""
        number = f"#{earlier} → #{later}" if earlier and later and earlier != later else f"#{later or earlier or '?'}"
        details = [values.get(field, "") for field in ("When", "Who", "Event")]
        return " · ".join([f"Event {number}"] + [part for part in details if part])
    if sheet == data.QUESTIONS_SHEET:
        return f"Question {values.get('Q') or '?'}"
    if sheet == data.ROUNDS_SHEET:
        return f"Process {values.get('Process') or '?'} / Round {values.get('Round') or '?'}"
    return f"Deal fact {values.get('Field') or '?'}"


class Workspace:
    def __init__(self, cockpit: data.Cockpit):
        self.cockpit = cockpit
        self.root = cockpit.repo_root.resolve()
        self.catalog_path = self.root / "_dev/cockpit/catalog.json"
        self.db_path = self.root / "_dev/cockpit/state/workspace.sqlite3"

    @property
    def available(self) -> bool:
        return self.catalog_path.is_file()

    def catalog(self) -> dict[str, Any]:
        try:
            doc = json.loads(self.catalog_path.read_text(encoding="utf-8"))
        except (OSError, ValueError) as exc:
            raise WorkspaceError("catalog cannot be read") from exc
        if doc.get("schema_version") != 1 or not isinstance(doc.get("deals"), dict):
            raise WorkspaceError("unsupported catalog")
        return doc

    def item(self, slug: str) -> dict[str, Any]:
        if not data.SLUG_RE.fullmatch(slug or ""):
            raise data.DealNotFound("unknown deal")
        item = self.catalog()["deals"].get(slug)
        if not isinstance(item, dict):
            raise data.DealNotFound("unknown deal")
        return item

    def _path(self, relative: str) -> Path:
        if not isinstance(relative, str) or not relative or Path(relative).is_absolute():
            raise WorkspaceError("invalid catalog path")
        path = (self.root / relative).resolve()
        if self.root not in path.parents or "ref" in path.relative_to(self.root).parts:
            raise WorkspaceError("catalog path outside allowed repository")
        return path

    def version(self, item: dict[str, Any], ident: str) -> dict[str, Any]:
        found = next((v for v in item.get("versions", []) if v.get("id") == ident), None)
        if not isinstance(found, dict):
            raise Missing("unknown version")
        path = self._path(found.get("path"))
        if not path.is_file() or _hash(path) != found.get("sha256"):
            raise Conflict("catalog workbook hash does not match")
        return found

    def base(self, item: dict[str, Any]) -> dict[str, Any]:
        ident = item.get("default_base")
        if not isinstance(ident, str) or not ident:
            raise WorkspaceError("catalog deal has no default_base")
        return self.version(item, ident)

    def _connect(self, write: bool = False) -> sqlite3.Connection | None:
        if not write and not self.db_path.is_file():
            return None
        if write:
            self.db_path.parent.mkdir(parents=True, exist_ok=True)
        conn = sqlite3.connect(self.db_path, timeout=30)
        conn.row_factory = sqlite3.Row
        if write:
            conn.execute("PRAGMA journal_mode=WAL")
            conn.execute("CREATE TABLE IF NOT EXISTS revisions (slug TEXT NOT NULL, revision INTEGER NOT NULL, base_id TEXT NOT NULL, base_sha256 TEXT NOT NULL, at TEXT NOT NULL, actor TEXT NOT NULL, reason TEXT NOT NULL, summary TEXT NOT NULL, changes TEXT NOT NULL, snapshot TEXT NOT NULL, PRIMARY KEY (slug, revision))")
            from cockpit import trace  # imported here: trace builds on this module
            try:
                trace.ensure_schema(conn)
            except Exception:
                conn.close()
                raise
        return conn

    def _latest(self, conn: sqlite3.Connection | None, slug: str, revision: int | None = None) -> sqlite3.Row | None:
        if conn is None:
            return None
        try:
            if revision is None:
                return conn.execute("SELECT * FROM revisions WHERE slug=? ORDER BY revision DESC LIMIT 1", (slug,)).fetchone()
            return conn.execute("SELECT * FROM revisions WHERE slug=? AND revision=?", (slug, revision)).fetchone()
        except sqlite3.OperationalError:
            return None

    def _base_state(self, slug: str, item: dict[str, Any], ver: dict[str, Any]) -> dict[str, Any]:
        wb = openpyxl.load_workbook(io.BytesIO(self._path(ver["path"]).read_bytes()), data_only=False)
        sheets = {}
        for sheet in SHEETS:
            ws = wb[sheet]
            columns = [data.display_value(cell.value) for cell in ws[1]]
            while columns and not columns[-1]: columns.pop()
            body = []
            for excel_row in range(2, ws.max_row + 1):
                cells = [ws.cell(excel_row, column + 1) for column in range(len(columns))]
                if all(check_lean.is_blank(cell.value) for cell in cells): continue
                body.append((excel_row, [_encode(cell.value) if cell.data_type != "f" else {"$type": "formula", "value": cell.value} for cell in cells]))
            sheets[sheet] = {"columns": columns, "rows": [
                {"uid": hashlib.sha256(f"{slug}|{ver['sha256']}|{sheet}|{excel_row}".encode()).hexdigest()[:24],
                 "excel_row": excel_row,
                 "source_row": excel_row,
                 "values": {col: value for col, value in zip(columns, values) if col}}
                for excel_row, values in body]}
        wb.close()
        return {"sheets": sheets, "findings": {}, "row_review": {}, "reference_warnings": []}

    def _state(self, slug: str, item: dict[str, Any], conn: sqlite3.Connection | None, revision: int | None = None) -> tuple[dict[str, Any], sqlite3.Row | None, dict[str, Any]]:
        row = self._latest(conn, slug, revision)
        if revision is not None and row is None and revision != 0:
            raise Missing("unknown revision")
        base = self.base(item)
        if row:
            if row["base_id"] != base["id"] or row["base_sha256"] != base["sha256"]:
                raise Conflict("working base differs from catalog; migration required")
            return json.loads(row["snapshot"]), row, base
        if revision == 0 and self._latest(conn, slug):
            raise Missing("unknown revision")
        return self._base_state(slug, item, base), None, base

    def _tables(self, state: dict[str, Any]) -> dict[str, Any]:
        return {sheet: (part["columns"], [(row.get("excel_row") or i + 2, [_decode(row["values"].get(col)) for col in part["columns"]]) for i, row in enumerate(part["rows"])]) for sheet, part in state["sheets"].items()}

    def _render_xlsx(self, base: dict[str, Any], state: dict[str, Any]) -> bytes:
        raw = self._path(base["path"]).read_bytes()
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        original = self._base_state("_style_only", {}, base)
        for sheet in SHEETS:
            ws = wb[sheet]
            part = state["sheets"][sheet]
            base_part = original["sheets"][sheet]
            if ([r["values"] for r in part["rows"]] == [r["values"] for r in base_part["rows"]]
                    and [r.get("excel_row") for r in part["rows"]] == [r.get("excel_row") for r in base_part["rows"]]):
                continue
            columns = part["columns"]
            original_cells = {}
            original_dims = {}
            for record in base_part["rows"]:
                source_row = record["excel_row"]
                original_cells[source_row] = [copy.copy(ws.cell(source_row, col)) for col in range(1, len(columns) + 1)]
                original_dims[source_row] = copy.copy(ws.row_dimensions[source_row])
            template = next(iter(original_cells.values()), None)
            if ws.max_row > 1:
                ws.delete_rows(2, ws.max_row - 1)
            for key in list(ws.row_dimensions):
                if key >= 2: del ws.row_dimensions[key]
            for row_number, record in enumerate(part["rows"], 2):
                source_row = record.get("source_row")
                source_cells = original_cells.get(source_row, template)
                if source_row in original_dims:
                    dim = copy.copy(original_dims[source_row])
                    dim.index = row_number
                    ws.row_dimensions[row_number] = dim
                for col_number, field in enumerate(columns, 1):
                    cell = ws.cell(row_number, col_number)
                    value = record["values"].get(field)
                    cell.value = _decode(value)
                    if isinstance(value, str) and value.startswith("="):
                        cell.data_type = "s"
                    if source_cells:
                        source = source_cells[col_number - 1]
                        cell._style = copy.copy(source._style)
                        cell.comment = copy.copy(source.comment)
                        cell.hyperlink = copy.copy(source.hyperlink)
                        cell.number_format = source.number_format
                    elif field in DATE_FIELDS:
                        cell.number_format = "yyyy-mm-dd"
            if ws.auto_filter.ref:
                ws.auto_filter.ref = f"A1:{openpyxl.utils.get_column_letter(len(columns))}{len(part['rows']) + 1}"
            if ws.freeze_panes is None:
                ws.freeze_panes = "A2"
        out = io.BytesIO()
        wb.save(out)
        wb.close()
        return out.getvalue()

    def _payload(self, slug: str, item: dict[str, Any], state: dict[str, Any], row: sqlite3.Row | None, base: dict[str, Any], selected: str) -> dict[str, Any]:
        _, filing_path, entry = self.cockpit.resolve(slug)
        filing = self.cockpit.filing(filing_path)
        if row is None:
            source_path = self._path(base["path"])
            report = self.cockpit.check(source_path, filing_path)
        else:
            content = self._render_xlsx(base, state)
            with tempfile.TemporaryDirectory(prefix="cockpit-check-") as folder:
                path = Path(folder) / "working.xlsx"
                path.write_bytes(content)
                report = check_lean.LeanChecker(path, filing_path).run()
        payload = data.build_deal_payload(slug, entry, self._tables(state), filing, report)
        payload["name"] = item.get("name", slug)
        for sheet, key in ((data.LEDGER_SHEET, "ledger"), (data.ROUNDS_SHEET, "rounds"), (data.QUESTIONS_SHEET, "questions")):
            for shown, stored in zip(payload[key]["rows"], state["sheets"][sheet]["rows"]):
                shown["uid"] = stored["uid"]
        for shown, stored in zip(payload["facts"], state["sheets"][data.FACTS_SHEET]["rows"]):
            shown["uid"] = stored["uid"]
        for shown in payload["ledger"]["rows"]:
            number = shown["id"]
            shown["has_references"] = self._references(state, number, shown["uid"])
        working_base = self.base(item)
        payload["versions"] = [{"id": "working", "label": "Working copy", "instruction_version": working_base.get("instruction_version"), "kind": "working", "sha256": working_base["sha256"], "review_status": "in_review" if row else working_base.get("review_status", "unreviewed")}] + [
            {key: version.get(key) for key in ("id", "label", "instruction_version", "kind", "sha256", "review_status")}
            for version in item.get("versions", [])]
        payload["workspace"] = {"revision": row["revision"] if row else 0, "base_version": working_base["id"], "base_sha256": working_base["sha256"], "updated_at": row["at"] if row else None, "updated_by": row["actor"] if row else None, "editable": selected == "working", "selected_version": selected}
        payload["workspace"]["reference_warnings"] = state.get("reference_warnings", [])
        payload["findings"] = [dict(finding, **state["findings"].get(finding["id"], {})) for finding in item.get("findings", [])]
        payload["documents"] = [{key: doc.get(key) for key in ("id", "label", "source_version", "kind")} for doc in item.get("documents", [])]
        payload["row_review"] = state["row_review"]
        payload["choices"] = {
            "Type": sorted(check_lean.TYPES), "Event": sorted(check_lean.EVENTS),
            "All cash": sorted(check_lean.ALL_CASH), "Formality": sorted(check_lean.FORMALITY),
            "Conditions": sorted(check_lean.CONDITIONS), "Exit reason": sorted(check_lean.EXIT_REASONS),
            "Finality": sorted(check_lean.FINALITY), "Deadline outcome": sorted(check_lean.DEADLINE_OUTCOMES),
            "Initiation": sorted(check_lean.INITIATION),
        }
        return payload

    def deal(self, slug: str, version: str = "working") -> dict[str, Any]:
        item = self.item(slug)
        conn = self._connect()
        try:
            if version == "working":
                state, row, base = self._state(slug, item, conn)
            elif SAFE_ID.fullmatch(version or ""):
                base = self.version(item, version)
                state = self._base_state(slug, item, base)
                working_state, _, _ = self._state(slug, item, conn)
                state["findings"] = copy.deepcopy(working_state["findings"])
                row = None
            else:
                raise Missing("unknown version")
            return self._payload(slug, item, state, row, base, version)
        finally:
            if conn: conn.close()

    def document(self, slug: str, ident: str) -> dict[str, str]:
        item = self.item(slug)
        if not SAFE_ID.fullmatch(ident or ""):
            raise Missing("unknown document")
        doc = next((d for d in item.get("documents", []) if d.get("id") == ident), None)
        if not doc:
            raise Missing("unknown document")
        path = self._path(doc.get("path"))
        if not path.is_file() or path.suffix.lower() not in (".md", ".txt"):
            raise Missing("document unavailable")
        return {"title": doc.get("label", ident), "text": path.read_text(encoding="utf-8")}

    def export(self, slug: str, version: str = "working") -> bytes:
        item = self.item(slug)
        if version != "working":
            return self._path(self.version(item, version)["path"]).read_bytes()
        conn = self._connect()
        try:
            state, row, base = self._state(slug, item, conn)
            return self._render_xlsx(base, state) if row else self._path(base["path"]).read_bytes()
        finally:
            if conn: conn.close()

    def history(self, slug: str) -> dict[str, Any]:
        self.item(slug)
        conn = self._connect()
        try:
            rows = [] if conn is None else conn.execute("SELECT revision,at,actor,reason,summary,changes FROM revisions WHERE slug=? ORDER BY revision DESC", (slug,)).fetchall()
            history = [{"revision": r["revision"], "at": r["at"], "actor": r["actor"], "reason": r["reason"], "summary": r["summary"], "changes": json.loads(r["changes"])} for r in rows]
            history.append({"revision": 0, "at": None, "actor": None, "reason": "Original working base", "summary": "Base workbook", "changes": []})
            return {"history": history}
        finally:
            if conn: conn.close()

    def changes(self, slug: str) -> dict[str, Any]:
        item = self.item(slug)
        conn = self._connect()
        try:
            state, _, base = self._state(slug, item, conn)
            original = self._base_state(slug, item, base)
            return {"base_label": base.get("label", base["id"]), "changes": self._diff(original, state)}
        finally:
            if conn: conn.close()

    def _diff(self, before: dict[str, Any], after: dict[str, Any]) -> list[dict[str, Any]]:
        changes = []
        for sheet in SHEETS:
            old = {r["uid"]: r for r in before["sheets"][sheet]["rows"]}
            new = {r["uid"]: r for r in after["sheets"][sheet]["rows"]}
            for uid in old.keys() | new.keys():
                a, b = old.get(uid), new.get(uid)
                row_label = record_label(sheet, a, b)
                if a is None or b is None:
                    changes.append({"sheet": sheet, "uid": uid, "record_label": row_label, "field": None, "before": None if a is None else {k: data.display_value(_decode(v)) for k,v in a["values"].items()}, "after": None if b is None else {k: data.display_value(_decode(v)) for k,v in b["values"].items()}, "type": "insert" if a is None else "delete"})
                else:
                    for field in after["sheets"][sheet]["columns"]:
                        va, vb = a["values"].get(field), b["values"].get(field)
                        if va != vb:
                            changes.append({"sheet": sheet, "uid": uid, "record_label": row_label, "field": field, "before": data.display_value(_decode(va)), "after": data.display_value(_decode(vb)), "type": "update"})
            old_order = [r["uid"] for r in before["sheets"][sheet]["rows"] if r["uid"] in new]
            new_order = [r["uid"] for r in after["sheets"][sheet]["rows"] if r["uid"] in old]
            if old_order != new_order:
                changes.append({"sheet": sheet, "uid": None, "record_label": f"{sheet} order", "field": "order", "before": [{"uid": uid, "label": record_label(sheet, old[uid], None)} for uid in old_order], "after": [{"uid": uid, "label": record_label(sheet, None, new[uid])} for uid in new_order], "type": "move"})
        for key in ("findings", "row_review"):
            for uid in before[key].keys() | after[key].keys():
                if before[key].get(uid) != after[key].get(uid):
                    if key == "findings":
                        row_label = f"Finding {uid}"
                    else:
                        row_label = next((record_label(sheet, next((r for r in before["sheets"][sheet]["rows"] if r["uid"] == uid), None), next((r for r in after["sheets"][sheet]["rows"] if r["uid"] == uid), None)) for sheet in SHEETS if any(r["uid"] == uid for r in before["sheets"][sheet]["rows"] + after["sheets"][sheet]["rows"])), f"Row {uid}")
                    changes.append({"sheet": key, "uid": uid, "record_label": row_label, "field": None, "before": before[key].get(uid), "after": after[key].get(uid), "type": "decision"})
        return changes

    def _coerce(self, sheet: str, field: str, value: Any) -> Any:
        if value is None or value == "": return None
        if isinstance(value, bool) or not isinstance(value, (str, int, float)):
            raise WorkspaceError("invalid cell value")
        if isinstance(value, str):
            if len(value.encode("utf-16-le")) // 2 > 32767:
                raise WorkspaceError("cell text exceeds Excel's 32767-character limit")
        if field in DATE_FIELDS:
            if not isinstance(value, str): raise WorkspaceError("date must be text")
            try:
                parsed = dt.datetime.fromisoformat(value) if "T" in value else dt.datetime.strptime(value, "%Y-%m-%d")
            except ValueError:
                try: parsed = dt.datetime.strptime(value, "%m/%d/%Y")
                except ValueError as exc: raise WorkspaceError("invalid date") from exc
            return _encode(parsed)
        if field == "Round" and sheet == data.LEDGER_SHEET and value == "post":
            return "post"
        if field in NUMBER_FIELDS:
            try:
                number = float(value)
            except (TypeError, ValueError) as exc: raise WorkspaceError("invalid number") from exc
            if not -1e12 < number < 1e12: raise WorkspaceError("number out of range")
            if field in {"#", "Process", "Round", "Count"}:
                if not number.is_integer(): raise WorkspaceError("whole number required")
                if field == "Round" and ((sheet == data.ROUNDS_SHEET and number < 1) or (sheet == data.LEDGER_SHEET and number < 0)):
                    raise WorkspaceError("invalid round number")
                return int(number)
            return number
        return value

    def _row(self, state: dict[str, Any], sheet: str, uid: str) -> dict[str, Any]:
        if sheet not in SHEETS: raise WorkspaceError("unknown sheet")
        return next((r for r in state["sheets"][sheet]["rows"] if r["uid"] == uid), None) or self._missing_row()

    @staticmethod
    def _missing_row() -> Any:
        raise WorkspaceError("unknown row uid")

    def _renumber(self, state: dict[str, Any], deleted: dict[str, str | None] | None = None, old_override: dict[int, str] | None = None) -> None:
        ledger = state["sheets"][data.LEDGER_SHEET]["rows"]
        old = old_override or {int(r["values"].get("#")): r["uid"] for r in ledger if str(r["values"].get("#", "")).isdigit()}
        new = {r["uid"]: i for i, r in enumerate(ledger, 1)}
        deleted = deleted or {}
        warnings = []
        def destination(uid: str) -> str | None:
            seen = set()
            while uid in deleted:
                if uid in seen: raise WorkspaceError("cyclic replacement references")
                seen.add(uid)
                uid = deleted[uid]
                if uid is None: return None
            return uid
        def mapref(match: re.Match) -> str:
            first, last = int(match.group(1)), int(match.group(2) or match.group(1))
            if abs(last - first) > 10000: raise WorkspaceError("ledger reference range is too large")
            numbers = range(first, last + (1 if last >= first else -1), 1 if last >= first else -1)
            mapped = []
            for number in numbers:
                uid = old.get(number)
                if uid is None:
                    mapped.append(number)
                    continue
                dest = destination(uid)
                if dest is None: raise WorkspaceError("remaining text references a deleted event without replacement")
                if new[dest] not in mapped: mapped.append(new[dest])
            return ", ".join(f"#{n}" for n in mapped) if mapped else match.group(0)
        for sheet in SHEETS:
            for row in state["sheets"][sheet]["rows"]:
                for field, val in list(row["values"].items()):
                    if field in ("#", data.QUOTE_COLUMN) or not isinstance(val, str): continue
                    if sheet == data.QUESTIONS_SHEET and field == "Rows affected":
                        refs, parseable = _affected(val)
                        if parseable and refs:
                            mapped = []
                            for number in sorted(refs):
                                uid = old.get(number)
                                if uid is None:
                                    mapped.append(number)
                                    continue
                                dest = destination(uid)
                                if dest is None: raise WorkspaceError("Rows affected references a deleted event without replacement")
                                if new[dest] not in mapped: mapped.append(new[dest])
                            row["values"][field] = ", ".join(f"#{n}" for n in sorted(mapped)) or None
                        elif val.strip():
                            warnings.append({"uid": row["uid"], "question": row["values"].get("Q"), "message": "Narrative Rows affected needs manual review after ledger reorder"})
                        continue
                    row["values"][field] = REF_EXPR.sub(mapref, val)
        for i, row in enumerate(ledger, 1): row["values"]["#"] = i
        state["reference_warnings"] = warnings

    def _apply(self, state: dict[str, Any], item: dict[str, Any], op: dict[str, Any], actor: str, deleted: dict[str, str | None]) -> str | None:
        typ = op.get("type")
        if typ == "restore": raise WorkspaceError("restore must be the only operation")
        if typ in ("update", "insert", "delete", "move"):
            sheet = op.get("sheet")
            if sheet not in SHEETS: raise WorkspaceError("unknown sheet")
            rows = state["sheets"][sheet]["rows"]
            columns = state["sheets"][sheet]["columns"]
            if typ in ("update", "insert"):
                values = op.get("values")
                if not isinstance(values, dict) or not values or any(k not in columns for k in values): raise WorkspaceError("invalid columns")
                if sheet == data.LEDGER_SHEET and "#" in values: raise WorkspaceError("ledger number is managed")
                if typ == "update":
                    row = self._row(state, sheet, op.get("uid"))
                    previous_q = row["values"].get("Q") if sheet == data.QUESTIONS_SHEET else None
                    row["values"].update({k: self._coerce(sheet,k,v) for k,v in values.items()})
                    if previous_q != row["values"].get("Q"):
                        self._rename_question(state, previous_q, row["values"].get("Q"))
                else:
                    after = op.get("after_uid")
                    if after is not None and not any(r["uid"] == after for r in rows): raise WorkspaceError("unknown after_uid")
                    row = {"uid": uuid.uuid4().hex, "source_row": None, "values": {k: None for k in columns}}
                    row["values"].update({k: self._coerce(sheet,k,v) for k,v in values.items()})
                    index = next((i + 1 for i,r in enumerate(rows) if r["uid"] == after), 0)
                    rows.insert(index, row)
                    return row["uid"]
            elif typ == "move":
                uid, after = op.get("uid"), op.get("after_uid")
                if uid == after: return
                row = self._row(state, sheet, uid)
                if after is not None and not any(r["uid"] == after for r in rows): raise WorkspaceError("unknown after_uid")
                rows.remove(row)
                index = next((i + 1 for i,r in enumerate(rows) if r["uid"] == after), 0)
                rows.insert(index, row)
            else:
                uid = op.get("uid")
                row = self._row(state, sheet, uid)
                if sheet == data.LEDGER_SHEET:
                    number = row["values"].get("#")
                    refs = self._references(state, number, uid)
                    if refs and not op.get("replacement_uid"): raise WorkspaceError("referenced row needs a live replacement_uid or prior reference edits")
                    replacement = op.get("replacement_uid")
                    if replacement is not None and (replacement == uid or not any(r["uid"] == replacement for r in rows)):
                        raise WorkspaceError("invalid replacement_uid")
                    rows.remove(row)
                    deleted[uid] = replacement
                else:
                    if sheet == data.QUESTIONS_SHEET and self._question_referenced(state, row["values"].get("Q"), uid):
                        raise WorkspaceError("question ID is referenced elsewhere")
                    rows.remove(row)
                state["row_review"].pop(uid, None)
            return
        if typ == "finding":
            ident = op.get("id")
            if ident not in {f.get("id") for f in item.get("findings", [])}: raise WorkspaceError("unknown finding")
            judgment, implementation, verification = op.get("judgment"), op.get("implementation"), op.get("verification")
            if judgment not in ("unreviewed", "supported", "rejected", "deferred") or implementation not in ("unassessed", "not_applied", "applied") or verification not in ("unchecked", "verified"):
                raise WorkspaceError("invalid finding decision")
            note = op.get("note", "")
            if not isinstance(note, str) or len(note) > 100000: raise WorkspaceError("invalid note")
            state["findings"][ident] = {"judgment": judgment, "implementation": implementation, "verification": verification, "note": note, "actor": actor, "at": _now()}
            return
        if typ == "review":
            uid = op.get("uid")
            if not any(uid == r["uid"] for part in state["sheets"].values() for r in part["rows"]): raise WorkspaceError("unknown row uid")
            if op.get("status") not in ("unreviewed", "reviewed", "needs_decision"): raise WorkspaceError("invalid review status")
            note = op.get("note", "")
            if not isinstance(note, str) or len(note) > 100000: raise WorkspaceError("invalid note")
            state["row_review"][uid] = {"status": op["status"], "note": note, "actor": actor, "at": _now()}
            return
        raise WorkspaceError("unknown operation")

    def _references(self, state: dict[str, Any], number: Any, uid: str) -> bool:
        if not isinstance(number, int): return False
        for sheet in SHEETS:
            for row in state["sheets"][sheet]["rows"]:
                if row["uid"] == uid: continue
                for field, value in row["values"].items():
                    if not isinstance(value, str) or field == data.QUOTE_COLUMN: continue
                    for match in REF_EXPR.finditer(value):
                        first, last = int(match.group(1)), int(match.group(2) or match.group(1))
                        if abs(last - first) > 10000: raise WorkspaceError("ledger reference range is too large")
                        if min(first, last) <= number <= max(first, last): return True
                    if sheet == data.QUESTIONS_SHEET and field == "Rows affected":
                        refs, parseable = _affected(value)
                        if parseable and number in refs: return True
        return False

    def _question_referenced(self, state: dict[str, Any], qid: Any, skip_uid: str | None = None) -> bool:
        if not isinstance(qid, str): return False
        for sheet in SHEETS:
            for row in state["sheets"][sheet]["rows"]:
                if row["uid"] == skip_uid: continue
                for field, value in row["values"].items():
                    if field in ("Q", data.QUOTE_COLUMN) or not isinstance(value, str): continue
                    if qid in QREF.findall(value): return True
        return False

    def _rename_question(self, state: dict[str, Any], old: Any, new: Any) -> None:
        if not isinstance(old, str) or not QID.fullmatch(old) or not isinstance(new, str) or not QID.fullmatch(new):
            raise WorkspaceError("question ID must be Q plus a positive number")
        for sheet in SHEETS:
            for row in state["sheets"][sheet]["rows"]:
                for field, value in row["values"].items():
                    if field in ("Q", data.QUOTE_COLUMN) or not isinstance(value, str): continue
                    row["values"][field] = re.sub(rf"(?<![\w]){re.escape(old)}\b", new, value)

    def _validate(self, state: dict[str, Any], removed_questions: set[str] | None = None) -> None:
        questions = [r["values"].get("Q") for r in state["sheets"][data.QUESTIONS_SHEET]["rows"]]
        if any(not isinstance(q, str) or not QID.fullmatch(q) for q in questions) or len(set(questions)) != len(questions):
            raise WorkspaceError("question IDs must be unique Q numbers")
        for row in state["sheets"][data.LEDGER_SHEET]["rows"]:
            flag = row["values"].get("Flag")
            if isinstance(flag, str) and any(q not in questions for q in QREF.findall(flag)):
                raise WorkspaceError("ledger Flag references missing question")
        for sheet in SHEETS:
            for row in state["sheets"][sheet]["rows"]:
                for field, value in row["values"].items():
                    if not isinstance(value, str): continue
                    if field not in ("Q", data.QUOTE_COLUMN) and removed_questions and removed_questions.intersection(QREF.findall(value)):
                        raise WorkspaceError("text still references a removed question ID")
                    for match in REF_EXPR.finditer(value):
                        if match.group(2) and abs(int(match.group(2)) - int(match.group(1))) > 10000:
                            raise WorkspaceError("ledger reference range is too large")
                    if sheet == data.QUESTIONS_SHEET and field == "Rows affected": _affected(value)

    def edit(self, slug: str, request: dict[str, Any], actor: str) -> dict[str, Any]:
        item = self.item(slug)
        if not isinstance(request, dict) or type(request.get("revision")) is not int or not isinstance(request.get("base_sha256"), str): raise WorkspaceError("revision and base_sha256 required")
        ops = request.get("operations")
        reason = request.get("reason")
        if not isinstance(ops, list) or not ops or len(ops) > 100 or any(not isinstance(op, dict) for op in ops): raise WorkspaceError("operations required")
        if not isinstance(reason, str) or not reason.strip() or len(reason) > 1000: raise WorkspaceError("save reason required")
        conn = self._connect(write=True)
        assert conn is not None
        try:
            conn.execute("BEGIN IMMEDIATE")
            current, row, base = self._state(slug, item, conn)
            revision = row["revision"] if row else 0
            if request["revision"] != revision or request["base_sha256"] != base["sha256"]: raise Conflict("stale revision or base hash")
            before = copy.deepcopy(current)
            original_questions = {r["values"].get("Q") for r in before["sheets"][data.QUESTIONS_SHEET]["rows"]}
            if any(op.get("type") == "restore" for op in ops):
                if len(ops) != 1: raise WorkspaceError("restore must be the only operation")
                target = ops[0].get("target_revision")
                if type(target) is not int or target < 0: raise WorkspaceError("invalid target_revision")
                if target == 0: current = self._base_state(slug, item, base)
                else:
                    target_row = self._latest(conn, slug, target)
                    if target_row is None: raise WorkspaceError("unknown target revision")
                    current = json.loads(target_row["snapshot"])
            else:
                client_uids: dict[str, str] = {}
                old_ledger = {int(r["values"].get("#")): r["uid"] for r in current["sheets"][data.LEDGER_SHEET]["rows"] if str(r["values"].get("#", "")).isdigit()}
                deleted: dict[str, str | None] = {}
                structural = False
                for source_op in ops:
                    op = dict(source_op)
                    for field in ("uid", "after_uid", "replacement_uid"):
                        if isinstance(op.get(field), str) and op[field] in client_uids: op[field] = client_uids[op[field]]
                    if op.get("type") == "insert" and "client_uid" in op:
                        client_uid = op["client_uid"]
                        if not isinstance(client_uid, str) or not re.fullmatch(r"new-[a-zA-Z0-9-]{1,80}", client_uid) or client_uid in client_uids:
                            raise WorkspaceError("invalid or duplicate client_uid")
                    if op.get("sheet") == data.LEDGER_SHEET and op.get("type") in ("insert", "delete", "move"):
                        structural = True
                    created = self._apply(current, item, op, actor, deleted)
                    if created is not None and "client_uid" in op:
                        client_uids[op["client_uid"]] = created
                if structural:
                    self._renumber(current, deleted, old_ledger)
            final_questions = {r["values"].get("Q") for r in current["sheets"][data.QUESTIONS_SHEET]["rows"]}
            self._validate(current, original_questions - final_questions)
            for sheet in SHEETS:
                old_rows = before["sheets"][sheet]["rows"]
                new_rows = current["sheets"][sheet]["rows"]
                if [(r["uid"], r["values"]) for r in old_rows] != [(r["uid"], r["values"]) for r in new_rows]:
                    for index, record in enumerate(new_rows, 2): record["excel_row"] = index
            changes = self._diff(before, current)
            if not changes: raise WorkspaceError("no changes to save")
            at = _now()
            summary = f"{len(changes)} change{'s' if len(changes) != 1 else ''}"
            conn.execute("INSERT INTO revisions VALUES (?,?,?,?,?,?,?,?,?,?)", (slug, revision + 1, base["id"], base["sha256"], at, actor, reason.strip(), summary, _dump(changes), _dump(current)))
            from cockpit import trace
            trace.record_revision(conn, slug, actor, revision + 1, reason.strip(), restore=any(op.get("type") == "restore" for op in ops), at=at)
            saved = self._latest(conn, slug)
            payload = self._payload(slug, item, current, saved, base, "working")
            conn.commit()
            return payload
        except Exception:
            conn.rollback()
            raise
        finally:
            conn.close()
