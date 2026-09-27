"""The analysis derive tool on synthetic Version 1 ledgers."""
from __future__ import annotations

import contextlib
import csv
import datetime as dt
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

import check_lean
import derive_analysis as derive
from test_check_lean import build_examples_fixture

D = dt.date(2020, 1, 1)
FACTS = {"Target": "Target Co", "Acquirer": "Alpha", "Auction screen": "Met (process 1): at least 3 parties",
         "Whole-company bids": "Yes", "Initiation": "target-led"}


def row(n, who, event, rnd=1, day=0, **values):
    date = D + dt.timedelta(days=day)
    return {"#": n, "When": date.strftime("%m/%d/%Y"), "Who": who, "Event": event, "Process": 1, "Round": rnd,
            "Sort date": date, "Date from": date, "Date to": date, **values}


def bid(n, who, low, high=None, rnd=1, day=0, formality="Informal", conditions="Unclear", **values):
    terms = {"Stock %": 0, "Due diligence": "Not stated", "Financing": "Not stated", "Regulatory": "Not stated", "Exclusivity": "Not stated"}
    return row(n, who, values.pop("event", "Bid"), rnd, day, Type=values.pop("Type", "Financial"), Count=values.pop("Count", 1),
               **{"Price low": low, "Price high": high if high is not None else low, "Formality": formality,
                  "Conditions": conditions, **terms, **values})


def rounds_line(rnd, outcome, finality="Not final", received="1: Alpha", who="2: Alpha, Beta", due="01/10/2020"):
    return {"Process": 1, "Round": rnd, "Opened": D, "How opened": "letter", "Who was in": who, "Due dates": due,
            "Deadline outcome": outcome, "Finality": finality, "Bids received": received, "How it ended": "signing"}


def write(path: Path, ledger, rounds, facts=FACTS, source=None, questions=()) -> Path:
    wb = Workbook()
    for title, header, records in (("Deal ledger", check_lean.LEDGER_COLUMNS, ledger), ("Rounds", check_lean.ROUND_COLUMNS, rounds),
                                   ("Questions", check_lean.QUESTION_COLUMNS, list(questions))):
        ws = wb.active if title == "Deal ledger" else wb.create_sheet(title)
        ws.title = title
        ws.append(header)
        for record in records:
            ws.append([record.get(c) for c in header])
    ws = wb.create_sheet("Deal facts")
    ws.append(check_lean.FACT_COLUMNS)
    for field, value in facts.items():
        ws.append([field, value])
    if source:
        ws = wb.create_sheet("Source")
        for label, value in source.items():
            ws.append([label, value])
        ws["B1"].hyperlink = source["EDGAR filing index"]
    wb.save(path)
    return path


def read(folder: Path, name: str) -> list[dict[str, str]]:
    with (folder / name).open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def quiet(fn, *args):
    with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(io.StringIO()):
        return fn(*args)


class DeriveTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.dir = Path(self.temp.name)

    def tearDown(self): self.temp.cleanup()

    def derive(self, ledger, rounds, name="alpha.xlsx", **kwargs):
        path = write(self.dir / name, ledger, rounds, **kwargs)
        out = self.dir / (name + "-out")
        manifest = derive.run(path, out)
        return out, manifest

    def test_every_deadline_outcome_value_gets_one_class_per_due_date(self):
        rounds = [rounds_line(1, "Extended (late bid accepted)"), rounds_line(2, "Extended; Extended (late bid accepted); Enforced"),
                  rounds_line(3, None), rounds_line(4, "No deadline stated", due="none stated"),
                  rounds_line(5, "Passed without action; Unclear")]
        ledger = [row(1, "Target", "Deadline", rnd=1)] + [row(2 + i, "Target", "Deadline", rnd=2) for i in range(3)] \
            + [row(6, "Target", "Deadline", rnd=5), row(7, "Target", "Deadline", rnd=5)]
        out, manifest = self.derive(ledger, rounds)
        classes = [(d["round"], d["value"], d["class"]) for d in manifest["deadline_outcomes"]]
        self.assertEqual(classes, [
            (1, "Extended (late bid accepted)", "extended"), (2, "Extended", "extended"),
            (2, "Extended (late bid accepted)", "extended"), (2, "Enforced", "hard"),
            (3, "", "not reached"), (4, "No deadline stated", "no deadline"),
            (5, "Passed without action", "soft"), (5, "Unclear", "missing")])
        self.assertEqual(manifest["incomplete"], [])
        by_round = {r["round"]: r for r in read(out, "rounds.csv")}
        self.assertEqual(by_round["2"]["deadline_classes"], "extended; extended; hard")
        self.assertEqual(by_round["3"]["deadline_classes"], "not reached")
        self.assertFalse([r for r in manifest["review"] if "Deadline row" in r["issue"]])

    def test_unknown_deadline_value_is_listed_and_the_manifest_is_incomplete(self):
        out, manifest = self.derive([row(1, "Target", "Deadline")], [rounds_line(1, "Sort of enforced")])
        self.assertIn("deadline_outcomes.class", manifest["incomplete"])
        self.assertTrue(any("matches no label" in r["issue"] for r in manifest["review"]))

    def test_readings_prices_and_all_cash(self):
        ledger = [
            bid(1, "Alpha", 10, 12, rnd=1, formality="Formal", conditions="Light"),
            bid(2, "Beta", 11, rnd=1, formality="Formal", conditions="Heavy", **{"Stock %": 40}),
            bid(3, "Gamma", 11, rnd=2, formality="Formal", conditions="Unclear", **{"Stock %": "Part stock"}),
            bid(4, "Delta", 9, rnd=2, formality="Informal", conditions="None", **{"Stock %": "Not stated"}),
            bid(5, "Eps", 9, rnd=2, formality="Unclear", conditions="Light", **{"Stock %": "50-75"}),
            bid(6, "Zeta", None, rnd=2, formality="Formal", conditions="Light", **{"Stock %": "0-50", "Price high": None}),
            bid(7, "Eta", 20, rnd=2, formality="Formal", conditions="Light", **{"CVR/earnout": "Y", "CVR/earnout value": 2.5}),
            bid(8, "Theta", 8, rnd=0, formality="Formal", conditions="None"),
            bid(9, "Iota", 8, rnd="post", formality="Formal", conditions="None", **{"Stock %": "Varies"}),
            bid(10, "Kappa", 7, rnd=2, **{"CVR/earnout": "Y"}),
            bid(11, "Lambda", 8, rnd=3, formality="Formal", conditions="None"),
            bid(12, "Mu", 8, rnd=2, formality="Formal", conditions=None),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced"), rounds_line(2, "Enforced", finality="Inferred final")])
        bids = {r["who"]: r for r in read(out, "bids.csv")}
        readings = {w: [bids[w][t] for t in ("T0", "T1", "T1u", "T2", "T3")] for w in bids}
        self.assertEqual(readings["Alpha"], ["Formal", "Formal", "Formal", "Informal", "Informal"])
        self.assertEqual(readings["Beta"], ["Formal", "Informal", "Informal", "Formal", "Informal"])
        self.assertEqual(readings["Gamma"], ["Formal", "Formal", "Informal", "Formal", "Formal"])
        self.assertEqual(readings["Delta"], ["Informal"] * 5)
        self.assertEqual(readings["Eps"], [""] * 5)  # recorded Unclear is missing under every reading
        # §7.10: outside a reading's test a Formal bid is Informal; only recorded Unclear is missing.
        self.assertEqual(readings["Zeta"], ["Formal", "Formal", "Formal", "Informal", "Formal"])  # no stated price: no point price
        self.assertEqual(readings["Theta"][4], "Informal")  # round 0 precedes any organized stage
        self.assertEqual(readings["Iota"][4], "Informal")  # post is not a final round
        self.assertEqual(readings["Lambda"][4], "Informal")  # a round with no Rounds line is not known to be final
        self.assertEqual(readings["Mu"][1:3], ["Formal", "Informal"])  # blank Conditions: not Heavy, and not None or Light
        issues = " | ".join(r["issue"] for r in manifest["review"])
        self.assertIn("round 3, which has no Rounds line", issues)
        self.assertIn("Conditions is blank on a Formal bid", issues)
        # A stated range (not a cell value) still means some stock, so "0-50" is 0.
        self.assertEqual({w: bids[w]["all_cash"] for w in bids},
                         {"Alpha": "1", "Beta": "0", "Gamma": "0", "Delta": "", "Eps": "0", "Zeta": "0", "Eta": "1", "Theta": "1", "Iota": "",
                          "Kappa": "1", "Lambda": "1", "Mu": "1"})
        for stock, cash in ((0, 1), (40, 0), ("0-50", 0), ("Part stock", 0), ("Not stated", None), ("Varies", None)):
            with self.subTest(stock=stock):
                self.assertEqual(derive.all_cash({"Stock %": stock}), cash)
        self.assertEqual((bids["Eta"]["price_low"], bids["Eta"]["package_low"], bids["Eta"]["package_high"]), ("20", "22.5", "22.5"))
        self.assertEqual((bids["Kappa"]["price_low"], bids["Kappa"]["package_low"]), ("7", ""))  # a CVR with no value
        self.assertEqual(bids["Kappa"]["package_basis"], "missing: CVR/earnout marked without a value")
        self.assertEqual(bids["Alpha"]["package_basis"], "upfront only (no CVR/earnout)")
        self.assertEqual((bids["Eps"]["stock_kind"], bids["Eps"]["stock_lo"], bids["Eps"]["stock_hi"]), ("range (not a Version 1 value)", "50", "75"))
        self.assertIsNone(manifest["readings"]["default"])

    def test_same_price_revision_marker_and_its_two_variants(self):
        ledger = [bid(1, "Alpha", 10, day=0), bid(2, "Alpha", 10, day=1), bid(3, "Alpha", 11, day=2),
                  bid(4, "Alpha", 11, day=3, event="Bid reaffirmed", formality="Formal")]
        out, _ = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Alpha")])
        rows = read(out, "bids.csv")
        self.assertEqual([r["same_price_revision"] for r in rows], ["", "Y", "", ""])
        self.assertEqual([r["price_obs__same_price_as_terms"] for r in rows], ["1", "0", "1", "1"])
        self.assertEqual({r["price_obs__same_price_as_new"] for r in rows}, {"1"})

    def test_count_bounds(self):
        cases = [(3, "", (3, 3, 3, "exact")), (None, "Count: at least 11; x", (11, None, None, "at least")),
                 (None, "Count: approximately 20.", (1, None, 20, "approximately")), (None, "Count: 11–14.", (11, 14, None, "range")),
                 (None, "Count: fewer than 5", (1, 4, None, "fewer than")), (None, "Count: unknown", (1, None, None, "unknown")),
                 (None, "no prefix", (1, None, None, "unknown (no Count: prefix)"))]
        for count, note, expected in cases:
            with self.subTest(note=note):
                self.assertEqual(derive.count_bounds(count, note), expected)

    def test_participation_follows_the_live_units_formula_with_bounds(self):
        ledger = [
            row(1, "Alpha", "Bid", rnd=0, Type="Financial", Count=1, **{"Price low": 10, "Price high": 10, "Formality": "Informal"}),
            row(2, "Target", "Round opened", day=1),
            row(3, "12 financial NDA signers", "NDA signed", day=2, Type="Financial", Count=12),
            row(4, "Beta", "NDA signed", day=2, Type="Financial", Count=1),
            bid(5, "Gamma", 11, day=3),  # first seen bidding after a cohort: maybe one of the 12
            row(6, "Non-submitters", "Did not submit", day=4, Type="Financial", Inferred="Y",
                **{"Exit reason": "Not stated", "Note": "Count: approximately 9 of the 12."}),
            row(7, "Beta", "Withdrew", day=5, Type="Financial", Count=1, **{"Exit reason": "Value below earlier offer"}),
            row(8, "Beta", "Re-entered", day=6, Type="Financial", Count=1),
            row(9, "Alpha", "Dropped by target", day=7, Type="Financial", Count=1, Inferred="Y",
                **{"Exit reason": "Not stated", "Note": "Inferred exit: a live rival when the target signed exclusivity."}),
            row(10, "Gamma", "Merger agreement signed", day=8, Type="Financial", Count=1),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Gamma", who="14: many")])
        rows = read(out, "participation.csv")
        self.assertEqual([r["change"] for r in rows], ["entry", "round opening", "entry", "entry", "entry (membership uncertain)",
                                                      "exit", "exit", "re-entry", "exit", "win"])
        self.assertEqual([r["live_point"] for r in rows], ["1", "1", "13", "14", "14", "5", "4", "5", "4", "3"])
        self.assertEqual([r["live_hi"] for r in rows][:5], ["1", "1", "13", "14", "15"])
        self.assertEqual(rows[5]["live_lo"], "0")  # an unbounded exit leaves no lower bound
        self.assertEqual((rows[6]["alex_drop_code"], rows[6]["exit_actor"]), ("DropBelowInf", "bidder"))
        self.assertEqual((rows[8]["alex_drop_code"], rows[8]["exit__inferred_as_dropout"], rows[8]["exit__inferred_as_censoring"]),
                         ("DropTarget", "dropout", "censored"))
        self.assertEqual((rows[5]["exit_inferred"], rows[5]["exit__inferred_as_censoring"]), ("inferred exit", "censored"))
        issues = [r["issue"] for r in manifest["review"]]
        self.assertTrue(any("no recorded entry" in i for i in issues))
        self.assertTrue(any("ends with 3 live unit(s)" in i for i in issues))
        self.assertEqual(read(out, "rounds.csv")[0]["live_open_point"], "1")

    def test_an_exit_after_the_win_is_listed_and_not_subtracted_again(self):
        ledger = [row(1, "Alpha", "NDA signed", Type="Financial", Count=1), row(2, "Beta", "NDA signed", Type="Financial", Count=1),
                  bid(3, "Alpha", 10, day=1), row(4, "Alpha", "Merger agreement signed", day=2, Type="Financial", Count=1),
                  row(5, "Alpha", "Withdrew", day=3, Type="Financial", Count=1, **{"Exit reason": "Not stated"})]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Alpha")])
        rows = read(out, "participation.csv")
        self.assertEqual([(r["change"], r["live_point"]) for r in rows], [("entry", "1"), ("entry", "2"), ("win", "1")])
        self.assertTrue(any("an exit after this unit's signing" in r["issue"] and r["row"] == 5 for r in manifest["review"]))

    def test_the_deal_slug_is_read_from_cockpit_file_names(self):
        known = {"meredith", "mac-gray"}
        cases = [("meredith-opus55-medium-20260924-2241-abc123.xlsx", None, "meredith", True),
                 ("meredith-working.xlsx", None, "meredith", False), ("mac-gray-working-r8.xlsx", None, "mac-gray", False),
                 ("meredith-version1-raw-with-source.xlsx", {"Version ID": "version1-raw"}, "meredith", False),
                 ("zz-unknown.xlsx", None, "zz-unknown", True)]
        for name, source, slug, warned in cases:
            with self.subTest(name=name):
                found, warning = derive.deal_from_name(Path(name), source, known)
                self.assertEqual((found, warning is not None), (slug, warned))
        path = write(self.dir / "meredith-opus55-medium-20260924-2241-abc123.xlsx", [bid(1, "Alpha", 10)],
                     [rounds_line(1, "Enforced")], facts={**FACTS, "Auction screen": "Met (process 1): 2 parties"})
        manifest = derive.run(path, self.dir / "m")
        self.assertEqual(manifest["deal"], "meredith")
        self.assertIn("read from the file name", manifest["warnings"][0])
        self.assertEqual(read(self.dir / "m", "deal.csv")[0]["descriptive_only"], "1")

    def test_other_scope_rows_stay_out_and_a_partial_only_candidate_is_counted_as_a_bound(self):
        ledger = [
            row(1, "Alpha", "NDA signed", Type="Strategic", Count=1),
            bid(2, "Alpha", 10, day=1),
            row(3, "Alpha", "Other-scope bid", day=2, Type="Strategic", Count=1, Formality="Informal", Conditions="Unclear"),
            row(4, "SegmentCo", "NDA signed", day=2, Type="Strategic", Count=1),
            row(5, "SegmentCo", "Other-scope bid", day=3, Type="Strategic", Count=1, Formality="Informal", Conditions="Unclear"),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Alpha; SegmentCo bid for a segment")])
        self.assertEqual([(r["row"], r["reason"]) for r in read(out, "other_scope.csv")],
                         [("3", "Other-scope bid"), ("4", "partial-only candidate (also in participation.csv)"), ("5", "Other-scope bid")])
        self.assertEqual([r["who"] for r in read(out, "bids.csv")], ["Alpha"])
        # No exit row: partial from the start, or a switch whose exit is missing. The entry is kept, as a bound.
        rows = read(out, "participation.csv")
        self.assertEqual([(r["who"], r["change"], r["live_lo"], r["live_hi"], r["live_point"]) for r in rows],
                         [("Alpha", "entry", "1", "1", "1"), ("SegmentCo", "entry (scope uncertain)", "1", "2", "1")])
        issue = next(r["issue"] for r in manifest["review"] if r["who"] == "segmentco")
        self.assertIn("partial-only candidate", issue)
        self.assertIn("no exit row", issue)
        self.assertEqual(read(out, "rounds.csv")[0]["bids_received_check"], "ok")

    def test_a_partial_only_candidate_does_not_initiate_the_process(self):
        # SegCo's interest comes first, but with no exit row it may never have been in the whole-company contest.
        ledger = [row(1, "SegCo", "Bidder interest", rnd=0, Type="Strategic", Count=1),
                  row(2, "Target", "Target sale decision", rnd=0, day=1),
                  row(3, "SegCo", "Other-scope bid", day=2, Type="Strategic", Count=1, Formality="Informal", Conditions="Unclear")]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="none")])
        deal = read(out, "deal.csv")[0]
        self.assertEqual((deal["initiation_first_event"], deal["initiation_first_row"]), ("target-led", "#2 Target sale decision"))
        self.assertTrue(any(r["row"] == 1 and "first initiating row belongs to a partial-only candidate" in r["issue"]
                            for r in manifest["review"]))

    def test_whole_company_entry_then_a_first_priced_offer_that_is_partial(self):
        # E1: a bidder that entered the whole-company contest and switches to a partial offer leaves it at the switch.
        # Its every bid row is Other-scope, yet its entry and exit are whole-company participation.
        ledger = [
            row(1, "Alpha", "NDA signed", Type="Strategic", Count=1),
            row(2, "SwitchCo", "NDA signed", Type="Strategic", Count=1),
            row(3, "SwitchCo", "Withdrew", day=4, Type="Strategic", Count=1,
                **{"Exit reason": "Other stated reason", "Note": "Stopped pursuing the whole company; talks continued on a partial basis."}),
            row(4, "SwitchCo", "Other-scope bid", day=5, Type="Strategic", Count=1, Formality="Informal", Conditions="Unclear"),
            bid(5, "Alpha", 10, day=6),
            row(6, "Alpha", "Merger agreement signed", day=8, Type="Strategic", Count=1),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Alpha; SwitchCo for a division")])
        rows = read(out, "participation.csv")
        self.assertEqual([(r["row"], r["who"], r["change"], r["live_point"]) for r in rows],
                         [("1", "Alpha", "entry", "1"), ("2", "SwitchCo", "entry", "2"), ("3", "SwitchCo", "exit", "1"), ("6", "Alpha", "win", "0")])
        self.assertEqual([r["live_lo"] for r in rows], ["1", "2", "1", "0"])
        self.assertEqual([(r["row"], r["reason"]) for r in read(out, "other_scope.csv")],
                         [("2", "partial-only candidate (also in participation.csv)"), ("3", "partial-only candidate (also in participation.csv)"),
                          ("4", "Other-scope bid")])
        issue = next(r["issue"] for r in manifest["review"] if r["who"] == "switchco")
        self.assertIn("switch to a partial offer", issue)
        self.assertNotIn("follows its first partial offer", issue)
        self.assertFalse([r for r in manifest["review"] if "ends with" in r["issue"]])
        # The same party whose exit row comes after its partial offer: still counted until the exit, and listed.
        late = [ledger[0], ledger[1], {**ledger[3], "#": 3}, {**ledger[2], "#": 4, "Sort date": D + dt.timedelta(days=6)}, *ledger[4:]]
        out, manifest = self.derive(late, [rounds_line(1, "Enforced", received="1: Alpha")], name="late.xlsx")
        self.assertEqual([r["change"] for r in read(out, "participation.csv")], ["entry", "entry", "exit", "win"])
        self.assertIn("follows its first partial offer #3", next(r["issue"] for r in manifest["review"] if r["who"] == "switchco"))



    def test_package_basis_is_the_stated_per_share_amount(self):
        def cvr(n, who, value, note, marker="Y"):
            return bid(n, who, 20, formality="Formal", conditions="Light", Note=note, **{"CVR/earnout": marker, "CVR/earnout value": value})
        ledger = [
            bid(1, "Plain", 20),
            cvr(2, "Maximum", 3, "Plus a CVR of up to $3.00 per share on FDA approval."),
            cvr(3, "Valued", 1.5, "Package $21.50 = $20.00 cash + $1.50 CVR (Valued Co's valuation)."),
            cvr(4, "Cohort", None, "Two members offered CVRs of $1 and $2.", marker="Varies"),
            cvr(5, "Unmarked", 1, "A $1.00 contingent payment.", marker=None),
        ]
        # E13 fixes the basis (the stated per-share amount, the maximum where several): the Note is not read.
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="5: all")])
        fixed = "upfront + CVR/earnout value (the stated per-share amount, the maximum where several; E13)"
        self.assertEqual({w: (b["package_low"], b["package_basis"]) for w, b in ((r["who"], r) for r in read(out, "bids.csv"))}, {
            "Plain": ("20", "upfront only (no CVR/earnout)"), "Maximum": ("23", fixed), "Valued": ("21.5", fixed),
            "Cohort": ("", "missing: CVR/earnout Varies on a cohort row (amounts in the Note)"),
            "Unmarked": ("21", "upfront + CVR/earnout value, with CVR/earnout not marked Y")})
        self.assertEqual({r["row"] for r in manifest["review"] if "CVR/earnout" in r["issue"]}, {5})

    def test_t2_needs_two_present_valid_equal_prices(self):
        ledger = [
            bid(1, "Point", 10, formality="Formal", conditions="Light"),
            # An unsplittable package (E13): both price cells blank, the package in the Note.
            bid(2, "Package", None, formality="Formal", conditions="Light",
                **{"Price high": None, "Note": "Package valued at $25.00 including a CVR; the parts cannot be separated."}),
            bid(3, "OneSided", 10, formality="Formal", conditions="Light", **{"Price high": None}),
            bid(4, "Text", "10", formality="Formal", conditions="Light", **{"Price high": "10"}),
            bid(5, "Range", 10, 12, formality="Formal", conditions="Light"),
            bid(6, "Zero", 0, formality="Formal", conditions="Light"),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="6: all")])
        bids = {r["who"]: r for r in read(out, "bids.csv")}
        self.assertEqual({w: (b["T0"], b["T2"]) for w, b in bids.items()},
                         {"Point": ("Formal", "Formal"), "Package": ("Formal", "Informal"), "OneSided": ("Formal", "Informal"),
                          "Text": ("Formal", "Informal"), "Range": ("Formal", "Informal"), "Zero": ("Formal", "Informal")})
        self.assertEqual((bids["Package"]["package_low"], bids["Package"]["package_basis"]), ("", "missing: no upfront price"))
        # An invalid price cell is shown as read, but gives no package; the review item says so.
        for who in ("Text", "Zero"):
            self.assertEqual((bids[who]["package_low"], bids[who]["package_high"], bids[who]["package_basis"]),
                             ("", "", "missing: a price cell is not a valid price"))
        self.assertEqual(bids["Text"]["price_low"], "10")
        self.assertEqual(sorted(r["row"] for r in manifest["review"] if "not a positive number" in r["issue"]), [4, 4, 6, 6])
        self.assertEqual(bids["OneSided"]["package_low"], "10")

    def test_rounds_sheet_disagreements_go_to_the_review_list(self):
        ledger = [bid(1, "Alpha", 10), bid(2, "Beta", 9), row(3, "Target", "Deadline")]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="3: Alpha, Beta, Gamma"),
                                             rounds_line(2, "Enforced", received="none")])
        checks = {r["round"]: r["bids_received_check"] for r in read(out, "rounds.csv")}
        self.assertEqual(checks, {"1": "mismatch", "2": "not compared"})
        issues = " | ".join(r["issue"] for r in manifest["review"])
        self.assertIn("Bids received says 3", issues)
        self.assertIn("round 2: 1 deadline outcome(s) but 0 Deadline row(s)", issues)

    def test_a_workbook_without_the_current_ledger_header_is_an_error(self):
        path = write(self.dir / "old.xlsx", [bid(1, "Alpha", 10)], [rounds_line(1, "Enforced")])
        wb = load_workbook(path)
        wb["Deal ledger"].delete_cols(check_lean.LEDGER_COLUMNS.index("Stock %") + 1, 10)
        wb.save(path)
        with self.assertRaisesRegex(derive.DeriveError, "not the Version 1 ledger"):
            derive.load(path)
        self.assertEqual(quiet(derive.main, [str(path), "--out", str(self.dir / "old-out")]), 2)
        broken = self.dir / "broken.xlsx"
        broken.write_text("not an xlsx", encoding="utf-8")
        with self.assertRaisesRegex(derive.DeriveError, "not a readable workbook"):
            derive.load(broken)

    def test_five_sheet_download_source_is_provenance_and_every_output_is_written(self):
        source = {"EDGAR filing index": "https://www.sec.gov/x-index.htm", "Working revision": 3, "Review status": "not set"}
        out, manifest = self.derive([bid(1, "Alpha", 10)], [rounds_line(1, "Enforced", received="1: Alpha")], source=source)
        self.assertEqual(manifest["input"]["kind"], "five-sheet cockpit download")
        self.assertEqual(manifest["input"]["source_sheet"]["Working revision"], 3)
        self.assertEqual(manifest["input"]["source_sheet"]["EDGAR filing index (link)"], "https://www.sec.gov/x-index.htm")
        self.assertEqual(sorted(p.name for p in out.iterdir()),
                         ["bids.csv", "deal.csv", "manifest.json", "other_scope.csv", "participation.csv", "rounds.csv"])
        self.assertEqual(derive.manifest_complete(json.loads((out / "manifest.json").read_text())), [])
        self.assertEqual(manifest["ledger_schema"], "Version 1")
        self.assertEqual({s["id"] for s in manifest["switches"]},
                         {"count_ranges", "unclear", "formality_reading", "same_price_revisions", "same_offer_restatements",
                          "inferred_exits", "estimation_price"})
        deal = read(out, "deal.csv")[0]
        self.assertEqual((deal["auction_status"], deal["auction_count_lo"], deal["auction_count_hi"], deal["estimation_sample"]),
                         ("Met", "3", "", "1"))

    def test_auction_screen_and_sample(self):
        screen = derive.auction_screen("Met (process 1): 3 parties; Not met (process 3): 1 party")
        self.assertEqual({p: (s["status"], s["lo"], s["hi"]) for p, s in screen.items()}, {1: ("Met", 3, 3), 3: ("Not met", 1, 1)})
        ledger = [bid(1, "Alpha", 10)]
        facts = {**FACTS, "Auction screen": "Met (process 1): 2 parties", "Whole-company bids": "Yes"}
        path = write(self.dir / "meredith.xlsx", ledger, [rounds_line(1, "Enforced")], facts=facts)
        derive.run(path, self.dir / "m")
        deal = read(self.dir / "m", "deal.csv")[0]
        self.assertEqual((deal["descriptive_only"], deal["estimation_sample"]), ("1", "0"))

    def test_out_folder_must_be_new_or_empty_and_outside_the_data(self):
        path = write(self.dir / "alpha.xlsx", [bid(1, "Alpha", 10)], [rounds_line(1, "Enforced")])
        full = self.dir / "full"
        full.mkdir()
        (full / "x").write_text("keep")
        self.assertEqual(quiet(derive.main, [str(path), "--out", str(full)]), 2)
        self.assertEqual((full / "x").read_text(), "keep")
        for forbidden in ("extraction/new", "raw_filing/new", "ref/new"):
            with self.subTest(forbidden=forbidden), self.assertRaises(derive.DeriveError):
                derive.check_out(derive.PROJECT / forbidden)
        self.assertEqual(quiet(derive.main, [str(path), "--out", str(self.dir / "new")]), 0)

    # ---- P1, same_offer_of and the E-rule readings -------------------------------------------------------------------

    def test_the_manifest_is_labelled_version_1_and_there_is_no_rules_option(self):
        path = write(self.dir / "alpha.xlsx", [bid(1, "Alpha", 10)], [rounds_line(1, "Enforced", received="1: Alpha")])
        self.assertEqual(derive.load(path)["schema"], "Version 1")
        manifest = derive.run(path, self.dir / "out")
        self.assertEqual((manifest["ledger_schema"], manifest["tool_version"], manifest["checker_version"]), ("Version 1", "Version 1", "Version 1"))
        self.assertNotIn("contract", manifest)
        with self.assertRaises(SystemExit):
            quiet(derive.parse_args, [str(path), "--out", str(self.dir / "bad"), "--rules", "legacy"])

    def test_upfront_price_kind_and_the_price_observation_flags(self):
        # P1: only valid numeric price information is a price observation.
        ledger = [
            bid(1, "Point", 10), bid(2, "Range", 10, 12), bid(3, "Lower", 10, **{"Price high": None}),
            bid(4, "Upper", None, 12), bid(5, "Blank", None, **{"Price high": None}),
            bid(6, "Text", "10", **{"Price high": "10"}), bid(7, "Zero", 0), bid(8, "Negative", -5),
            bid(9, "Reversed", 12, 10),
        ]
        kinds = {"Point": "point", "Range": "range", "Lower": "lower_bound", "Upper": "upper_bound", "Blank": "not_available",
                 "Text": "invalid", "Zero": "invalid", "Negative": "invalid", "Reversed": "invalid"}
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="9: all")])
        bids = {r["who"]: r for r in read(out, "bids.csv")}
        self.assertEqual({w: b["upfront_price_kind"] for w, b in bids.items()}, kinds)
        usable = {w for w, k in kinds.items() if k not in ("not_available", "invalid")}
        for who, b in bids.items():
            flags = (b["price_obs__same_price_as_new"], b["price_obs__same_price_as_terms"])
            self.assertEqual(flags, ("1", "1") if who in usable else ("0", "0"), who)
        reviewed = {r["row"] for r in manifest["review"] if "not a positive number" in r["issue"] or "reversed range" in r["issue"]}
        self.assertEqual(reviewed, {6, 7, 8, 9})  # every invalid kind is a review item
        self.assertEqual((bids["Reversed"]["package_low"], bids["Reversed"]["package_basis"]),
                         ("", "missing: the price range is reversed"))
        self.assertEqual(derive.upfront_price_kind({"Price low": 10.0, "Price high": 10}), "point")
        self.assertEqual(derive.upfront_price_kind({"Price low": True, "Price high": None}), "invalid")

    def test_a_blank_price_commitment_bid_is_kept_but_is_no_price_observation(self):
        # P1 and H4: the event and its terms stay; the price is not available, which is not "no economic price".
        ledger = [
            bid(1, "Alpha", 10, day=0), bid(2, "Alpha", 10, day=1),
            bid(3, "Alpha", None, day=2, formality="Formal", conditions="Heavy", Financing="Contingent",
                **{"Price high": None, "Note": "H1: financing not committed. Sponsor liability cap $50m proposed."}),
            bid(4, "Alpha", None, day=3, **{"Price high": None, "Stock %": "Not stated",
                                            "Note": "Package valued at $25.00 including a CVR; the parts cannot be separated.",
                                            "CVR/earnout": "Y"}),
            bid(5, "Alpha", 10, day=4),
            bid(6, "Alpha", 10, day=5, event="Bid reaffirmed", formality="Formal"),
        ]
        out, _ = self.derive(ledger, [rounds_line(1, "Enforced", received="1: Alpha")])
        rows = read(out, "bids.csv")
        self.assertEqual([r["row"] for r in rows], ["1", "2", "3", "4", "5", "6"])  # every bid event is kept
        self.assertEqual([r["upfront_price_kind"] for r in rows],
                         ["point", "point", "not_available", "not_available", "point", "point"])
        # The documented comparison: the previous row of the unit, whatever it holds. A blank-price row in
        # between breaks it (no unseen price continuity), and Bid reaffirmed is never marked, only told apart by Event.
        self.assertEqual([r["same_price_revision"] for r in rows], ["", "Y", "", "", "", ""])
        self.assertEqual([r["price_obs__same_price_as_new"] for r in rows], ["1", "1", "0", "0", "1", "1"])
        self.assertEqual([r["price_obs__same_price_as_terms"] for r in rows], ["1", "0", "0", "0", "1", "1"])
        self.assertEqual((rows[2]["conditions"], rows[2]["financing"], rows[2]["formality"]), ("Heavy", "Contingent", "Formal"))
        self.assertEqual((rows[3]["cvr_earnout"], rows[3]["package_low"], rows[3]["package_basis"]),
                         ("Y", "", "missing: no upfront price"))
        self.assertIn("Package valued at $25.00", rows[3]["note"])
        self.assertEqual(rows[5]["event"], "Bid reaffirmed")

    def test_the_instructions_examples_give_same_offer_of_and_no_price_observation_for_priceless_rows(self):
        workbook, _ = build_examples_fixture(self.dir / "examples")
        original = workbook.read_bytes()
        out = self.dir / "examples-out"
        manifest = derive.run(workbook, out, "examples")
        self.assertEqual(manifest["ledger_schema"], "Version 1")
        bids = {r["row"]: r for r in read(out, "bids.csv")}
        # Example 2: "Same as #5" copies #5's price; it is a restatement, not a new same-price revision.
        self.assertEqual((bids["14"]["same_offer_of"], bids["14"]["same_price_revision"]), ("5", ""))
        self.assertEqual((bids["14"]["price_low"], bids["14"]["upfront_price_kind"]), ("30", "point"))
        self.assertEqual({r for r, b in bids.items() if b["same_offer_of"]}, {"14"})
        # Party G's priceless proposal (row 15) and Example 5's liability cap (row 18) are kept, but are no price observations.
        for number in ("15", "18"):
            self.assertEqual((bids[number]["upfront_price_kind"], bids[number]["price_obs__same_price_as_new"],
                              bids[number]["price_obs__same_price_as_terms"]), ("not_available", "0", "0"))
        self.assertFalse([r for r in manifest["review"] if "Same as" in r["issue"] or "Same-offer" in r["issue"]])
        self.assertEqual(workbook.read_bytes(), original)  # the input is never changed
        deal = read(out, "deal.csv")[0]
        self.assertEqual((deal["initiation_first_event"], deal["initiation_check"]), ("target-led", "agrees"))
        # Without the prefix the same row repeats #5's price and would be a same-price revision.
        wb = load_workbook(workbook)
        ws = wb["Deal ledger"]
        ws.cell(15, check_lean.LEDGER_COLUMNS.index("Note") + 1).value = "H1: financing not committed."
        wb.save(self.dir / "unprefixed.xlsx")
        derive.run(self.dir / "unprefixed.xlsx", self.dir / "unprefixed", "examples")
        row14 = next(r for r in read(self.dir / "unprefixed", "bids.csv") if r["row"] == "14")
        self.assertEqual((row14["same_offer_of"], row14["same_price_revision"], row14["price_obs__same_price_as_terms"]), ("", "Y", "0"))

    def test_a_same_offer_row_that_points_elsewhere_is_listed(self):
        ledger = [bid(1, "Alpha", 10), bid(2, "Beta", 11), bid(3, "Alpha", 10, day=1, Note="Same as #2."),
                  bid(4, "Beta", 12, day=1, Note="Same as #2.")]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="2: Alpha, Beta")])
        self.assertEqual([r["same_offer_of"] for r in read(out, "bids.csv")], ["", "", "2", "2"])
        issues = {r["row"]: r["issue"] for r in manifest["review"] if "Same" in r["issue"]}
        self.assertIn("does not point to an earlier whole-company bid row of this bidder", issues[3])
        self.assertIn("prices differ from #2", issues[4])

    def test_inferred_exits_are_read_from_inferred_alone(self):
        ledger = [row(1, "Beta", "NDA signed", Type="Financial", Count=1),
                  row(2, "Beta", "Withdrew", day=3, Type="Financial", Count=1, Inferred="Y",
                      **{"Exit reason": "Terms or process", "Note": "Date inferred: the filing reports the withdrawal without a day."}),
                  row(3, "Gamma", "NDA signed", Type="Financial", Count=1),
                  row(4, "Gamma", "Did not submit", day=4, Type="Financial", Count=1, Inferred="Y", When="by 01/05/2020",
                      **{"Exit reason": "Not stated"}),
                  row(5, "Eps", "NDA signed", Type="Financial", Count=1),
                  row(6, "Eps", "Withdrew", day=5, Type="Financial", Count=1, **{"Exit reason": "Not stated"})]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="none")])
        found = {r["who"]: (r["exit_inferred"], r["exit__inferred_as_dropout"], r["exit__inferred_as_censoring"])
                 for r in read(out, "participation.csv") if r["change"] == "exit"}
        self.assertEqual(found, {"Beta": ("inferred exit", "dropout", "censored"), "Gamma": ("inferred exit", "dropout", "censored"),
                                 "Eps": ("", "dropout", "dropout")})
        self.assertFalse([r for r in manifest["review"] if "Inferred = Y" in r["issue"]])

    def test_initiation_follows_d5_and_is_checked_against_deal_facts(self):
        activist_first = [row(1, "Alpha", "Bidder interest", rnd=0, Type="Financial", Count=1),
                          row(2, "Fund X", "Activist", rnd=0, day=1, Note="Demands sale: seek a buyer."), row(3, "Target", "Round opened", day=2)]
        activist_late = [row(1, "Target", "Target sale decision", rnd=0), row(2, "Target", "Round opened", day=1),
                         row(3, "Fund X", "Activist", day=2)]
        bidder_first = [row(1, "Alpha", "Bid", rnd=0, Type="Financial", Count=1, **{"Price low": 10, "Price high": 10}),
                        row(2, "Target", "Target sale decision", rnd=0, day=1), row(3, "Target", "Round opened", day=2)]
        cases = [(activist_first, "activist-influenced", "#2 Activist"),  # a sale demand before the target's first step wins
                 (activist_late, "target-led", "#1 Target sale decision"),  # a later Activist row does not count
                 (bidder_first, "mixed", "#1 Bid")]
        for i, (ledger, derived, first_row) in enumerate(cases):
            with self.subTest(case=i):
                out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="none")], name=f"init{i}.xlsx")
                deal = read(out, "deal.csv")[0]
                self.assertEqual((deal["initiation_first_event"], deal["initiation_first_row"]), (derived, first_row))
                self.assertEqual(deal["initiation_check"], "agrees" if derived == "target-led" else "differs")
                issues = [r for r in manifest["review"] if r["sheet"] == "Deal facts" and "Initiation" in r["issue"]]
                self.assertEqual(len(issues), 0 if derived == "target-led" else 1)
        out, _ = self.derive(activist_late, [rounds_line(1, "Enforced", received="none")], name="unrecorded.xlsx",
                             facts={k: v for k, v in FACTS.items() if k != "Initiation"})
        self.assertEqual(read(out, "deal.csv")[0]["initiation_check"], "not recorded")

    def test_sale_one_option_does_not_override_mixed_initiation(self):
        ledger = [row(1, "Fund X", "Activist", rnd=0, Note="Sale one option: explore alternatives."),
                  row(2, "Target", "Target interest", rnd=0, day=1),
                  bid(3, "Alpha", 10, rnd=0, day=2), row(4, "Target", "Round opened", day=3)]
        out, _ = self.derive(ledger, [rounds_line(1, "Enforced")], facts={**FACTS, "Initiation": "mixed"})
        deal = read(out, "deal.csv")[0]
        self.assertEqual((deal["initiation_first_event"], deal["initiation_check"]), ("mixed", "agrees"))

    def test_signing_outside_contest_does_not_subtract_or_win(self):
        ledger = [row(1, "Alpha", "NDA signed", Type="Financial", Count=1),
                  row(2, "Outside buyer", "Merger agreement signed", day=1)]
        out, _ = self.derive(ledger, [rounds_line(1, "No deadline stated")])
        participation = read(out, "participation.csv")
        self.assertFalse([r for r in participation if r["change"] == "win"])
        self.assertEqual(participation[-1]["live_point"], "1")

    def test_opening_live_is_after_same_day_exits_in_previous_round(self):
        ledger = [row(1, "Alpha", "NDA signed", Count=1), row(2, "Beta", "NDA signed", Count=1),
                  row(3, "Target", "Round opened", rnd=2, day=1),
                  row(4, "Gamma", "NDA signed", rnd=2, day=1, Count=1),
                  row(5, "Beta", "Dropped by target", rnd=1, day=1, Count=1,
                      **{"Exit reason": "Not stated"})]
        out, _ = self.derive(ledger, [rounds_line(1, "No deadline stated"), rounds_line(2, "No deadline stated")])
        rounds = {r["round"]: r for r in read(out, "rounds.csv")}
        self.assertEqual(rounds["2"]["live_open_point"], "1")

    def test_opening_live_carries_pending_exits_from_earlier_same_day_openings(self):
        ledger = [row(1, "Alpha", "NDA signed", Count=1), row(2, "Beta", "NDA signed", Count=1),
                  row(3, "Gamma", "NDA signed", Count=1),
                  row(4, "Target", "Round opened", rnd=2, day=1),
                  row(5, "Target", "Round opened", rnd=3, day=1),
                  row(6, "Gamma", "Dropped by target", rnd=2, day=1, Count=1,
                      **{"Exit reason": "Not stated"}),
                  row(7, "Beta", "Dropped by target", rnd=1, day=1, Count=1,
                      **{"Exit reason": "Not stated"})]
        out, _ = self.derive(ledger, [rounds_line(1, "No deadline stated"), rounds_line(2, "No deadline stated"),
                                      rounds_line(3, "No deadline stated")])
        rounds = {r["round"]: r for r in read(out, "rounds.csv")}
        self.assertEqual((rounds["2"]["live_open_point"], rounds["3"]["live_open_point"]), ("2", "1"))

    def test_confirmation_by_documents_is_a_same_offer_restatement(self):
        ledger = [bid(1, "Alpha", 10),
                  bid(2, "Alpha", None, day=1, Note="Sponsor liability cap reduced."),
                  bid(3, "Alpha", 10, day=2, event="Bid reaffirmed", formality="Formal",
                      Note="Same as #1. Revised markup submitted.")]
        out, manifest = self.derive(ledger, [rounds_line(1, "No deadline stated")])
        bids = read(out, "bids.csv")
        self.assertEqual((bids[2]["same_offer_of"], bids[2]["same_price_revision"]), ("1", ""))
        self.assertEqual(bids[2]["price_obs__same_price_as_terms"], "1")
        self.assertEqual([b["row"] for b in bids if not b["same_offer_of"]], ["1", "2"])
        self.assertFalse([r for r in manifest["review"] if "Same as #1" in r["issue"]])

    def test_auction_screen_has_no_uncertain(self):
        facts = {**FACTS, "Auction screen": "Uncertain (process 1): count unknown"}
        out, manifest = self.derive([bid(1, "Alpha", 10)], [rounds_line(1, "Enforced", received="1: Alpha")], facts=facts)
        deal = read(out, "deal.csv")[0]
        self.assertEqual((deal["auction_status"], deal["estimation_sample"]), ("", ""))  # unsupported entry: unread, and listed
        self.assertIn("process 1: no Auction screen entry could be parsed", manifest["warnings"])
        self.assertEqual(derive.auction_screen("Met (process 1): 3 parties; Not met (process 2): 1 party"),
                         {1: {"status": "Met", "lo": 3, "hi": 3, "text": "Met (process 1): 3 parties"},
                          2: {"status": "Not met", "lo": 1, "hi": 1, "text": "Not met (process 2): 1 party"}})

    def test_who_was_in_count_and_merger_of_equals_rows(self):
        rounds = [rounds_line(1, "Enforced", received="1: Alpha", who="2: Alpha, Beta; Gamma still being received, not admitted")]
        ledger = [bid(1, "Alpha", 10), row(2, "Target", "Other material event", Note="Merger of equals talks with Delta ended.")]
        questions = [{"Q": "Q1", "Question": "Was the merger of equals with Delta a sale attempt?"}]
        out, manifest = self.derive(ledger, rounds, questions=questions)
        line = read(out, "rounds.csv")[0]
        self.assertEqual(line["who_was_in_count"], "2")  # the stage's own count is read
        self.assertNotIn("not_admitted", line)
        self.assertEqual(read(out, "deal.csv")[0]["merger_of_equals_rows"], "#2")  # rows only; Questions are not read
        self.assertFalse(any("merger-of-equals" in w for w in manifest["warnings"]))
        self.assertFalse({"eligible_unadmitted", "process_initiator", "merger_of_equals"} & {s["id"] for s in manifest["switches"]})

    def test_stock_ranges_and_count_ranges(self):
        ledger = [
            row(1, "Target", "Round opened"),
            bid(2, "Alpha", 10, **{"Stock %": "Part stock", "Note": "Stock 40–60% by election; exchange ratio 0.5."}),
            bid(3, "Beta", 10, **{"Stock %": "50-75"}),
            bid(4, "Gamma", 10, **{"Stock %": "Part stock", "Note": "Exchange ratio 0.5."}),
            row(5, "Others", "NDA signed", Type="Unknown", Note="Count: 11–14."),
        ]
        out, manifest = self.derive(ledger, [rounds_line(1, "Enforced", received="3: all")])
        stock = {r["who"]: (r["stock_kind"], r["stock_lo"], r["stock_hi"], r["all_cash"]) for r in read(out, "bids.csv")}
        self.assertEqual(stock, {"Alpha": ("part stock (range in the Note)", "40", "60", "0"),
                                 "Beta": ("range (not a Version 1 value)", "50", "75", "0"),
                                 "Gamma": ("part stock", "", "", "0")})
        issues = {r["row"]: r["issue"] for r in manifest["review"]}
        self.assertIn("E13 records Part stock", issues[3])
        self.assertIn("numeric Count range", issues[5])
        # The range is read as bounds; a qualifier is read as before.
        cohort = next(r for r in read(out, "participation.csv") if r["who"] == "Others")
        self.assertEqual((cohort["count_kind"], cohort["count_lo"], cohort["count_hi"]), ("range", "11", "14"))
        self.assertEqual(derive.count_bounds(None, "Count: more than ten; 12 named."), (1, None, None, "unknown (unparsed Count: prefix)"))
        self.assertEqual(derive.count_bounds(None, "Count: more than 10."), (11, None, None, "more than"))


if __name__ == "__main__":
    unittest.main()
