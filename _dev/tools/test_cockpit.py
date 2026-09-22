#!/usr/bin/env python3
"""Synthetic-fixture tests for the read-only review cockpit (no real filings)."""

from __future__ import annotations

import datetime as dt
import hashlib
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

from openpyxl import Workbook, load_workbook

import check_lean
from cockpit import data, server

SPLIT_QUOTE = "On March 1, 2020, Alpha submitted a proposal of $10 per share"
CROSS_QUOTE = "The board met on March 2, 2020. It rejected the proposal"
SPACED_QUOTE = "Beta Capital signed a confidentiality agreement"
WRONG_PAGE_QUOTE = "On March 9, 2020, Gamma withdrew"
MISSING_QUOTE = "This sentence is nowhere in the filing"
NBSP_QUOTE = "Delta Partners said “no” to the offer"
ROW_QUOTE = "Epsilon Group $12.00"
REPEATED_QUOTE = "The special committee reconvened by telephone"


def narrative_page(page: int) -> str:
    """Filler text plus special passages for the pages the tests quote."""

    body = [
        "<h5><a href='#toc'>Table of Contents</a></h5>",
        f"<p>Filler paragraph on page {page} describing ordinary business matters at length.</p>",
    ]
    if page == 1:
        body.append(
            "<table><tr><td>The Merger</td><td>2</td></tr>"
            "<tr><td>Background of the Merger</td><td>3</td></tr></table>"
        )
    if page == 3:
        body.append("<p><b>Background of the Merger</b></p>")
        body.append(
            "<p>On March 1, 2020, Alpha sub<font>mitted</font> a proposal of $<b>10</b> per share.</p>"
        )
        body.append("<p>The board met on March 2, 2020.</p><p>It rejected the proposal and asked for more.</p>")
        body.append("<p><font>Beta </font><font>Capital</font> signed a confidentiality agreement.</p>")
        body.append("<p>Delta&nbsp;Partners said &#8220;no&#8221; to the offer.</p>")
        body.append("<table><tr><td>Bidder</td><td>Epsilon Group</td><td>$12.00</td></tr></table>")
        body.append(f"<p>{REPEATED_QUOTE}.</p>")
        body.append("<p>" + " ".join(["The parties continued their discussions."] * 20) + "</p>")
    if page == 5:
        body.append(f"<p>{REPEATED_QUOTE} again.</p>")
    if page == 4:
        body.append("<p>On March 9, 2020, Gamma withdrew from the process.</p>")
        body.append("<p><b>Opinion of the Financial Adviser</b></p>")
    return "".join(body)


def build_filing(pages: int = 12, numbered: bool = True) -> str:
    parts = ["<html><head><style>p {}</style><script>var x = 1;</script></head><body>"]
    parts.append("<p>COVER PAGE</p><p>Proxy statement cover without a folio</p>")
    for page in range(1, pages + 1):
        parts.append("<p style='page-break-before:always'><hr></p>")
        parts.append(narrative_page(page))
        if numbered:
            parts.append(f"<p align='center'><font size='2'>{page}</font></p>")
    parts.append("</body></html>")
    return "".join(parts)


LEDGER_HEADER = list(check_lean.LEDGER_COLUMNS)


def ledger_row(number: int, quote: str | None, page: str, flag: str | None = None) -> list:
    row = [None] * len(LEDGER_HEADER)
    values = {
        "#": number,
        "When": "03/01/2020",
        "Who": "Alpha",
        "Type": "Financial",
        "Event": "Bid",
        "Process": 1,
        "Round": 1,
        "Price low": 10.0,
        "Price high": 10.5,
        "Quote and page": None if quote is None else f"“{quote}” (p. {page})",
        "Flag": flag,
        "Sort date": dt.datetime(2020, 3, number),
    }
    for column, value in values.items():
        row[LEDGER_HEADER.index(column)] = value
    return row


def build_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal ledger"
    ws.append(LEDGER_HEADER)
    ws.append(ledger_row(1, SPLIT_QUOTE, "3", "Q1"))
    ws.append(ledger_row(2, CROSS_QUOTE, "3"))
    ws.append(ledger_row(3, SPACED_QUOTE, "3"))
    ws.append(ledger_row(4, WRONG_PAGE_QUOTE, "5"))
    ws.append(ledger_row(5, MISSING_QUOTE, "3"))
    ws.append([None] * len(LEDGER_HEADER))  # blank row is skipped
    ws.append(ledger_row(7, None, "3"))
    rounds = wb.create_sheet("Rounds")
    rounds.append(list(check_lean.ROUND_COLUMNS))
    rounds.append([1, 1, dt.datetime(2020, 3, 1), "Outreach", "Alpha", "none", "No deadline stated", "Not final", "1", "Signed"])
    questions = wb.create_sheet("Questions")
    questions.append(list(check_lean.QUESTION_COLUMNS))
    questions.append(["Q1", "Is Alpha's bid formal?", "No", "Informal (p. 3)", "#1, #2", "Formality flips", None])
    facts = wb.create_sheet("Deal facts")
    facts.append(["Field", "Value"])
    facts.append(["Target", "Target Co."])
    facts.append(["Acquirer", "Alpha"])
    wb.save(path)


def build_repo(root: Path, slug: str = "alpha-deal", filing_html: str | None = None) -> Path:
    (root / "extraction" / "archive").mkdir(parents=True)
    (root / "raw_filing").mkdir()
    build_workbook(root / "extraction" / f"{slug}.xlsx")
    build_workbook(root / "extraction" / "archive" / "old.xlsx")  # ignored
    filing_name = f"{slug}_2020-04-01_DEFM14A.htm"
    (root / "raw_filing" / filing_name).write_text(filing_html or build_filing(), encoding="utf-8")
    (root / "raw_filing" / "MANIFEST.csv").write_text(
        "file,deal,form_type,date_filed,source_url,document,fetched_utc,bytes,sha256\n"
        f"{filing_name},{slug},DEFM14A,2020-04-01,https://example.invalid/x.txt,x.htm,,,\n"
        "missing.htm,ghost,DEFM14A,2020-01-01,,,,,\n",
        encoding="utf-8",
    )
    return root


def add_corrupt_deal(root: Path, slug: str) -> None:
    """A workbook that is not a zip file, beside a readable filing."""

    (root / "extraction" / f"{slug}.xlsx").write_text("not a zip", encoding="utf-8")
    filing_name = f"{slug}_2020-04-01_DEFM14A.htm"
    (root / "raw_filing" / filing_name).write_text(build_filing(), encoding="utf-8")
    with (root / "raw_filing" / "MANIFEST.csv").open("a", encoding="utf-8") as handle:
        handle.write(f"{filing_name},{slug},DEFM14A,2020-04-01,,,,,\n")


def tree_digest(root: Path) -> dict[str, str]:
    return {
        str(path.relative_to(root)): hashlib.sha256(path.read_bytes()).hexdigest()
        for path in sorted(root.rglob("*"))
        if path.is_file()
    }


def text_between(filing: data.Filing, start: dict, end: dict) -> str:
    if start["block"] == end["block"]:
        return filing.texts[start["block"]][start["offset"] : end["offset"]]
    pieces = [filing.texts[start["block"]][start["offset"] :]]
    pieces += filing.texts[start["block"] + 1 : end["block"]]
    pieces.append(filing.texts[end["block"]][: end["offset"]])
    return " ".join(pieces)


class FilingTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.html = build_filing()
        cls.filing = data.Filing(cls.html.encode("utf-8"))

    def test_renderings_match_checker(self) -> None:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(self.html.encode("utf-8"), "html.parser")
        for element in soup(["script", "style", "noscript"]):
            element.decompose()
        for sep in (" ", ""):
            expected = check_lean.normalize_contiguous(soup.get_text(sep))
            self.assertEqual(self.filing.renderings[sep][0], expected)

    def test_quote_split_across_font_nodes_uses_joined_rendering(self) -> None:
        located = self.filing.locate(SPLIT_QUOTE, "3")
        self.assertTrue(located["located"])
        self.assertEqual(located["rendering"], "")
        self.assertEqual(located["start"]["block"], located["end"]["block"])
        shown = text_between(self.filing, located["start"], located["end"])
        self.assertEqual(shown, SPLIT_QUOTE)
        self.assertEqual(located["found_page"], "3")
        self.assertEqual(located["occurrences"], 1)

    def test_quote_across_two_blocks(self) -> None:
        located = self.filing.locate(CROSS_QUOTE, "3")
        self.assertTrue(located["located"])
        self.assertEqual(located["rendering"], " ")
        self.assertEqual(located["end"]["block"], located["start"]["block"] + 1)
        self.assertEqual(text_between(self.filing, located["start"], located["end"]), CROSS_QUOTE)

    def test_quote_with_space_between_inline_nodes(self) -> None:
        located = self.filing.locate(SPACED_QUOTE, "3")
        self.assertTrue(located["located"])
        self.assertEqual(text_between(self.filing, located["start"], located["end"]), SPACED_QUOTE)

    def test_nbsp_and_curly_quotes(self) -> None:
        located = self.filing.locate(NBSP_QUOTE, "3")
        self.assertTrue(located["located"])
        self.assertEqual(text_between(self.filing, located["start"], located["end"]), NBSP_QUOTE)

    def test_quote_across_table_cells(self) -> None:
        located = self.filing.locate(ROW_QUOTE, "3")
        self.assertTrue(located["located"])
        self.assertEqual(self.filing.kinds[located["start"]["block"]], "row")
        self.assertEqual(text_between(self.filing, located["start"], located["end"]), "Epsilon Group | $12.00")

    def test_repeated_quote_prefers_cited_page(self) -> None:
        for cited, expected in (("5", "5"), ("3", "3"), (None, "3")):
            located = self.filing.locate(REPEATED_QUOTE, cited)
            self.assertEqual(located["occurrences"], 2)
            self.assertEqual(located["found_page"], expected, cited)

    def test_overlapping_repeats_are_counted(self) -> None:
        filing = data.Filing(b"<p>aaa</p>")
        self.assertEqual(filing.locate("aa", None)["occurrences"], 2)

    def test_toc_back_link_is_not_a_heading(self) -> None:
        links = [kind for kind, text in zip(self.filing.kinds, self.filing.texts) if text == "Table of Contents"]
        self.assertEqual(len(links), 12)
        self.assertEqual(set(links), {"p"})

    def test_missing_quote(self) -> None:
        located = self.filing.locate(MISSING_QUOTE, "3")
        self.assertFalse(located["located"])
        self.assertIsNone(located["start"])
        self.assertEqual(located["occurrences"], 0)

    def test_every_rendering_substring_is_locatable(self) -> None:
        for sep in (" ", ""):
            text = self.filing.renderings[sep][0]
            for start in range(0, len(text) - 20, 97):
                passage = text[start : start + 20].strip()
                if passage:
                    self.assertTrue(self.filing.locate(passage, None)["located"], passage)

    def test_page_detection(self) -> None:
        self.assertTrue(self.filing.pages_reliable)
        self.assertEqual([page["page"] for page in self.filing.pages], [str(n) for n in range(1, 13)])
        payload = self.filing.blocks_payload()
        cover = next(block for block in payload if block["text"] == "COVER PAGE")
        self.assertIsNone(cover["page"])
        gamma = next(block for block in payload if block["text"].startswith("On March 9"))
        self.assertEqual(gamma["page"], "4")
        for page in self.filing.pages:
            self.assertEqual(payload[page["block"]]["page"], page["page"])

    def test_unnumbered_filing_is_not_reliable(self) -> None:
        filing = data.Filing(build_filing(numbered=False).encode("utf-8"))
        self.assertFalse(filing.pages_reliable)
        self.assertEqual(filing.pages, [])
        self.assertTrue(all(block["page"] is None for block in filing.blocks_payload()))

    def test_background_heading_skips_table_of_contents(self) -> None:
        index = self.filing.background_block
        self.assertIsNotNone(index)
        block = self.filing.blocks_payload()[index]
        self.assertEqual(block["text"], "Background of the Merger")
        self.assertEqual(block["kind"], "h")
        self.assertEqual(block["page"], "3")

    def test_table_rows_are_row_blocks(self) -> None:
        rows = [block for block in self.filing.blocks_payload() if block["kind"] == "row"]
        self.assertIn("Background of the Merger | 3", [row["text"] for row in rows])

    def test_script_and_style_text_hidden(self) -> None:
        self.assertFalse(any("var x" in text for text in self.filing.texts))


DISPLAY_HTML = (
    "<html><body>"
    "<dl><dt>&#149;</dt><dd>the first risk</dd>"
    "<dt><b><i>Q:</i></b></dt><dd><b>What is the effect?</b></dd><dt>A:</dt><dd>It merges.</dd></dl>"
    "<div style='float:left'>&#8226;<br></div><div>Bullet in a div.</div><div>&#8203;</div>"
    "<table><tr><td>&#8203;</td><td>Agreement</td><td>&#8203;</td><td>&#8203;</td><td>A-1</td></tr>"
    "<tr><td>Revenue</td><td>$</td><td>(40.9</td><td>)</td><td>12</td><td>%</td></tr></table>"
    "</body></html>"
)


class DisplayTests(unittest.TestCase):
    """Display-only joins (bullets, labels, spacer and currency cells)."""

    @classmethod
    def setUpClass(cls) -> None:
        cls.filing = data.Filing(DISPLAY_HTML.encode("cp1252"))

    def test_labels_join_their_text(self) -> None:
        texts = self.filing.texts
        self.assertIn("\u2022 the first risk", texts)
        self.assertIn("Q: What is the effect?", texts)
        self.assertIn("A: It merges.", texts)
        self.assertIn("\u2022 Bullet in a div.", texts)
        self.assertNotIn("\u2022", texts)

    def test_table_cells_join_for_display(self) -> None:
        rows = [text for kind, text in zip(self.filing.kinds, self.filing.texts) if kind == "row"]
        self.assertEqual([row.replace("\u200b", "") for row in rows], ["Agreement | A-1", "Revenue | $(40.9) | 12%"])

    def test_offsets_still_match_renderings(self) -> None:
        from bs4 import BeautifulSoup

        soup = BeautifulSoup(DISPLAY_HTML.encode("cp1252"), "html.parser")
        for sep in (" ", ""):
            text, starts, blocks, offsets = self.filing.renderings[sep]
            self.assertEqual(text, check_lean.normalize_contiguous(soup.get_text(sep)))
            ends = list(starts[1:]) + [len(text)]
            for start, end, block, offset in zip(starts, ends, blocks, offsets):
                token = text[start:end].rstrip(" ")
                self.assertEqual(self.filing.texts[block][offset : offset + len(token)], token)


class HelperTests(unittest.TestCase):
    def test_page_label(self) -> None:
        cases = {
            "24": "24", " 7 ": "7", "-ii-": "ii", "- 7 -": "7", "\u2013 12 \u2013": "12", "(iii)": "iii",
            "A-1": "A-1", "B-2-B": "B-2-B", "ii\u200b": "ii",
            "(1)": None, "Table": None, "The Merger": None, "12345": None,
        }
        for text, expected in cases.items():
            self.assertEqual(data.page_label(text), expected, text)

    def test_display_value(self) -> None:
        self.assertEqual(data.display_value(None), "")
        self.assertEqual(data.display_value(10.0), "10")
        self.assertEqual(data.display_value(10.5), "10.5")
        self.assertEqual(data.display_value(3), "3")
        self.assertEqual(data.display_value(dt.datetime(2020, 3, 1, 0, 0)), "2020-03-01")
        self.assertEqual(data.display_value("  text "), "  text ")

    def test_cited_pages(self) -> None:
        self.assertEqual(data.cited_pages("24"), {"24"})
        self.assertEqual(data.cited_pages("30–31"), {"30", "31"})
        self.assertEqual(data.cited_pages("130-31"), {"130", "131"})
        self.assertIn("A-1", data.cited_pages("A-1"))
        self.assertEqual(data.cited_pages("30, 32"), {"30", "32"})

    def test_reader_mapping(self) -> None:
        self.assertEqual(server.reader_for("A.Gorbenko@ucl.ac.uk"), "alex")
        self.assertEqual(server.reader_for("junyu.li.24@ucl.ac.uk"), "austin")
        self.assertEqual(server.reader_for(None), "local")
        self.assertEqual(server.reader_for("someone@example.com"), "local")


class CockpitTests(unittest.TestCase):
    def setUp(self) -> None:
        self.tmp = tempfile.TemporaryDirectory()
        self.root = build_repo(Path(self.tmp.name))
        self.cockpit = data.Cockpit(self.root)

    def tearDown(self) -> None:
        self.tmp.cleanup()

    def test_read_only(self) -> None:
        before = tree_digest(self.root)
        self.cockpit.list_deals()
        self.cockpit.deal("alpha-deal")
        self.cockpit.filing_payload("alpha-deal")
        self.assertEqual(tree_digest(self.root), before)

    def test_list_deals(self) -> None:
        deals = self.cockpit.list_deals()
        self.assertEqual([deal["slug"] for deal in deals], ["alpha-deal"])
        deal = deals[0]
        self.assertEqual(deal["target"], "Target Co.")
        self.assertEqual(deal["form_type"], "DEFM14A")
        self.assertEqual(deal["rows"], 6)
        self.assertEqual(deal["quotes_total"], 5)
        self.assertEqual(deal["quotes_located"], 4)
        self.assertEqual(deal["questions"], 1)
        self.assertEqual(deal["rounds"], 1)
        self.assertIn("errors", deal["check"])

    def test_deal_payload(self) -> None:
        payload = self.cockpit.deal("alpha-deal")
        rows = {row["id"]: row for row in payload["ledger"]["rows"]}
        self.assertEqual(sorted(rows), [1, 2, 3, 4, 5, 7])
        self.assertEqual(rows[1]["excel_row"], 2)
        self.assertEqual(rows[7]["excel_row"], 8)
        self.assertEqual(rows[1]["cells"]["Price low"], "10")
        self.assertEqual(rows[1]["cells"]["Sort date"], "2020-03-01")
        self.assertEqual(rows[1]["cells"]["Reviewer note"], "")
        self.assertEqual(rows[1]["quote"]["text"], SPLIT_QUOTE)
        self.assertEqual(rows[1]["quote"]["cited_page"], "3")
        self.assertIsNone(rows[7]["quote"])
        self.assertFalse(rows[5]["quote"]["located"])
        self.assertIsNone(rows[1]["page_hint"])
        self.assertIn("page hint", rows[4]["page_hint"])
        self.assertEqual(rows[4]["quote"]["found_page"], "4")
        self.assertEqual(payload["questions"]["rows"][0]["id"], "Q1")
        self.assertEqual(payload["facts"][0], {"field": "Target", "value": "Target Co."})
        self.assertTrue(payload["pages_reliable"])
        self.assertIsNotNone(payload["background_block"])
        self.assertEqual(payload["filing"]["form_type"], "DEFM14A")
        self.assertIn("checker_version", payload["check"])
        json.dumps(payload)

    def test_checker_issue_maps_to_row(self) -> None:
        payload = self.cockpit.deal("alpha-deal")
        rows = {row["id"]: row for row in payload["ledger"]["rows"]}
        codes = [issue["code"] for issue in rows[5]["issues"]]
        self.assertIn("quote.not_contiguous_in_filing", codes)
        self.assertNotIn("quote.not_contiguous_in_filing", [i["code"] for i in rows[1]["issues"]])
        for issue in payload["check"]["other_issues"]:
            self.assertFalse(
                issue.get("sheet") == "Deal ledger" and issue.get("row") in {r["excel_row"] for r in rows.values()}
            )

    def test_issue_mapping_by_sheet(self) -> None:
        workbook, filing_path, entry = self.cockpit.resolve("alpha-deal")
        tables = self.cockpit.workbook(workbook)
        filing = self.cockpit.filing(filing_path)
        report = {
            "status": "fail",
            "summary": {"errors": 4, "warnings": 1, "information": 0, "total_issues": 5},
            "scope_note": "note",
            "checker_version": "x",
            "issues": [
                {"severity": "error", "code": "a", "sheet": "Deal ledger", "row": 3, "column": "Who", "message": "m"},
                {"severity": "warning", "code": "b", "sheet": "Rounds", "row": 2, "column": None, "message": "m"},
                {"severity": "error", "code": "c", "sheet": "Questions", "row": 2, "column": "Q", "message": "m"},
                {"severity": "error", "code": "d", "sheet": "Deal facts", "row": 2, "column": "Value", "message": "m"},
                {"severity": "error", "code": "e", "sheet": "Deal ledger", "row": 1, "column": None, "message": "m"},
            ],
        }
        payload = data.build_deal_payload("alpha-deal", entry, tables, filing, report)
        ledger = {row["excel_row"]: row for row in payload["ledger"]["rows"]}
        self.assertEqual([i["code"] for i in ledger[3]["issues"]], ["a"])
        self.assertEqual(ledger[3]["issues"][0], {"severity": "error", "code": "a", "column": "Who", "message": "m"})
        self.assertEqual([i["code"] for i in payload["rounds"]["rows"][0]["issues"]], ["b"])
        self.assertEqual([i["code"] for i in payload["questions"]["rows"][0]["issues"]], ["c"])
        self.assertEqual(sorted(i["code"] for i in payload["check"]["other_issues"]), ["d", "e"])

    def test_slug_validation(self) -> None:
        for slug in ("", "../extraction/alpha-deal", "Alpha-deal", "-alpha", "alpha deal", "ghost", "old", "archive"):
            with self.assertRaises(data.DealNotFound, msg=slug):
                self.cockpit.resolve(slug)
        with self.assertRaises(data.DealNotFound):
            self.cockpit.deal("nope")

    def test_payload_is_cached(self) -> None:
        first = self.cockpit.deal("alpha-deal")
        self.assertIs(self.cockpit.deal("alpha-deal"), first)

    def test_cache_follows_workbook_changes(self) -> None:
        self.cockpit.deal("alpha-deal")
        path = self.root / "extraction" / "alpha-deal.xlsx"
        wb = load_workbook(path)
        wb["Deal ledger"]["C2"] = "Omega"
        wb.save(path)
        stat = path.stat()
        os.utime(path, ns=(stat.st_atime_ns, stat.st_mtime_ns + 10**9))
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        self.assertEqual(rows[0]["cells"]["Who"], "Omega")

    def test_unreadable_workbook_is_listed_not_fatal(self) -> None:
        add_corrupt_deal(self.root, "bad")
        deals = {deal["slug"]: deal for deal in self.cockpit.list_deals()}
        self.assertEqual(sorted(deals), ["alpha-deal", "bad"])
        self.assertEqual(deals["alpha-deal"]["quotes_located"], 4)
        self.assertEqual(deals["bad"]["check"]["status"], "unreadable")
        self.assertNotIn(str(self.root), deals["bad"]["error"])
        with self.assertRaises(data.DealUnavailable) as caught:
            self.cockpit.deal("bad")
        self.assertNotIn(str(self.root), str(caught.exception))


class ServerTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls) -> None:
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = build_repo(Path(cls.tmp.name))
        add_corrupt_deal(cls.root, "bad")
        cls.httpd = server.make_server(0, data.Cockpit(cls.root), quiet=True)
        cls.port = cls.httpd.server_address[1]
        cls.thread = threading.Thread(target=cls.httpd.serve_forever, daemon=True)
        cls.thread.start()

    @classmethod
    def tearDownClass(cls) -> None:
        cls.httpd.shutdown()
        cls.httpd.server_close()
        cls.tmp.cleanup()

    def request(self, path: str, method: str = "GET", headers: dict | None = None):
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", method=method, headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=30) as response:
                return response.status, dict(response.headers), response.read()
        except urllib.error.HTTPError as exc:
            return exc.code, dict(exc.headers), exc.read()

    def test_binds_loopback(self) -> None:
        self.assertEqual(self.httpd.server_address[0], "127.0.0.1")

    def test_api_endpoints(self) -> None:
        status, headers, body = self.request("/api/deals")
        self.assertEqual(status, 200)
        self.assertEqual(headers["Cache-Control"], "no-store")
        self.assertEqual(json.loads(body)[0]["slug"], "alpha-deal")
        status, _, body = self.request("/api/deal/alpha-deal", headers={"Cf-Access-Authenticated-User-Email": "a.gorbenko@ucl.ac.uk"})
        self.assertEqual(status, 200)
        payload = json.loads(body)
        self.assertEqual(payload["reader"], "local")
        self.assertEqual(payload["slug"], "alpha-deal")
        status, _, body = self.request("/api/filing/alpha-deal")
        self.assertEqual(status, 200)
        filing = json.loads(body)
        self.assertEqual(set(filing), {"blocks", "pages"})
        self.assertEqual(set(filing["blocks"][0]), {"kind", "text", "page"})

    def test_unknown_and_invalid_slugs_are_404(self) -> None:
        for path in ("/api/deal/nope", "/api/deal/..%2Fref", "/api/filing/ghost", "/api/deal/", "/api/other"):
            status, headers, body = self.request(path)
            self.assertEqual(status, 404, path)
            self.assertIn("error", json.loads(body))
            self.assertEqual(headers["Cache-Control"], "no-store")

    def test_pages_and_static(self) -> None:
        for path in ("/", "/deal/alpha-deal"):
            status, headers, _ = self.request(path)
            self.assertEqual(status, 200, path)
            self.assertTrue(headers["Content-Type"].startswith("text/html"))
        self.assertEqual(self.request("/static/server.py")[0], 404)
        self.assertEqual(self.request("/static/../data.py")[0], 404)

    def test_other_methods_rejected(self) -> None:
        for method in ("POST", "PUT", "DELETE", "PATCH", "OPTIONS", "HEAD", "TRACE", "PROPFIND", "FOO"):
            status, headers, body = self.request("/api/deal/alpha-deal", method=method)
            self.assertEqual(status, 405, method)
            self.assertEqual(headers["Cache-Control"], "no-store", method)
            if method != "HEAD":
                self.assertIn("error", json.loads(body), method)

    def test_unreadable_deal(self) -> None:
        status, _, body = self.request("/api/deals")
        self.assertEqual(status, 200)
        deals = {deal["slug"]: deal for deal in json.loads(body)}
        self.assertEqual(deals["alpha-deal"]["quotes_located"], 4)
        self.assertEqual(deals["bad"]["check"]["status"], "unreadable")
        status, headers, body = self.request("/api/deal/bad")
        self.assertEqual(status, 409)
        self.assertEqual(headers["Cache-Control"], "no-store")
        self.assertIn("error", json.loads(body))
        self.assertNotIn(str(self.root).encode(), body)


if __name__ == "__main__":
    unittest.main()
