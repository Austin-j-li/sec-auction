#!/usr/bin/env python3
"""Set a ledger beside Alex's hand coding of the same deal: a review aid, never a target.

    python3 _dev/tools/compare_alex.py LEDGER.xlsx --out DIR [--deal SLUG] [--alex XLSX] [--seed CSV]

The deal is joined to `ref/deal_details_Alex_2026.xlsx` through `ref/seed.csv`'s deal_number. Each
of Alex's labelled bids (bid_type Formal or Informal) is aligned to one whole-company Bid or Bid
reaffirmed row by bidder, price (upfront, or upfront + CVR/earnout value) and nearest date. For
aligned bids the output reports agreement of bid_type with each Formality reading of
derive_analysis.py (T0, T1, T1u, T2, T3), of all_cash and of the per-share value. The ledger must be
a Version 1 workbook. Every other coded row
(bid_note) is checked against the ledger's Event and Exit reason through audit D §3.3's code map.
Each Alex row is marked as carrying his red-font correction or the earlier Chicago coding; only
nine deals carry corrections.

Nothing here may reach an extraction run or be used to tune the instruction. On a
held-out deal, run it only at Austin's request. It writes only to the --out folder: alex_bids.csv,
alex_events.csv and summary.json.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path
from typing import Any

import openpyxl

import derive_analysis as derive
from derive_analysis import cell, text

ALEX = derive.PROJECT / "ref/deal_details_Alex_2026.xlsx"
SEED = derive.PROJECT / "ref/seed.csv"
READINGS = ("T0", "T1", "T1u", "T2", "T3")
TOLERANCE = 0.005
RED = {"FFFF0000", "FF9C0006"}
WHOLE_ROW_RED = 30
COMMENTS = {"comments_1", "comments_2", "comments_3"}
CORRECTED = {"imprivata", "mac-gray", "medivation", "penford", "petsmart", "providence-worcester", "saks", "stec", "zep"}
BID_COLUMNS = [
    "deal", "alex_row", "bidder_id", "bidder_name", "alex_date", "bid_type", "bid_value_lower",
    "bid_value_upper", "bid_value_pershare", "all_cash_alex", "bid_note", "provenance", "red_cells",
    "ledger_row", "ledger_who", "ledger_event", "ledger_round", "alignment", "price_basis", "date_gap_days",
    *READINGS, *(f"agree_{reading}" for reading in READINGS), "all_cash_ledger", "agree_all_cash",
    "ledger_upfront", "ledger_package", "ledger_package_basis", "per_share_match",
]
EVENT_COLUMNS = [
    "deal", "alex_row", "bidder_id", "bidder_name", "alex_date", "bid_note", "provenance", "red_cells",
    "expected", "status", "ledger_row", "ledger_who", "ledger_event", "ledger_exit_reason", "ledger_round",
    "round_finality", "date_gap_days",
]
# Alex's bid_note codes and the corresponding whole-company ledger events.
CODES = {
    "NA": ({"Bid", "Bid reaffirmed"}, None),
    "IB": ({"Adviser"}, None), "IB Terminated": ({"Adviser ended"}, None),
    "NDA": ({"NDA signed"}, None),
    "Bidder Interest": ({"Bidder interest"}, None), "Bidder Sale": ({"Bid"}, None),
    "Target Interest": ({"Target interest"}, None), "Target Sale": ({"Target sale decision"}, None),
    "Target Sale Public": ({"Target sale decision", "Sale process announced"}, None),
    "Sale Press Release": ({"Sale process announced"}, None),
    "Bid Press Release": ({"Bid announced"}, None),
    "Activist Sale": ({"Activist"}, None),
    "Terminated": ({"Process terminated"}, None), "Restarted": ({"Process restarted"}, None),
    "Drop": ({"Withdrew"}, None), "DropM": ({"Withdrew"}, None),
    "DropBelowM": ({"Withdrew"}, None), "DropBelowInf": ({"Withdrew"}, None),
    "DropAtInf": ({"Withdrew"}, None), "DropTarget": ({"Dropped by target"}, None),
    "Executed": ({"Merger agreement signed"}, None),
    "Final Round Ann": ({"Deadline set", "Round opened"}, True),
    "Final Round Inf Ann": ({"Deadline set", "Round opened"}, False),
    "Final Round": ({"Deadline"}, True),
    "Final Round Inf": ({"Deadline"}, False),
    "Final Round Ext Ann": ({"Deadline set", "Round opened", "Deadline revised"}, True),
    "Final Round Ext": ({"Deadline", "Deadline revised"}, True),
    "Final Round Inf Ext Ann": ({"Deadline set", "Round opened", "Deadline revised"}, False),
    "Final Round Inf Ext": ({"Deadline", "Deadline revised"}, False),
}


def na(value: Any) -> Any:
    return None if value is None or (isinstance(value, str) and value.strip() in ("", "NA")) else value


def as_date(value: Any) -> dt.date | None:
    return derive.as_date(na(value))


def names(value: Any) -> list[str]:
    """The name and its '/'-separated parts; "Party E/F" also gives "party f"."""
    full = derive.unit_key(value)
    parts = [p.strip() for p in full.split("/") if p.strip()]
    prefix = parts[0].rsplit(" ", 1)[0] + " " if parts and " " in parts[0] else ""
    return [full] + [prefix + p if len(p) <= 2 and prefix else p for p in parts]


def name_strength(alex: Any, who: Any, count: int | None = None) -> int:
    """3 same name, 2 a shared part or one name inside the other, 1 cohorts of the same size, 0 none."""
    if na(alex) is None:
        return 0
    a, b = names(alex), names(who)
    if a[0] == b[0]:
        return 3
    if set(a) & set(b) or any(len(x) >= 3 and re.search(rf"(?<!\w){re.escape(x)}(?!\w)", b[0]) for x in a) \
            or any(len(x) >= 3 and re.search(rf"(?<!\w){re.escape(x)}(?!\w)", a[0]) for x in b):
        return 2
    size = derive.leading_integer(a[0])
    return 1 if size is not None and size > 1 and size in (derive.leading_integer(b[0]), count) else 0


def close(a: float | None, b: float | None) -> bool:
    return a is not None and b is not None and abs(a - b) <= TOLERANCE + 1e-9


def price_basis(alex: tuple[float | None, float | None], bid: dict[str, Any]) -> str | None:
    lo, hi = alex
    lo, hi = lo if lo is not None else hi, hi if hi is not None else lo
    def fill(a, b): return (a if a is not None else b, b if b is not None else a)
    upfront, package = fill(bid["price_low"], bid["price_high"]), fill(bid["package_low"], bid["package_high"])
    if lo is None:
        return "undisclosed" if upfront == (None, None) else None
    if close(lo, upfront[0]) and close(hi, upfront[1]):
        return "upfront"
    if close(lo, package[0]) and close(hi, package[1]):
        return "package"
    return None


def gap(day: dt.date | None, row: dict[str, Any]) -> int | None:
    """Days from Alex's date to the ledger row's window (0 inside it)."""
    ends = [derive.as_date(row.get(k)) for k in ("date_from", "sort_date", "date_to")]
    ends = [d for d in ends if d]
    if day is None or not ends:
        return None
    first, last = min(ends), max(ends)
    return 0 if first <= day <= last else min(abs((day - first).days), abs((day - last).days))


def seed_number(seed: Path, deal: str) -> str:
    with seed.open(newline="", encoding="utf-8") as handle:
        for row in csv.DictReader(handle):
            if row.get("deal") == deal:
                return str(row["deal_number"]).strip()
    raise derive.DeriveError(f"{deal}: not in {seed}")


def alex_rows(path: Path, number: str) -> list[dict[str, Any]]:
    """Alex's rows for one DealNumber, with the columns in red font."""
    wb = openpyxl.load_workbook(path)
    try:
        ws = wb["deal_details"] if "deal_details" in wb.sheetnames else wb.active
        header = [c.value for c in ws[1]]
        at = header.index("DealNumber")
        out = []
        for row in ws.iter_rows(min_row=2):
            value = row[at].value
            if value is None or str(int(value) if isinstance(value, float) else value).strip() != number:
                continue
            red = [str(h) if h is not None else "(index)" for h, c in zip(header, row)
                   if c.font is not None and c.font.color is not None and c.font.color.type == "rgb" and c.font.color.rgb in RED]
            record = {str(h): c.value for h, c in zip(header, row) if h is not None}
            record.update(excel_row=row[0].row, red=red)
            out.append(record)
        return out
    finally:
        wb.close()


def provenance(deal: str, red: list[str]) -> str:
    if len(red) >= WHOLE_ROW_RED:
        return "Alex: whole row in red (added or rewritten)"
    if set(red) - COMMENTS:
        return "Alex: red-font correction"
    if red:
        return "Alex: red comments only (coding as Chicago)"
    return "Chicago coding" if deal in CORRECTED else "Chicago coding (deal has no corrections)"


def bid_type(value: Any) -> str | None:
    value = text(na(value))
    return {"formal": "Formal", "informal": "Informal", "informsl": "Informal"}.get(value.casefold())


def alex_price(row: dict[str, Any]) -> tuple[float | None, float | None]:
    lo, hi = derive.number(na(row.get("bid_value_lower"))), derive.number(na(row.get("bid_value_upper")))
    if lo is None and hi is None:
        per = derive.number(na(row.get("bid_value_pershare")))
        return per, per
    return lo, hi


def align(bids: list[dict[str, Any]], labelled: list[dict[str, Any]]) -> dict[int, tuple[dict[str, Any], str, str, int | None]]:
    """Alex row index -> (bid, alignment, price basis, days); one ledger row per Alex row."""
    candidates = []
    for i, row in enumerate(labelled):
        day = as_date(row.get("bid_date_precise")) or as_date(row.get("bid_date_rough"))
        rough = as_date(row.get("bid_date_rough"))
        for j, bid in enumerate(bids):
            basis = price_basis(alex_price(row), bid)
            if basis is None:
                continue
            strength, days = name_strength(row.get("BidderName"), bid["who"]), gap(day, bid)
            if strength == 0 and (days is None or days > 7):
                continue
            candidates.append((-strength, 10**6 if days is None else days, gap(rough, bid) or 0, i, j, basis, strength, days))
    done: dict[int, tuple[dict[str, Any], str, str, int | None]] = {}
    used: set[int] = set()
    labels = {3: "bidder, price, date", 2: "bidder (part of name), price, date", 1: "cohort size, price, date", 0: "price and date only"}
    for *_, i, j, basis, strength, days in sorted(candidates):
        if i in done or j in used:
            continue
        done[i] = (bids[j], labels[strength], basis, days)
        used.add(j)
    return done


def final_code(code: str) -> tuple[set[str], bool] | None:
    """The ledger events and required round finality for an Alex final-round code."""
    if code.startswith("Final Round") and code in CODES:
        events, final = CODES[code]
        return events, bool(final)
    return None


def code_mapping(code: str) -> tuple[set[str], bool | None] | None:
    return CODES.get(code) or (({"Exclusivity changed"}, None) if code.startswith("Exclusivity ") else None)


def match_code(row: dict[str, Any], code: str, ledger: list[dict[str, Any]], finality: dict[tuple, str]) -> dict[str, Any]:
    """Match a coded non-bid row to a nearby source event; keep disagreements visible."""
    mapped = code_mapping(code)
    if mapped is None:
        return {"expected": "(no code in the map)", "status": "not compared"}
    events, required_final = mapped
    expected = " or ".join(sorted(events))
    days = [as_date(row.get(key)) for key in ("bid_date_precise", "bid_date_rough")]
    days = [day for day in days if day is not None]
    name = na(row.get("BidderName"))

    def date_gap(r):
        gaps = [distance for day in days if (distance := gap(day, r)) is not None]
        return min(gaps) if gaps else None

    def candidates(test):
        found=[]
        for r in ledger:
            if not test(r): continue
            strength = name_strength(name, r["who"], r.get("count_point")) if name is not None else 0
            if name is not None and not strength: continue
            distance = date_gap(r)
            if distance is None or distance > 31: continue
            found.append((distance, -strength, r["row"], r))
        return [x[-1] for x in sorted(found)]

    def describe(r, status):
        return {"expected": expected, "status": status, "ledger_row": r["row"], "ledger_who": r["who"],
                "ledger_event": r["event"], "ledger_exit_reason": r["exit_reason"], "ledger_round": r["round"],
                "round_finality": finality.get((r["process"], r["round"]), ""), "date_gap_days": date_gap(r)}

    matched = candidates(lambda r: r["event"] in events)
    if matched:
        r=matched[0]
        if required_final is not None:
            recorded_finality = finality.get((r["process"], r["round"]), "")
            is_final = recorded_finality in derive.FINAL
            if is_final != required_final:
                return describe(r, "event agrees, round finality differs")
        reason = r["exit_reason"]
        desired = {"DropM": {"Value below market price", "Value at or below market price"},
                   "DropBelowM": {"Value below market price", "Value at or below market price"},
                   "DropBelowInf": {"Value below earlier offer"}, "DropAtInf": {"Value at earlier offer"}}
        if code in desired and reason not in desired[code]:
            return describe(r, "exit, other label or reason")
        return describe(r, "agree")
    if code.startswith("Drop"):
        found=candidates(lambda r: r["event"] in {"Withdrew", "Did not submit", "Not selected at signing", "Dropped by target"})
        if found:
            return describe(found[0], "exit, other label or reason")
    return {"expected": expected, "status": "unaligned"}


def compare(workbook: Path, deal: str, alex: Path, seed: Path) -> dict[str, Any]:
    ledger = derive.load(workbook)
    result = derive.derive(ledger, deal)
    number = seed_number(seed, deal)
    rows = alex_rows(alex, number)
    if not rows:
        raise derive.DeriveError(f"{deal}: no rows for DealNumber {number} in {alex}")
    bids = result["bids"]
    finality = {(r["process"], r["round"]): r["finality"] for r in result["rounds"]}
    all_rows = [{**derive.base_cols(deal, r, *derive.row_round(r)), "exit_reason": text(r.get("Exit reason")),
                 "count_point": derive.count_bounds(r.get("Count"), r.get("Note"))[2]} for r in ledger["ledger"]]
    labelled = [r for r in rows if bid_type(r.get("bid_type"))]
    aligned = align(bids, labelled)

    bid_out = []
    for i, row in enumerate(labelled):
        lo, hi = alex_price(row)
        record = {"deal": deal, "alex_row": row["excel_row"], "bidder_id": cell(na(row.get("BidderID"))),
                  "bidder_name": text(na(row.get("BidderName"))),
                  "alex_date": cell(as_date(row.get("bid_date_precise")) or as_date(row.get("bid_date_rough"))),
                  "bid_type": bid_type(row.get("bid_type")), "bid_value_lower": lo, "bid_value_upper": hi,
                  "bid_value_pershare": derive.number(na(row.get("bid_value_pershare"))),
                  "all_cash_alex": derive.number(na(row.get("all_cash"))), "bid_note": text(na(row.get("bid_note"))) or "NA",
                  "provenance": provenance(deal, row["red"]), "red_cells": "; ".join(row["red"]), "alignment": "unaligned"}
        if i in aligned:
            bid, how, basis, days = aligned[i]
            record.update(ledger_row=bid["row"], ledger_who=bid["who"], ledger_event=bid["event"], ledger_round=bid["round"],
                          alignment=how, price_basis=basis, date_gap_days=days, all_cash_ledger=bid["all_cash"],
                          ledger_upfront=bid["price_low"], ledger_package=bid["package_low"], ledger_package_basis=bid["package_basis"])
            for reading in READINGS:
                record[reading] = bid[reading]
                record[f"agree_{reading}"] = "" if not bid[reading] else "Y" if bid[reading] == record["bid_type"] else "N"
            if record["all_cash_alex"] is not None and bid["all_cash"] is not None:
                record["agree_all_cash"] = "Y" if int(record["all_cash_alex"]) == bid["all_cash"] else "N"
            per = record["bid_value_pershare"]
            record["per_share_match"] = ("" if per is None else "upfront" if close(per, bid["price_low"]) else
                                         "package" if close(per, bid["package_low"]) else "differs")
        bid_out.append(record)

    event_out = []
    for row in rows:
        code = text(na(row.get("bid_note"))) or "NA"
        record = {"deal": deal, "alex_row": row["excel_row"], "bidder_id": cell(na(row.get("BidderID"))),
                  "bidder_name": text(na(row.get("BidderName"))),
                  "alex_date": cell(as_date(row.get("bid_date_precise")) or as_date(row.get("bid_date_rough"))),
                  "bid_note": code, "provenance": provenance(deal, row["red"]), "red_cells": "; ".join(row["red"])}
        if code == "NA" and not bid_type(row.get("bid_type")):
            record.update(expected="(no event: an unlabelled row)", status="not compared")
        elif bid_type(row.get("bid_type")):
            # A labelled bid is checked on the ledger row it was aligned to.
            i = next(k for k, r in enumerate(labelled) if r is row)
            bid = aligned.get(i, (None,))[0]
            events = CODES["NA"][0] if code == "NA" else (code_mapping(code) or (set(),))[0]
            record.update(expected=" or ".join(sorted(events)) or "(no code in the map)",
                          status="unaligned" if not bid else "agree" if bid["event"] in events else "disagree",
                          **({"ledger_row": bid["row"], "ledger_who": bid["who"], "ledger_event": bid["event"],
                              "ledger_round": bid["round"], "round_finality": bid["round_finality"],
                              "date_gap_days": aligned[i][3]} if bid else {}))
        else:
            record.update(match_code(row, code, all_rows, finality))
        event_out.append(record)

    def tally(column: str) -> dict[str, int]:
        values = [r.get(column, "") for r in bid_out if r["alignment"] != "unaligned"]
        return {"agree": values.count("Y"), "compared": values.count("Y") + values.count("N"), "missing": values.count("") + values.count(None)}

    statuses: dict[str, int] = {}
    for r in event_out:
        statuses[r["status"]] = statuses.get(r["status"], 0) + 1
    summary = {
        "note": "Review aid only: never a target for the instruction, and never shown to an extraction run.",
        "deal": deal, "deal_number": number, "workbook": str(workbook), "workbook_sha256": derive.sha256(workbook),
        "ledger_schema": ledger["schema"], "alex_file": str(alex), "alex_sha256": derive.sha256(alex),
        "deal_has_alex_corrections": deal in CORRECTED,
        "alex_rows": len(rows), "alex_labelled_bids": len(labelled),
        "aligned": sum(1 for r in bid_out if r["alignment"] != "unaligned"),
        "alignment": {k: sum(1 for r in bid_out if r["alignment"] == k) for k in sorted({r["alignment"] for r in bid_out})},
        "readings": {reading: tally(f"agree_{reading}") for reading in READINGS},
        "all_cash": tally("agree_all_cash"),
        "per_share": {k: sum(1 for r in bid_out if r.get("per_share_match") == k) for k in ("upfront", "package", "differs")},
        "bid_note_codes": statuses,
        "provenance": {k: sum(1 for r in rows if provenance(deal, r["red"]) == k) for k in sorted({provenance(deal, r["red"]) for r in rows})},
        "labelled_bids_with_red_bid_type": sum(1 for r in labelled if "bid_type" in r["red"]),
    }
    return {"bids": bid_out, "events": event_out, "summary": summary}


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--deal", help="deal slug (default: from the file name)")
    parser.add_argument("--out", type=Path, required=True, help="new or empty output folder")
    parser.add_argument("--alex", type=Path, default=ALEX)
    parser.add_argument("--seed", type=Path, default=SEED)
    args = parser.parse_args(argv)
    deal, guessed = (args.deal, None) if args.deal else derive.deal_from_name(args.workbook, None, derive.known_deals())
    try:
        derive.check_out(args.out)
        result = compare(args.workbook.resolve(), deal, args.alex, args.seed)
        if guessed:
            result["summary"]["warning"] = guessed
        args.out.mkdir(parents=True, exist_ok=True)
        derive.write_csv(args.out / "alex_bids.csv", BID_COLUMNS, result["bids"])
        derive.write_csv(args.out / "alex_events.csv", EVENT_COLUMNS, result["events"])
        (args.out / "summary.json").write_text(json.dumps(result["summary"], indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    except (derive.DeriveError, OSError, KeyError, ValueError) as exc:
        print(f"compare_alex: {exc}", file=sys.stderr)
        return 2
    s = result["summary"]
    print(f"{deal}: {s['aligned']} of {s['alex_labelled_bids']} labelled bids aligned; "
          + ", ".join(f"{k} {v['agree']}/{v['compared']}" for k, v in s["readings"].items()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
