"""Catalog deals are found from catalog.json, and verify_catalog checks them read-only.

A synthetic nine-deal repository stands in for the checkout, so the workbooks can be
relocated as D25 will relocate the real ones.
"""
from __future__ import annotations

import contextlib
import hashlib
import io
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import fetch_filing
from cockpit import data, provenance, verify_catalog
from cockpit import import_results as imp
from test_cockpit import build_filing, build_workbook, tree_digest

RELOCATED = f"{imp.REEXTRACT}/workbooks"


def sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def nine_deal_repo(root: Path, folder: str = "extraction") -> data.Cockpit:
    """The nine catalog deals, each one version whose workbook is in `folder`, as import_results writes them."""
    rows = ["file,deal,form_type,date_filed,source_url,document,fetched_utc,bytes,sha256"]
    catalog = {"schema_version": 1, "deals": {}}
    (root / "raw_filing").mkdir(parents=True)
    for slug in imp.DEALS:
        filing = f"{slug}_2020-04-01_DEFM14A.htm"
        (root / "raw_filing" / filing).write_text(build_filing(), encoding="utf-8")
        rows.append(f"{filing},{slug},DEFM14A,2020-04-01,https://example.invalid/{slug}.txt,{slug}.htm,,,")
        path = root / folder / f"{slug}.xlsx"
        path.parent.mkdir(parents=True, exist_ok=True)
        build_workbook(path)
        catalog["deals"][slug] = dict(
            name=imp.NAMES.get(slug, slug.title()), default_base=imp.VERSION_ID,
            versions=[dict(id=imp.VERSION_ID, label="Opus 5.5 medium extraction", path=f"{folder}/{slug}.xlsx",
                           sha256=imp.sha256(path), instruction_version="v1.13.2", kind="raw",
                           review_status="unreviewed; Austin review pending")],
            findings=[], documents=[])
    (root / "raw_filing/MANIFEST.csv").write_text("\n".join(rows) + "\n", encoding="utf-8")
    (root / "_dev/cockpit").mkdir(parents=True)
    (root / "_dev/cockpit/catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    return data.Cockpit(root)


def relocate(root: Path, slugs=imp.DEALS, folder: str = RELOCATED) -> None:
    """D25: move the workbooks, bytes unchanged, and repoint only their catalog paths."""
    path = root / "_dev/cockpit/catalog.json"
    catalog = json.loads(path.read_text(encoding="utf-8"))
    for slug in slugs:
        for version in catalog["deals"][slug]["versions"]:
            source, target = root / version["path"], root / folder / f"{slug}.xlsx"
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.move(source, target)
            version["path"] = f"{folder}/{slug}.xlsx"
    path.write_text(json.dumps(catalog), encoding="utf-8")


def edit(cockpit: data.Cockpit, slug: str, price: str = "99") -> dict:
    payload = cockpit.deal(slug)
    row = payload["ledger"]["rows"][0]
    return cockpit.workspace.edit(slug, {"revision": payload["workspace"]["revision"],
                                         "base_sha256": payload["workspace"]["base_sha256"], "reason": "Review edit",
                                         "operations": [{"type": "update", "sheet": "Deal ledger", "uid": row["uid"],
                                                         "values": {"Price low": price}}]}, "austin")


def import_version(cockpit: data.Cockpit, slug: str, ident: str = "run1") -> bytes:
    """A version imported from a cockpit run, as the worker records it: a changed copy of the deal's workbook."""
    folder = cockpit.repo_root / "_dev/cockpit/state/versions" / slug / ident
    folder.mkdir(parents=True)
    book = load_workbook(cockpit.workspace.root / cockpit.workspace.item(slug)["versions"][0]["path"])
    book["Deal facts"]["B2"] = "Run target"
    book.save(folder / f"{slug}.xlsx")
    content = (folder / f"{slug}.xlsx").read_bytes()
    conn = cockpit.workspace._connect(write=True)
    conn.execute("INSERT INTO versions (slug, id, label, path, sha256, kind, engine, model, effort, instruction_version, instruction_sha256, filing_sha256, started_by, started_at, finished_at, receipts, checker) VALUES (?,?,?,?,?,'raw','Opus 5.5','claude-opus-5-5','medium','v1.14','i','f','austin','2026-09-26T10:00:00+00:00','2026-09-26T10:10:00+00:00',?,'{}')",
                 (slug, ident, "Opus 5.5 · medium", f"_dev/cockpit/state/versions/{slug}/{ident}/{slug}.xlsx", sha(content), f"_dev/cockpit/state/versions/{slug}/{ident}"))
    conn.commit(); conn.close()
    return content


def shown(payload: dict) -> tuple:
    """What a reader sees of a working copy: its revision, base, and every cell."""
    return (payload["workspace"]["revision"], payload["workspace"]["base_version"], payload["workspace"]["base_sha256"],
            [[(r["uid"], r["cells"]) for r in payload[key]["rows"]] for key in ("ledger", "rounds", "questions")], payload["facts"])


class CatalogLookupTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit = nine_deal_repo(self.root)

    def tearDown(self): self.temp.cleanup()

    def test_nothing_changes_while_the_files_are_in_place(self):
        self.assertEqual(self.cockpit.slugs(), sorted(imp.DEALS))
        workbook, filing, entry = self.cockpit.resolve("kraton")
        self.assertEqual(workbook, (self.root / "extraction/kraton.xlsx").resolve())
        self.assertEqual((filing, entry["deal"]), (self.root / "raw_filing/kraton_2020-04-01_DEFM14A.htm", "kraton"))
        # A workbook in extraction/ that the catalog does not list is still found there, as before.
        extra = self.root / "extraction/zeta-deal.xlsx"
        build_workbook(extra)
        (self.root / "raw_filing/zeta-deal_2020-04-01_DEFM14A.htm").write_text(build_filing(), encoding="utf-8")
        with (self.root / "raw_filing/MANIFEST.csv").open("a", encoding="utf-8") as handle:
            handle.write("zeta-deal_2020-04-01_DEFM14A.htm,zeta-deal,DEFM14A,2020-04-01,,,,,\n")
        self.assertIn("zeta-deal", self.cockpit.slugs())
        self.assertEqual(self.cockpit.resolve("zeta-deal")[0], extra)
        # A catalog deal without its filing is not listed, as before.
        (self.root / "raw_filing/stec_2020-04-01_DEFM14A.htm").unlink()
        self.assertNotIn("stec", self.cockpit.slugs())
        with self.assertRaises(data.DealNotFound):
            self.cockpit.resolve("stec")

    def test_relocated_deal_lists_renders_and_exports(self):
        edit(self.cockpit, "kraton")
        before = {slug: shown(self.cockpit.deal(slug)) for slug in ("kraton", "datalink")}
        before_tables = data.read_workbook(self._export_file("kraton"))
        relocate(self.root)
        shutil.rmtree(self.root / "extraction")
        cockpit = data.Cockpit(self.root)
        self.assertEqual(cockpit.slugs(), sorted(imp.DEALS))
        listed = {deal["slug"]: deal for deal in cockpit.list_deals()}
        self.assertEqual(set(listed), set(imp.DEALS))
        self.assertFalse([slug for slug, deal in listed.items() if deal.get("error")])
        self.assertEqual(cockpit.resolve("kraton")[0], (self.root / RELOCATED / "kraton.xlsx").resolve())
        # The working copies resolve on the same base id and hash, at the same revision, with the same cells.
        for slug, seen in before.items():
            self.assertEqual(shown(cockpit.deal(slug)), seen)
        self.assertEqual(before["kraton"][:2], (1, imp.VERSION_ID))
        self.assertEqual(data.read_workbook(self._export_file("kraton", cockpit)), before_tables)
        base = cockpit.workspace.item("datalink")["versions"][0]
        self.assertEqual(sha(cockpit.workspace.export("datalink")), base["sha256"])
        self.assertEqual(sha(cockpit.workspace.export("datalink", imp.VERSION_ID)), base["sha256"])
        self.assertTrue(cockpit.filing_payload("kraton")["blocks"])
        outcome = verify_catalog.verify(self.root)
        self.assertEqual((outcome["deals"]["kraton"]["edited"], outcome["deals"]["kraton"]["path"]),
                         (True, f"{RELOCATED}/kraton.xlsx"))

    def test_relocated_deal_downloads_with_its_source_sheet(self):
        """S3 with S5: the download's Source sheet takes the manifest row through the catalog-based resolve."""
        submission = "https://www.sec.gov/Archives/edgar/data/77/0000000077-20-000001.txt"
        index = fetch_filing.index_link(submission)
        filing_hash = sha((self.root / "raw_filing/kraton_2020-04-01_DEFM14A.htm").read_bytes())
        manifest = self.root / "raw_filing/MANIFEST.csv"
        manifest.write_text(manifest.read_text(encoding="utf-8").replace(
            "https://example.invalid/kraton.txt,kraton.htm,,,", f"{submission},kraton.htm,,,{filing_hash}"), encoding="utf-8")
        (self.root / "ref").mkdir()
        (self.root / "ref/seed.csv").write_text(f"deal,target_name,form_type,date_filed,index_url,status\nkraton,KRATON,DEFM14A,2020-04-01,{index},ok\n", encoding="utf-8")
        edit(self.cockpit, "kraton")
        base_hash = self.cockpit.workspace.item("kraton")["versions"][0]["sha256"]
        relocate(self.root)
        shutil.rmtree(self.root / "extraction")
        cockpit = data.Cockpit(self.root)

        content, name = provenance.download(cockpit, "kraton")
        self.assertEqual(name, "kraton-working-r1.xlsx")
        book = load_workbook(io.BytesIO(content))
        self.assertEqual(book.sheetnames, ["Deal ledger", "Rounds", "Questions", "Deal facts", "Source"])
        values = {field: value for field, value in book["Source"].iter_rows(min_row=2, values_only=True)}
        book.close()
        self.assertEqual((values["EDGAR filing index"], values["Complete submission (.txt)"], values["Filing SHA-256"]), (index, submission, filing_hash))
        self.assertEqual((values["Version ID (working-copy base)"], values["Raw workbook SHA-256"], values["Instruction"]), (imp.VERSION_ID, base_hash, "v1.13.2"))
        self.assertEqual((values["Working revision"], values["Review status"]), ("1", "not set"))

        four, four_name = provenance.download(cockpit, "kraton", source=False)
        self.assertEqual(four_name, "kraton-working-r1.xlsx")
        path = self.root / "_download-four.xlsx"
        path.write_bytes(four)
        self.assertEqual(load_workbook(path).sheetnames, ["Deal ledger", "Rounds", "Questions", "Deal facts"])
        self.assertEqual(data.read_workbook(path), data.read_workbook(self._export_file("kraton", cockpit)))
        self.assertEqual(sha(provenance.download(cockpit, "kraton", "rev:0", source=False)[0]), base_hash)  # the relocated bytes

    def test_verify_passes_in_place(self):
        outcome = verify_catalog.verify(self.root)
        self.assertEqual({deal["path"] for deal in outcome["deals"].values()},
                         {f"extraction/{slug}.xlsx" for slug in imp.DEALS})

    def test_a_missing_catalog_workbook_is_reported_not_hidden(self):
        (self.root / "extraction/penford.xlsx").unlink()
        self.assertIn("penford", self.cockpit.slugs())
        listed = {deal["slug"]: deal for deal in self.cockpit.list_deals()}
        self.assertIn("hash does not match", listed["penford"]["error"])

    def _export_file(self, slug: str, cockpit: data.Cockpit | None = None) -> Path:
        path = self.root / f"_export-{slug}-{len(list(self.root.glob('_export-*')))}.xlsx"
        path.write_bytes((cockpit or self.cockpit).workspace.export(slug))
        return path


class VerifyCatalogTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit = nine_deal_repo(self.root, RELOCATED)

    def tearDown(self): self.temp.cleanup()

    def main(self, *args: str) -> tuple[int, str]:
        err = io.StringIO()
        with contextlib.redirect_stdout(io.StringIO()), contextlib.redirect_stderr(err):
            code = verify_catalog.main(["--root", str(self.root), *args])
        return code, err.getvalue()

    def test_unedited_and_edited_working_copies(self):
        edit(self.cockpit, "mac-gray")
        edit(self.cockpit, "mac-gray", "98")
        import_version(self.cockpit, "petsmart")
        db = self.cockpit.workspace.db_path
        before = sha(db.read_bytes())
        outcome = verify_catalog.verify(self.root)
        self.assertEqual(sha(db.read_bytes()), before)  # read-only
        self.assertEqual(set(outcome["deals"]), set(imp.DEALS))
        mac = outcome["deals"]["mac-gray"]
        self.assertEqual((mac["edited"], mac["revision"], mac["base_version"], mac["fresh_check_summary"]),
                         (True, 2, imp.VERSION_ID, None))
        stec = outcome["deals"]["stec"]
        self.assertEqual((stec["edited"], stec["revision"], stec["versions"]), (False, 0, ["working", imp.VERSION_ID]))
        self.assertEqual(stec["check_summary"], stec["fresh_check_summary"])
        self.assertEqual(outcome["deals"]["petsmart"]["versions"], ["working", imp.VERSION_ID, "run1"])

    def test_catalog_workbook_must_match_its_hash(self):
        book = load_workbook(self.root / RELOCATED / "synacor.xlsx")
        book["Deal facts"]["B2"] = "Another target"
        book.save(self.root / RELOCATED / "synacor.xlsx")
        with self.assertRaisesRegex(RuntimeError, "missing or differs from its hash: synacor"):
            verify_catalog.verify(self.root)
        (self.root / RELOCATED / "synacor.xlsx").unlink()
        with self.assertRaisesRegex(RuntimeError, "missing or differs from its hash: synacor"):
            verify_catalog.verify(self.root)

    def test_main_writes_a_new_file_and_never_overwrites(self):
        earlier = self.root / imp.REEXTRACT / "catalog-verification.json"
        earlier.parent.mkdir(parents=True, exist_ok=True)
        earlier.write_text("22 September evidence\n", encoding="utf-8")
        self.assertEqual(self.main(), (0, ""))
        written = sorted((self.root / imp.REEXTRACT).glob("catalog-verification-*.json"))
        self.assertEqual(len(written), 1)
        self.assertRegex(written[0].name, r"^catalog-verification-\d{8}T\d{6}Z\.json$")
        self.assertEqual(set(json.loads(written[0].read_text())["deals"]), set(imp.DEALS))
        self.assertEqual(earlier.read_text(encoding="utf-8"), "22 September evidence\n")
        before = tree_digest(self.root)
        for target in (earlier, written[0]):
            code, err = self.main("--output", str(target))
            self.assertEqual(code, 2)
            self.assertIn("never overwrites", err)
        self.assertEqual(tree_digest(self.root), before)
        with self.assertRaises(FileExistsError):
            verify_catalog.write_new(earlier, {})
        self.assertEqual(earlier.read_text(encoding="utf-8"), "22 September evidence\n")
        self.assertEqual(tree_digest(self.root), before)


if __name__ == "__main__":
    unittest.main()
