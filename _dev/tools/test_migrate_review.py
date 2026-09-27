"""Register, aligner, triage and port batch on a disposable cockpit root; nothing live is read."""

from __future__ import annotations

import collections
import contextlib
import csv
import datetime as dt
import hashlib
import io
import json
import re
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from openpyxl import Workbook

import check_lean
import diff_workbooks
import migrate_review as mr
from cockpit import data

SENTENCES = [f"On the reported date event number {n} took place as the filing describes it." for n in range(1, 13)]


def quote(n: int) -> str:
    return f"“{SENTENCES[n - 1]}” (p. 3)"


def event(n, who, name, day, **extra):
    values = {"#": n, "When": day.strftime("%m/%d/%Y"), "Who": who, "Event": name, "Process": 1, "Round": 1,
              "Quote and page": quote(n), "Sort date": day, "Date from": day, "Date to": day}
    values.update({key.replace("_", " "): value for key, value in extra.items()})
    return values


D = dt.datetime
BASE = [
    event(1, "Bank X", "Adviser", D(2020, 1, 5)),
    event(2, "Target", "Round opened", D(2020, 2, 1)),
    event(3, "Party A", "NDA signed", D(2020, 2, 3), Type="Strategic", Count=1),
    event(4, "Party B", "NDA signed", D(2020, 2, 3), Type="Strategic", Count=1),
    event(5, "Target", "Deadline", D(2020, 3, 1)),
    event(6, "Party A", "Bid", D(2020, 3, 1), Type="Strategic", Price_low=10, Price_high=10, All_cash="Yes", Formality="Informal", Conditions="Unclear", Count=1),
    event(7, "Party B", "Bid", D(2020, 3, 1), Type="Strategic", Price_low=12, Price_high=12, All_cash="Yes", Formality="Informal", Conditions="Heavy", Count=1, Note="Includes a $1.00 CVR."),
    event(8, "Party C", "Did not submit", D(2020, 3, 1), Type="Financial", Count=1, Exit_reason="Not stated"),
    event(9, "Party A", "Bid", D(2020, 3, 10), Type="Strategic", Price_low=11, Price_high=11, All_cash="No", Formality="Formal", Conditions="Light", Count=1),
    event(10, "Party D", "Other-scope bid", D(2020, 3, 12), Type="Strategic", Price_low=5, Price_high=5, All_cash="Not stated", Formality="Informal", Conditions="Unclear", Count=1, Note="Division offer, see #10."),
    event(11, "Party B", "Merger agreement signed", D(2020, 4, 1)),
]
TERMS = {"CVR/earnout": None, "Due diligence": "Not stated", "Financing": "Not stated", "Regulatory": "Not stated", "Exclusivity": "Not stated"}
RUN = [
    event(1, "Bank X", "Adviser", D(2020, 1, 5)),
    event(2, "Target", "Round opened", D(2020, 2, 1)),
    event(3, "Party A", "NDA signed", D(2020, 2, 3), Type="Strategic", Count=1),
    event(4, "Party B", "NDA signed", D(2020, 2, 3), Type="Financial", Count=1),
    event(5, "Party E", "Contact", D(2020, 2, 15), Type="Financial", Count=1),
    event(6, "Target", "Deadline", D(2020, 3, 1)),
    event(7, "Party A", "Bid", D(2020, 3, 1), Type="Strategic", Price_low=10, Price_high=10, Formality="Informal", Conditions="Unclear", Count=1, **{"Stock %": 0}, **TERMS),
    event(8, "Party B", "Bid", D(2020, 3, 1), Type="Strategic", Price_low=11, Price_high=11, Formality="Informal", Conditions="Heavy", Count=1,
          **{**TERMS, "Stock %": 0, "CVR/earnout": "Y", "CVR/earnout value": 1}),
    event(9, "Party A", "Bid", D(2020, 3, 10), Type="Strategic", Price_low=11, Price_high=11, Formality="Informal", Conditions="Light", Count=1, **{"Stock %": 20}, **TERMS),
    event(10, "Party D", "Other-scope bid", D(2020, 3, 12), Type="Strategic", Formality="Informal", Conditions="Unclear", Count=1, **{"Stock %": "Not stated"}, **TERMS),
    event(11, "Party B", "Merger agreement signed", D(2020, 4, 2)),
]
ROUND_BASE = [1, 1, D(2020, 2, 1), "Outreach", "Party A, Party B and Party C", "03/01/2020", "Late bids accepted", "Not final", "3", "Signed"]
ROUND_RUN = [1, 1, D(2020, 2, 1), "Outreach", "Party A and Party B", "03/01/2020", "Extended (late bid accepted)", "Not final", "2", "Signed"]
FACTS_BASE = [("Initiation", "Target-led"), ("Auction screen", "Yes"), ("Whole-company bids", "Yes"), ("Number of processes", "1")]
FACTS_RUN = [("Initiation", "target-led"), ("Auction screen", "Yes"), ("Whole-company bids", "No (Party D bid for a division)"), ("Number of processes", "1")]
README = """# Audit

## 1 Verdict

| Id | The choice | Why | Rows affected | Recommendation |
| --- | --- | --- | --- | --- |
| A1 | Not in section 2 | x | Demo #1 | x |

## 2 Questions that genuinely need Alex

| Id | The choice | Why a research choice | Rows affected | Recommendation |
| --- | --- | --- | --- | --- |
| A7 | Exclusivity and Conditions | E12 | Demo #7 | Unclear |
| A15 | Confirmations | decided | Demo #11, Rounds R1 | Yes |

## 3 Next
"""


def workbook(path: Path, ledger_columns, ledger, round_row, facts) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Deal ledger"
    ws.append(list(ledger_columns))
    for values in ledger:
        ws.append([values.get(column) for column in ledger_columns])
    rounds = wb.create_sheet("Rounds")
    rounds.append(list(check_lean.ROUND_COLUMNS))
    rounds.append(round_row)
    questions = wb.create_sheet("Questions")
    questions.append(list(check_lean.QUESTION_COLUMNS))
    questions.append(["Q1", "Round 1 deadline outcome?", "Late bids accepted", "Why (p. 3)", "#5", "The count", None])
    sheet = wb.create_sheet("Deal facts")
    sheet.append(["Field", "Value"])
    for field, value in facts:
        sheet.append([field, value])
    wb.save(path)
    return path.read_bytes()


def build_root(root: Path) -> dict:
    """A cockpit root: a v1.13.2 base, a v1.14 run, a filing and the audit's files for deal 'demo'."""
    (root / "extraction").mkdir()
    (root / "runs").mkdir()
    (root / "raw_filing").mkdir()
    (root / "_dev/cockpit").mkdir(parents=True)
    base = workbook(root / "extraction/demo.xlsx", check_lean.LEDGER_COLUMNS, BASE, ROUND_BASE, FACTS_BASE)
    run = workbook(root / "runs/demo-v114.xlsx", check_lean.LEDGER_COLUMNS_V114, RUN, ROUND_RUN, FACTS_RUN)
    body = "".join(f"<p>{sentence}</p>" for sentence in SENTENCES)
    (root / "raw_filing/demo_2020-04-01_DEFM14A.htm").write_text(
        "<html><body><p>COVER</p><p style='page-break-before:always'><hr></p>" + body + "<p align='center'><font size='2'>3</font></p></body></html>", encoding="utf-8")
    (root / "raw_filing/MANIFEST.csv").write_text("file,deal,form_type,date_filed,source_url,document,fetched_utc,bytes,sha256\n"
                                                  "demo_2020-04-01_DEFM14A.htm,demo,DEFM14A,2020-04-01,https://example.invalid/x.txt,x.htm,,,\n", encoding="utf-8")
    shas = {"base": hashlib.sha256(base).hexdigest(), "run": hashlib.sha256(run).hexdigest()}
    catalog = {"schema_version": 1, "deals": {"demo": {"name": "Demo", "default_base": "base", "findings": [], "documents": [], "versions": [
        {"id": "base", "label": "Base", "path": "extraction/demo.xlsx", "sha256": shas["base"], "instruction_version": "v1.13.2", "kind": "raw", "review_status": "unreviewed"},
        {"id": "run14", "label": "Run", "path": "runs/demo-v114.xlsx", "sha256": shas["run"], "instruction_version": "v1.14", "kind": "raw", "review_status": "unreviewed"}]}}}
    (root / "_dev/cockpit/catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    (root / "SEC_Deal_Ledger_Extraction_Instruction.md").write_text("# Synthetic\n\n**Revision of 1 January 2026, v1.13.2.**\n", encoding="utf-8")
    uids = [row["uid"] for row in mr.load_run(root / "extraction/demo.xlsx", "demo")["sheets"][mr.LEDGER]["rows"]]
    audit = root / "audit"
    (audit / "inventory").mkdir(parents=True)
    (audit / "deals").mkdir()
    with (audit / "inventory/demo.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["change_id", "sheet", "uid", "original_ref", "final_ref", "change_type", "field", "substantive", "original_value", "final_value", "revisions", "attribution"])
        writer.writerow(["demo:1", "ledger", uids[5], 6, 6, "update", "Conditions", "yes", "Heavy", "Unclear", "r1", "H"])
        writer.writerow(["demo:2", "ledger", uids[8], 9, 9, "update", "Formality", "yes", "Informal", "Formal", "r1", "A"])
        writer.writerow(["demo:3", "ledger", uids[9], 10, 10, "update", "Note", "text", "Division offer, see #9.", "Changed.", "r1", "A"])
    with (audit / "deals/demo-verdicts.csv").open("w", newline="", encoding="utf-8") as stream:
        writer = csv.writer(stream)
        writer.writerow(["change_id", "verdict", "pre_challenge_verdict", "basis", "filing_cite", "rule_or_decision", "comment"])
        writer.writerow(["demo:1", "supported", "", "existing-rule correction (E12)", "p.3 b2", "E12", ""])
        writer.writerow(["demo:2", "overstated", "", "convention change without a recorded decision", "p.3 b9", "E11", ""])
        writer.writerow(["demo:3", "incorrect", "", "existing-rule correction", "p.3 b10", "D1", ""])
    (audit / "README.md").write_text(README, encoding="utf-8")
    return {"shas": shas, "base_uids": uids}


def issue_codes(payload) -> collections.Counter:
    rows = [row for key in ("ledger", "rounds", "questions") for row in payload[key]["rows"]]
    return collections.Counter(issue["code"] for issue in [i for row in rows for i in row["issues"]] + payload["check"]["other_issues"])


class MigrateReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.info = build_root(self.root)
        self.cockpit = data.Cockpit(self.root)
        self.ws = self.cockpit.workspace
        uids = self.info["base_uids"]
        marks = [{"type": "review", "uid": uid, "status": "needs_decision" if n == 7 else "reviewed"} for n, uid in enumerate(uids, 1)]
        self.ws.edit("demo", {"revision": 0, "base_sha256": self.info["shas"]["base"], "reason": "Review", "operations": marks}, "austin")
        self.db = self.root / "_dev/cockpit/state/workspace.sqlite3"
        self.patch = mock.patch.dict(mr.AUDIT_DEALS, {"demo": "demo"})
        self.patch.start()

    def tearDown(self):
        self.patch.stop()
        self.temp.cleanup()

    def register(self):
        conn = mr.connect(self.db)
        try:
            return mr.build_register("demo", conn, self.root / "audit", self.root / "raw_filing")
        finally:
            conn.close()

    def facts(self, register, row=None, sheet=mr.LEDGER):
        return {fact["kind"]: fact for fact in register["facts"] if fact["sheet"] == sheet and (row is None or fact["row"] == row)}

    def triage(self, register=None):
        register = register or self.register()
        return register, mr.triage(register, mr.load_run(self.root / "runs/demo-v114.xlsx", "demo"), "run14", seed=3, sample=2)

    # -- read-only access -------------------------------------------------------------------

    def test_sqlite_is_opened_only_read_only(self):
        before = hashlib.sha256(self.db.read_bytes()).hexdigest()
        with mock.patch("migrate_review.sqlite3.connect", wraps=sqlite3.connect) as spy:
            self.register()
            args = mr.parser().parse_args(["register", "--state-db", str(self.db), "--audit-dir", str(self.root / "audit"),
                                           "--filings", str(self.root / "raw_filing"), "--out", str(self.root / "out")])
            with contextlib.redirect_stdout(io.StringIO()):
                args.func(args)
        self.assertGreaterEqual(spy.call_count, 2)
        for call in spy.call_args_list:
            self.assertTrue(call.kwargs.get("uri"))
            self.assertTrue(call.args[0].startswith("file:") and call.args[0].endswith("?mode=ro"))
        conn = mr.connect(self.db)
        with self.assertRaises(sqlite3.OperationalError):
            conn.execute("CREATE TABLE written (x)")
        conn.close()
        with self.assertRaises(sqlite3.OperationalError):
            mr.connect(self.root / "missing.sqlite3")  # never creates a database
        self.assertFalse((self.root / "missing.sqlite3").exists())
        self.assertEqual(hashlib.sha256(self.db.read_bytes()).hexdigest(), before)
        source = Path(mr.__file__).read_text(encoding="utf-8")
        self.assertEqual(source.count("sqlite3.connect("), 1)

    # -- register ---------------------------------------------------------------------------

    def test_row_list_carries_every_row_and_its_mark(self):
        register = self.register()
        self.assertEqual([row["#"] for row in register["rows"]], list(range(1, 12)))
        self.assertEqual(register["marks"], {"reviewed": 10, "needs_decision": 1, "unreviewed": 0})
        self.assertEqual(register["rows"][6]["mark"], "needs_decision")
        self.assertEqual(register["source"]["revision"], 1)
        out = self.root / "out"
        mr.write_register(register, out)
        args = mr.parser().parse_args(["verify", "--state-db", str(self.db), "--out", str(out)])
        with contextlib.redirect_stdout(io.StringIO()) as shown:
            args.func(args)
        self.assertIn("demo: marks match", shown.getvalue())
        with (out / "demo/rows.csv").open(newline="", encoding="utf-8") as stream:
            rows = list(csv.DictReader(stream))
        rows[0]["mark"] = "needs_decision"
        with (out / "demo/rows.csv").open("w", newline="", encoding="utf-8") as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        with contextlib.redirect_stdout(io.StringIO()), self.assertRaises(SystemExit):
            args.func(args)

    def test_facts_follow_part_a_and_carry_key_value_basis_status_and_impact(self):
        register = self.register()
        sections = [mr.SECTIONS.index(fact["section"]) for fact in register["facts"]]
        self.assertEqual(sections, sorted(sections))
        for fact in register["facts"]:
            self.assertIn(fact["basis"], ("reported", "inference", "convention"))
            self.assertIn(fact["status"], ("supported", "reverted", "unresolved"))
            self.assertIn(fact["impact"]["class"], ("unchanged", "re-judge", "new column"))
        bid = self.facts(register, 6)
        self.assertEqual(set(bid), {"assignment", "price", "stock", "formality", "conditions", "terms", "order"})
        self.assertEqual(bid["price"]["filing_key"]["page"], "3")
        self.assertRegex(bid["price"]["filing_key"]["block"], r"^b\d+$")
        self.assertEqual(bid["price"]["filing_key"]["quote"], SENTENCES[5])
        self.assertEqual(bid["price"]["value"], {"Price low": "10", "Price high": "10"})
        self.assertEqual((bid["conditions"]["status"], bid["conditions"]["basis"], bid["conditions"]["audit"]), ("supported", "reported", ["demo:1"]))
        self.assertEqual(bid["stock"]["crosswalk"], {"All cash": "Yes", "Stock %": 0})
        self.assertEqual(bid["terms"]["impact"]["class"], "new column")
        self.assertEqual(self.facts(register, 8)["participation"]["impact"]["decisions"], ["D10", "R5"])
        self.assertIn("D18", self.facts(register, 10)["price"]["impact"]["decisions"])
        revision = self.facts(register, 9)
        self.assertEqual(revision["formality"]["impact"]["decisions"], ["D9", "v1.14.1 D1", "express incorporation removed", "Formality routes"])
        self.assertEqual((revision["formality"]["status"], revision["formality"]["basis"]), ("unresolved", "convention"))
        self.assertEqual(revision["price"]["status"], "supported")
        self.assertIn("H1–H3", revision["conditions"]["impact"]["decisions"])
        rounds = self.facts(register, sheet=mr.ROUNDS)
        self.assertEqual(rounds["deadline"]["impact"]["decisions"], ["D11"])
        self.assertEqual(rounds["round"]["impact"]["decisions"], ["D8", "v1.14.1 D6", "not invited: Dropped by target"])
        facts = {fact["label"]: fact for fact in register["facts"] if fact["kind"] == "deal fact"}
        self.assertEqual([facts[f]["section"] for f in ("Initiation", "Auction screen", "Number of processes")],
                         ["participation and exits", "participation and exits", "round map and deadline outcomes"])
        self.assertEqual(facts["Initiation"]["impact"]["decisions"], ["Initiation from the first row"])
        self.assertEqual(facts["Number of processes"]["impact"]["decisions"], ["E5 process test"])
        self.assertEqual(self.facts(register, 1)["order"]["impact"]["class"], "unchanged")  # the target's adviser, a reported day
        self.assertIn("Extended (late bid accepted)", rounds["deadline"]["impact"]["note"])
        self.assertEqual(rounds["round"]["filing_key"]["quote"], SENTENCES[1])  # the Round opened row
        self.assertEqual(rounds["deadline"]["filing_key"]["quote"], SENTENCES[4])

    def test_status_from_marks_audit_verdicts_and_open_conventions(self):
        register = self.register()
        other = self.facts(register, 10)
        self.assertEqual({fact["status"] for fact in other.values()}, {"reverted"})  # incorrect Note restored, renumbered
        held = self.facts(register, 7)["conditions"]
        self.assertEqual((held["status"], held["basis"], held["basis_refs"]), ("unresolved", "convention", ["A7"]))
        self.assertIn("row marked needs decision", held["status_reasons"])
        self.assertEqual(self.facts(register, 7)["order"]["status"], "unresolved")  # needs decision alone
        self.assertEqual(self.facts(register, 11)["order"]["basis_refs"], ["A15"])
        self.assertEqual(self.facts(register, 1)["order"]["status"], "supported")  # section 1's A1 is not an open question
        self.assertEqual(self.facts(register, sheet=mr.ROUNDS)["round"]["status"], "unresolved")

    def test_audit_convention_table_parsing(self):
        audit = self.root / "parse"
        audit.mkdir()
        (audit / "README.md").write_text("## 2 Questions\n\n| Id | a | b | Rows affected | c |\n| --- | --- | --- | --- | --- |\n"
                                        "| A1 | x | y | Kraton #47, #49, Rounds R2/R3; Meredith #23, #26–#28, Rounds; sTec Deal facts | z |\n"
                                        "| A11 | x | y | Meredith 26 LMG rows, #37, Deal facts | z |\n"
                                        "| SYN-A; SYN-B | x | y | Synacor #9; #24 | z |\n| A17 (optional) | x | y | Kraton Q1's 15 rows | z |\n", encoding="utf-8")
        found = mr.audit_conventions(audit)
        self.assertEqual(found["kraton"][0]["rows"], {47, 49})
        self.assertEqual(found["kraton"][0]["rounds"], {2, 3})
        self.assertEqual(found["meredith"][0]["rows"], {23, 26, 27, 28})
        self.assertEqual(found["meredith"][0]["rounds"], {"all"})
        self.assertTrue(found["stec"][0]["facts"])
        self.assertEqual([ref["rows"] for ref in found["synacor"]], [{9}, {24}])  # the deal name carries over
        self.assertEqual((found["kraton"][1]["id"], found["kraton"][1]["questions"]), ("A17", {"Q1"}))
        self.assertEqual(found["kraton"][1]["unnamed"], [])
        lmg = found["meredith"][1]
        self.assertEqual((lmg["rows"], lmg["facts"], lmg["unnamed"]), ({37}, True, ["Meredith 26 LMG rows"]))

    def test_v1141_changes_mark_the_facts_they_touch(self):
        """R1–R6, V1141_SPEC D1–D6 and the structural cuts: a touched fact is re-judged, never accepted on agreement."""
        def tags(kind, context=None, **values):
            return mr._impact(kind, {key.replace("_", " "): value for key, value in values.items()}, context or {})["decisions"]
        day, later = dt.date(2020, 3, 1), dt.date(2020, 3, 31)
        exact = {"Date_from": day, "Date_to": day}
        # R1 and the removed express incorporation: a same-price revision and a Bid reaffirmed.
        same = {"revision": True, "same_price": True}
        self.assertIn("R1", tags("price", same, Event="Bid"))
        self.assertIn("R1", tags("price", Event="Bid reaffirmed"))
        self.assertEqual(tags("formality", same, Event="Bid", Formality="Informal"), ["D9", "v1.14.1 D1", "express incorporation removed", "R1"])
        self.assertEqual(tags("formality", Event="Bid reaffirmed", Formality="Formal"),
                         ["D9", "v1.14.1 D1", "express incorporation removed", "R1", "Formality routes"])
        self.assertEqual(tags("formality", Event="Bid", Formality="Informal"), [])  # a first informal bid: no narrower route moves it
        self.assertEqual(tags("formality", Event="Bid", Formality="Unclear"), ["Formality routes"])
        conditions = tags("conditions", {"revision": True}, Event="Bid", Conditions="Light", Note="Two weeks of confirmatory diligence")
        self.assertEqual(conditions, ["H1–H3", "R2", "express incorporation removed", "R4", "R6", "v1.14.1 D3"])
        self.assertNotIn("express incorporation", conditions)
        self.assertIn("v1.14.1 D2", tags("conditions", Event="Bid", Conditions="Heavy", Note="Will not proceed unless the standstill is waived"))
        self.assertEqual(tags("conditions", Event="Bid", Conditions="None"), ["H1–H3", "R2"])
        # R3 and R5: cohort closure and not-invited exits, in participation and in the order of events.
        self.assertEqual(tags("participation", Event="Did not submit", Count=12, Inferred="Y"), ["D10", "R3", "R5"])
        self.assertEqual(tags("participation", Event="Dropped by target", Count=1), ["R5"])
        self.assertEqual(tags("participation", Event="Re-entered", Count=1), ["R5"])
        self.assertEqual(tags("participation", Event="Withdrew", Count=1), [])
        self.assertEqual(tags("participation", Event="Not selected at signing", Count=1, Inferred="Y"), ["R5"])
        self.assertEqual(tags("participation", Event="NDA signed", Count="11–14"), ["R3"])  # a v1.13.2 count range
        self.assertEqual(tags("order", Event="Did not submit", Count=12, Inferred="Y", **exact), ["D10", "R5", "R3", "Sort-date ladder"])
        # Contact versus interest around round 1.
        for name in ("Target interest", "Bidder interest", "Contact"):
            self.assertEqual(tags("participation", Event=name, Count=1), ["Contact vs interest"])
        self.assertEqual(tags("participation", Event="NDA signed", Count=1), [])
        # The Sort-date ladder: a row without a single reported day.
        self.assertEqual(tags("order", Event="NDA signed", **exact), [])
        self.assertEqual(tags("order", Event="NDA signed", Date_from=day, Date_to=later), ["Sort-date ladder"])
        self.assertEqual(tags("order", Event="NDA signed", Date_to=day), ["Sort-date ladder"])
        # The closed Other material event list, and a condition on proceeding recorded as one (v1.14.1 D2).
        self.assertEqual(tags("order", Event="Other material event", Note="Price feedback to Party A", **exact), ["Other material event list"])
        self.assertEqual(tags("order", Event="Other material event", Note="Party B will not proceed unless waived", **exact),
                         ["v1.14.1 D2", "Other material event list"])
        # Round map (v1.14.1 D6), processes (E5), and the target's advisers only.
        self.assertEqual(tags("assignment", Event="Bid"), ["D8", "v1.14.1 D6"])
        self.assertEqual(tags("order", Event="Round opened", **exact), ["D8", "v1.14.1 D6"])
        self.assertEqual(tags("order", Event="Process restarted", Inferred="Y", **exact), ["E5 process test", "Sort-date ladder"])
        self.assertEqual(tags("order", {"names_bidder": True}, Event="Adviser", **exact), ["bidders' advisers to Notes"])
        self.assertEqual(tags("order", {"names_bidder": False}, Event="Adviser", **exact), [])
        self.assertEqual(tags("deal fact", Field="Initiation"), ["Initiation from the first row"])
        # Every tag written is either a V114_SPEC decision (D1–D27, H1–H3) or a documented v1.14.1 change.
        written = set()
        for kind in mr.FACT_FIELDS:
            for event_name in list(check_lean.EVENTS) + [""]:
                for context in ({}, {"revision": True, "same_price": True, "names_bidder": True, "partial": True}):
                    written.update(tags(kind, context, Event=event_name, Count=3, Inferred="Y", Formality="Unclear", Conditions="Light",
                                        Note="partial financing exclusivity diligence nda unless commitment", Field="Initiation"))
        self.assertTrue(written)
        undocumented = {t for t in written if not re.fullmatch(r"D\d+|H1–H3|new column: .+", t)} - set(mr.IMPACT_TAGS_V1141)
        self.assertEqual(undocumented, set())
        self.assertEqual(set(mr.IMPACT_TAGS_V1141) - written, set())  # and every documented tag is reachable

    def test_bidder_adviser_rows_are_found_from_the_ledger(self):
        def adviser_tags():
            register = self.register()
            return next(fact for fact in register["facts"] if fact["kind"] == "order" and fact["row"] == 1)["impact"]["decisions"]
        self.assertEqual(adviser_tags(), [])  # Bank X, the target's adviser
        uid = self.info["base_uids"][0]
        for revision, note in enumerate(("Financial adviser to Party A", "Financing source for the buyer"), 1):
            self.ws.edit("demo", {"revision": revision, "base_sha256": self.info["shas"]["base"], "reason": "Adviser",
                                  "operations": [{"type": "update", "sheet": mr.LEDGER, "uid": uid, "values": {"Note": note}}]}, "austin")
            self.assertEqual(adviser_tags(), ["bidders' advisers to Notes"], note)

    def test_a_run_is_read_under_the_rules_asked_for(self):
        path = self.root / "runs/demo-v114.xlsx"
        self.assertEqual(mr.load_run(path, "demo")["schema"], check_lean.SCHEMA_V1141)
        self.assertEqual(mr.load_run(path, "demo", check_lean.SCHEMA_V114)["schema"], check_lean.SCHEMA_V114)
        self.assertEqual(mr.load_run(self.root / "extraction/demo.xlsx", "demo", check_lean.SCHEMA_V114)["schema"], check_lean.SCHEMA_V1132)
        args = mr.parser().parse_args(["triage", "demo", str(path), "--run-id", "run14", "--rules", "v1.14", "--state-db", str(self.db),
                                       "--audit-dir", str(self.root / "audit"), "--filings", str(self.root / "raw_filing"), "--out", str(self.root / "out")])
        with contextlib.redirect_stdout(io.StringIO()):
            args.func(args)
        result = json.loads((self.root / "out/demo/triage-run14.json").read_text(encoding="utf-8"))
        self.assertEqual(result["run_schema"], check_lean.SCHEMA_V114)
        with contextlib.redirect_stderr(io.StringIO()), self.assertRaises(SystemExit):
            mr.parser().parse_args(["triage", "demo", str(path), "--run-id", "x", "--rules", "v1.13.2"])

    # -- crosswalk, parties and aligner ---------------------------------------------------------

    def test_stock_crosswalk(self):
        cases = [("Yes", 0, True), ("Yes", 20, False), ("No", 20, True), ("No", "20-40", True), ("No", "Part stock", True),
                 ("No", 0, False), ("Not stated", "Not stated", True), ("Not stated", 0, False), ("Yes", "Varies", None)]
        cases += [("No", 150, False), (None, None, True), ("Maybe", 0, False)]  # S7: above 100 is no figure; blank pairs agree
        for all_cash, stock, expected in cases:
            self.assertIs(mr.stock_agrees(all_cash, stock), expected, (all_cash, stock))
            if expected is not None:  # the one mapping: S7's crosswalk
                self.assertEqual(expected, diff_workbooks.crosswalk("Stock %", all_cash, stock, {})[0] == "same")
        self.assertEqual([mr.stock_from_all_cash(v) for v in ("Yes", "No", "Not stated")], [0, None, "Not stated"])
        for stock in (0, "Not stated"):
            self.assertEqual(diff_workbooks.all_cash_for_stock(mr.stock_from_all_cash(diff_workbooks.all_cash_for_stock(stock))),
                             diff_workbooks.all_cash_for_stock(stock))

    def test_parties_are_matched_on_normalized_names(self):
        self.assertTrue(mr.same_party("Genesee & Wyoming Inc. (G&W)", "G&W"))
        self.assertTrue(mr.same_party("Moab Partners, L.P. (Moab)", "Moab Partners, L.P."))
        self.assertTrue(mr.same_party("PWRR (target)", "Providence and Worcester (target)"))
        self.assertTrue(mr.same_party("9 IOI bidders (unnamed)", "9 IOI bidders"))
        self.assertFalse(mr.same_party("Party A", "Party B"))
        self.assertFalse(mr.same_party("2 lowest IOI bidders (unnamed)", "9 IOI bidders (unnamed)"))
        self.assertFalse(mr.same_party("Party L", "Party P"))
        self.assertFalse(mr.same_party("Party A", "Party C (a strategic buyer)"))  # the article is not the letter
        self.assertFalse(mr.same_party("Party A", "Party B, a financial sponsor"))
        self.assertTrue(mr.same_party("Party A", "Party A (a strategic buyer)"))
        self.assertTrue(mr.same_party("Moab / CSC/Pamplona", "Moab Partners, L.P. / CSC/Pamplona"))

    def test_aligner_accepts_unique_matches_only(self):
        def record(uid, who, name, day, **extra):
            return {"uid": uid, "values": event(0, who, name, day, **extra)}
        old = [record("a", "Party A", "Bid", D(2020, 3, 1), Price_low=10, Price_high=10),
               record("b", "Party B", "NDA signed", D(2020, 2, 1)), record("c", "Party B", "NDA signed", D(2020, 2, 1)),
               record("d", "Party C", "Exclusivity changed", D(2020, 4, 1))]
        new = [record("x", "Party A", "Bid", D(2020, 3, 1), Price_low=9, Price_high=9, **{"CVR/earnout value": 1}),
               record("y", "Party B", "NDA signed", D(2020, 2, 1)), record("z", "Party C", "Bid", D(2020, 4, 1))]
        result = mr.align_ledger(old, new)
        self.assertEqual(result["matches"]["a"]["uid"], "x")
        self.assertEqual(result["matches"]["a"]["price"], "with CVR")
        self.assertNotIn("b", result["matches"])  # two reviewed rows fit one run row: no match
        self.assertEqual(sorted(result["omitted"]), ["b", "c", "d"])
        self.assertEqual(result["inserted"], ["y", "z"])  # never across event families
        entry = [record("e", "Party C", "Contact", D(2020, 1, 1), Date_to=D(2020, 1, 2)), record("f", "Party C", "NDA signed", D(2020, 1, 2))]
        matched = mr.align_ledger(entry, [record("g", "Party C", "Contact", D(2020, 1, 1)), record("h", "Party C", "NDA signed", D(2020, 1, 2))])["matches"]
        self.assertEqual({uid: m["uid"] for uid, m in matched.items()}, {"e": "g", "f": "h"})  # the Event itself decides first

    def test_rounds_align_by_opening_date_and_members(self):
        vocabulary = {("party", "a"), ("party", "b"), ("party", "c"), ("party", "d")}
        def line(uid, opened, who):
            return {"uid": uid, "values": {"Process": 1, "Round": 1, "Opened": opened, "Who was in": who}}
        old = [line("r1", dt.date(2020, 2, 1), "Party A and Party B"), line("r2", dt.date(2020, 3, 1), "Party A, Party B and Party C"),
               line("r3", dt.date(2020, 6, 1), "Party C and Party D"), line("r4", dt.date(2020, 8, 1), "Party A")]
        new = [line("n1", dt.date(2020, 2, 3), "Party A and Party B"),  # two days later, same members
               line("n2", dt.date(2020, 3, 1), "Party D"),  # same date, other members
               line("n3", dt.date(2020, 5, 27), "Party C and Party D"),  # five days earlier, same members
               line("n4", dt.date(2020, 9, 1), "Party A"), line("n5", dt.date(2020, 10, 1), "Party A")]  # two fit r4: none taken
        result = mr.align_rounds(old, new, vocabulary)
        self.assertEqual({uid: (m["uid"], m["pass"]) for uid, m in result["matches"].items()},
                         {"r1": ("n1", 1), "r2": ("n2", 2), "r3": ("n3", 3)})
        self.assertEqual((result["omitted"], result["inserted"]), (["r4"], ["n4", "n5"]))
        self.assertEqual(mr.align_rounds([line("x", dt.date(2020, 2, 1), "Party A")], [line("y", dt.date(2020, 2, 8), "Party B")], vocabulary)["matches"], {})

    def test_run_uids_are_the_ones_the_cockpit_gives_a_base(self):
        payload = self.ws.deal("demo", "run14")
        run = mr.load_run(self.root / "runs/demo-v114.xlsx", "demo")
        self.assertEqual([row["uid"] for row in payload["ledger"]["rows"]], [row["uid"] for row in run["sheets"][mr.LEDGER]["rows"]])
        self.assertEqual([row["uid"] for row in payload["rounds"]["rows"]], [row["uid"] for row in run["sheets"][mr.ROUNDS]["rows"]])
        self.assertEqual(run["schema"], check_lean.SCHEMA_V1141)

    # -- triage and port batch ------------------------------------------------------------------

    def test_triage_puts_every_fact_in_one_of_four_buckets(self):
        register, result = self.triage()
        where = {item["fact"]: bucket for bucket, items in result["buckets"].items() for item in items if "fact" in item}
        ids = {fact["id"]: fact for fact in register["facts"]}
        self.assertTrue(all(result["summary"][bucket]["facts"] > 0 for bucket in mr.BUCKETS))
        for bucket, items in result["buckets"].items():
            self.assertEqual(result["summary"][bucket]["facts"], len(items))
        by_row = lambda n, kind: next(f["id"] for f in register["facts"] if f["row"] == n and f["kind"] == kind)
        self.assertEqual(where[by_row(6, "price")], "agrees")
        self.assertEqual(where[by_row(6, "stock")], "agrees")
        self.assertEqual(where[by_row(9, "stock")], "agrees")  # No and 20
        self.assertEqual(where[by_row(7, "price")], "differs, v1.14.1 changed the rule")  # 11 plus a CVR of 1
        self.assertEqual(where[by_row(9, "formality")], "differs, v1.14.1 changed the rule")  # D9
        self.assertEqual(where[by_row(10, "price")], "differs, v1.14.1 changed the rule")  # D18
        self.assertEqual(where[by_row(4, "participation")], "differs, rule unchanged")  # Type
        self.assertEqual(where[by_row(11, "order")], "differs, rule unchanged")  # a day later, matched with slack
        order = next(f for f in register["facts"] if f["id"] == by_row(3, "order"))
        same_days = event(3, "Party A", "NDA signed", D(2020, 2, 3))
        self.assertTrue(mr._compare(order, same_days, {**same_days, "When": "on 02/03/2020"}, {"in_order": True}, set())[0])  # wording of When alone
        self.assertEqual(where[by_row(8, "participation")], "omitted or inserted")
        rounds = {f["kind"]: f["id"] for f in register["facts"] if f["sheet"] == mr.ROUNDS}
        self.assertEqual(where[rounds["deadline"]], "differs, v1.14.1 changed the rule")
        self.assertEqual(where[rounds["round"]], "differs, v1.14.1 changed the rule")  # Party C left the members
        facts = {f["label"]: f["id"] for f in ids.values() if f["kind"] == "deal fact"}
        self.assertEqual(where[facts["Whole-company bids"]], "differs, v1.14.1 changed the rule")  # D7
        self.assertEqual(where[facts["Auction screen"]], "agrees")
        inserted = [item for item in result["buckets"]["omitted or inserted"] if item.get("side") == "inserted in the run"]
        self.assertEqual([item["label"] for item in inserted], ["#5 Contact · Party E"])
        self.assertNotIn(by_row(6, "terms"), where)
        self.assertEqual(len(result["new_columns"]), 4)
        self.assertEqual(result["spot_check"], self.triage()[1]["spot_check"])  # seeded
        self.assertEqual(len(result["spot_check"]), 2)
        mr.write_triage(result, mr.port_batch(register, result), self.root / "out")
        text = (self.root / "out/demo/triage-run14.md").read_text(encoding="utf-8")
        for bucket in mr.BUCKETS:
            self.assertIn(f"## {bucket} ({result['summary'][bucket]['facts']})", text)

    def test_an_agreement_accepts_only_supported_facts_whose_rule_is_unchanged(self):
        register, result = self.triage()
        ids = {fact["id"]: fact for fact in register["facts"]}
        by_row = lambda n, kind: next(f["id"] for f in register["facts"] if f["row"] == n and f["kind"] == kind)
        agrees = {item["fact"]: item for item in result["buckets"]["agrees"]}
        # A repeated unresolved fact: #7's Conditions (Needs decision, open question A7) is Heavy in both.
        repeated = agrees[by_row(7, "conditions")]
        self.assertEqual((repeated["run"], repeated["list"]), ({"Conditions": "Heavy"}, "review"))
        self.assertIn("status unresolved", repeated["hold"][0])
        self.assertIn("row marked needs decision", repeated["hold"][0])
        self.assertEqual(agrees[by_row(10, "order")]["list"], "review")  # reverted: restored, not confirmed
        self.assertEqual(agrees[by_row(6, "stock")]["list"], "review")  # a new column inherits no acceptance
        self.assertTrue(any("new column: Stock %" in reason for reason in agrees[by_row(6, "stock")]["hold"]))
        self.assertEqual(agrees[by_row(1, "order")]["list"], "accept")
        for item in agrees.values():
            fact = ids[item["fact"]]
            settled = fact["status"] == "supported" and fact["impact"]["class"] == "unchanged"
            self.assertEqual(item["list"] == "accept", settled, item["fact"])
            self.assertEqual(bool(item["hold"]), not settled)
        self.assertTrue(set(result["spot_check"]) <= {f for f, item in agrees.items() if item["list"] == "accept"})
        summary = result["summary"]["agrees"]
        self.assertEqual(summary["accept"]["facts"] + summary["review"]["facts"], summary["facts"])
        self.assertEqual(summary["review"]["facts"], sum(1 for item in agrees.values() if item["list"] == "review"))
        # Even when a reviewer accepts the repeated fact, its row is never ported as reviewed.
        batch = mr.port_batch(register, result, {by_row(7, "conditions")})
        run_uid = {p["reviewed_row"]: p["run_uid"] for p in result["alignment"]["pairs"]}
        marks = {op["uid"]: op["status"] for request in batch["requests"] for op in request["operations"] if op["type"] == "review"}
        self.assertNotEqual(marks.get(run_uid[7]), "reviewed")
        mr.write_triage(result, batch, self.root / "out")
        text = (self.root / "out/demo/triage-run14.md").read_text(encoding="utf-8")
        self.assertIn(f"### accept, after the seeded spot check ({summary['accept']['facts']})", text)
        self.assertIn(f"### review: agrees, but not settled ({summary['review']['facts']})", text)

    def test_port_batch_is_an_edit_request_that_applies_after_a_rebase(self):
        register, result = self.triage()
        by_row = lambda n, kind: next(f["id"] for f in register["facts"] if f["row"] == n and f["kind"] == kind)
        deadline = next(f["id"] for f in register["facts"] if f["kind"] == "deadline")
        accepted = {by_row(4, "participation"), by_row(11, "order"), by_row(7, "price"), deadline}
        batch = mr.port_batch(register, result, accepted)
        self.assertFalse(batch["applied"])
        self.assertEqual({item["fact"] for item in batch["skipped"]}, {by_row(7, "price"), deadline})
        operations = [op for request in batch["requests"] for op in request["operations"]]
        updates = [op for op in operations if op["type"] == "update"]
        self.assertEqual([op["values"] for op in updates], [{"Type": "Strategic"}, {"When": "04/01/2020", "Sort date": "2020-04-01", "Date from": "2020-04-01", "Date to": "2020-04-01"}])
        marks = {op["uid"]: op["status"] for op in operations if op["type"] == "review"}
        run_uid = {p["reviewed_row"]: p["run_uid"] for p in result["alignment"]["pairs"]}
        # Reviewed only where every fact is supported and unchanged by v1.14: never a bid row (new condition columns),
        # nor Round opened (D8) or Deadline (D11).
        self.assertEqual(marks, {**{run_uid[n]: "reviewed" for n in (1, 3, 4)}, **{run_uid[n]: "needs_decision" for n in (2, 5, 6, 11)}})
        with mock.patch.object(mr, "MAX_OPERATIONS", 4):
            self.assertEqual([len(r["operations"]) for r in mr.port_batch(register, result, accepted)["requests"]], [4, 4, 1])
        # Rebase the disposable working copy onto the run, then apply the batch as the edit API receives it.
        self.ws.edit("demo", {"revision": 1, "base_sha256": self.info["shas"]["base"], "reason": "Rebase", "operations": [{"type": "rebase", "target_version": "run14"}]}, "austin")
        before = issue_codes(self.ws.deal("demo"))
        request = dict(batch["requests"][0], revision=2)
        payload = self.ws.edit("demo", request, "austin")
        self.assertEqual(payload["workspace"]["revision"], 3)
        self.assertEqual({uid: mark["status"] for uid, mark in payload["row_review"].items()}, marks)
        self.assertEqual(payload["ledger"]["rows"][3]["cells"]["Type"], "Strategic")
        self.assertEqual(issue_codes(payload) - before, collections.Counter())  # the ported values add no checker issue

    def test_port_batch_records_accepted_facts_it_cannot_place(self):
        register, result = self.triage()
        omitted = next(f["id"] for f in register["facts"] if f["row"] == 8 and f["kind"] == "participation")
        batch = mr.port_batch(register, result, {omitted})
        self.assertEqual([item["fact"] for item in batch["skipped"]], [omitted])
        self.assertIn("insert must be prepared by hand", batch["skipped"][0]["reason"])
        self.assertFalse(any(op["type"] == "update" for request in batch["requests"] for op in request["operations"]))


if __name__ == "__main__":
    unittest.main()
