#!/usr/bin/env python3
"""Turn one Version 1 ledger workbook into estimation tables.

    python3 _dev/tools/derive_analysis.py LEDGER.xlsx --out DIR [--deal SLUG]

The tool is mechanical: it never corrects the ledger, and every research choice awaiting Alex is a
switch with no default, emitted as side-by-side columns. It writes only into the --out folder, which
must be new or empty, and never under `extraction/`, `raw_filing/` or `ref/`.

Outputs: bids.csv, other_scope.csv, rounds.csv, participation.csv, deal.csv and manifest.json. The
manifest's "review" list holds what a person should look at: disagreements with the Rounds sheet,
partial-only parties, counts that cannot be parsed, and the like.

The workbook must have the Version 1 29-column Deal ledger; any other header is an error. An exit's
Inferred = Y is an inferred exit, and Initiation is derived by D5's rule and checked against Deal
facts. Every bid gets upfront_price_kind (a blank or invalid price is no price observation) and
same_offer_of (a Note beginning "Same as #n"). A five-sheet cockpit download's Source sheet is read
for provenance.
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import hashlib
import json
import math
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl

import check_lean

TOOL_VERSION = "Version 1"
PROJECT = Path(__file__).resolve().parents[2]
FORBIDDEN_OUT = ("extraction", "raw_filing", "ref")
# Settled: Alex keeps Meredith for descriptive work only.
DESCRIPTIVE_ONLY = {"meredith"}

WHOLE_BIDS = {"Bid", "Bid reaffirmed"}
ENTRY_EVENTS = {"NDA signed", "Bid", "Bid reaffirmed"}
# A party's own contact row: after the cohort entries, it marks a new entrant, not a cohort member (E3).
CONTACT_EVENTS = {"Contact", "Bidder interest", "Target interest"}
# Entries whose unit takes back exactly what it added when it exits with the same inexact count.
OWN_ENTRIES = {"entry", "re-entry", "entry (late contact)"}
# A cohort Note's "less X" or "excluding X" clause names parties outside the cohort.
COHORT_EXCLUSION_RE = re.compile(r"\b(?:less|excluding|except|other than|not including)\b[^;.]*", re.IGNORECASE)
FINAL = {"Announced as final", "Inferred final"}
# One class per due date. Longest labels first for prefix matching.
DEADLINE_CLASSES = {
    "Extended (late bid accepted)": "extended",
    "Passed without action": "soft",
    check_lean.NO_DEADLINE: "no deadline",
    "Enforced": "hard",
    "Extended": "extended",
    "Unclear": "missing",
}
EXIT_ACTOR = {"Dropped by target": "target", "Withdrew": "bidder", "Did not submit": "bidder", "Not selected at signing": "target"}
# Alex's drop codes (audit D §3.3). The audit pairs DropM and DropBelowM with the two market-price
# reasons without saying which is which, so both codes are given for either reason.
MARKET_CODES = "DropM|DropBelowM"
WITHDREW_CODES = {
    "Value below market price": MARKET_CODES,
    "Value at or below market price": MARKET_CODES,
    "Value below earlier offer": "DropBelowInf",
    "Value at earlier offer": "DropAtInf",
}
INITIATION_EVENTS = {"Target interest": "target-led", "Target sale decision": "target-led",
                     "Bidder interest": "bidder-led", "Bid": "bidder-led", "Activist": "activist-influenced"}
# P1: what the two price cells give, from validated values only.
USABLE_PRICE_KINDS = {"point", "range", "lower_bound", "upper_bound"}
# E13 leaves a Stock % range to the Note ("Part stock", the range in the Note); a percentage range read there.
NOTE_STOCK_RANGE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*%?\s*(?:[-–]|to)\s*(\d+(?:\.\d+)?)\s*%")
NUMBER_WORDS = {word: n for n, word in enumerate(
    "zero one two three four five six seven eight nine ten eleven twelve thirteen fourteen fifteen sixteen "
    "seventeen eighteen nineteen twenty".split())}
NUMBER_WORDS.update({"thirty": 30, "forty": 40, "fifty": 50, "sixty": 60, "seventy": 70, "eighty": 80, "ninety": 90})
# (lo, hi, point) for each qualifier check_lean.COUNT_QUALIFIER_RE accepts, given the figure after it. "Approximately
# N is N" is the presumed point; "several" needs no figure and, being plural, is at least two.
COUNT_QUALIFIER_BOUNDS = {
    "at least": lambda n: (n, None, None),
    "more than": lambda n: (n + 1, None, None), "over": lambda n: (n + 1, None, None),
    "at most": lambda n: (1, n, None), "up to": lambda n: (1, n, None),
    "fewer than": lambda n: (1, n - 1, None), "less than": lambda n: (1, n - 1, None),
    "approximately": lambda n: (1, None, n), "about": lambda n: (1, None, n), "around": lambda n: (1, None, n),
    "some": lambda n: (1, None, n),
    "nearly": lambda n: (1, n, n),
    "several": lambda n: (2, None, None),
}

SWITCHES = [
    {"id": "count_ranges", "source": "D22; Alex Q1 / Decision 3",
     "question": "How estimation uses counts that are ranges or bounds.",
     "variants": {"bounds": "live_lo, live_hi (and count_lo, count_hi)",
                  "presumed ordinary sequence": "live_point (a party first seen bidding after cohort entries is presumed a member unless its own contact row follows them and no cohort Note names it or cites its row; approximately N is N)",
                  "one bound": "live_lo or live_hi alone"}},
    {"id": "unclear", "source": "D22",
     "question": "How Unclear and Not stated values are treated.",
     "variants": {"missing": "Formality Unclear is missing under every reading; deadline Unclear has class missing",
                  "Conditions Unclear as not Heavy": "T1", "Conditions Unclear as Heavy": "T1u"}},
    {"id": "formality_reading", "source": "D22; Alex Decision 3b",
     "question": "Which Formality reading is primary.",
     "variants": {"T0": "T0", "T1": "T1", "T1u": "T1u", "T2": "T2", "T3": "T3"}},
    {"id": "same_price_revisions", "source": "D22; D13; Alex Decision 3b; P1",
     "question": "Whether a same-price revision of terms is a new price observation.",
     "variants": {"new observation": "price_obs__same_price_as_new", "change of terms only": "price_obs__same_price_as_terms"}},
    {"id": "same_offer_restatements", "source": "E10 (R1); questionnaire 3.3(b)",
     "question": "Whether a Same-offer row (the bidder says its offer stands or confirms it by documents) is kept as a bid observation.",
     "variants": {"kept": "every bids.csv row", "dropped": "bids.csv rows with same_offer_of blank"}},
    {"id": "inferred_exits", "source": "D22; Alex Decision 3b",
     "question": "Whether inferred exits enter as dropouts or as censoring.",
     "variants": {"dropout": "exit__inferred_as_dropout",
                  "censoring": "exit__inferred_as_censoring"}},
    {"id": "estimation_price", "source": "audit D §4; spec §7.10",
     "question": "Upfront or package (upfront + CVR/earnout value) as the estimation price.",
     "variants": {"upfront": "price_low, price_high", "package": "package_low, package_high, compared only within one package_basis"}},
]


BID_COLUMNS = [
    "deal", "process", "round", "row", "when", "sort_date", "date_from", "date_to", "who", "unit", "type", "event",
    "count", "count_lo", "count_hi", "price_low", "price_high", "upfront_price_kind", "cvr_earnout", "cvr_value", "package_low",
    "package_high", "package_basis",
    "stock_pct", "stock_kind", "stock_lo", "stock_hi", "all_cash", "formality", "conditions", "due_diligence", "financing",
    "regulatory", "antitrust", "exclusivity", "inferred", "flag", "same_offer_of", "same_price_revision",
    "price_obs__same_price_as_new", "price_obs__same_price_as_terms", "round_finality", "T0", "T1", "T1u", "T2", "T3",
    "live_lo", "live_hi", "live_point", "note",
]
OTHER_COLUMNS = ["deal", "process", "round", "row", "when", "sort_date", "who", "type", "event", "reason", "count",
                 "price_low", "price_high", "stock_pct", "formality", "conditions", "note"]
ROUND_COLUMNS = [
    "deal", "process", "round", "opened", "how_opened", "finality", "who_was_in", "who_was_in_count",
    "due_dates", "due_dates_reached", "deadline_outcome", "deadline_values", "deadline_classes",
    "deadline_rows", "bids_received", "bids_received_count", "bidders_bid_lo", "bidders_bid_hi", "bidders_bid",
    "bids_received_check", "live_open_lo", "live_open_hi", "live_open_point", "live_max_hi", "how_it_ended",
]
PARTICIPATION_COLUMNS = [
    "deal", "process", "round", "row", "when", "sort_date", "date_from", "date_to", "who", "unit", "type", "event",
    "change", "inferred", "exit_inferred", "exit_actor", "exit_reason", "alex_drop_code", "exit__inferred_as_dropout",
    "exit__inferred_as_censoring", "count_kind", "count_lo", "count_hi", "count_point", "delta_lo", "delta_hi",
    "delta_point", "live_lo", "live_hi", "live_point", "note",
]
DEAL_COLUMNS = [
    "deal", "process", "target", "acquirer", "acquirer_type", "agreed_price", "merger_signed", "merger_announced",
    "filing", "background_pages", "number_of_processes", "earlier_approaches", "currency_units", "auction_screen",
    "auction_status", "auction_count_lo", "auction_count_hi", "auction_met", "whole_company_bids", "whole_company",
    "descriptive_only", "estimation_sample", "initiation_recorded", "initiation_first_event", "initiation_first_row",
    "initiation_check",
    "merger_of_equals_rows",
]
MANIFEST_KEYS = ("tool", "tool_version", "generated_at", "deal", "input", "ledger_schema", "checker_version", "readings",
                 "switches", "deadline_outcomes", "outputs", "review", "warnings")


class DeriveError(ValueError):
    pass


# ---- values ------------------------------------------------------------------------------------

def text(value: Any) -> str:
    return "" if check_lean.is_blank(value) else check_lean.normalize_contiguous(value)


def number(value: Any) -> float | None:
    if isinstance(value, bool) or check_lean.is_blank(value):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).strip())
    except ValueError:
        return None


def cell(value: Any) -> Any:
    """A CSV cell: blank for None, ISO for dates, integers without a decimal point."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        value = value.date()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value, 6)
    return value


def as_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, str):
        try:
            return dt.date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def unit_key(who: Any) -> str:
    """The name a bidder unit is followed by: Who without parentheticals, case-folded (shared with check_lean.py)."""
    return check_lean.unit_key(who)


def cohort_includes(note: Any, key: str, number: int | None) -> bool:
    """Whether an entry-cohort Note counts this unit inside the cohort (E3): it names the unit or cites its row
    outside a "less ..." clause. A list after the plural also names it: "Parties E and F" names party e."""
    inside = COHORT_EXCLUSION_RE.sub(" ", text(note)).casefold()
    if number is not None and "#" in inside and number in check_lean.parse_affected_rows(inside)[0]:
        return True
    if re.search(rf"(?<!\w){re.escape(key)}(?!\w)", inside):
        return True
    m = re.fullmatch(r"(.+?)\s+(\w{1,3})", key)
    if not m:
        return False
    base, label = m.groups()
    plural = base[:-1] + "ies" if base.endswith("y") else base + "s"
    for found in re.finditer(rf"(?<!\w){re.escape(plural)}\s+(\w{{1,3}}(?:\s*(?:,|&|\band\b)\s*\w{{1,3}})+)", inside):
        if label in re.split(r"\s*(?:,|&|\band\b)\s*", found.group(1)):
            return True
    return False


def leading_figure(value: str) -> int | None:
    """The whole number at the start of value, in digits or words ("ten", "twenty-five"); else None."""
    m = re.match(r"\s*(?:(\d+)|([a-z]+)(?:-([a-z]+))?)\b", value.casefold())
    if not m:
        return None
    if m.group(1):
        return int(m.group(1))
    first, second = NUMBER_WORDS.get(m.group(2)), NUMBER_WORDS.get(m.group(3) or "")
    if m.group(3) is None:
        return first
    return first + second if first and first >= 20 and first % 10 == 0 and second and second < 10 else None


def count_bounds(count: Any, note: Any) -> tuple[int | None, int | None, int | None, str]:
    """(lo, hi, point, kind) for a row's bidder units: Count, else the Note's "Count: ..." prefix, read with the
    checker's patterns (a range, a qualifier in check_lean.COUNT_QUALIFIER_RE, unknown or not stated, a figure)."""
    exact = check_lean.as_integer(count)
    if exact is not None and exact > 0:
        return exact, exact, exact, "exact"
    note = text(note)
    found = re.search(r"\bCount:\s*[\"\u201c\u2018']?", note, re.IGNORECASE)
    if not found:
        return 1, None, None, "unknown (no Count: prefix)"
    if (m := check_lean.COUNT_RANGE_RE.search(note)):
        lo, hi = (int(n) for n in re.findall(r"\d+", m.group(0)))
        return lo, hi, None, "range"
    if (m := check_lean.COUNT_QUALIFIER_RE.search(note)):
        word = re.sub(r"(?i)^Count:\s*[\"\u201c\u2018']?", "", m.group(0)).casefold()
        n = leading_figure(note[m.end():])
        if n is None and word != "several":
            return 1, None, None, "unknown (unparsed Count: prefix)"
        return COUNT_QUALIFIER_BOUNDS[word](n) + (word,)
    if check_lean.COUNT_UNKNOWN_RE.search(note):
        return 1, None, None, "unknown"
    if (n := leading_figure(note[found.end():])) is not None:
        return 1, None, n, "figure in Note"
    return 1, None, None, "unknown (unparsed Count: prefix)"


def stock_bounds(value: Any, note: Any = None) -> tuple[str, float | None, float | None]:
    """(kind, lo, hi) for a Stock % cell. E13 records a stated range as Part stock with the range in the Note, so a
    Part stock row takes its bounds from a percentage range in the Note."""
    if check_lean.is_blank(value):
        return "blank", None, None
    if (n := number(value)) is not None:
        return "exact", n, n
    raw = text(value)
    if (m := check_lean.STOCK_RANGE_RE.fullmatch(raw)):
        return "range (not a Version 1 value)", float(m.group(1)), float(m.group(2))
    if raw == "Part stock" and (m := NOTE_STOCK_RANGE_RE.search(text(note))):
        lo, hi = float(m.group(1)), float(m.group(2))
        if 0 <= lo <= hi <= 100:
            return "part stock (range in the Note)", lo, hi
    return {"Part stock": "part stock", "Not stated": "not stated", "Varies": "varies"}.get(raw, "unparsed"), None, None


def all_cash(record: dict[str, Any]) -> int | None:
    """1 if all cash (Stock % 0), 0 if any stock, None if unknown."""
    stock = record.get("Stock %")
    if isinstance(stock, bool) or check_lean.is_blank(stock):
        return None
    if isinstance(stock, (int, float)):
        return 1 if stock == 0 else 0 if 0 < stock <= 100 else None
    value = str(stock).strip()
    if value == "Part stock":
        return 0
    match = check_lean.STOCK_RANGE_RE.fullmatch(value)
    return 0 if match and float(match.group(2)) > 0 else None


def valid_price(value: Any) -> float | None:
    """A price cell the checker accepts as a price (bid.price_type): a finite number above 0; else None."""
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0:
        return None
    return float(value)


def upfront_price_kind(record: dict[str, Any]) -> str:
    """P1: what Price low and Price high give, from validated values only. A point needs two valid equal prices, a
    range two valid ordered unequal ones; a filled cell that is not a valid price, or a reversed range, is invalid."""
    cells = [record.get("Price low"), record.get("Price high")]
    filled = [not check_lean.is_blank(c) for c in cells]
    low, high = (valid_price(c) for c in cells)
    if (filled[0] and low is None) or (filled[1] and high is None):
        return "invalid"
    if low is not None and high is not None:
        return "point" if low == high else "range" if low < high else "invalid"
    return "lower_bound" if low is not None else "upper_bound" if high is not None else "not_available"


def package_basis(record: dict[str, Any], low: float | None, high: float | None,
                  cvr: float | None) -> tuple[str, str | None]:
    """(what package_low and package_high are, a review issue or None). E13 fixes the basis: the CVR/earnout value is
    the stated per-share amount, the maximum where several, and a package stated only as a whole leaves the price
    cells blank."""
    marker = text(record.get("CVR/earnout"))
    if low is None and high is None:
        return "missing: no upfront price", None
    if any(not check_lean.is_blank(record.get(c)) and valid_price(record.get(c)) is None for c in ("Price low", "Price high")):
        return "missing: a price cell is not a valid price", None
    if upfront_price_kind(record) == "invalid":
        return "missing: the price range is reversed", None
    if marker == "Varies":
        return "missing: CVR/earnout Varies on a cohort row (amounts in the Note)", None
    if marker == "Y" and cvr is None:
        return "missing: CVR/earnout marked without a value", None
    if cvr is None:
        return "upfront only (no CVR/earnout)", None
    if marker != "Y":
        return ("upfront + CVR/earnout value, with CVR/earnout not marked Y",
                "a CVR/earnout value without CVR/earnout Y: the package adds it anyway; check the row")
    return "upfront + CVR/earnout value (the stated per-share amount, the maximum where several; E13)", None


def deadline_class(value: str) -> str | None:
    """The class of one Deadline outcome value; None if it matches no label."""
    folded = value.casefold()
    for label, klass in DEADLINE_CLASSES.items():
        if folded == label.casefold() or (folded.startswith(label.casefold()) and not value[len(label)].isalnum()):
            return klass
    return None


def split_outcomes(value: Any) -> list[str]:
    return [part.strip() for part in text(value).split(";") if part.strip()]


def leading_integer(value: Any) -> int | None:
    m = re.match(r"\s*(\d+)\b", text(value))
    return int(m.group(1)) if m else None


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


# ---- workbook ----------------------------------------------------------------------------------

def read_sheet(ws: Any) -> list[dict[str, Any]]:
    rows = ws.iter_rows(values_only=True)
    header = [text(h) for h in next(rows, ())]
    out = []
    for values in rows:
        if all(check_lean.is_blank(v) for v in values):
            continue
        out.append({h: v for h, v in zip(header, values) if h})
    return out


def read_source(ws: Any) -> dict[str, Any]:
    """A Source sheet's label/value pairs, with any hyperlink targets, whatever its layout."""
    found: dict[str, Any] = {}
    for row in ws.iter_rows():
        cells = [c for c in row if not check_lean.is_blank(c.value)]
        if len(cells) < 2:
            continue
        label, value = text(cells[0].value), cells[1]
        found[label] = cell(value.value) if not isinstance(value.value, str) else text(value.value)
        if value.hyperlink is not None and value.hyperlink.target:
            found[label + " (link)"] = value.hyperlink.target
    return found


def load(path: Path) -> dict[str, Any]:
    """Read the current ledger workbook; any other Deal ledger header is an error."""
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as exc:
        raise DeriveError(f"{path}: not a readable workbook ({type(exc).__name__}: {exc})") from exc
    try:
        missing = [s for s in check_lean.SHEETS if s not in wb.sheetnames]
        if missing:
            raise DeriveError(f"{path}: missing sheets {', '.join(missing)}")
        header = [text(h) for h in next(wb["Deal ledger"].iter_rows(max_row=1, values_only=True), ())]
        while header and not header[-1]:
            header.pop()
        if header != check_lean.LEDGER_COLUMNS:
            raise DeriveError(f"{path}: the Deal ledger header is not the Version 1 ledger's {len(check_lean.LEDGER_COLUMNS)} columns")
        return {"schema": check_lean.LEDGER_SCHEMA, "sheets": list(wb.sheetnames), "header": header,
                "ledger": read_sheet(wb["Deal ledger"]), "rounds": read_sheet(wb["Rounds"]), "questions": read_sheet(wb["Questions"]),
                "facts": {text(r.get("Field")): r.get("Value") for r in read_sheet(wb["Deal facts"])},
                "source": read_source(wb["Source"]) if "Source" in wb.sheetnames else None}
    finally:
        wb.close()


def fact(facts: dict[str, Any], options: set[str] | str) -> str:
    options = {options} if isinstance(options, str) else options
    return next((text(facts[k]) for k in facts if k in options), "")


def auction_screen(value: str) -> dict[int, dict[str, Any]]:
    """Per process: status (Met or Not met; E1 has no Uncertain) and the supported count as bounds."""
    marks = list(re.finditer(r"\b(Met|Not met)\b\s*(?:\(\s*process\s*(\d+)\s*\))?\s*:?", value, re.IGNORECASE))
    out: dict[int, dict[str, Any]] = {}
    for i, m in enumerate(marks):
        body = value[m.end(): marks[i + 1].start() if i + 1 < len(marks) else len(value)]
        process = int(m.group(2)) if m.group(2) else (1 if len(marks) == 1 else None)
        if process is None or process in out:
            continue
        status = {"met": "Met", "not met": "Not met"}[m.group(1).casefold()]
        lo = hi = None
        if (r := re.search(r"(\d+)\s*[-–]\s*(\d+)", body)):
            lo, hi = int(r.group(1)), int(r.group(2))
        elif (r := re.search(r"at least\s+(\d+)", body, re.IGNORECASE)):
            lo = int(r.group(1))
        elif (r := re.search(r"(\d+)\s+(?:independent\s+)?(?:parties|party|acquirers?|bidders?)", body, re.IGNORECASE)):
            lo = hi = int(r.group(1))
        out[process] = {"status": status, "lo": lo, "hi": hi, "text": text(m.group(0) + body).strip("; ")}
    return out


# ---- derivation --------------------------------------------------------------------------------

def reading(t0: str, formal_if: bool) -> str:
    """Formal where T0 is Formal and the reading's test holds, otherwise Informal; missing only where T0 is (§7.10)."""
    if not t0:
        return ""
    return "Formal" if t0 == "Formal" and formal_if else "Informal"


def derive(ledger: dict[str, Any], deal: str) -> dict[str, Any]:
    review: list[dict[str, Any]] = []
    warnings: list[str] = []
    rows = ledger["ledger"]

    def note_review(row: dict[str, Any] | None, issue: str, sheet: str = "Deal ledger") -> None:
        review.append({"sheet": sheet, "row": cell(row.get("#")) if row else "", "who": text(row.get("Who")) if row else "", "issue": issue})

    # Scope: Other-scope bid rows leave the whole-company contest. A Who whose every bid row is Other-scope is only a
    # candidate for a partial-only party: a bidder may enter the whole-company contest and switch to a partial offer,
    # which closes its participation at the switch with an exit row (E1). Its other rows stay in the counts as the
    # ledger records them; with no exit row its entry counts as a bound (partial from the start, or a missing exit).
    bid_events = defaultdict(set)
    for r in rows:
        if text(r.get("Event")) in check_lean.BID_EVENTS:
            bid_events[unit_key(r.get("Who"))].add(text(r.get("Event")))
    partial_only = {k for k, events in bid_events.items() if events == {"Other-scope bid"}}
    other_scope, whole = [], []
    for r in rows:
        event, key = text(r.get("Event")), unit_key(r.get("Who"))
        if event == "Other-scope bid":
            other_scope.append((r, "Other-scope bid"))
            continue
        if key in partial_only:
            other_scope.append((r, "partial-only candidate (also in participation.csv)"))
        whole.append(r)
    scope_uncertain = set()
    for key in sorted(partial_only):
        own = [r for r in rows if unit_key(r.get("Who")) == key]
        numbers = ", ".join(f"#{cell(r.get('#'))}" for r in own)
        first_partial = next(r for r in own if text(r.get("Event")) == "Other-scope bid")
        exits = [r for r in own if text(r.get("Event")) in check_lean.EXIT_EVENTS | {"Re-entered"}]
        if exits:
            issue = (f"partial-only candidate (every bid row is Other-scope), but the ledger records whole-company participation "
                     f"ending at #{cell(exits[0].get('#'))} {text(exits[0].get('Event'))}: read as a switch to a partial offer (E1); "
                     "its entry and exit stay in the whole-company counts")
            if (as_date(exits[0].get("Sort date")) or dt.date.min) > (as_date(first_partial.get("Sort date")) or dt.date.max):
                issue += f"; the exit follows its first partial offer #{cell(first_partial.get('#'))}, and E1 closes participation at the switch"
        else:
            scope_uncertain.add(key)
            issue = ("partial-only candidate (every bid row is Other-scope) with no exit row: partial from the start (no entry, E1, E3) "
                     "or a switch to a partial offer whose exit row is missing; its entry counts as [0, n] with point 0 in participation.csv")
        review.append({"sheet": "Deal ledger", "row": numbers, "who": key, "issue": issue})

    rounds = {}
    for line in ledger["rounds"]:
        key = (check_lean.as_integer_or_text(line.get("Process")), check_lean.as_integer_or_text(line.get("Round")))
        rounds[key] = line

    # Participation: E14's formula, per process, as bounds and a presumed point.
    participation, live_at, changed = [], {}, set()
    # hi_open and point_open hold what leaves a bound unknown: the key of a unit's own entry, or "" for any other row.
    state = defaultdict(lambda: {"lo": 0, "hi": 0, "point": 0, "hi_open": set(), "point_open": set()})
    status: dict[tuple[Any, str], str] = {}
    cohort_seen: set[Any] = set()
    cohort_notes: dict[Any, list[Any]] = defaultdict(list)
    last_cohort_at: dict[Any, int] = {}
    entered: dict[tuple[Any, str], tuple[int, int | None, int | None]] = {}
    closed: set[Any] = set()
    open_units: dict[Any, dict[str, int]] = defaultdict(dict)
    # A round's opening state: live at its Round opened row, plus same-day exits from the round before and
    # same-day re-entries into it, wherever the ledger places those rows.
    openings: dict[tuple[Any, Any], dict[str, Any]] = {}
    opened_on: dict[Any, tuple[tuple[Any, Any], dt.date | None]] = {}

    def apply(s: dict[str, Any], d_lo: int | None, d_hi: int | None, d_point: int | None, unit: str = "") -> None:
        """Add a signed change in live units; None is unbounded (d_lo below, d_hi above) or unknown."""
        s["lo"] = 0 if d_lo is None else max(0, s["lo"] + d_lo)
        if d_hi is None:
            s["hi_open"].add(unit)
        s["hi"] = max(0, s["hi"] + (d_hi or 0))
        if d_point is None:
            s["point_open"].add(unit)
        s["point"] += d_point or 0

    def take_back(s: dict[str, Any], unit: str, added: tuple[int, int | None, int | None]) -> None:
        """Remove exactly what a unit's own entry added; its exit repeats the entry's inexact count."""
        a_lo, a_hi, a_point = added
        s["lo"] = max(0, s["lo"] - a_lo)
        if a_hi is None:
            s["hi_open"].discard(unit)
        s["hi"] = max(0, s["hi"] - (a_hi or 0))
        if a_point is None:
            s["point_open"].discard(unit)
        s["point"] -= a_point or 0

    def live_of(s: dict[str, Any]) -> tuple[int | None, int | None, int | None]:
        return s["lo"], None if s["hi_open"] else s["hi"], None if s["point_open"] else max(0, s["point"])

    def live(process: Any) -> tuple[int | None, int | None, int | None]:
        return live_of(state[process])

    def at_opening(event: str, row_round: Any, opening_round: Any) -> bool:
        return (isinstance(row_round, int) and isinstance(opening_round, int)
                and (event in check_lean.EXIT_EVENTS and row_round < opening_round
                     or event == "Re-entered" and row_round == opening_round))

    for position, r in enumerate(whole):
        process, rnd = row_round(r)
        event, key = text(r.get("Event")), unit_key(r.get("Who"))
        lo, hi, point, kind = count_bounds(r.get("Count"), r.get("Note"))
        change, delta, extra, own = "", None, {}, None
        prior = status.get((process, key))
        if event in ENTRY_EVENTS or event == "Re-entered":
            if event == "Re-entered":
                if prior != "exited":
                    note_review(r, "Re-entered without a recorded exit for this unit")
                change, delta = "re-entry", (lo, hi, point)
                status[(process, key)] = "live"
            elif prior == "live" or prior == "won":
                pass
            elif prior == "exited":
                if event in WHOLE_BIDS:
                    note_review(r, "bid after this unit's exit with no Re-entered row")
            elif key in scope_uncertain:
                change, delta = "entry (scope uncertain)", (0, hi, 0)
                status[(process, key)] = "live"
            elif event == "NDA signed" or process not in cohort_seen:
                change, delta = "entry", (lo, hi, point)
                status[(process, key)] = "live"
            else:
                # A party first seen bidding after cohort entries may be one of those cohorts' members (E3). A cohort
                # Note that names it settles that; its own contact row after the cohorts makes it a new entrant.
                contact = next((c for c in whole[last_cohort_at[process] + 1:position]
                                if text(c.get("Event")) in CONTACT_EVENTS and unit_key(c.get("Who")) == key
                                and row_round(c)[0] == process), None)
                if any(cohort_includes(note, key, check_lean.as_integer(r.get("#"))) for note in cohort_notes[process]):
                    change, delta = "entry (cohort member)", (0, 0, 0)
                elif contact is not None:
                    change, delta = "entry (late contact)", (lo, hi, point)
                    note_review(r, f"first seen bidding after cohort entries; its own {text(contact.get('Event'))} row "
                                f"#{cell(contact.get('#'))} follows them and no cohort Note names it: counted as a new entrant (E3)")
                else:
                    change, delta = "entry (membership uncertain)", (0, hi, 0)
                    if (lo, hi) == (1, 1):
                        note_review(r, "first seen bidding after cohort entries, and no cohort Note names it or cites its row: "
                                    "live_point presumes it a member (E3), live_hi allows a new entrant")
                status[(process, key)] = "live"
            if change in OWN_ENTRIES:
                entered[(process, key)] = (lo, hi, point)
            if change.startswith("entry") and change != "entry (cohort member)" and (lo, hi) != (1, 1):
                cohort_seen.add(process)
                cohort_notes[process].append(r.get("Note"))
                last_cohort_at[process] = position
            if change:
                open_units[process][key] = r.get("#")
        elif event in check_lean.EXIT_EVENTS or (event == "Merger agreement signed"
                                                and check_lean.as_integer(r.get("Count")) == 1):
            signing = event == "Merger agreement signed"
            if signing and prior is None:
                note_review(r, "the signing party has no recorded entry")
            if prior in ("exited", "won"):
                note_review(r, ("an exit after this unit's signing" if prior == "won" else "a second exit for this unit")
                            + " with no Re-entered row; not subtracted again")
            elif prior is None and process not in cohort_seen:
                # No entry of its own and no earlier cohort to belong to: it was never counted, so nothing is subtracted.
                if not signing:
                    note_review(r, "exit of a Who with no recorded entry and no earlier cohort entry in its process; not subtracted")
                change = "win (no recorded entry)" if signing else "exit (no recorded entry)"
                status[(process, key)] = "won" if signing else "exited"
            else:
                if prior is None and not signing:
                    note_review(r, "exit of a Who with no recorded entry (a residual or members of an earlier cohort); subtracted")
                change, delta = ("win" if signing else "exit"), (lo, hi, point)
                status[(process, key)] = "won" if signing else "exited"
                open_units[process].pop(key, None)
                own = entered.pop((process, key), None)
                if own is not None and own == (lo, hi, point) and (lo != hi or point is None):
                    # The same inexact count as its own entry ("several"): the whole unit leaves.
                    change += " (takes back its entry)"
                else:
                    own = None
            if event in check_lean.EXIT_EVENTS:
                # Inferred = Y marks an inferred event (Part B), so on an exit it is an inferred exit.
                inference = "inferred exit" if text(r.get("Inferred")) == "Y" else ""
                extra = {"exit_inferred": inference, "exit_actor": EXIT_ACTOR[event], "exit_reason": text(r.get("Exit reason")),
                         "alex_drop_code": ("DropTarget" if event == "Dropped by target" else
                                            WITHDREW_CODES.get(text(r.get("Exit reason")), "Drop") if event == "Withdrew" else ""),
                         "exit__inferred_as_dropout": "dropout",
                         "exit__inferred_as_censoring": "censored" if inference else "dropout"}
        elif event == "Bidding group changed":
            m = re.search(r"(\d+)\s+(?:\w+\s+)?units?\s+becomes?\s+(\d+)", text(r.get("Note")), re.IGNORECASE)
            if m:
                d = int(m.group(2)) - int(m.group(1))
                change, delta = "group change", (d, d, d)
            else:
                change, delta = "group change", (-1, 1, None)
                note_review(r, "group change whose units before and after cannot be read from the Note; live counts widened by one each way")
            if prior is None:
                # E4: this row counts the resulting unit, so its later bids are not new entries.
                status[(process, key)] = "live"
                open_units[process][key] = r.get("#")
        elif event in ("Process terminated", "Process restarted"):
            target = process if event == "Process terminated" else (process - 1 if isinstance(process, int) else None)
            if target is not None and target not in closed:
                closed.add(target)
                before = live(target)
                state[target].update(lo=0, hi=0, point=0, hi_open=set(), point_open=set())
                for k in list(open_units[target]):
                    status[(target, k)] = "exited"
                open_units[target].clear()
                recorded = check_lean.as_integer(r.get("Count"))
                if recorded is not None and before[2] is not None and recorded != before[2]:
                    note_review(r, f"closure Count {recorded} differs from the {before[2]} units the ledger leaves open in process {target}")
                participation.append({**base_cols(deal, r, process, rnd), "change": "process closure", "count_kind": kind,
                                      "count_lo": lo, "count_hi": hi, "count_point": point, "live_lo": 0, "live_hi": 0, "live_point": 0,
                                      "note": f"closes process {target}; live before: {before}"})
            live_at[id(r)] = live(process)
            continue
        elif event == "Round opened":
            change, (lo, hi, point, kind) = "round opening", (None, None, None, "")
        if delta is not None:
            if own is not None:
                delta = (-own[0], None if own[1] is None else -own[1], None if own[2] is None else -own[2])
            elif change in ("exit", "win"):
                lo_, hi_, point_ = delta
                delta = (None if hi_ is None else -hi_, -(lo_ or 0), None if point_ is None else -point_)
            targets = [state[process]]
            opening = opened_on.get(process)
            if opening and as_date(r.get("Sort date")) == opening[1] and at_opening(event, rnd, opening[0][1]):
                targets.append(openings[opening[0]])
            for s in targets:
                if own is not None:
                    take_back(s, key, own)
                else:
                    apply(s, *delta, unit=key if change in OWN_ENTRIES else "")
        if change:
            if event != "Round opened":
                changed.add(id(r))
            live_lo, live_hi, live_point = live(process)
            participation.append({
                **base_cols(deal, r, process, rnd), "change": change, "inferred": text(r.get("Inferred")), **extra,
                "count_kind": kind, "count_lo": lo, "count_hi": hi, "count_point": point,
                **(dict(zip(("delta_lo", "delta_hi", "delta_point"), delta)) if delta else {}),
                "live_lo": live_lo, "live_hi": live_hi, "live_point": live_point, "note": text(r.get("Note"))})
        live_at[id(r)] = live(process)
        if event == "Round opened":
            s = state[process]
            openings[(process, rnd)] = {**s, "hi_open": set(s["hi_open"]), "point_open": set(s["point_open"])}
            opened_on[process] = ((process, rnd), as_date(r.get("Sort date")))
        if kind == "range":
            note_review(r, "the Note gives a numeric Count range, which the ledger does not create (B, E3): read as bounds")
        if kind.startswith("unknown (") and (event in ENTRY_EVENTS | check_lean.EXIT_EVENTS | {"Re-entered"}):
            note_review(r, f"Count is blank and the Note gives no parseable 'Count: ...' ({kind})")
    for process, units in open_units.items():
        end_lo, end_hi, end_point = live(process)
        if process not in closed and end_point != 0:
            names = ", ".join(sorted(units)) or "none under their own names"
            review.append({"sheet": "Deal ledger", "row": "", "who": "", "issue": f"process {process} ends with {end_point if end_point is not None else 'an unknown number of'} live unit(s) "
                           f"(bounds {end_lo}-{end_hi if end_hi is not None else 'open'}): open at the filing cutoff, or an exit is missing. Open under their own names: {names}"})

    # Bids.
    bids, previous, bid_rows = [], {}, {}
    for r in whole:
        event = text(r.get("Event"))
        if event not in WHOLE_BIDS:
            continue
        process, rnd = row_round(r)
        key = unit_key(r.get("Who"))
        low, high = number(r.get("Price low")), number(r.get("Price high"))
        cvr = number(r.get("CVR/earnout value"))
        # A Same-offer row (E10: "Same as #n") copies #n's price; it is a restatement, never a same-price revision.
        same_as = check_lean.SAME_AS_RE.match(text(r.get("Note")))
        same_offer_of = int(same_as.group(1)) if same_as else None
        if same_offer_of is not None:
            source, number_here = bid_rows.get(same_offer_of), check_lean.as_integer(r.get("#"))
            if source is None or unit_key(source.get("Who")) != key or number_here is None or same_offer_of >= number_here:
                note_review(r, f"'Same as #{same_offer_of}' does not point to an earlier whole-company bid row of this bidder (E10)")
            elif (number(source.get("Price low")), number(source.get("Price high"))) != (low, high):
                note_review(r, f"a Same-offer row whose prices differ from #{same_offer_of}'s (E10 copies the price)")
        same = ""
        prior = previous.get((process, key))
        if event == "Bid" and same_offer_of is None and prior and (low, high) != (None, None) and (low, high, cvr) == prior:
            same = "Y"
        previous[(process, key)] = (low, high, cvr)
        if check_lean.as_integer(r.get("#")) is not None:
            bid_rows[check_lean.as_integer(r.get("#"))] = r
        # P1: a blank or invalid price is no price observation under either variant (H4).
        price_kind = upfront_price_kind(r)
        usable = price_kind in USABLE_PRICE_KINDS
        if price_kind == "invalid" and not any(not check_lean.is_blank(r.get(c)) and valid_price(r.get(c)) is None
                                               for c in ("Price low", "Price high")):
            note_review(r, "Price low is above Price high (a reversed range): upfront_price_kind invalid, no price observation")
        kind, s_lo, s_hi = stock_bounds(r.get("Stock %"), r.get("Note"))
        if kind == "range (not a Version 1 value)":
            note_review(r, "Stock % holds a range; E13 records Part stock with the range in the Note: read as bounds")
        lo, hi, point, _ = count_bounds(r.get("Count"), r.get("Note"))
        formality, level = text(r.get("Formality")), text(r.get("Conditions"))
        t0 = formality if formality in ("Formal", "Informal") else ""
        finality = text(rounds.get((process, rnd), {}).get("Finality"))
        if isinstance(rnd, int) and rnd >= 1 and (process, rnd) not in rounds:
            note_review(r, f"bid in process {process} round {rnd}, which has no Rounds line (T3 reads it as not final)")
        if t0 == "Formal" and level not in check_lean.CONDITIONS:
            note_review(r, (f"Conditions {level!r} is not a listed value" if level else "Conditions is blank")
                        + " on a Formal bid: T1 reads it as not Heavy (Formal), T1u as not None or Light (Informal)")
        # Package = upfront + CVR/earnout value; missing where a CVR is marked but its value is not given. package_basis
        # says what the sum is, so packages on different bases are not read as comparable (E13).
        # A price cell the checker rejects gives no package (price_low/high still show the cell as read).
        cvr_unvalued = text(r.get("CVR/earnout")) in {"Y", "Varies"} and cvr is None
        invalid = [c for c in ("Price low", "Price high") if not check_lean.is_blank(r.get(c)) and valid_price(r.get(c)) is None]
        package = (lambda p: None if p is None or cvr_unvalued or invalid or price_kind == "invalid" else p + (cvr or 0))
        basis, basis_issue = package_basis(r, low, high, cvr)
        if basis_issue:
            note_review(r, basis_issue)
        for column in invalid:
            note_review(r, f"{column} {r.get(column)!r} is not a positive number: no point price for T2 and no package; "
                        "price_low/price_high show the cell as read")
        point_price = valid_price(r.get("Price low")) is not None and valid_price(r.get("Price low")) == valid_price(r.get("Price high"))
        live_lo, live_hi, live_point = live_at.get(id(r), (None, None, None))
        bids.append({
            **base_cols(deal, r, process, rnd), "count": cell(r.get("Count")), "count_lo": lo, "count_hi": hi,
            "price_low": low, "price_high": high, "upfront_price_kind": price_kind, "cvr_earnout": text(r.get("CVR/earnout")), "cvr_value": cvr,
            "package_low": package(low), "package_high": package(high), "package_basis": basis,
            "stock_pct": cell(r.get("Stock %")), "stock_kind": kind, "stock_lo": s_lo, "stock_hi": s_hi,
            "all_cash": all_cash(r), "formality": formality, "conditions": level,
            **{c.lower().replace(" ", "_"): text(r.get(c)) for c in ("Due diligence", "Financing", "Regulatory", "Antitrust", "Exclusivity")},
            "inferred": text(r.get("Inferred")), "flag": text(r.get("Flag")), "same_offer_of": same_offer_of,
            "same_price_revision": same,
            # L253: the price a Same-offer row copies is not a new price observation under either variant.
            "price_obs__same_price_as_new": 1 if usable and same_offer_of is None else 0,
            "price_obs__same_price_as_terms": 1 if usable and same_offer_of is None and not same else 0,
            "round_finality": finality, "T0": t0,
            "T1": reading(t0, level != "Heavy"),
            "T1u": reading(t0, level in ("None", "Light")),
            # Two present, valid and equal prices make a point price; a blank or invalid cell does not (an unsplittable
            # package leaves both blank, E13). Round 0, post and a round with no Rounds line are not final.
            "T2": reading(t0, point_price),
            "T3": reading(t0, finality in FINAL and isinstance(rnd, int) and rnd >= 1),
            "live_lo": live_lo, "live_hi": live_hi, "live_point": live_point, "note": text(r.get("Note"))})

    # Rounds.
    deadline_log, round_rows = [], []
    deadline_rows = Counter(row_round(r) for r in whole if text(r.get("Event")) == "Deadline")
    opening_live = {key: live_of(s) for key, s in openings.items()}
    # The most live units at any point in a round: its opening state, then the state after each row that changes a
    # count. A row placed before a same-day exit at the opening would overstate it, and other rows change nothing.
    max_hi = defaultdict(int, {key: live[1] for key, live in opening_live.items()})
    for r in whole:
        key = row_round(r)
        if key in opening_live and id(r) not in changed:
            continue
        hi_now = live_at.get(id(r), (None, None, None))[1]
        max_hi[key] = None if hi_now is None or max_hi[key] is None else max(max_hi[key], hi_now)
    for (process, rnd), line in rounds.items():
        values = split_outcomes(line.get("Deadline outcome"))
        classes = []
        if not values:
            deadline_log.append({"process": process, "round": rnd, "position": 1, "value": "", "class": "not reached"})
        for i, value in enumerate(values, 1):
            found = deadline_class(value)
            if found is None:
                note_review(line, f"Deadline outcome value {value!r} matches no label; class missing", "Rounds")
                found = "unmapped"
            classes.append(found)
            deadline_log.append({"process": process, "round": rnd, "position": i, "value": value, "class": found})
        reached = []
        for part in re.split(r"\s*(?:→|->)\s*", text(line.get("Due dates"))):
            if re.search(r"superseded|future at filing", part, re.IGNORECASE):
                continue
            reached += re.findall(r"\d{1,2}/\d{1,2}/\d{4}", part)[:1]
        dated = [c for c in classes if c not in ("no deadline",)]
        if len(dated) != deadline_rows[(process, rnd)]:
            note_review(line, f"process {process} round {rnd}: {len(dated)} deadline outcome(s) but {deadline_rows[(process, rnd)]} Deadline row(s) in the ledger", "Rounds")
        units = defaultdict(lambda: (0, 0))
        for b in bids:
            if (b["process"], b["round"]) == (process, rnd):
                lo, hi = units[b["unit"]]
                units[b["unit"]] = (max(lo, b["count_lo"] or 0), None if b["count_hi"] is None or hi is None else max(hi, b["count_hi"]))
        bid_lo = sum(lo for lo, _ in units.values())
        bid_hi = None if any(hi is None for _, hi in units.values()) else sum(hi for _, hi in units.values())
        received = leading_integer(line.get("Bids received"))
        if received is None:
            check = "not compared"
        elif received >= bid_lo and (bid_hi is None or received <= bid_hi):
            check = "ok"
        else:
            check = "mismatch"
            note_review(line, f"process {process} round {rnd}: Bids received says {received}; the ledger's whole-company bid rows give {bid_lo}-{bid_hi if bid_hi is not None else 'open'}", "Rounds")
        admitted = leading_integer(line.get("Who was in"))
        open_lo, open_hi, open_point = opening_live.get((process, rnd)) or (None, None, None)
        if admitted is not None and max_hi.get((process, rnd)) is not None and admitted > max_hi[(process, rnd)]:
            note_review(line, f"process {process} round {rnd}: Who was in says {admitted}, more than the ledger's live units at any point in the round (at most {max_hi[(process, rnd)]})", "Rounds")
        round_rows.append({
            "deal": deal, "process": process, "round": rnd, "opened": cell(as_date(line.get("Opened"))),
            "how_opened": text(line.get("How opened")), "finality": text(line.get("Finality")),
            "who_was_in": text(line.get("Who was in")), "who_was_in_count": admitted,
            "due_dates": text(line.get("Due dates")), "due_dates_reached": "; ".join(reached),
            "deadline_outcome": text(line.get("Deadline outcome")), "deadline_values": "; ".join(values) or "(blank)",
            "deadline_classes": "; ".join(classes) or "not reached",
            "deadline_rows": deadline_rows[(process, rnd)], "bids_received": text(line.get("Bids received")),
            "bids_received_count": received, "bidders_bid_lo": bid_lo, "bidders_bid_hi": bid_hi,
            "bidders_bid": "; ".join(sorted(units)), "bids_received_check": check,
            "live_open_lo": open_lo, "live_open_hi": open_hi, "live_open_point": open_point,
            "live_max_hi": max_hi.get((process, rnd)), "how_it_ended": text(line.get("How it ended"))})
    ledger_rounds = {row_round(r) for r in whole if isinstance(row_round(r)[1], int) and row_round(r)[1] >= 1}
    for key in sorted(ledger_rounds - set(rounds), key=str):
        review.append({"sheet": "Rounds", "row": "", "who": "", "issue": f"process {key[0]} round {key[1]} has ledger rows but no Rounds line"})

    # Deal facts, per process.
    facts = ledger["facts"]
    screen = auction_screen(fact(facts, check_lean.AUCTION_SCREEN_FIELDS))
    whole_bids = fact(facts, check_lean.WHOLE_COMPANY_FIELDS)
    processes = sorted({p for p, _ in rounds} | {row_round(r)[0] for r in whole} | set(screen), key=str)
    moe = [r for r in rows if re.search(r"merger of equals", text(r.get("Note")), re.IGNORECASE)]
    recorded_initiation = fact(facts, "Initiation")
    deal_rows = []
    for process in processes:
        s = screen.get(process, {})
        process_rows = [r for r in whole if row_round(r)[0] == process]
        opening_index = next((i for i, r in enumerate(process_rows) if text(r.get("Event")) == "Round opened"
                              and row_round(r)[1] == 1), len(process_rows))
        # A partial-only candidate with no exit row may never have been in the whole-company contest, so its rows do
        # not initiate it; where one would have come first, a reviewer decides. The round-1 Round opened row
        # initiates only as the target's step, as check_lean.py reads it (D5; a bilateral opening is not).
        target_opening = check_lean.target_opened_round_one(process_rows)
        initiating = [r for i, r in enumerate(process_rows)
                      if text(r.get("Event")) in INITIATION_EVENTS or i == target_opening]
        # D5: only a demand for sale before the target's first sale step makes the process activist-influenced.
        # A target-side first step and a bidder's own Bid before round 1 make it mixed.
        preopening = process_rows[:opening_index]
        target_steps = [r for r in preopening if text(r.get("Event")) in {"Target interest", "Target sale decision"}]
        own_bids = [r for r in preopening if text(r.get("Event")) == "Bid"
                    and unit_key(r.get("Who")) not in scope_uncertain]
        first_target_index = check_lean.first_target_step(process_rows)
        activist = next((r for r in process_rows[:first_target_index] if text(r.get("Event")) == "Activist"
                         and text(r.get("Note")).startswith("Demands sale")), None)
        initiating = [activist] if activist else [r for r in initiating if text(r.get("Event")) != "Activist"]
        first = next((r for r in initiating if unit_key(r.get("Who")) not in scope_uncertain), None)
        if initiating and initiating[0] is not first:
            note_review(initiating[0], f"process {process}: the first initiating row belongs to a partial-only candidate with no exit row; "
                        "initiation_first_event uses " + (f"#{cell(first.get('#'))} {text(first.get('Event'))}" if first else "no row"))
        derived = ("activist-influenced" if activist else "mixed" if target_steps and own_bids
                   else INITIATION_EVENTS.get(text(first.get("Event")), "target-led") if first else "")
        check = ""
        if process == 1:
            # D5 decides Initiation from process 1; the Deal facts value is checked against the rule.
            check = ("not recorded" if not recorded_initiation else "not derivable (no initiating row)" if not derived
                     else "agrees" if recorded_initiation.casefold().startswith(derived) else "differs")
            if check == "differs":
                review.append({"sheet": "Deal facts", "row": "", "who": "", "issue":
                               f"Initiation is {recorded_initiation!r}, but D5's rule gives {derived} from "
                               f"#{cell(first.get('#'))} {text(first.get('Event'))} in process 1"})
        met = {"Met": 1, "Not met": 0}.get(s.get("status"))
        whole_flag = 1 if whole_bids.startswith("Yes") else 0 if whole_bids.startswith("No") else None
        descriptive = 1 if deal in DESCRIPTIVE_ONLY else 0
        deal_rows.append({
            "deal": deal, "process": process, "target": fact(facts, "Target"), "acquirer": fact(facts, "Acquirer"),
            "acquirer_type": fact(facts, "Acquirer type"), "agreed_price": fact(facts, "Agreed price and consideration"),
            "merger_signed": fact(facts, "Merger agreement signed"), "merger_announced": fact(facts, "Merger announced"),
            "filing": fact(facts, "Filing type and date"), "background_pages": fact(facts, "Background pages"),
            "number_of_processes": fact(facts, "Number of processes"),
            "earlier_approaches": fact(facts, check_lean.EARLIER_APPROACHES_FIELDS),
            "currency_units": fact(facts, "Currency and units of bid prices"),
            "auction_screen": s.get("text", ""), "auction_status": s.get("status", ""),
            "auction_count_lo": s.get("lo"), "auction_count_hi": s.get("hi"), "auction_met": met,
            "whole_company_bids": whole_bids, "whole_company": whole_flag, "descriptive_only": descriptive,
            "estimation_sample": None if met is None or whole_flag is None else int(met == 1 and whole_flag == 1 and not descriptive),
            "initiation_recorded": recorded_initiation, "initiation_first_event": derived,
            "initiation_first_row": f"#{cell(first.get('#'))} {text(first.get('Event'))}" if first else "",
            "initiation_check": check,
            "merger_of_equals_rows": "; ".join(f"#{cell(r.get('#'))}" for r in moe)})
        if not s:
            warnings.append(f"process {process}: no Auction screen entry could be parsed")

    others = [{**base_cols(deal, r, *row_round(r)), "reason": reason, "count": cell(r.get("Count")),
               "price_low": number(r.get("Price low")), "price_high": number(r.get("Price high")),
               "stock_pct": cell(r.get("Stock %")), "formality": text(r.get("Formality")),
               "conditions": text(r.get("Conditions")), "note": text(r.get("Note"))} for r, reason in other_scope]
    return {"bids": bids, "other_scope": others, "rounds": round_rows, "participation": participation,
            "deal": deal_rows, "deadline_outcomes": deadline_log, "review": review, "warnings": warnings}


def row_round(r: dict[str, Any]) -> tuple[Any, Any]:
    """(Process, Round) as integers where they are integers; Round may be "post"."""
    rnd = check_lean.as_integer_or_text(r.get("Round"))
    return check_lean.as_integer_or_text(r.get("Process")), rnd if rnd is not None else text(r.get("Round"))


def base_cols(deal: str, r: dict[str, Any], process: Any, rnd: Any) -> dict[str, Any]:
    return {"deal": deal, "process": process, "round": rnd, "row": cell(r.get("#")), "when": text(r.get("When")),
            "sort_date": cell(as_date(r.get("Sort date"))), "date_from": cell(as_date(r.get("Date from"))),
            "date_to": cell(as_date(r.get("Date to"))), "who": text(r.get("Who")), "unit": unit_key(r.get("Who")),
            "type": text(r.get("Type")), "event": text(r.get("Event"))}


# ---- output ------------------------------------------------------------------------------------

def check_outside_data(path: Path, option: str = "--out") -> None:
    """Refuse an output path in or under extraction/, raw_filing/ or ref/."""
    path = path.resolve()
    for part in FORBIDDEN_OUT:
        forbidden = (PROJECT / part).resolve()
        if path == forbidden or forbidden in path.parents:
            raise DeriveError(f"{option} must not be under {PROJECT / part}/")


def check_out(out: Path) -> None:
    check_outside_data(out)
    out = out.resolve()
    if out.exists() and (not out.is_dir() or any(out.iterdir())):
        raise DeriveError(f"--out {out} exists and is not an empty folder; choose a new one")


def write_csv(path: Path, columns: list[str], rows: list[dict[str, Any]]) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({c: cell(row.get(c)) for c in columns})


def manifest_complete(manifest: dict[str, Any]) -> list[str]:
    """The manifest keys that are missing or empty where a value is required."""
    missing = [k for k in MANIFEST_KEYS if k not in manifest]
    for key in ("path", "sha256", "kind"):
        if not manifest.get("input", {}).get(key):
            missing.append(f"input.{key}")
    if not manifest.get("ledger_schema"):
        missing.append("ledger_schema")
    if any(d.get("class") in (None, "", "unmapped") for d in manifest.get("deadline_outcomes", [])):
        missing.append("deadline_outcomes.class")
    return missing


def known_deals() -> set[str]:
    """Deal slugs of the seed, to read a slug from a file name."""
    known = set(DESCRIPTIVE_ONLY)
    try:
        with (PROJECT / "ref/seed.csv").open(newline="", encoding="utf-8") as handle:
            known |= {row["deal"] for row in csv.DictReader(handle) if row.get("deal")}
    except (OSError, KeyError):
        pass
    return known


def deal_from_name(workbook: Path, source: dict[str, Any] | None, known: set[str]) -> tuple[str, str | None]:
    """The deal slug from a file name, as the cockpit names downloads, and a warning when it is a guess.

    Cockpit downloads are "{slug}-{version}.xlsx", "{slug}-{version}-with-source.xlsx" and
    "{slug}-working-r{N}.xlsx"; the Source sheet's version ID, when present, is stripped as well.
    """
    stem = re.sub(r"-with-source$", "", workbook.stem)
    stem = re.sub(r"-working(?:-r\d+)?$", "", stem)
    version = next((text(v) for k, v in (source or {}).items() if k.startswith("Version ID") and text(v)), "")
    if version and stem.endswith("-" + version):
        stem = stem[: -len(version) - 1]
    if stem in known:
        return stem, None
    prefixes = [k for k in known if stem.startswith(k + "-")]
    if prefixes:
        slug = max(prefixes, key=len)
        return slug, f"deal slug {slug!r} read from the file name {workbook.name}; give --deal if that is wrong"
    return stem, f"deal slug {stem!r} from the file name {workbook.name} is not a known deal (seed.csv); give --deal"


def run(workbook: Path, out: Path, deal: str | None = None) -> dict[str, Any]:
    ledger = load(workbook)
    guessed = None
    if not deal:
        deal, guessed = deal_from_name(workbook, ledger["source"], known_deals())
    result = derive(ledger, deal)
    if guessed:
        result["warnings"].insert(0, guessed)
    out.mkdir(parents=True, exist_ok=True)
    tables = {"bids.csv": (BID_COLUMNS, result["bids"]), "other_scope.csv": (OTHER_COLUMNS, result["other_scope"]),
              "rounds.csv": (ROUND_COLUMNS, result["rounds"]), "participation.csv": (PARTICIPATION_COLUMNS, result["participation"]),
              "deal.csv": (DEAL_COLUMNS, result["deal"])}
    for name, (columns, rows) in tables.items():
        write_csv(out / name, columns, rows)
    kind = "five-sheet cockpit download" if ledger["source"] is not None else "four-sheet workbook"
    manifest = {
        "tool": "derive_analysis.py", "tool_version": TOOL_VERSION,
        "generated_at": dt.datetime.now(dt.timezone.utc).isoformat(timespec="seconds"), "deal": deal,
        "input": {"path": str(workbook), "sha256": sha256(workbook), "kind": kind, "sheets": ledger["sheets"],
                  "source_sheet": ledger["source"]},
        "ledger_schema": ledger["schema"], "checker_version": check_lean.CHECKER_VERSION,
        "readings": {"computed": ["T0", "T1", "T1u", "T2", "T3"], "all_cash_from": "Stock %",
                     "package": "Price + CVR/earnout value; package_basis says what each sum is (E13)",
                     "default": None},
        "switches": SWITCHES,
        "deadline_outcomes": result["deadline_outcomes"],
        "outputs": {name: len(rows) for name, (_, rows) in tables.items()},
        "review": result["review"], "warnings": result["warnings"],
    }
    manifest["incomplete"] = manifest_complete(manifest)
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=False, default=cell) + "\n", encoding="utf-8")
    return manifest


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("workbook", type=Path, help="a Version 1 ledger workbook (raw version or cockpit download)")
    parser.add_argument("--deal", help="deal slug (default: from the file name)")
    parser.add_argument("--out", type=Path, required=True, help="new or empty output folder")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    try:
        check_out(args.out)
        args.out.mkdir(parents=True, exist_ok=True)
        manifest = run(args.workbook.resolve(), args.out, args.deal)
    except (DeriveError, OSError) as exc:
        print(f"derive_analysis: {exc}", file=sys.stderr)
        return 2
    print(f"{manifest['deal']} ({manifest['ledger_schema']}): " + ", ".join(f"{k} {v}" for k, v in manifest["outputs"].items())
          + f"; {len(manifest['review'])} review item(s)")
    return 0 if not manifest["incomplete"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
