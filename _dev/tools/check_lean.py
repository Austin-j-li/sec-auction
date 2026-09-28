#!/usr/bin/env python3
"""Mechanical validator for the Version 1 deal-ledger workbook.

This checker intentionally does not decide whether events, bidders, rounds, or
classifications are substantively correct. In particular, it does not sum
``Count`` into a bidder population. It checks workbook structure and internal
consistency, then tests each quoted passage as one contiguous normalized
substring of the full filing parsed by BeautifulSoup.

The checker knows one 29-column Deal ledger. A workbook whose ledger
header differs fails with a schema.columns error.

Mechanical readings of rules the instruction states in words:

- A cohort row is a row whose Count is above 1, or whose Who names a group (it starts
  with a number, or names parties, bidders, signers and the like in the plural).
- The process Question (Part F) is the first Question whose Question text begins
  "Process:" or whose Rows affected cites a Process terminated or Process restarted row
  by #. It does not count toward the five-Question cap; one is expected when the ledger
  has more than one process or a Round opened row with Inferred = Y. Rounds opened by
  trigger (d) also call for it, but they cannot be recognised mechanically, so they are
  not checked.
- A bidder unit is followed by its Who without parentheticals, case-folded (unit_key),
  as derive_analysis.py follows it.
- The target's first sale step (D5) is its first Target interest or Target sale decision
  row, or the first round-1 Round opened row unless that row shares its Sort date with an
  NDA signed or Bid row of a party that approached the target before it (E6's bilateral
  fallback), whatever the opening row's Who.
- On Merger agreement signed, Count is blank or 1. Who is live at signing is not
  reconstructed: a blank Count beside a Bid or Bid reaffirmed row of the signer's unit,
  or a Count of 1 where every bid row of that unit is an Other-scope bid, is a warning (D1).
- Count is required on Process terminated, Process restarted and Bidding group changed
  rows. A process marker that closes no open participation leaves Count blank, so a
  blank Count on a process marker is a warning, not an error.

Usage:
    python3 check_lean.py --workbook candidate.xlsx --filing filing.htm \
        --output mechanical_report.json

Exit status is 0 when there are no errors, 1 when validation errors are found,
and 2 when an input cannot be opened or the JSON report cannot be written.
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import math
import os
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl
from bs4 import BeautifulSoup
from openpyxl.utils.cell import coordinate_to_tuple, range_boundaries


CHECKER_VERSION = "Version 1"
CHECKER_REVISION = "Checks the Version 1 ledger rules."
LEDGER_SCHEMA = "Version 1"

SHEETS = ["Deal ledger", "Rounds", "Questions", "Deal facts"]

LEDGER_COLUMNS = [
    "#",
    "When",
    "Who",
    "Type",
    "Event",
    "Process",
    "Round",
    "Price low",
    "Price high",
    "Stock %",
    "CVR/earnout",
    "CVR/earnout value",
    "Formality",
    "Conditions",
    "Due diligence",
    "Financing",
    "Regulatory",
    "Antitrust",
    "Exclusivity",
    "Count",
    "Exit reason",
    "Inferred",
    "Note",
    "Quote and page",
    "Flag",
    "Reviewer note",
    "Sort date",
    "Date from",
    "Date to",
]

ROUND_COLUMNS = [
    "Process",
    "Round",
    "Opened",
    "How opened",
    "Who was in",
    "Due dates",
    "Deadline outcome",
    "Finality",
    "Bids received",
    "How it ended",
]

QUESTION_COLUMNS = [
    "Q",
    "Question",
    "Recommended answer",
    "Why, with page",
    "Rows affected",
    "What changes if answered differently",
    "Reviewer note",
]

FACT_COLUMNS = ["Field", "Value"]

FACT_FIELDS = [
    "Target",
    "Acquirer",
    "Acquirer type",
    "Agreed price and consideration",
    "Merger agreement signed",
    "Merger announced",
    "Filing type and date",
    "Background pages",
    "Initiation",
    "Number of processes",
    "Earlier approaches",
    "Auction screen",
    "Whole-company bids",
    "Currency and units of bid prices",
    "Target financial advisers",
    "Target legal advisers",
    "Account",
]
ACCOUNT_FIELDS = {"Account", "Account (five or six plain sentences)"}
AUCTION_SCREEN_FIELDS = {"Auction screen", "Auction screen (E1)"}
EARLIER_APPROACHES_FIELDS = {
    "Earlier approaches", "Earlier approaches (E5)",
    "Earlier approaches (E5; “None reported” if none)",  # the field as D5 prints it
}
WHOLE_COMPANY_FIELDS = {
    "Whole-company bids",
    "Whole-company bids (Yes, or No with what was bid for)",
}
FACT_FIELD_OPTIONS = [{field} for field in FACT_FIELDS]
FACT_FIELD_OPTIONS[10] = EARLIER_APPROACHES_FIELDS
FACT_FIELD_OPTIONS[11] = AUCTION_SCREEN_FIELDS
FACT_FIELD_OPTIONS[12] = WHOLE_COMPANY_FIELDS
FACT_FIELD_OPTIONS[16] = ACCOUNT_FIELDS

EVENTS = {
    "Target interest",
    "Bidder interest",
    "Target sale decision",
    "Activist",
    "Adviser",
    "Adviser ended",
    "Contact",
    "NDA signed",
    "Round opened",
    "Deadline set",
    "Deadline revised",
    "Deadline",
    "Exclusivity changed",
    "Other material event",
    "Bid",
    "Bid reaffirmed",
    "Other-scope bid",
    "Bidding group changed",
    "Dropped by target",
    "Withdrew",
    "Did not submit",
    "Not selected at signing",
    "Re-entered",
    "Sale process announced",
    "Bid announced",
    "Merger announced",
    "Merger agreement signed",
    "Go-shop changed",
    "Process terminated",
    "Process restarted",
}

TYPES = {"Strategic", "Financial", "Mixed", "Unknown"}
BID_EVENTS = {"Bid", "Bid reaffirmed", "Other-scope bid"}
EXIT_EVENTS = {
    "Dropped by target",
    "Withdrew",
    "Did not submit",
    "Not selected at signing",
}
NO_BIDDER_COUNT_EVENTS = {
    "Target sale decision",
    "Activist",
    "Adviser",
    "Adviser ended",
    "Round opened",
    "Deadline set",
    "Deadline revised",
    "Deadline",
    "Sale process announced",
    "Bid announced",
    "Merger announced",
    "Go-shop changed",
}
FORMALITY = {"Formal", "Informal", "Unclear"}
CONDITIONS = {"None", "Light", "Heavy", "Unclear"}
STOCK_CODES = {"Part stock", "Not stated", "Varies"}
STOCK_RANGE_RE = re.compile(r"(\d+(?:\.\d+)?)\s*[-\u2013]\s*(\d+(?:\.\d+)?)")
MARKER = {"Y"}  # D1 cols 11 and 18: Y or blank; a cohort split goes in the Note.
DUE_DILIGENCE = {"Complete", "Incomplete", "Not begun", "Not stated", "Varies"}
FINANCING = {"Not needed", "Committed", "Contingent", "Not stated", "Varies"}
REGULATORY = {"No concern", "Concern", "Not stated", "Varies"}
EXCLUSIVITY = {"Required", "Requested", "Not stated", "Varies"}
CONDITION_COLUMNS = {
    "Due diligence": DUE_DILIGENCE,
    "Financing": FINANCING,
    "Regulatory": REGULATORY,
    "Exclusivity": EXCLUSIVITY,
}
TERM_COLUMNS = (
    "Stock %",
    "CVR/earnout",
    "CVR/earnout value",
    "Due diligence",
    "Financing",
    "Regulatory",
    "Antitrust",
    "Exclusivity",
)
EXIT_REASONS = {
    "Value below market price",
    "Value at or below market price",
    "Value below earlier offer",
    "Value at earlier offer",
    "Would not improve earlier offer",
    "Lower offer than rivals",
    "Terms or process",
    "Other stated reason",
    "Not stated",
}
FINALITY = {"Announced as final", "Inferred final", "Not final"}
INITIATION = {"target-led", "bidder-led", "activist-influenced", "mixed"}
# E9: an overdue required response the target still considered is an extension.
DEADLINE_OUTCOMES = {
    "Enforced",
    "Extended",
    "Extended (late bid accepted)",
    "Passed without action",
    "Unclear",
}
NO_DEADLINE = "No deadline stated"
# Inferred = Y marks only these row events (Part B, D1 col 22).
INFERRED_EVENTS = EXIT_EVENTS | {"Round opened", "Process restarted"}
PROCESS_MARKERS = {"Process terminated", "Process restarted"}
COHORT_WHO_RE = re.compile(
    r"^\s*\d+\b|\b(?:parties|bidders|signers|others|participants|buyers|sponsors|firms|companies|"
    r"acquirers|investors|recipients|members|cohort|group)\b",
    re.IGNORECASE,
)
SAME_AS_RE = re.compile(r"^\s*Same as #\s*(\d+)\b\.?\s*", re.IGNORECASE)
HEAVY_TRIGGER_RE = re.compile(r"^H[123]:")
COUNT_QUALIFIER_RE = re.compile(
    r"\bCount:\s*(?:at least|more than|over|at most|up to|fewer than|less than|approximately|about|"
    r"around|nearly|some|several)\b",
    re.IGNORECASE,
)
COUNT_UNKNOWN_RE = re.compile(r"\bCount:\s*(?:unknown|not stated)\b", re.IGNORECASE)
COUNT_RANGE_RE = re.compile(r"\bCount:\s*\d+\s*(?:[-\u2013\u2014]|to)\s*\d+", re.IGNORECASE)
QUESTION_CAP = 5
QUESTION_WORDS = 60

DATE_COLUMNS = ("Sort date", "Date from", "Date to")
PAGE_REF_RE = re.compile(
    r"(?:\(\s*pp?\.\s*(?P<paren>[^()]+?)\s*\)|pp?\.\s*(?P<bare>[A-Za-z0-9]+(?:\s*[-\u2013,]\s*[A-Za-z0-9]+)*))\s*$",
    re.IGNORECASE,
)
REVIEW_ID_RE = re.compile(r"[QR][1-9]\d*")


def choice_lists() -> dict[str, list[str]]:
    """The editor's controlled values for the current ledger format."""
    return {
        "Type": sorted(TYPES), "Event": sorted(EVENTS), "Formality": sorted(FORMALITY),
        "Conditions": sorted(CONDITIONS), "Exit reason": sorted(EXIT_REASONS),
        "Finality": sorted(FINALITY),
        "Deadline outcome": sorted(DEADLINE_OUTCOMES | {NO_DEADLINE}),
        "Initiation": sorted(INITIATION),
        "CVR/earnout": sorted(MARKER), "Antitrust": sorted(MARKER),
        **{field: sorted(values) for field, values in CONDITION_COLUMNS.items()},
    }


def normalize_contiguous(value: Any) -> str:
    """Normalize Unicode and whitespace without splitting or joining excerpts."""

    text = unicodedata.normalize("NFC", str(value)).replace("\u00a0", " ")
    return re.sub(r"\s+", " ", text).strip()


def word_count(value: Any) -> int:
    text = normalize_contiguous(value)
    return len(text.split()) if text else 0


def strip_outer_quotes(value: str) -> str:
    text = value.strip()
    if len(text) >= 2 and ((text[0], text[-1]) in {("\u201c", "\u201d"), ('"', '"')}):
        return text[1:-1].strip()
    return text


def parse_quote_and_page(value: Any) -> tuple[str, str] | None:
    """Separate one passage from its terminal p./pp. reference.

    Outer quotation marks are optional and may surround the passage alone or
    the passage plus its page reference. The returned passage remains whole;
    no ellipsis/bracket splitting or piecewise matching is performed.
    """

    if is_blank(value):
        return None
    text = normalize_contiguous(value)
    text_without_outer_quotes = strip_outer_quotes(text)
    match = PAGE_REF_RE.search(text_without_outer_quotes)
    if match is None:
        return None
    passage = strip_outer_quotes(text_without_outer_quotes[: match.start()].strip())
    page = normalize_contiguous(match.group("paren") or match.group("bare"))
    return passage, page


def is_blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def is_integer(value: Any) -> bool:
    return isinstance(value, int) and not isinstance(value, bool) or (
        isinstance(value, float) and math.isfinite(value) and value.is_integer()
    )


def as_integer(value: Any) -> int | None:
    return int(value) if is_integer(value) else None


def as_integer_or_text(value: Any) -> int | None:
    integer = as_integer(value)
    if integer is not None:
        return integer
    if isinstance(value, str) and re.fullmatch(r"[+-]?\d+", value.strip()):
        return int(value.strip())
    return None


def starts_with_canonical(value: Any, allowed: set[str]) -> bool:
    text = normalize_contiguous(value)
    folded = text.casefold()
    return any(
        folded == label.casefold()
        or (
            folded.startswith(label.casefold())
            and len(text) > len(label)
            and not text[len(label)].isalnum()
        )
        for label in allowed
    )


def parse_flag_ids(value: Any) -> set[str] | None:
    """Parse one or more Q or R ids; return None if other content is present."""

    if is_blank(value):
        return set()
    text = normalize_contiguous(value)
    ids = REVIEW_ID_RE.findall(text)
    remainder = REVIEW_ID_RE.sub("", text)
    remainder = re.sub(r"\band\b", "", remainder, flags=re.IGNORECASE)
    if not ids or re.fullmatch(r"[\s,;/&]*", remainder) is None:
        return None
    return set(ids)


def as_date(value: Any) -> dt.date | None:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    return None


def date_has_time(value: Any) -> bool:
    return isinstance(value, dt.datetime) and value.time() != dt.time()


def has_mmddyyyy_format(cell: Any) -> bool:
    return str(cell.number_format or "").strip().lower() == "mm/dd/yyyy"


def nonempty_rows(ws: Any, width: int) -> list[int]:
    scan_width = max(width, ws.max_column)
    return [
        row
        for row in range(2, ws.max_row + 1)
        if any(not is_blank(ws.cell(row, col).value) for col in range(1, scan_width + 1))
    ]


def parse_affected_rows(value: Any) -> tuple[set[int], bool]:
    """Return explicit ledger row references and whether the syntax was parseable.

    Accepted forms include ``#2, #4-#6`` and ``Rows 2, 4-6``. Without ``#``, semicolons separate
    segments, and a segment that is not a row list (a source event the ledger omits, D4) is skipped,
    so ``Rows 4, 6; June 5 board meeting, omitted`` cites rows 4 and 6. Narrative/global
    descriptions with no row list are left unparsed so the checker does not invent references.
    """

    if is_blank(value):
        return set(), False
    text = normalize_contiguous(value).replace("\u2013", "-").replace("\u2014", "-")
    hash_refs = {int(number) for number in re.findall(r"#\s*(\d+)", text)}
    for start, end in re.findall(r"#\s*(\d+)\s*-\s*#?\s*(\d+)", text):
        lo, hi = sorted((int(start), int(end)))
        hash_refs.update(range(lo, hi + 1))
    if "#" in text:
        return hash_refs, bool(hash_refs)

    refs: set[int] = set()
    parsed = False
    for segment in text.split(";"):
        stripped = re.sub(r"(?i)^\s*rows?\s*:?\s*", "", segment).strip()
        if not re.fullmatch(r"\d+(?:\s*-\s*\d+)?(?:\s*,\s*\d+(?:\s*-\s*\d+)?)*", stripped):
            continue
        parsed = True
        for token in re.split(r"\s*,\s*", stripped):
            if "-" in token:
                start, end = (int(part.strip()) for part in token.split("-", 1))
                lo, hi = sorted((start, end))
                refs.update(range(lo, hi + 1))
            else:
                refs.add(int(token))
    return refs, parsed



def is_process_question(question: Any, rows_affected: Any, marker_rows: set[int]) -> bool:
    """Part F's process Question: its Question begins "Process:" or its Rows affected cites, by #, a
    Process terminated or Process restarted row (marker_rows holds those rows' numbers)."""
    text = normalize_contiguous(question).lower() if not is_blank(question) else ""
    refs, _ = parse_affected_rows(rows_affected)
    return text.startswith("process:") or bool(refs & marker_rows)


def unit_key(who: Any) -> str:
    """The name a bidder unit is followed by: Who without parentheticals, case-folded."""
    name = re.sub(r"\([^()]*\)", " ", "" if is_blank(who) else normalize_contiguous(who))
    return re.sub(r"\s+", " ", name).strip(" .,;:").casefold()


def row_event(record: dict[str, Any]) -> str:
    return "" if is_blank(record.get("Event")) else normalize_contiguous(record.get("Event"))


def target_opened_round_one(records: list[dict[str, Any]]) -> int | None:
    """Index of the first round-1 Round opened row in one process's rows when D5 reads it as the target's step
    (a round "opened by the target's outreach"), else None.

    The test: the opening is the target's step unless it falls on the Sort date of an NDA signed or Bid row of a
    party that approached the target earlier in the process (a Bidder interest or Bid row before the opening). That
    is E6's bilateral fallback, which opens round 1 at the first NDA or price negotiation with such a party. The
    opening row's Who does not decide, since D1 gives the target as Who on process-wide rows."""
    approached: set[str] = set()
    for index, record in enumerate(records):
        if row_event(record) == "Round opened" and as_integer_or_text(record.get("Round")) == 1:
            day = as_date(record.get("Sort date"))
            bilateral = day is not None and any(
                unit_key(other.get("Who")) in approached and as_date(other.get("Sort date")) == day
                and row_event(other) in {"NDA signed", "Bid", "Bid reaffirmed"} for other in records)
            return None if bilateral else index
        if row_event(record) in {"Bidder interest", "Bid"} and unit_key(record.get("Who")):
            approached.add(unit_key(record.get("Who")))
    return None


def first_target_step(records: list[dict[str, Any]]) -> int:
    """Index of the target's first sale step (D5) in one process's rows, else len(records): its first Target
    interest or Target sale decision row, or its round 1 if the target opened it (target_opened_round_one)."""
    opening = target_opened_round_one(records)
    return next((index for index, record in enumerate(records)
                 if row_event(record) in {"Target interest", "Target sale decision"} or index == opening),
                len(records))


def is_cohort(record: dict[str, Any]) -> bool:
    """A row standing for several bidders: Count above 1, or a Who naming a group."""
    count = as_integer(record.get("Count"))
    who = normalize_contiguous(record.get("Who")) if not is_blank(record.get("Who")) else ""
    return (count is not None and count > 1) or bool(COHORT_WHO_RE.search(who))


class LeanChecker:
    def __init__(self, workbook_path: Path, filing_path: Path) -> None:
        self.workbook_path = workbook_path
        self.filing_path = filing_path
        self.issues: list[dict[str, Any]] = []
        self.fatal = False
        self.wb: Any = None
        self.filing_texts: tuple[str, ...] = ()

    def add(
        self,
        severity: str,
        code: str,
        message: str,
        *,
        sheet: str | None = None,
        row: int | None = None,
        column: str | None = None,
    ) -> None:
        self.issues.append(
            {
                "severity": severity,
                "code": code,
                "sheet": sheet,
                "row": row,
                "column": column,
                "message": message,
            }
        )

    def load_inputs(self) -> bool:
        try:
            self.wb = openpyxl.load_workbook(self.workbook_path, data_only=False)
        except Exception as exc:  # openpyxl exposes several format/XML exceptions
            self.fatal = True
            self.add(
                "error",
                "input.workbook_unreadable",
                f"Could not open workbook: {type(exc).__name__}: {exc}",
            )
            return False

        try:
            raw = self.filing_path.read_bytes()
            soup = BeautifulSoup(raw, "html.parser")
            for element in soup(["script", "style", "noscript"]):
                element.decompose()
            # SEC HTML frequently divides a visually contiguous phrase among
            # inline FONT/SPAN nodes. Both renderings preserve DOM order and
            # contiguous matching; neither permits excerpt-by-excerpt splicing.
            renderings = (
                normalize_contiguous(soup.get_text(" ")),
                normalize_contiguous(soup.get_text("")),
            )
            self.filing_texts = tuple(dict.fromkeys(text for text in renderings if text))
            if not self.filing_texts:
                raise ValueError("parsed filing contains no text")
        except Exception as exc:
            self.fatal = True
            self.add(
                "error",
                "input.filing_unreadable",
                f"Could not parse filing: {type(exc).__name__}: {exc}",
            )
            return False
        return True

    def check_schema(self, ws: Any, expected: list[str]) -> bool:
        width = max(ws.max_column, len(expected))
        actual = [ws.cell(1, col).value for col in range(1, width + 1)]
        while actual and is_blank(actual[-1]):
            actual.pop()
        if actual != expected:
            self.add(
                "error",
                "schema.columns",
                f"Expected columns {expected!r}; found {actual!r}."
                + (" The checker reads only the current ledger." if ws.title == "Deal ledger" else ""),
                sheet=ws.title,
                row=1,
            )
            return False
        return True

    def check_presentation(self, ws: Any, expected_width: int) -> None:
        if ws.sheet_state != "visible":
            self.add(
                "error",
                "presentation.hidden_sheet",
                "Required sheet is not visible.",
                sheet=ws.title,
            )

        if ws.merged_cells.ranges:
            self.add(
                "error",
                "presentation.merged_cells",
                "Merged cells are prohibited: "
                + ", ".join(str(item) for item in list(ws.merged_cells.ranges)[:10]),
                sheet=ws.title,
            )

        freeze = ws.freeze_panes
        try:
            coordinate = freeze.coordinate if hasattr(freeze, "coordinate") else str(freeze)
            freeze_row, _ = coordinate_to_tuple(coordinate)
        except Exception:
            freeze_row = None
        if freeze_row != 2:
            self.add(
                "error",
                "presentation.freeze_header",
                f"Header row is not frozen at row 1 (freeze_panes={freeze!r}).",
                sheet=ws.title,
                row=1,
            )

        filter_refs: list[str] = []
        if ws.auto_filter.ref:
            filter_refs.append(str(ws.auto_filter.ref))
        filter_refs.extend(str(table.ref) for table in ws.tables.values())
        covered = False
        for ref in filter_refs:
            try:
                min_col, min_row, max_col, _ = range_boundaries(ref)
                covered = covered or (min_col == 1 and min_row == 1 and max_col >= expected_width)
            except ValueError:
                continue
        if not covered:
            self.add(
                "error",
                "presentation.filters",
                "No filter range starts at A1 and covers every prescribed column.",
                sheet=ws.title,
                row=1,
            )

        unwrapped: list[str] = []
        for row in range(1, ws.max_row + 1):
            for col in range(1, expected_width + 1):
                cell = ws.cell(row, col)
                if not is_blank(cell.value) and cell.alignment.wrap_text is not True:
                    unwrapped.append(cell.coordinate)
        if unwrapped:
            self.add(
                "error",
                "presentation.wrap_text",
                f"{len(unwrapped)} nonempty prescribed cells do not have wrap text enabled; "
                f"first: {', '.join(unwrapped[:12])}.",
                sheet=ws.title,
                row=ws[unwrapped[0]].row,
            )

    def check_bid_terms(self, ws: Any, excel_row: int, record: dict[str, Any]) -> None:
        """Consideration and condition columns on one bid row (E12, E13)."""

        def issue(severity: str, code: str, message: str, column: str) -> None:
            self.add(severity, code, message, sheet=ws.title, row=excel_row, column=column)

        stock = record["Stock %"]
        stock_ok = False
        reported = False
        if isinstance(stock, bool) or is_blank(stock):
            pass
        elif isinstance(stock, (int, float)):
            stock_ok = math.isfinite(float(stock)) and 0 <= float(stock) <= 100
        elif isinstance(stock, str):
            text = stock.strip()
            match = STOCK_RANGE_RE.fullmatch(text)
            if text in STOCK_CODES:
                stock_ok = True
            elif match:
                issue("error", "controlled.stock_pct", f"Stock % {stock!r} is a range; a stated range is Part stock, "
                      "with the range in the Note (E13).", "Stock %")
                reported = True
            else:
                try:
                    float(text)
                except ValueError:
                    pass
                else:
                    issue("error", "controlled.stock_pct", f"Stock % {stock!r} is text; store a single figure as a number.", "Stock %")
                    reported = True
        if not stock_ok and not reported:
            issue(
                "error",
                "controlled.stock_pct",
                f"Stock % must be a number from 0 to 100, or Part stock, Not stated or Varies (E13); found {stock!r}.",
                "Stock %",
            )

        for marker in ("CVR/earnout", "Antitrust"):
            if not is_blank(record[marker]) and record[marker] not in MARKER:
                code = "controlled." + marker.lower().replace("/", "_")
                issue("error", code, f"{marker} must be Y or blank; a cohort's split goes in the Note (D1); "
                      f"found {record[marker]!r}.", marker)

        cvr_value = record["CVR/earnout value"]
        if not is_blank(cvr_value):
            if isinstance(cvr_value, bool) or not isinstance(cvr_value, (int, float)) or not math.isfinite(float(cvr_value)) or float(cvr_value) <= 0:
                issue("error", "bid.cvr_value_type", f"CVR/earnout value must be a positive number or blank; found {cvr_value!r}.", "CVR/earnout value")
            if record["CVR/earnout"] != "Y":
                issue("error", "bid.cvr_value_marker", "CVR/earnout value is filled, so CVR/earnout must be Y.", "CVR/earnout value")

        if record["Event"] == "Other-scope bid":
            for column in ("Price low", "Price high", "CVR/earnout value"):
                if not is_blank(record[column]):
                    issue(
                        "error",
                        "bid.other_scope_per_share",
                        f"{column} must be blank on Other-scope bid rows; give the amount, units and scope in the Note (D2).",
                        column,
                    )
            if is_blank(record["Note"]):
                issue("warning", "bid.other_scope_note", "Other-scope bid row has no Note; give the amount, units and scope there (D1).", "Note")

        if record["Antitrust"] == "Y" and record["Regulatory"] != "Concern":
            issue(
                "error",
                "conditions.antitrust_regulatory",
                f"Antitrust is Y but Regulatory is {record['Regulatory']!r}; Antitrust Y needs Regulatory Concern (E12).",
                "Antitrust",
            )

        if as_integer(record["Count"]) == 1:
            for column in TERM_COLUMNS:
                if record[column] == "Varies":
                    issue("error", "bid.varies_single", f"{column} is Varies on a row for one bidder; Varies is for cohort rows whose members differ or for which the filing reports the term for only some members (E12).", column)

        level = record["Conditions"]
        diligence, financing, regulatory = record["Due diligence"], record["Financing"], record["Regulatory"]
        note = normalize_contiguous(record["Note"]) if not is_blank(record["Note"]) else ""
        after_same = SAME_AS_RE.sub("", note, count=1)
        stated_heavy_trigger = bool(HEAVY_TRIGGER_RE.match(after_same))
        if financing == "Contingent" and level in CONDITIONS and level != "Heavy":
            issue("error", "conditions.financing_heavy", f"Financing is Contingent, so Conditions must be Heavy (E12, H1); found {level!r}.", "Conditions")
        if (stated_heavy_trigger and level in CONDITIONS and level != "Heavy"
                and not (level == "Unclear" and is_cohort(record))):
            issue("error", "conditions.heavy_trigger_level",
                  "The Note states an H1, H2 or H3 trigger, so Conditions must be Heavy (E12).", "Conditions")
        fully_supported = (diligence == "Complete" and financing in {"Committed", "Not needed"}
                           and regulatory in {"No concern", "Not stated"})
        silent_formal = (record["Formality"] == "Formal"
                         and diligence in {"Complete", "Not stated"}
                         and financing in {"Committed", "Not needed", "Not stated"}
                         and regulatory in {"No concern", "Not stated"})
        if level == "None" and (not (fully_supported or silent_formal) or record["Exclusivity"] == "Required"):
            issue(
                "error",
                "conditions.none_support",
                "Conditions None needs supported absence of remaining conditions (a Formal bid may be silent), "
                "and Exclusivity must not be Required (E12); found "
                f"{diligence!r}, {financing!r}, {regulatory!r}, {record['Exclusivity']!r}.",
                "Conditions",
            )
        if (level == "Unclear" and (fully_supported or silent_formal) and record["Exclusivity"] != "Required"
                and not stated_heavy_trigger and not is_cohort(record)):
            # E12 takes None before Light and Unclear; a silent Formal bid with no H trigger meets it.
            issue(
                "warning",
                "conditions.none_expected",
                "Conditions is Unclear, but E12 gives None: no H trigger, Exclusivity not Required, and Due diligence, "
                "Financing and Regulatory meet None's tests (a Formal bid may be silent on them); found "
                f"{record['Formality']!r}, {diligence!r}, {financing!r}, {regulatory!r}, {record['Exclusivity']!r}.",
                "Conditions",
            )
        if level == "Light" and diligence not in {"Complete", "Incomplete"} and record["Exclusivity"] != "Required":
            # A warning: E12's "only documentation remains" route can be read with diligence Not stated.
            issue(
                "warning",
                "conditions.light_support",
                f"Conditions Light needs Due diligence Complete or Incomplete (E12); found {diligence!r}.",
                "Conditions",
            )
        if record["Exclusivity"] == "Required" and level == "Unclear" and not is_cohort(record):
            issue("error", "conditions.exclusivity_level",
                  "Required exclusivity makes Conditions at least Light for a single bidder (E12).", "Conditions")
        if record["Formality"] == "Unclear" and not is_cohort(record):
            issue(
                "error",
                "bid.formality_unclear",
                "Formality Unclear is only for a cohort row whose members differ (D1); code Formal or Informal (E11).",
                "Formality",
            )
        if level == "Heavy" and not HEAVY_TRIGGER_RE.match(after_same):
            issue(
                "warning",
                "conditions.heavy_trigger",
                "A Heavy row's Note begins with its trigger, 'H1:', 'H2:' or 'H3:', after any 'Same as #n. ' (E12, D1).",
                "Note",
            )
        if record["Event"] == "Bid reaffirmed":
            if record["Formality"] != "Formal":
                issue("error", "bid.reaffirmed_formal", f"Bid reaffirmed is Formal (E11 route 3); found {record['Formality']!r}.", "Formality")
            if not SAME_AS_RE.match(note):
                issue("error", "bid.reaffirmed_same_as", "Bid reaffirmed is a Same-offer row; its Note begins 'Same as #n' (E10).", "Note")

    def check_review_sequences(self, ws: Any, rows: list[tuple[int, dict[str, Any]]]) -> None:
        """Review warnings that compare ledger rows (E10, E14)."""

        def who(record: dict[str, Any]) -> str:
            return normalize_contiguous(record["Who"]) if not is_blank(record["Who"]) else ""

        requests: dict[tuple[str, dt.date], Any] = {}
        for _, record in rows:
            sort_date = as_date(record["Sort date"])
            if record["Event"] in BID_EVENTS and record["Exclusivity"] in {"Requested", "Required"} and sort_date and who(record):
                requests.setdefault((who(record), sort_date), record["#"])
        for excel_row, record in rows:
            if record["Event"] != "Exclusivity changed":
                continue
            bid_number = requests.get((who(record), as_date(record["Sort date"])))
            if bid_number is not None:
                self.add(
                    "warning",
                    "ledger.exclusivity_duplicate",
                    f"Exclusivity changed has the same Who and Sort date as bid row #{bid_number}, which codes an "
                    "exclusivity request; if this row repeats the bid's own request, remove it (E10); a grant, "
                    "execution or extension is its own event.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Event",
                )

        def numbered_round(record: dict[str, Any]) -> int | None:
            return None if record["Round"] == "post" else as_integer(record["Round"])

        def day(record: dict[str, Any]) -> dt.date | None:
            return as_date(record["Date to"]) or as_date(record["Sort date"])

        # E14: a bidder asked to bid by a due date that has not bid by then exits at the due date.
        due_days: dict[tuple[Any, Any], set[dt.date]] = defaultdict(set)
        for _, record in rows:
            if record["Event"] == "Deadline" and numbered_round(record) is not None and day(record):
                due_days[(as_integer(record["Process"]), numbered_round(record))].add(day(record))
        for excel_row, record in rows:
            key = (as_integer(record["Process"]), numbered_round(record))
            exit_day = day(record)
            if record["Event"] == "Did not submit" and exit_day and due_days.get(key) and exit_day not in due_days[key]:
                self.add(
                    "warning",
                    "exit.did_not_submit_date",
                    f"Did not submit is dated {exit_day:%m/%d/%Y}, but no Deadline row of Process {key[0]} Round "
                    f"{key[1]} falls on that day; a bidder asked to bid by a due date that has not bid by then exits "
                    "at the due date (E14).",
                    sheet=ws.title,
                    row=excel_row,
                    column="Date to",
                )

        # Rounds whose Deadline outcome records that the target considered a late bid (E9).
        late_bid_rounds: set[tuple[Any, Any]] = set()
        if "Rounds" in self.wb:
            summary = self.wb["Rounds"]
            outcome_column = ROUND_COLUMNS.index("Deadline outcome") + 1
            for summary_row in nonempty_rows(summary, len(ROUND_COLUMNS)):
                outcome = summary.cell(summary_row, outcome_column).value
                if not is_blank(outcome) and "Extended (late bid accepted)" in [
                        part.strip() for part in normalize_contiguous(outcome).split(";")]:
                    late_bid_rounds.add((as_integer(summary.cell(summary_row, 1).value),
                                         as_integer(summary.cell(summary_row, 2).value)))

        exited: dict[tuple[Any, str], tuple[Any, Any, int | None]] = {}
        for excel_row, record in rows:
            key = (as_integer(record["Process"]), who(record))
            event = record["Event"]
            if not key[1]:
                continue
            if event in EXIT_EVENTS:
                exited[key] = (record["#"], event, numbered_round(record))
            elif event in {"Re-entered", "Bid", "Bid reaffirmed", "NDA signed", "Bidding group changed"} and key in exited:
                number, exit_event, exit_round = exited.pop(key)
                same_round_late_bid = (event in {"Re-entered", "Bid", "Bid reaffirmed"} and exit_event == "Did not submit"
                                       and exit_round is not None and numbered_round(record) == exit_round)
                if same_round_late_bid and (event == "Re-entered" or (key[0], exit_round) in late_bid_rounds):
                    self.add(
                        "warning",
                        "exit.late_bid_in_round",
                        f"{record['Who']} did not submit at #{number}, and this {event} row follows in the same round"
                        + ("" if event == "Re-entered" else ", whose Deadline outcome includes Extended (late bid accepted)")
                        + "; review the exit against the rules for a late bid (E14 closing events; E9).",
                        sheet=ws.title,
                        row=excel_row,
                        column="Event",
                    )
                elif event != "Re-entered" and not same_round_late_bid:
                    self.add(
                        "warning",
                        "exit.activity_without_reentry",
                        f"{record['Who']} left at #{number}, and this {event} row follows in the same process "
                        "with no Re-entered between them; add Re-entered if the bidder came back, or review the exit (E14).",
                        sheet=ws.title,
                        row=excel_row,
                        column="Event",
                    )

    def check_same_as_and_round_zero(self, ws: Any, rows: list[tuple[int, dict[str, Any]]]) -> None:
        """Checks that compare ledger rows: 'Same as #n' (E10) and Round 0 (E6)."""

        def who(record: dict[str, Any]) -> str:
            return normalize_contiguous(record["Who"]).casefold() if not is_blank(record["Who"]) else ""

        by_number = {as_integer(record["#"]): record for _, record in rows if as_integer(record["#"]) is not None}
        for excel_row, record in rows:
            match = SAME_AS_RE.match(normalize_contiguous(record["Note"])) if not is_blank(record["Note"]) else None
            if not match:
                continue
            target, number = by_number.get(int(match.group(1))), as_integer(record["#"])
            problem = None
            if record["Event"] not in BID_EVENTS:
                problem = "only a bid row is a Same-offer row"
            elif target is None or number is None or int(match.group(1)) >= number:
                problem = f"#{match.group(1)} is not an earlier row"
            elif target["Event"] not in BID_EVENTS:
                problem = f"#{match.group(1)} is a {target['Event']!r} row, not a bid row"
            elif who(target) != who(record):
                problem = f"#{match.group(1)} is {target['Who']!r}'s row, not this bidder's"
            if problem:
                self.add("error", "ledger.same_as", f"'Same as #{match.group(1)}' must point to an earlier bid row with the "
                         f"same Who (E10); {problem}.", sheet=ws.title, row=excel_row, column="Note")

        opened: set[int] = set()
        for excel_row, record in rows:
            process = as_integer(record["Process"])
            if record["Event"] == "Round opened" and as_integer(record["Round"]) == 1 and record["Round"] != "post":
                opened.add(process)
            elif (process in opened and as_integer(record["Round"]) == 0 and record["Round"] != "post"
                  and record["Event"] not in EXIT_EVENTS):
                # An exit row carries the round being left (E6), and the exits an opening causes follow its
                # Round opened row (E8, E14 rule 1), so a round-0 bidder's exit may follow round 1's opening.
                self.add("error", "round.zero_after_opening", f"Round 0 is only for rows before round 1 of Process {process} "
                         "opens (D1, E6); this row follows its Round opened row.", sheet=ws.title, row=excel_row, column="Round")

    def check_ledger(self) -> dict[str, Any]:
        ws = self.wb["Deal ledger"]
        ledger_columns = LEDGER_COLUMNS
        schema_ok = self.check_schema(ws, ledger_columns)
        self.check_presentation(ws, len(ledger_columns))
        if not schema_ok:
            return {
                "rows": [],
                "by_number": {},
                "rounds": defaultdict(list),
                "round_openings": defaultdict(list),
                "deadlines": Counter(),
                "deadline_rows": defaultdict(list),
                "flags": {},
                "processes": set(),
            }

        columns = {name: index + 1 for index, name in enumerate(ledger_columns)}
        rows = nonempty_rows(ws, len(ledger_columns))
        records: list[dict[str, Any]] = []
        by_number: dict[int, int] = {}
        round_rows: dict[tuple[int, int], list[int]] = defaultdict(list)
        round_openings: dict[tuple[int, int], list[tuple[int, dt.date | None]]] = defaultdict(list)
        deadline_counts: Counter[tuple[int, int]] = Counter()
        deadline_rows: dict[tuple[int, int], list[tuple[str, dt.date | None]]] = defaultdict(list)
        flags: dict[int, set[str]] = {}
        processes: set[int] = set()
        previous_sort: tuple[dt.date, int] | None = None

        # D1: On Merger agreement signed, Who is the signing acquirer, never the target.
        target_names = {"target", "the target", "company", "the company"}
        if "Deal facts" in self.wb:
            facts_sheet = self.wb["Deal facts"]
            for fact_row in nonempty_rows(facts_sheet, len(FACT_COLUMNS)):
                if facts_sheet.cell(fact_row, 1).value == "Target" and unit_key(facts_sheet.cell(fact_row, 2).value):
                    target_names.add(unit_key(facts_sheet.cell(fact_row, 2).value))

        def value(row: int, column: str) -> Any:
            return ws.cell(row, columns[column]).value

        for ordinal, excel_row in enumerate(rows, start=1):
            record = {name: value(excel_row, name) for name in ledger_columns}
            records.append(record)
            row_number = as_integer(record["#"])
            if row_number != ordinal:
                self.add(
                    "error",
                    "ledger.sequence",
                    f"Event number must be {ordinal}; found {record['#']!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="#",
                )
            elif row_number in by_number:
                self.add(
                    "error",
                    "ledger.duplicate_number",
                    f"Duplicate event number {row_number}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="#",
                )
            else:
                by_number[row_number] = excel_row

            if is_blank(record["Who"]):
                self.add(
                    "error",
                    "ledger.who_blank",
                    "Who must never be blank.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Who",
                )

            type_value = record["Type"]
            if not is_blank(type_value) and type_value not in TYPES:
                self.add(
                    "error",
                    "controlled.type",
                    f"Type {type_value!r} is not allowed.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Type",
                )

            event = record["Event"]
            if event not in EVENTS:
                self.add(
                    "error",
                    "controlled.event",
                    f"Event {event!r} is not an allowed label.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Event",
                )
            if event == "Activist" and not str(record["Note"] or "").startswith(("Demands sale", "Sale one option")):
                self.add("warning", "activist.note_prefix",
                         "An Activist Note begins 'Demands sale' or 'Sale one option' (D2).",
                         sheet=ws.title, row=excel_row, column="Note")

            process = as_integer(record["Process"])
            if process is None or process < 1:
                self.add(
                    "error",
                    "ledger.process",
                    f"Process must be a positive integer; found {record['Process']!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Process",
                )
            else:
                processes.add(process)

            round_value = record["Round"]
            numbered_round = as_integer(round_value)
            if round_value == "post":
                numbered_round = None
            elif numbered_round is None or numbered_round < 0:
                self.add(
                    "error",
                    "ledger.round",
                    f"Round must be a nonnegative integer or 'post'; found {round_value!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Round",
                )

            if process is not None and process >= 1 and numbered_round is not None and numbered_round > 0:
                key = (process, numbered_round)
                round_rows[key].append(excel_row)
                if event == "Round opened":
                    round_openings[key].append((excel_row, as_date(record["Sort date"])))
                if event == "Deadline":
                    deadline_counts[key] += 1
                if event in {"Deadline", "Deadline set", "Deadline revised"}:
                    deadline_rows[key].append((event, as_date(record["Sort date"])))
            elif event == "Round opened":
                self.add(
                    "error",
                    "round.opening_key",
                    "Round opened must carry a positive numbered round.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Round",
                )

            count = record["Count"]
            if not is_blank(count) and (not is_integer(count) or int(count) <= 0):
                self.add(
                    "error",
                    "ledger.count",
                    f"Count must be a positive integer or blank; found {count!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Count",
                )
            if event in NO_BIDDER_COUNT_EVENTS and not is_blank(count):
                self.add(
                    "error",
                    "ledger.count_nonbidder",
                    f"Count must be blank on {event} rows.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Count",
                )
            if event in PROCESS_MARKERS and is_blank(count):
                self.add(
                    "warning",
                    "ledger.count_marker",
                    f"Count is blank on a {event} row; it holds the number of open participations the marker "
                    "closes, and stays blank only when it closes none (E5).",
                    sheet=ws.title,
                    row=excel_row,
                    column="Count",
                )
            if event in BID_EVENTS | {"Re-entered", "Bidding group changed"} and is_blank(count):
                count_note = normalize_contiguous(record["Note"])
                if COUNT_QUALIFIER_RE.search(count_note):
                    pass  # a qualified figure keeps its qualifier, with Count blank (E3)
                elif COUNT_RANGE_RE.search(count_note) or COUNT_UNKNOWN_RE.search(count_note):
                    self.add(
                        "warning",
                        "ledger.count_range",
                        "Count is blank and the Note gives a range or an unknown population; keep the filing's own "
                        "figure or qualifier ('Count: more than ten') and create no range (B, E3).",
                        sheet=ws.title,
                        row=excel_row,
                        column="Count",
                    )
                else:
                    self.add(
                        "error",
                        "ledger.count_bidder",
                        f"Count is required on {event} rows unless the filing's figure is qualified, with the Note "
                        "beginning 'Count: more than ten' (E3, E4).",
                        sheet=ws.title,
                        row=excel_row,
                        column="Count",
                    )

            for price_column in ("Price low", "Price high"):
                price = record[price_column]
                if not is_blank(price) and (
                    isinstance(price, bool)
                    or not isinstance(price, (int, float))
                    or not math.isfinite(float(price))
                    or float(price) <= 0
                ):
                    self.add(
                        "error",
                        "bid.price_type",
                        f"{price_column} must be a positive number or blank; found {price!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column=price_column,
                    )
            low, high = record["Price low"], record["Price high"]
            if isinstance(low, (int, float)) and not isinstance(low, bool) and isinstance(high, (int, float)) and not isinstance(high, bool) and low > high:
                self.add(
                    "error",
                    "bid.price_order",
                    f"Price low {low!r} exceeds Price high {high!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Price low",
                )
            if event in BID_EVENTS:
                controlled = [("Formality", FORMALITY), ("Conditions", CONDITIONS), *CONDITION_COLUMNS.items()]
                for field, allowed in controlled:
                    field_value = record[field]
                    if field_value not in allowed:
                        self.add(
                            "error",
                            f"controlled.{field.lower().replace(' ', '_')}",
                            f"{field} must be filled with an allowed value on {event} rows; found {field_value!r}.",
                            sheet=ws.title,
                            row=excel_row,
                            column=field,
                        )
                self.check_bid_terms(ws, excel_row, record)
                if is_blank(low) != is_blank(high):
                    note = normalize_contiguous(record["Note"]).lower()
                    markers = ("at least", "no more than", "floor", "ceiling", "\u2265", "\u2264")
                    if not any(marker in note for marker in markers):
                        self.add(
                            "warning",
                            "bid.one_sided_price",
                            "Only one price endpoint is filled, but the Note does not visibly identify a floor or ceiling.",
                            sheet=ws.title,
                            row=excel_row,
                            column="Note",
                        )
            else:
                bid_fields = ("Price low", "Price high", "Formality", "Conditions", *TERM_COLUMNS)
                for field in bid_fields:
                    if not is_blank(record[field]):
                        self.add(
                            "error",
                            "bid.fields_on_nonbid",
                            f"{field} must be blank because {event!r} is not a bid row.",
                            sheet=ws.title,
                            row=excel_row,
                            column=field,
                        )

            exit_reason = record["Exit reason"]
            if event in EXIT_EVENTS:
                if exit_reason not in EXIT_REASONS:
                    self.add(
                        "error",
                        "controlled.exit_reason",
                        f"Exit reason must be filled with an allowed value on {event} rows; found {exit_reason!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Exit reason",
                    )
            elif not is_blank(exit_reason):
                self.add(
                    "error",
                    "ledger.exit_reason_nonexit",
                    "Exit reason is permitted only on exit rows.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Exit reason",
                )

            inferred = record["Inferred"]
            if inferred not in (None, "", "Y"):
                self.add(
                    "error",
                    "controlled.inferred",
                    f"Inferred must be 'Y' or blank; found {inferred!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Inferred",
                )
            if inferred == "Y" and event not in INFERRED_EVENTS:
                self.add(
                    "error",
                    "ledger.inferred_event",
                    f"Inferred = Y marks only an inferred exit, Round opened or Process restarted row (B); "
                    f"this is a {event!r} row. Coding a column is never an inference.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Inferred",
                )
            needs_note = event == "Process restarted" or (event == "Did not submit" and is_cohort(record))
            if inferred == "Y" and needs_note and is_blank(record["Note"]):
                self.add(
                    "warning",
                    "ledger.inference_note",
                    "An inferred Process restarted row gives the last reported acquirer contact in its Note (E5), and "
                    "an inferred cohort Did not submit row gives its Count arithmetic (E3).",
                    sheet=ws.title,
                    row=excel_row,
                    column="Note",
                )

            if word_count(record["Note"]) > 40:
                self.add(
                    "error",
                    "ledger.note_length",
                    f"Note has {word_count(record['Note'])} words; the limit is 40 (D1).",
                    sheet=ws.title,
                    row=excel_row,
                    column="Note",
                )

            quote_cell = record["Quote and page"]
            parsed_quote = parse_quote_and_page(quote_cell)
            if parsed_quote is None:
                self.add(
                    "error",
                    "quote.format",
                    "Quote and page must end with a p. or pp. printed-page reference.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Quote and page",
                )
            else:
                quote, _ = parsed_quote
                if not quote:
                    self.add(
                        "error",
                        "quote.empty",
                        "The quoted passage is empty.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Quote and page",
                    )
                elif word_count(quote) > 30:
                    self.add(
                        "error",
                        "quote.length",
                        f"Quotation has {word_count(quote)} words; maximum is 30.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Quote and page",
                    )
                if quote and not any(quote in filing_text for filing_text in self.filing_texts):
                    self.add(
                        "error",
                        "quote.not_contiguous_in_filing",
                        "Normalized quotation does not occur as one contiguous passage in the parsed full filing.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Quote and page",
                    )

            flag = record["Flag"]
            if not is_blank(flag):
                flag_ids = parse_flag_ids(flag)
                if flag_ids is None:
                    self.add(
                        "error",
                        "controlled.flag",
                        f"Flag must contain only Question or Review ids such as Q1, R1; found {flag!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Flag",
                    )
                elif row_number is not None:
                    flags[row_number] = flag_ids

            if not is_blank(record["Reviewer note"]):
                self.add(
                    "error",
                    "reviewer_note.not_empty",
                    "Reviewer note must be empty.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Reviewer note",
                )

            parsed_dates: dict[str, dt.date | None] = {}
            for date_column in DATE_COLUMNS:
                date_value = record[date_column]
                parsed = as_date(date_value)
                parsed_dates[date_column] = parsed
                cell = ws.cell(excel_row, columns[date_column])
                if date_column == "Sort date" and is_blank(date_value):
                    self.add(
                        "error",
                        "date.sort_blank",
                        "Sort date must always be filled.",
                        sheet=ws.title,
                        row=excel_row,
                        column=date_column,
                    )
                elif not is_blank(date_value) and parsed is None:
                    self.add(
                        "error",
                        "date.not_excel_date",
                        f"{date_column} must be a real Excel date; found {date_value!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column=date_column,
                    )
                elif parsed is not None:
                    if date_has_time(date_value):
                        self.add(
                            "error",
                            "date.has_time",
                            f"{date_column} contains a time component.",
                            sheet=ws.title,
                            row=excel_row,
                            column=date_column,
                        )
                    if not has_mmddyyyy_format(cell):
                        self.add(
                            "error",
                            "date.format",
                            f"{date_column} must use Excel format MM/DD/YYYY; found {cell.number_format!r}.",
                            sheet=ws.title,
                            row=excel_row,
                            column=date_column,
                        )

            sort_date = parsed_dates["Sort date"]
            date_from = parsed_dates["Date from"]
            date_to = parsed_dates["Date to"]
            if date_from and date_to and date_from > date_to:
                self.add(
                    "error",
                    "date.window_reversed",
                    "Date from is later than Date to.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Date from",
                )
            when = record["When"]
            if isinstance(when, str) and re.fullmatch(r"\d{2}/\d{2}/\d{4}", when.strip()):
                try:
                    reported_day = dt.datetime.strptime(when.strip(), "%m/%d/%Y").date()
                except ValueError:
                    reported_day = None
                for date_column in DATE_COLUMNS:
                    if reported_day and parsed_dates[date_column] not in (None, reported_day):
                        self.add(
                            "error",
                            "date.exact_day_mismatch",
                            f"When reports the day {when.strip()}, so {date_column} must equal it (E8).",
                            sheet=ws.title,
                            row=excel_row,
                            column=date_column,
                        )
            if sort_date and date_from and sort_date < date_from:
                self.add(
                    "error",
                    "date.sort_outside_window",
                    "Sort date falls before Date from.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Sort date",
                )
            if sort_date and date_to and sort_date > date_to:
                self.add(
                    "error",
                    "date.sort_outside_window",
                    "Sort date falls after Date to.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Sort date",
                )
            if sort_date and previous_sort and sort_date < previous_sort[0]:
                self.add(
                    "error",
                    "date.sort_decreases",
                    f"Sort date {sort_date:%m/%d/%Y} is earlier than the prior populated row's "
                    f"{previous_sort[0]:%m/%d/%Y} (Excel row {previous_sort[1]}).",
                    sheet=ws.title,
                    row=excel_row,
                    column="Sort date",
                )
            if sort_date:
                previous_sort = (sort_date, excel_row)

        # D1: on Merger agreement signed, Count is 1 for a whole-company signer, otherwise blank. Only the
        # signer's bid rows are compared; who is live at signing is not reconstructed.
        bid_events_of: dict[str, set[str]] = defaultdict(set)
        for record in records:
            if record["Event"] in BID_EVENTS and unit_key(record["Who"]):
                bid_events_of[unit_key(record["Who"])].add(record["Event"])
        for excel_row, record in zip(rows, records):
            if record["Event"] != "Merger agreement signed":
                continue
            signer = unit_key(record["Who"])
            if signer in target_names:
                self.add("warning", "ledger.signing_who",
                         "On Merger agreement signed, Who is the signing acquirer, not the target (D1).",
                         sheet=ws.title, row=excel_row, column="Who")
            count, signer_bids = record["Count"], bid_events_of.get(signer, set())
            if not is_blank(count) and as_integer(count) != 1:
                self.add("error", "ledger.count_signing",
                         f"Count on Merger agreement signed is 1 or blank (D1); found {count!r}.",
                         sheet=ws.title, row=excel_row, column="Count")
            elif is_blank(count) and signer_bids & {"Bid", "Bid reaffirmed"}:
                self.add("warning", "ledger.count_signing_scope",
                         "Count is blank, but the signer has a whole-company Bid row; a whole-company signer has "
                         "Count 1 (D1).", sheet=ws.title, row=excel_row, column="Count")
            elif as_integer(count) == 1 and signer_bids == {"Other-scope bid"}:
                self.add("warning", "ledger.count_signing_scope",
                         "Count is 1, but every bid row of the signer is an Other-scope bid; a signer outside the "
                         "whole-company contest has blank Count (D1).", sheet=ws.title, row=excel_row, column="Count")

        self.check_review_sequences(ws, list(zip(rows, records)))
        self.check_same_as_and_round_zero(ws, list(zip(rows, records)))

        expected_processes = list(range(1, max(processes) + 1)) if processes else []
        if sorted(processes) != expected_processes:
            self.add(
                "error",
                "ledger.process_sequence",
                f"Processes must be consecutive from 1; found {sorted(processes)}.",
                sheet=ws.title,
            )
        for process in sorted(processes):
            used = sorted(round_number for proc, round_number in round_rows if proc == process)
            if used and used != list(range(1, max(used) + 1)):
                self.add(
                    "error",
                    "round.sequence",
                    f"Numbered rounds for Process {process} must be consecutive from 1; found {used}.",
                    sheet=ws.title,
                )
        for key, key_rows in round_rows.items():
            openings = round_openings.get(key, [])
            if len(openings) != 1:
                self.add(
                    "error",
                    "round.opening_count",
                    f"Process {key[0]} Round {key[1]} has {len(openings)} Round opened rows; expected exactly one.",
                    sheet=ws.title,
                    row=key_rows[0],
                    column="Event",
                )
            elif openings[0][0] != min(key_rows):
                self.add(
                    "error",
                    "round.opening_order",
                    f"Round opened is not the first ledger row assigned to Process {key[0]} Round {key[1]}.",
                    sheet=ws.title,
                    row=openings[0][0],
                    column="Event",
                )

        return {
            "rows": rows,
            "records": records,
            "by_number": by_number,
            "rounds": round_rows,
            "round_openings": round_openings,
            "deadlines": deadline_counts,
            "deadline_rows": deadline_rows,
            "flags": flags,
            "processes": processes,
        }

    def check_rounds(self, ledger: dict[str, Any]) -> None:
        ws = self.wb["Rounds"]
        schema_ok = self.check_schema(ws, ROUND_COLUMNS)
        self.check_presentation(ws, len(ROUND_COLUMNS))
        if not schema_ok:
            return

        rows = nonempty_rows(ws, len(ROUND_COLUMNS))
        keys: list[tuple[int, int]] = []
        opened: dict[tuple[int, int], tuple[int, dt.date | None]] = {}

        def none_stated(text: str) -> bool:
            # Compared by prefix, so "none stated (…)" still reads as no due date.
            return text.lower().startswith("none stated")

        for excel_row in rows:
            values = {name: ws.cell(excel_row, col + 1).value for col, name in enumerate(ROUND_COLUMNS)}
            process, round_number = as_integer(values["Process"]), as_integer(values["Round"])
            if process is None or process < 1 or round_number is None or round_number < 1:
                self.add(
                    "error",
                    "rounds.key",
                    f"Process and Round must be positive integers; found {values['Process']!r}, {values['Round']!r}.",
                    sheet=ws.title,
                    row=excel_row,
                )
                continue
            key = (process, round_number)
            keys.append(key)
            if key in opened:
                self.add(
                    "error",
                    "rounds.duplicate",
                    f"Duplicate Process {process} Round {round_number} summary row.",
                    sheet=ws.title,
                    row=excel_row,
                )

            for field in ROUND_COLUMNS[2:]:
                if field == "Deadline outcome":
                    # Future or superseded deadlines can leave this cell blank;
                    # validate it against the reached Deadline rows below.
                    continue
                if is_blank(values[field]):
                    self.add(
                        "error",
                        "rounds.required",
                        f"{field} must be filled on every Rounds row.",
                        sheet=ws.title,
                        row=excel_row,
                        column=field,
                    )

            opened_date = as_date(values["Opened"])
            opened[key] = (excel_row, opened_date)
            opened_cell = ws.cell(excel_row, ROUND_COLUMNS.index("Opened") + 1)
            if opened_date is None and not is_blank(values["Opened"]):
                self.add(
                    "error",
                    "rounds.opened_type",
                    f"Opened must be a real Excel date; found {values['Opened']!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Opened",
                )
            elif opened_date is not None and not has_mmddyyyy_format(opened_cell):
                self.add(
                    "error",
                    "date.format",
                    f"Opened must use Excel format MM/DD/YYYY; found {opened_cell.number_format!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Opened",
                )

            if values["Finality"] not in FINALITY:
                self.add(
                    "error",
                    "controlled.finality",
                    f"Finality {values['Finality']!r} is not allowed.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Finality",
                )

            due_dates = normalize_contiguous(values["Due dates"])
            outcome = (
                "" if is_blank(values["Deadline outcome"])
                else normalize_contiguous(values["Deadline outcome"])
            )
            deadline_count = ledger["deadlines"].get(key, 0)
            if outcome == "No deadline stated":
                if not none_stated(due_dates):
                    self.add(
                        "error",
                        "rounds.no_deadline_pair",
                        "Deadline outcome is 'No deadline stated' but Due dates is not 'none stated'.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Due dates",
                    )
                if deadline_count:
                    self.add(
                        "error",
                        "rounds.deadline_count",
                        f"Rounds says no deadline, but the ledger has {deadline_count} Deadline row(s) for this round.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Deadline outcome",
                    )
            elif outcome:
                outcomes = [part.strip() for part in outcome.split(";")]
                bad = [part for part in outcomes if part not in DEADLINE_OUTCOMES]
                if bad:
                    self.add(
                        "error",
                        "controlled.deadline_outcome",
                        f"Unallowed deadline outcome value(s): {bad!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Deadline outcome",
                    )
                if deadline_count != len(outcomes):
                    self.add(
                        "error",
                        "rounds.deadline_count",
                        f"Ledger has {deadline_count} Deadline row(s), but Rounds lists {len(outcomes)} outcome(s).",
                        sheet=ws.title,
                        row=excel_row,
                        column="Deadline outcome",
                    )
                else:
                    # Outcomes follow the round's Deadline rows in ledger order.
                    round_rows = ledger["deadline_rows"].get(key, [])
                    reached = [day for event, day in round_rows if event == "Deadline"]
                    new_dates = [day for event, day in round_rows if event != "Deadline" and day]
                    for position, (part, day) in enumerate(zip(outcomes, reached, strict=True), start=1):
                        if part == "Extended" and day and not any(new_day >= day for new_day in new_dates):
                            self.add(
                                "warning",
                                "rounds.extended_without_new_date",
                                f"Outcome {position} is Extended, but the ledger has no Deadline set or Deadline revised "
                                f"row in this round on or after that deadline ({day:%m/%d/%Y}); record the new due date "
                                "or review the outcome (E9).",
                                sheet=ws.title,
                                row=excel_row,
                                column="Deadline outcome",
                            )
            elif deadline_count:
                self.add(
                    "error",
                    "rounds.deadline_count",
                    f"Ledger has {deadline_count} Deadline row(s), but Deadline outcome is blank.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Deadline outcome",
                )
            elif none_stated(due_dates):
                self.add(
                    "error",
                    "rounds.no_deadline_pair",
                    "Due dates is 'none stated', so Deadline outcome must be 'No deadline stated'.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Deadline outcome",
                )

        if keys != sorted(keys):
            self.add(
                "error",
                "rounds.order",
                "Rounds rows must be ordered by Process and Round.",
                sheet=ws.title,
            )
        ledger_keys = set(ledger["rounds"])
        sheet_keys = set(keys)
        for key in sorted(ledger_keys - sheet_keys):
            self.add(
                "error",
                "rounds.missing_summary",
                f"Process {key[0]} Round {key[1]} appears in the ledger but not in Rounds.",
                sheet=ws.title,
            )
        for key in sorted(sheet_keys - ledger_keys):
            row = opened[key][0] if key in opened else None
            self.add(
                "error",
                "rounds.orphan_summary",
                f"Process {key[0]} Round {key[1]} appears in Rounds but not in the ledger.",
                sheet=ws.title,
                row=row,
            )
        for key in sorted(ledger_keys & sheet_keys):
            ledger_openings = ledger["round_openings"].get(key, [])
            if len(ledger_openings) == 1 and key in opened:
                ledger_date = ledger_openings[0][1]
                summary_date = opened[key][1]
                if ledger_date is not None and summary_date is not None and ledger_date != summary_date:
                    self.add(
                        "error",
                        "rounds.opened_mismatch",
                        f"Opened {summary_date:%m/%d/%Y} does not equal the Round opened row's Sort date "
                        f"{ledger_date:%m/%d/%Y}.",
                        sheet=ws.title,
                        row=opened[key][0],
                        column="Opened",
                    )

    def check_questions(self, ledger: dict[str, Any]) -> None:
        ws = self.wb["Questions"]
        schema_ok = self.check_schema(ws, QUESTION_COLUMNS)
        self.check_presentation(ws, len(QUESTION_COLUMNS))
        if not schema_ok:
            return

        rows = nonempty_rows(ws, len(QUESTION_COLUMNS))
        review_rows: dict[str, int] = {}
        affected_by_id: dict[str, tuple[set[int], bool]] = {}
        marker_rows = {
            as_integer(record["#"]) for record in ledger.get("records", [])
            if record.get("Event") in PROCESS_MARKERS and as_integer(record.get("#")) is not None
        }
        process_question: int | None = None
        q_count = r_count = 0
        for excel_row in rows:
            values = {name: ws.cell(excel_row, col + 1).value for col, name in enumerate(QUESTION_COLUMNS)}
            qid = str(values["Q"]).strip() if not is_blank(values["Q"]) else ""
            is_review = qid.startswith("R")
            if is_review:
                r_count += 1
            else:
                q_count += 1
            expected = f"R{r_count}" if is_review else f"Q{q_count}"
            if qid != expected:
                self.add(
                    "error",
                    "questions.sequence",
                    f"Question id must be {expected}; found {values['Q']!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Q",
                )
            if qid in review_rows:
                self.add(
                    "error",
                    "questions.duplicate",
                    f"Duplicate question id {qid!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Q",
                )
            elif REVIEW_ID_RE.fullmatch(qid):
                review_rows[qid] = excel_row

            required_fields = QUESTION_COLUMNS[1:-1] if not is_review else QUESTION_COLUMNS[1:-2]
            for field in required_fields:
                if is_blank(values[field]):
                    self.add(
                        "error",
                        "questions.required",
                        f"{field} must be filled.",
                        sheet=ws.title,
                        row=excel_row,
                        column=field,
                    )
            if not is_blank(values["Reviewer note"]):
                self.add(
                    "error",
                    "reviewer_note.not_empty",
                    "Reviewer note must be empty.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Reviewer note",
                )

            refs, parseable = parse_affected_rows(values["Rows affected"])
            if qid:
                affected_by_id[qid] = (refs, parseable)
            if not is_review and not is_blank(values["Rows affected"]) and not parseable:
                self.add(
                    "warning",
                    "questions.rows_unparsed",
                    "Rows affected is narrative/global, so its row links could not be checked mechanically.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Rows affected",
                )
            for ref in sorted(refs):
                if ref not in ledger["by_number"]:
                    self.add(
                        "error",
                        "questions.unknown_row",
                        f"Rows affected cites nonexistent ledger row #{ref}.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Rows affected",
                    )
                elif qid not in ledger["flags"].get(ref, set()):
                    self.add(
                        "warning",
                        "questions.flag_mismatch",
                        f"Rows affected cites #{ref}, but that ledger row's focused Flag does not include {qid!r}.",
                        sheet=ws.title,
                        row=excel_row,
                        column="Rows affected",
                    )

            total_words = sum(word_count(values[field]) for field in QUESTION_COLUMNS[1:-1])
            if not is_review and process_question is None and is_process_question(values["Question"], values["Rows affected"], marker_rows):
                process_question = excel_row
            if total_words > QUESTION_WORDS:
                self.add(
                    "warning",
                    "questions.length",
                    f"Question entry has {total_words} words across its fields; the limit is {QUESTION_WORDS} (D4).",
                    sheet=ws.title,
                    row=excel_row,
                )

        for row_number, qids in ledger["flags"].items():
            excel_row = ledger["by_number"].get(row_number)
            for qid in sorted(qids):
                if qid not in review_rows:
                    self.add(
                        "error",
                        "questions.unknown_flag",
                        f"Flag {qid} has no Questions row.",
                        sheet="Deal ledger",
                        row=excel_row,
                        column="Flag",
                    )
                    continue
                refs, parseable = affected_by_id.get(qid, (set(), False))
                if parseable and row_number not in refs:
                    self.add(
                        "warning",
                        "questions.reverse_link",
                        f"Ledger row #{row_number} is flagged {qid}, but {qid}'s broad Rows affected range does not cite it.",
                        sheet="Deal ledger",
                        row=excel_row,
                        column="Flag",
                    )

        counted = q_count - (process_question is not None)
        if counted > QUESTION_CAP:
            self.add(
                "warning",
                "questions.count",
                f"{counted} Questions besides the process Question; raise at most {QUESTION_CAP} (F).",
                sheet=ws.title,
            )
        inferred_openings = [record.get("#") for record in ledger.get("records", [])
                             if record.get("Event") == "Round opened" and record.get("Inferred") == "Y"]
        reasons = (["the ledger has more than one process"] if len(ledger["processes"]) > 1 else []) + (
            [f"a round opened by inference (#{', #'.join(str(n) for n in inferred_openings)})"] if inferred_openings else [])
        if reasons and process_question is None:
            because = " and ".join(reasons)
            self.add(
                "warning",
                "questions.process_missing",
                f"{because[0].upper() + because[1:]}, but no Question is recognisably the process Question "
                "(its Question begins 'Process:' or its Rows affected cites a Process terminated or restarted row) (F).",
                sheet=ws.title,
            )

    def check_facts(self, ledger: dict[str, Any]) -> None:
        ws = self.wb["Deal facts"]
        schema_ok = self.check_schema(ws, FACT_COLUMNS)
        self.check_presentation(ws, len(FACT_COLUMNS))
        if not schema_ok:
            return

        rows = nonempty_rows(ws, len(FACT_COLUMNS))
        field_options = FACT_FIELD_OPTIONS
        actual_fields = [ws.cell(row, 1).value for row in rows]
        fields_match = len(actual_fields) == len(field_options) and all(
            actual in allowed for actual, allowed in zip(actual_fields, field_options, strict=True)
        )
        if not fields_match:
            self.add(
                "error",
                "facts.fields",
                f"Deal facts fields must appear exactly in the prescribed order; found {actual_fields!r}.",
                sheet=ws.title,
                row=rows[0] if rows else 2,
                column="Field",
            )
        values: dict[str, Any] = {}
        for excel_row in rows:
            field = ws.cell(excel_row, 1).value
            value = ws.cell(excel_row, 2).value
            if isinstance(field, str):
                values[field] = value
            if is_blank(value):
                self.add(
                    "error",
                    "facts.value_blank",
                    f"Value is blank for field {field!r}.",
                    sheet=ws.title,
                    row=excel_row,
                    column="Value",
                )

        acquirer_type = values.get("Acquirer type")
        if not is_blank(acquirer_type) and not starts_with_canonical(acquirer_type, TYPES):
            self.add(
                "error",
                "controlled.acquirer_type",
                f"Acquirer type {acquirer_type!r} is not allowed.",
                sheet=ws.title,
                row=self._fact_row(rows, ws, "Acquirer type"),
                column="Value",
            )
        initiation = values.get("Initiation")
        if not is_blank(initiation) and not starts_with_canonical(initiation, INITIATION):
            self.add(
                "error",
                "controlled.initiation",
                f"Initiation {initiation!r} is not allowed.",
                sheet=ws.title,
                row=self._fact_row(rows, ws, "Initiation"),
                column="Value",
            )
        elif starts_with_canonical(initiation, {"activist-influenced"}):
            ledger_rows = [record for record in ledger.get("records", []) if as_integer(record.get("Process")) == 1]
            first_target = first_target_step(ledger_rows)
            if not any(i < first_target and record.get("Event") == "Activist"
                       and str(record.get("Note") or "").startswith("Demands sale")
                       for i, record in enumerate(ledger_rows)):
                self.add("error", "initiation.activist_support",
                         "Activist-influenced needs a 'Demands sale' Activist row before the target's first sale step (D5).",
                         sheet=ws.title, row=self._fact_row(rows, ws, "Initiation"), column="Value")
        process_count = values.get("Number of processes")
        expected_process_count = len(ledger["processes"])
        parsed_process_count = as_integer_or_text(process_count)
        if not is_blank(process_count) and parsed_process_count != expected_process_count:
            self.add(
                "error",
                "facts.process_count",
                f"Number of processes must equal the ledger's {expected_process_count}; found {process_count!r}.",
                sheet=ws.title,
                row=self._fact_row(rows, ws, "Number of processes"),
                column="Value",
            )
        auction_field = next((field for field in AUCTION_SCREEN_FIELDS if field in values), None)
        auction = normalize_contiguous(values.get(auction_field)) if auction_field else ""
        if auction and (re.match(r"^(Met|Not met)(?:\b|:)", auction) is None or re.search(r"\b\d+\b", auction) is None
                        or re.search(r"\b(?:Uncertain|count unknown)\b", auction, re.IGNORECASE)):
            self.add(
                "error",
                "facts.auction_screen",
                "Auction screen entries start with Met or Not met and give the number of parties (E1), "
                "for example 'Met (process 1): 3 parties'.",
                sheet=ws.title,
                row=self._fact_row(rows, ws, auction_field),
                column="Value",
            )
        whole_company_field = next((field for field in WHOLE_COMPANY_FIELDS if field in values), None)
        whole_company = normalize_contiguous(values.get(whole_company_field)) if whole_company_field else ""
        if whole_company and not starts_with_canonical(whole_company, {"Yes", "No"}):
            self.add(
                "error",
                "facts.whole_company_bids",
                "Whole-company bids must be 'Yes', or start with 'No' and say what was bid for.",
                sheet=ws.title,
                row=self._fact_row(rows, ws, whole_company_field),
                column="Value",
            )
        account_field = next((field for field in ACCOUNT_FIELDS if field in values), None)
        account = normalize_contiguous(values.get(account_field)) if account_field else ""
        if account:
            sentence_count = len(re.findall(r"(?:[.!?](?:[\"\u201d']*)\s+|[.!?](?:[\"\u201d']*)$)", account))
            if sentence_count not in (5, 6):
                self.add(
                    "warning",
                    "facts.account_sentences",
                    f"Account appears to have {sentence_count} sentences; the instruction asks for five or six.",
                    sheet=ws.title,
                    row=self._fact_row(rows, ws, account_field),
                    column="Value",
                )

    @staticmethod
    def _fact_row(rows: list[int], ws: Any, field: str) -> int | None:
        return next((row for row in rows if ws.cell(row, 1).value == field), None)

    def run(self) -> dict[str, Any]:
        if self.load_inputs():
            actual_sheets = self.wb.sheetnames
            if actual_sheets != SHEETS:
                self.add(
                    "error",
                    "schema.sheets",
                    f"Workbook must contain exactly {SHEETS!r} in that order; found {actual_sheets!r}.",
                )
            available = set(actual_sheets)
            ledger: dict[str, Any] = {
                "rows": [],
                "by_number": {},
                "rounds": defaultdict(list),
                "round_openings": defaultdict(list),
                "deadlines": Counter(),
                "deadline_rows": defaultdict(list),
                "flags": {},
                "processes": set(),
            }
            if "Deal ledger" in available:
                ledger = self.check_ledger()
            if "Rounds" in available:
                self.check_rounds(ledger)
            if "Questions" in available:
                self.check_questions(ledger)
            if "Deal facts" in available:
                self.check_facts(ledger)

        counts = Counter(issue["severity"] for issue in self.issues)
        if self.fatal:
            status = "error"
        elif counts["error"]:
            status = "fail"
        elif counts["warning"]:
            status = "pass_with_warnings"
        else:
            status = "pass"
        return {
            "checker": "lean-mechanical",
            "checker_version": CHECKER_VERSION,
            "checker_revision": CHECKER_REVISION,
            "ledger_schema": LEDGER_SCHEMA,
            "workbook": str(self.workbook_path),
            "filing": str(self.filing_path),
            "status": status,
            "scope_note": (
                "Mechanical checks only. Quote occurrence proves source occurrence, not semantic support. "
                "No bidder-population arithmetic or substantive extraction grading is performed."
            ),
            "summary": {
                "errors": counts["error"],
                "warnings": counts["warning"],
                "information": counts["info"],
                "total_issues": len(self.issues),
            },
            "issues": [{**issue, "basis": "mechanical"} for issue in self.issues],
        }


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path, help="Workbook to check")
    parser.add_argument("--filing", required=True, type=Path, help="Full SEC filing in HTML")
    parser.add_argument("--output", required=True, type=Path, help="JSON report path")
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    checker = LeanChecker(args.workbook, args.filing)
    report = checker.run()
    try:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(report, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    except Exception as exc:
        print(f"Could not write JSON report {args.output}: {type(exc).__name__}: {exc}", file=sys.stderr)
        return 2
    print(
        f"{report['status']}: {report['summary']['errors']} error(s), "
        f"{report['summary']['warnings']} warning(s)"
        + f"; report={args.output}"
    )
    if report["status"] == "error":
        return 2
    return 1 if report["summary"]["errors"] else 0


if __name__ == "__main__":
    raise SystemExit(main())
