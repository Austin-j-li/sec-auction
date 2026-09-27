"""Synthetic regression checks for workbook comparisons and revision findings."""

import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest

import datetime as dt

from openpyxl import Workbook, load_workbook

import check_lean
import diff_workbooks
import findings_text

PROJECT = Path(__file__).resolve().parents[2]
# The v1.13.2 Mac-Gray workbook: in extraction/ until the release, then relocated (D25).
MAC_GRAY_V1132 = [PROJECT / "_dev" / "reviews" / "2026-09-22-opus55-reextraction" / "workbooks" / "mac-gray.xlsx",
                  PROJECT / "extraction" / "mac-gray.xlsx"]
STOCK_FOR_ALL_CASH = {"Yes": 0, "No": "Part stock", "Not stated": "Not stated"}


def render_v114(source, dest):
    """The same ledger in v1.14 columns: All cash mapped to Stock % by the crosswalk, new columns blank."""
    book = load_workbook(source)
    ledger = book["Deal ledger"]
    rows = list(ledger.iter_rows(values_only=True))
    ledger.delete_rows(1, ledger.max_row)
    ledger.append(check_lean.LEDGER_COLUMNS_V114)
    for values in rows[1:]:
        row = dict(zip(rows[0], values))
        row["Stock %"] = STOCK_FOR_ALL_CASH.get(row.pop("All cash"))
        ledger.append([row.get(column) for column in check_lean.LEDGER_COLUMNS_V114])
    book.save(dest)


def diff_output(before, after, **options):
    output = io.StringIO()
    with contextlib.redirect_stdout(output):
        diff_workbooks.main(before, after, **options)
    return output.getvalue()


def ledger_book(header, rows, rounds=()):
    book = Workbook()
    book.active.title = "Deal ledger"
    book.active.append(header)
    for row in rows:
        book.active.append([row.get(column) for column in header])
    sheet = book.create_sheet("Rounds")
    sheet.append(check_lean.ROUND_COLUMNS)
    for line in rounds:
        sheet.append([line.get(column) for column in check_lean.ROUND_COLUMNS])
    return book


class ReviewHelperTests(unittest.TestCase):
    def test_workbook_diff_reports_number_to_text_change(self):
        with tempfile.TemporaryDirectory() as tmp:
            before, after = Path(tmp) / "before.xlsx", Path(tmp) / "after.xlsx"
            wb = Workbook()
            wb.active.append(["Round"])
            wb.active.append([1])
            wb.save(before)
            wb.active["A2"] = "1"
            wb.save(after)
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                diff_workbooks.main(before, after)
            self.assertIn("1 [number]", output.getvalue())
            self.assertIn("1 [text]", output.getvalue())
            self.assertIn("Total: 1 change(s)", output.getvalue())

    def test_workbook_diff_reports_an_added_trailing_column(self):
        with tempfile.TemporaryDirectory() as tmp:
            before, after = Path(tmp) / "before.xlsx", Path(tmp) / "after.xlsx"
            wb = Workbook()
            wb.active.append(["Round"])
            wb.active.append([1])
            wb.save(before)
            wb.active["B1"] = "New field"
            wb.active["B2"] = "New value"
            wb.save(after)
            output = io.StringIO()
            with contextlib.redirect_stdout(output):
                diff_workbooks.main(before, after)
            # Columns are aligned by heading: an added column is listed once, not as a change in every row.
            self.assertIn("column only in after: New field", output.getvalue())
            self.assertNotIn("New value", output.getvalue())
            self.assertIn("Total: 0 change(s) in rows; 1 column(s) in one workbook only", output.getvalue())

    def test_crosswalk_rules(self):
        crosswalk = diff_workbooks.crosswalk
        bid = {"Event": "Bid"}
        for old, new in (("Yes", 0), ("No", 40), ("No", "50-75"), ("No", "Part stock"), ("Not stated", "Not stated"), (None, None)):
            with self.subTest(old=old, new=new):
                self.assertEqual(crosswalk("Stock %", old, new, bid), ("same", None))
        for old, new in (("Yes", 40), ("No", 0), ("Yes", "Not stated"), ("Yes", None)):
            with self.subTest(old=old, new=new):
                self.assertEqual(crosswalk("Stock %", old, new, bid)[0], "shown")
        self.assertEqual(crosswalk("Stock %", "No", "Varies", bid), ("shown", "v1.14 Varies has no v1.13.2 equivalent"))
        self.assertEqual(diff_workbooks.all_cash_for_stock("Varies"), None)
        self.assertIsNone(crosswalk("Price low", 10, 11, bid))
        cvr = crosswalk("Price low", 21.5, 19, {"Event": "Bid", "CVR/earnout": "Y", "CVR/earnout value": 2.5})
        self.assertEqual(cvr[0], "shown")
        self.assertIn("value 2.5; not equated", cvr[1])
        self.assertEqual(crosswalk("Price high", 5, None, {"Event": "Other-scope bid"})[0], "suppressed")
        self.assertIsNone(crosswalk("Price high", None, 5, {"Event": "Other-scope bid"}))
        self.assertEqual(crosswalk("Deadline outcome", "Late bids accepted", "Extended (late bid accepted)", {}),
                         ("shown", "D11 (wider): shown, not equated"))
        self.assertEqual(crosswalk("Conditions", "Heavy", "Light", bid), ("shown", "E12 rules changed"))
        self.assertEqual(crosswalk("Conditions", "Heavy", "Heavy", bid), ("same", None))
        self.assertIsNone(crosswalk("Note", "a", "b", bid))

    def test_equal_content_across_schemas_reports_no_row_change(self):
        source = next((path for path in MAC_GRAY_V1132 if check_lean.ledger_schema(path) == check_lean.SCHEMA_V1132), None)
        if source is None:
            self.skipTest("no v1.13.2 mac-gray.xlsx in the relocated workbooks or in extraction/")
        with tempfile.TemporaryDirectory() as tmp:
            rendered = Path(tmp) / "mac-gray-v114.xlsx"
            render_v114(source, rendered)
            self.assertEqual(check_lean.ledger_schema(rendered), check_lean.SCHEMA_V1141)
            text = diff_output(source, rendered, crosswalk_on=True)
            plain = diff_output(source, rendered)
        self.assertIn("Total: 0 change(s) in rows; 9 column(s) in one workbook only", text)
        self.assertIn("## Deal ledger: no change; rows by Event, Who, Sort date: 53 matched, 0 removed, 0 added", text)
        self.assertIn("column only in after: Due diligence [added in v1.14]", text)
        self.assertIn("column only in before: All cash [crosswalk: compared with Stock %]", text)
        self.assertIn("Conditions: compared [crosswalk: E12 rules changed]", text)
        self.assertIn("  Bid: 13 -> 13", text)
        # Without the crosswalk the rows still match by key; only the columns differ.
        self.assertIn("Total: 0 change(s) in rows; 9 column(s) in one workbook only", plain)

    def test_keyed_rows_and_labelled_schema_differences(self):
        header_old = check_lean.LEDGER_COLUMNS
        header_new = check_lean.LEDGER_COLUMNS_V114
        day = dt.datetime(2020, 1, 2)
        old_rows = [
            {"#": 1, "Event": "NDA signed", "Who": "A", "Sort date": day},
            {"#": 2, "Event": "Bid", "Who": "A", "Sort date": day, "Price low": 10, "Price high": 12, "All cash": "Yes",
             "Conditions": "Heavy"},
            {"#": 3, "Event": "Other-scope bid", "Who": "B", "Sort date": day, "Price low": 2, "Price high": 2, "All cash": "Yes"},
        ]
        new_rows = [
            {"#": 1, "Event": "NDA signed", "Who": "A", "Sort date": day},
            {"#": 2, "Event": "Exclusivity changed", "Who": "A", "Sort date": day},
            {"#": 3, "Event": "Bid", "Who": "A", "Sort date": day, "Price low": 10, "Price high": 12, "Stock %": 0,
             "CVR/earnout": "Y", "CVR/earnout value": 1, "Conditions": "Light"},
            {"#": 4, "Event": "Other-scope bid", "Who": "B", "Sort date": day, "Stock %": 0},
        ]
        old_round = [{"Process": 1, "Round": 1, "Deadline outcome": "Late bids accepted"}]
        new_round = [{"Process": 1, "Round": 1, "Deadline outcome": "Extended (late bid accepted)"}]
        with tempfile.TemporaryDirectory() as tmp:
            old, new = Path(tmp) / "old.xlsx", Path(tmp) / "new.xlsx"
            ledger_book(header_old, old_rows, old_round).save(old)
            book = ledger_book(header_new, new_rows, new_round)
            book.create_sheet("Source").append(["EDGAR link", "https://www.sec.gov/"])
            book.save(new)
            text = diff_output(old, new, crosswalk_on=True)
            reverse = diff_output(new, old, crosswalk_on=True)
            with_source = diff_output(old, new, include_source=True)
        # The inserted row is one added row; the renumbered "#" is not a change.
        self.assertIn("rows by Event, Who, Sort date: 3 matched, 0 removed, 1 added", text)
        self.assertIn("added row 3 [Exclusivity changed | A | 2020-01-02]", text)
        self.assertNotIn(", #:", text)
        self.assertIn("  Exclusivity changed: 0 -> 1", text)
        self.assertIn("E13 basis: v1.14 price with CVR/earnout Y, value 1; not equated", text)
        self.assertIn("crosswalk: E12 rules changed", text)
        self.assertIn("crosswalk: D11 (wider): shown, not equated", text)
        self.assertIn("2 Other-scope price cell(s) blank in v1.14 not shown (D18)", text)
        self.assertNotIn("All cash -> Stock %", text)  # Yes and 0 are equivalent
        self.assertIn("## Source: skipped (--include-source compares it)", text)
        self.assertIn("## Source: sheet added", with_source)
        # Either order works: the v1.13.2 side is the old one.
        self.assertIn("E13 basis", reverse)
        self.assertIn("2 Other-scope price cell(s) blank in v1.14 not shown (D18)", reverse)

    def test_by_quote_pairs_rows_whose_key_changed(self):
        header = check_lean.LEDGER_COLUMNS
        quote = "the Board authorized management to explore a possible sale of the company (p. 5)"
        before = [{"#": 1, "Event": "Other material event", "Who": "Moab", "Sort date": dt.datetime(2020, 1, 2), "Quote and page": quote}]
        after = [{"#": 1, "Event": "Other material event", "Who": "Moab Partners", "Sort date": dt.datetime(2020, 1, 2),
                  "Quote and page": '"' + quote.replace(" (p. 5)", '" (p. 5)')}]
        with tempfile.TemporaryDirectory() as tmp:
            old, new = Path(tmp) / "old.xlsx", Path(tmp) / "new.xlsx"
            ledger_book(header, before).save(old)
            ledger_book(header, after).save(new)
            keyed = diff_output(old, new)
            quoted = diff_output(old, new, by_quote=True)
        self.assertIn("0 matched, 1 removed, 1 added", keyed)
        self.assertIn("0 matched (+1 by quote), 0 removed, 0 added", quoted)
        self.assertIn("matched by quote; key now Other material event | Moab Partners | 2020-01-02), Who:", quoted)

    def test_findings_keep_warnings_as_review_leads(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, output = Path(tmp) / "report.json", Path(tmp) / "findings.md"
            source.write_text(json.dumps({"checker_version": "1.7", "ledger_schema": "v1.14", "issues": [
                {"severity": "warning", "basis": "mechanical", "message": "Required note may exceed target"},
                {"severity": "error", "basis": "mechanical", "message": "Invalid event label"},
            ]}))
            with contextlib.redirect_stdout(io.StringIO()):
                findings_text.main(source, output)
            text = output.read_text()
            self.assertEqual(text.splitlines()[0], "# Checker findings (checker 1.7; v1.14 ledger rules)")
            self.assertLess(text.index("Invalid event label"), text.index("Required note"))
            self.assertIn("warning; mechanical review lead", text)
            self.assertNotIn("certain", text)

    def test_findings_from_an_older_report_say_the_version_is_not_recorded(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, output = Path(tmp) / "report.json", Path(tmp) / "findings.md"
            source.write_text(json.dumps({"issues": []}))
            with contextlib.redirect_stdout(io.StringIO()):
                findings_text.main(source, output)
            first = output.read_text().splitlines()[0]
        self.assertEqual(first, "# Checker findings (checker version not recorded; ledger rules not recorded)")


if __name__ == "__main__":
    unittest.main()
