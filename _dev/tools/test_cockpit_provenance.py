"""The Source sheet of the cockpit's Excel download, on a fixture repository with a real workspace."""
from __future__ import annotations

import hashlib
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

import fetch_filing
from cockpit import provenance
from cockpit.workspace import Conflict
from test_cockpit_workspace import fixture

SUBMISSION = "https://www.sec.gov/Archives/edgar/data/42/0000000042-21-000007.txt"
INDEX = "https://www.sec.gov/Archives/edgar/data/42/0000000042-21-000007-index.htm"
FOUR = ["Deal ledger", "Rounds", "Questions", "Deal facts"]


def sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def source(content: bytes) -> dict[str, str]:
    wb = load_workbook(io.BytesIO(content))
    try:
        return {field: value for field, value in wb["Source"].iter_rows(min_row=2, values_only=True)}
    finally:
        wb.close()


class SourceSheetTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit, self.original = fixture(self.root)
        self.ws = self.cockpit.workspace

    def tearDown(self): self.temp.cleanup()

    def seed(self, index_url: str) -> None:
        (self.root / "ref").mkdir(exist_ok=True)
        (self.root / "ref/seed.csv").write_text(f"deal,target_name,deal_number,form_type,date_filed,index_url,status\nalpha-deal,ALPHA DEAL INC,1,DEFM14A,2020-04-01,{index_url},ok\n", encoding="utf-8")

    def manifest(self, source_url: str, digest: str) -> None:
        path = self.root / "raw_filing/MANIFEST.csv"
        path.write_text(path.read_text(encoding="utf-8").replace("https://example.invalid/x.txt,x.htm,,,", f"{source_url},x.htm,,,{digest}"), encoding="utf-8")

    def test_catalog_deal_takes_the_seed_link_and_the_published_instruction_hash(self):
        catalog = self.root / "_dev/cockpit/catalog.json"
        doc = json.loads(catalog.read_text(encoding="utf-8"))
        doc["deals"]["alpha-deal"]["versions"][0]["instruction_version"] = "v1.13.2"
        catalog.write_text(json.dumps(doc), encoding="utf-8")
        instruction = self.cockpit.instructions.list()["items"][0]  # imports the repository instruction as published v1.13.2
        self.manifest(SUBMISSION, "ab" * 32)
        self.seed(INDEX)
        content, name = provenance.download(self.cockpit, "alpha-deal")
        self.assertEqual(name, "alpha-deal-working-r0.xlsx")
        self.assertEqual(load_workbook(io.BytesIO(content)).sheetnames, [*FOUR, "Source"])
        values = source(content)
        self.assertEqual(list(values), ["EDGAR filing index", "Complete submission (.txt)", "Background pages", "Filing SHA-256", "Instruction",
                                        "Instruction SHA-256", "Version ID (working-copy base)", "Raw workbook SHA-256", "Working revision", "Review status", "Exported at"])
        self.assertEqual((values["EDGAR filing index"], values["Complete submission (.txt)"]), (INDEX, SUBMISSION))
        self.assertEqual(values["Filing SHA-256"], "ab" * 32)
        self.assertEqual((values["Instruction"], values["Instruction SHA-256"]), ("v1.13.2", instruction["sha256"]))
        self.assertEqual((values["Version ID (working-copy base)"], values["Raw workbook SHA-256"]), ("v1132-raw", sha(self.original)))
        self.assertEqual((values["Background pages"], values["Working revision"], values["Review status"]), ("not recorded", "0", "not set"))

        self.seed(INDEX.replace("000007", "000008"))  # a seed that disagrees with the filing's own link is not shown
        self.assertEqual(source(provenance.download(self.cockpit, "alpha-deal")[0])["EDGAR filing index"], "not recorded")
        four, name = provenance.download(self.cockpit, "alpha-deal", source=False)
        self.assertEqual((four, name), (self.original, "alpha-deal-working-r0.xlsx"))
        self.assertEqual(provenance.download(self.cockpit, "alpha-deal", "v1132-raw"), (self.original, "alpha-deal-v1132-raw.xlsx"))

    def test_review_status_background_pages_and_formula_like_text(self):
        facts = self.cockpit.deal("alpha-deal")["facts"]
        base = sha(self.original)
        self.ws.edit("alpha-deal", {"revision": 0, "base_sha256": base, "reason": "Pages", "operations": [
            {"type": "insert", "sheet": "Deal facts", "after_uid": facts[-1]["uid"], "values": {"Field": "Background pages", "Value": "=pp. 20-31"}}]}, "austin")
        values = source(provenance.download(self.cockpit, "alpha-deal")[0])
        self.assertEqual((values["Background pages"], values["Working revision"], values["Review status"]), ("=pp. 20-31", "1", "not set"))
        wb = load_workbook(io.BytesIO(provenance.download(self.cockpit, "alpha-deal")[0]))
        self.assertEqual(wb["Source"]["B4"].data_type, "s")  # text from the workbook, never a formula
        self.cockpit.trace.set_deal_review("alpha-deal", {"status": "in_review", "revision": 1}, "austin")
        self.assertEqual(source(provenance.download(self.cockpit, "alpha-deal")[0])["Review status"], "In review at revision 1")
        self.ws.edit("alpha-deal", {"revision": 1, "base_sha256": base, "reason": "Later", "operations": [
            {"type": "update", "sheet": "Deal facts", "uid": facts[0]["uid"], "values": {"Value": "Target Later"}}]}, "alex")
        content, name = provenance.download(self.cockpit, "alpha-deal")
        self.assertEqual(name, "alpha-deal-working-r2.xlsx")
        self.assertEqual(source(content)["Review status"], "In review at revision 1; edited since (working revision 2)")
        self.assertEqual(provenance.download(self.cockpit, "alpha-deal", source=False)[0], self.ws.export("alpha-deal"))

    def test_raw_version_with_source_keeps_its_sheets_and_names_no_working_state(self):
        content, name = provenance.download(self.cockpit, "alpha-deal", "v1132-raw", source=True)
        self.assertEqual(name, "alpha-deal-v1132-raw-with-source.xlsx")
        wb, raw = load_workbook(io.BytesIO(content)), load_workbook(io.BytesIO(self.original))
        self.assertEqual(wb.sheetnames, [*FOUR, "Source"])
        for sheet in FOUR:
            self.assertEqual(list(wb[sheet].values), list(raw[sheet].values))
        values = source(content)
        self.assertEqual((values["Version ID"], values["Raw workbook SHA-256"]), ("v1132-raw", sha(self.original)))
        self.assertEqual(values["Working revision"], provenance.RAW_ONLY)
        self.assertEqual(values["Review status"], provenance.RAW_ONLY)
        self.assertEqual((values["EDGAR filing index"], values["Complete submission (.txt)"], values["Filing SHA-256"]), ("not recorded",) * 3)

    def test_added_deal_takes_its_recorded_link_and_the_run_receipts(self):
        slug, file, filing = "beta-holdings", "beta-holdings_2021-03-04_DEFM14A.htm", b"<h2>Background of the Merger</h2><p>Beta met Party A.</p>"
        folder = self.root / "_dev/cockpit/state/filings" / slug
        folder.mkdir(parents=True)
        (folder / file).write_bytes(filing)
        version = self.root / "_dev/cockpit/state/versions" / slug / "run1"
        version.mkdir(parents=True)
        (version / f"{slug}.xlsx").write_bytes(self.original)
        conn = self.ws._connect(write=True)
        conn.execute("INSERT INTO added_deals VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                     (slug, "Beta Holdings", "DEFM14A", "2021-03-04", file, "link", None, fetch_filing.index_link(SUBMISSION), SUBMISSION,
                      "beta-proxy.htm", "2026-09-23T10:00:00+00:00", len(filing), sha(filing), "alex", "2026-09-23T10:00:05+00:00"))
        conn.execute("INSERT INTO versions (slug, id, label, path, sha256, kind, engine, model, effort, instruction_version, instruction_sha256, filing_sha256, started_by, started_at, finished_at, receipts, checker, instruction_id) VALUES (?,?,?,?,?,'raw','Opus 5.5','claude-opus-5-5','medium','v1.14',?,?,'alex','2026-09-23T10:00:00+00:00','2026-09-23T10:10:00+00:00',?,'{}','abcdef012345')",
                     (slug, "run1", "Opus 5.5 · medium", str((version / f"{slug}.xlsx").relative_to(self.root)), sha(self.original), "c" * 64, sha(filing), str(version.relative_to(self.root))))
        conn.commit(); conn.close()
        content, name = provenance.download(self.cockpit, slug)
        self.assertEqual(name, f"{slug}-working-r0.xlsx")
        values = source(content)
        self.assertEqual((values["EDGAR filing index"], values["Complete submission (.txt)"]), (INDEX, SUBMISSION))
        self.assertEqual(values["Filing SHA-256"], sha(filing))
        self.assertEqual((values["Instruction"], values["Instruction SHA-256"]), ("v1.14 (cockpit instruction abcdef012345)", "c" * 64))
        self.assertEqual((values["Version ID (working-copy base)"], values["Raw workbook SHA-256"]), ("run1", sha(self.original)))

    def test_a_workbook_that_already_has_a_source_sheet_is_refused(self):
        wb = load_workbook(io.BytesIO(self.original))
        wb.create_sheet("Source")
        out = io.BytesIO()
        wb.save(out)
        with self.assertRaises(Conflict):
            provenance.add_sheet(out.getvalue(), [("Field", "value", False)])


if __name__ == "__main__":
    unittest.main()
