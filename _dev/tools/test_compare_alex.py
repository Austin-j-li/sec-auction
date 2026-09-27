"""Observed compare_alex alignment, coding, provenance, and CLI behavior on a synthetic deal."""
from __future__ import annotations

import contextlib
import datetime as dt
import io
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

import compare_alex
import derive_analysis as derive
from test_derive_analysis import D, FACTS, bid, row, rounds_line, write

HEADERS = ["DealNumber", "BidderID", "BidderName", "bid_date_precise", "bid_date_rough", "bid_type",
           "bid_value_lower", "bid_value_upper", "bid_value_pershare", "all_cash", "bid_note",
           "comments_1", "comments_2", "comments_3", "additional_note", *(f"extra_{n}" for n in range(20))]


def coded(number, name, when, kind=None, low=None, high=None, cash=None, note="NA", red=()):
    return {"DealNumber": number, "BidderID": 1, "BidderName": name, "bid_date_precise": when,
            "bid_type": kind, "bid_value_lower": low, "bid_value_upper": high,
            "bid_value_pershare": low if low == high else None, "all_cash": cash, "bid_note": note,
            "_red": red}


def alex_book(path: Path, records: list[dict]) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "deal_details"
    ws.append(HEADERS)
    for values in records:
        red = values.get("_red", ())
        for col, key in enumerate(HEADERS, 1):
            cell = ws.cell(ws.max_row + 1 if col == 1 else ws.max_row, col)
            cell.value = values.get(key)
            if red == "all" or key in red:
                cell.font = Font(color="FFFF0000")
    wb.save(path)
    return path


class CompareAlexTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        d = lambda n: D + dt.timedelta(days=n)
        ledger = [
            row(1, "Bank", "Adviser", day=0),
            bid(2, "Party A", 10, 12, day=1),
            bid(3, "5 other parties", 8, 9, day=1, Count=5),
            row(4, "Party C", "Did not submit", day=2, Type="Financial", Count=1, **{"Exit reason": "Not stated"}),
            row(5, "Target", "Deadline", rnd=2, day=10),
            bid(6, "Party B", 19, rnd=2, day=10, formality="Formal", conditions="Heavy",
                **{"CVR/earnout": "Y", "CVR/earnout value": 2.5}),
            bid(7, "Party E", 23.81, rnd=2, day=11, conditions="Heavy"),
            row(8, "Party D", "Withdrew", rnd=2, day=12, Type="Financial", Count=1,
                **{"Exit reason": "Value below market price"}),
            row(9, "Party A", "Merger agreement signed", rnd=2, day=20),
        ]
        self.workbook = write(self.root / "alpha.xlsx", ledger,
                              [rounds_line(1, "Enforced"), rounds_line(2, "Enforced", finality="Announced as final")],
                              facts=FACTS)
        self.alex = alex_book(self.root / "alex.xlsx", [
            coded(111, "Party A", d(1), "Informal", 10, 12, 1),
            coded(111, "Party B", d(10), "Formal", 21.5, 21.5, 0, red=("bid_type",)),
            coded(111, "Party E/F", d(11), "Informsl", 23.81, 23.81),
            coded(111, "5 parties", d(1), "Informal", 8, 9),
            coded(111, "Party Z", d(5), "Formal", 99, 99, red="all"),
            coded(111, "Bank", d(0), note="IB"),
            coded(111, "Party C", d(2), note="Drop", red=("comments_1",)),
            coded(111, "Party D", d(12), note="DropM"),
            coded(111, "NA", d(10), note="Final Round"),
            coded(111, "NA", d(10), note="Final Round Inf"),
            coded(111, "Party A", d(20), note="Executed"),
            coded(111, "Party Q", d(3), note="NDA"),
            coded(222, "Party A", d(1), "Formal", 10, 12),
        ])
        self.seed = self.root / "seed.csv"
        self.seed.write_text("deal,target_name,deal_number,form_type,date_filed,index_url,status\n"
                             "alpha,ALPHA,111,DEFM14A,2020-02-01,https://example.invalid,ok\n")

    def tearDown(self):
        self.tmp.cleanup()

    def test_alignment_readings_and_red_provenance(self):
        result = compare_alex.compare(self.workbook, "alpha", self.alex, self.seed)
        rows = {r["bidder_name"]: r for r in result["bids"]}
        self.assertEqual({name: (r.get("ledger_row"), r["alignment"], r.get("price_basis")) for name, r in rows.items()}, {
            "Party A": (2, "bidder, price, date", "upfront"),
            "Party B": (6, "bidder, price, date", "package"),
            "Party E/F": (7, "bidder (part of name), price, date", "upfront"),
            "5 parties": (3, "cohort size, price, date", "upfront"),
            "Party Z": (None, "unaligned", None),
        })
        self.assertEqual(result["summary"]["readings"]["T0"], {"agree": 4, "compared": 4, "missing": 0})
        self.assertEqual(result["summary"]["readings"]["T1"], {"agree": 3, "compared": 4, "missing": 0})
        self.assertEqual((rows["Party B"]["agree_all_cash"], rows["Party B"]["per_share_match"]), ("N", "package"))
        self.assertEqual(rows["Party B"]["provenance"], "Alex: red-font correction")
        self.assertEqual(rows["Party Z"]["provenance"], "Alex: whole row in red (added or rewritten)")
        self.assertEqual(rows["Party A"]["provenance"], "Chicago coding (deal has no corrections)")
        self.assertEqual(result["summary"]["labelled_bids_with_red_bid_type"], 2)
        self.assertEqual(result["summary"]["ledger_schema"], "v0")

    def test_note_codes_and_seed_join(self):
        events = {(r["bidder_name"], r["bid_note"]): (r["status"], r.get("ledger_row"))
                  for r in compare_alex.compare(self.workbook, "alpha", self.alex, self.seed)["events"]}
        self.assertEqual({k: v for k, v in events.items() if k[1] != "NA"}, {
            ("Bank", "IB"): ("agree", 1), ("Party C", "Drop"): ("exit, other label or reason", 4),
            ("Party D", "DropM"): ("agree", 8), ("", "Final Round"): ("agree", 5),
            ("", "Final Round Inf"): ("event agrees, round finality differs", 5),
            ("Party A", "Executed"): ("agree", 9), ("Party Q", "NDA"): ("unaligned", None),
        })
        with self.assertRaisesRegex(derive.DeriveError, "not in"):
            compare_alex.compare(self.workbook, "beta", self.alex, self.seed)

    def test_final_round_announcement_accepts_inferred_finality(self):
        day = D + dt.timedelta(days=10)
        coded_row = {"BidderName": None, "bid_date_precise": day}
        ledger_row = {"row": 5, "who": "Target", "event": "Round opened", "process": 1, "round": 2,
                      "sort_date": day, "date_from": day, "date_to": day, "exit_reason": "", "count_point": None}
        finality = {(1, 2): "Inferred final"}
        self.assertEqual(compare_alex.match_code(coded_row, "Final Round Ann", [ledger_row], finality)["status"],
                         "agree")
        self.assertEqual(compare_alex.match_code(coded_row, "Final Round Inf Ann", [ledger_row], finality)["status"],
                         "event agrees, round finality differs")
        self.assertEqual(compare_alex.match_code(coded_row, "Final Round", [ledger_row], finality)["status"], "unaligned")
        self.assertEqual(compare_alex.match_code(coded_row, "Final Round", [{**ledger_row, "event": "Deadline"}], finality)["status"], "agree")

    def test_providence_worcester_pilot_row_6045(self):
        day = dt.date(2016, 7, 27)
        coded_row = {"BidderName": None, "bid_date_precise": day}
        ledger_row = {"row": 35, "who": "Providence and Worcester (target)", "event": "Round opened",
                      "process": 1, "round": 3, "sort_date": day, "date_from": day, "date_to": day,
                      "exit_reason": "", "count_point": None}
        matched = compare_alex.match_code(coded_row, "Final Round Ann", [ledger_row],
                                          {(1, 3): "Inferred final"})
        self.assertEqual((matched["status"], matched["ledger_row"]), ("agree", 35))

    def test_event_code_map_and_final_round_variants(self):
        day = D + dt.timedelta(days=10)
        coded = {"BidderName": None, "bid_date_precise": day}
        base = {"row": 5, "who": "Target", "process": 1, "round": 2, "sort_date": day,
                "date_from": day, "date_to": day, "exit_reason": "", "count_point": None}
        cases = {
            "Bidder Sale": "Bid", "Sale Press Release": "Sale process announced",
            "Target Sale Public": "Sale process announced", "Bid Press Release": "Bid announced",
            "Terminated": "Process terminated", "Restarted": "Process restarted",
            "Exclusivity 30 days": "Exclusivity changed", "Final Round Ann": "Deadline set",
            "Final Round Ext Ann": "Deadline revised", "Final Round Ext": "Deadline revised",
            "Final Round": "Deadline", "Final Round Inf Ann": "Round opened",
            "Final Round Inf Ext Ann": "Deadline revised", "Final Round Inf Ext": "Deadline revised",
            "Final Round Inf": "Deadline",
        }
        for code, event_name in cases.items():
            with self.subTest(code=code):
                finality = "Not final" if " Inf" in code else "Announced as final"
                row = {**base, "event": event_name}
                self.assertEqual(compare_alex.match_code(coded, code, [row], {(1, 2): finality})["status"], "agree")
        self.assertEqual(compare_alex.match_code(coded, "Final Round Ann", [{**base, "event": "Deadline"}],
                                                 {(1, 2): "Announced as final"})["status"], "unaligned")

    def test_event_match_uses_both_dates_with_31_day_limit(self):
        day = D + dt.timedelta(days=10)
        base = {"row": 5, "who": "Party A", "event": "NDA signed", "process": 1, "round": 1,
                "sort_date": day, "date_from": day, "date_to": day, "exit_reason": "", "count_point": None}
        def result(code, event_name, offset, precise=None, rough=None):
            row = {**base, "event": event_name, "sort_date": day + dt.timedelta(days=offset),
                   "date_from": day + dt.timedelta(days=offset), "date_to": day + dt.timedelta(days=offset)}
            coded = {"BidderName": "Party A", "bid_date_precise": precise, "bid_date_rough": rough}
            return compare_alex.match_code(coded, code, [row], {})
        self.assertEqual(result("NDA", "NDA signed", 31, day)["date_gap_days"], 31)
        self.assertEqual(result("NDA", "NDA signed", 32, day)["status"], "unaligned")
        self.assertEqual(result("NDA", "NDA signed", 365, day)["status"], "unaligned")
        self.assertEqual(result("Drop", "Withdrew", 20, day)["status"], "agree")
        self.assertEqual(result("Drop", "Withdrew", 32, day)["status"], "unaligned")
        self.assertEqual(result("NDA", "NDA signed", 31, day - dt.timedelta(days=300), day)["status"], "agree")
        self.assertEqual(result("NDA", "NDA signed", 0)["status"], "unaligned")

    def test_half_cent_price_and_correction_provenance(self):
        self.assertTrue(compare_alex.close(10, 10.004))
        self.assertTrue(compare_alex.close(10, 10.005))
        self.assertFalse(compare_alex.close(10, 10.006))
        self.assertFalse(compare_alex.close(10, 10.01))
        self.assertEqual(compare_alex.provenance("mac-gray", ["comments_1"]), "Alex: red comments only (coding as Chicago)")
        self.assertEqual(compare_alex.provenance("mac-gray", ["additional_note"]), "Alex: red-font correction")
        self.assertEqual(compare_alex.provenance("mac-gray", ["comments_1"] * 29), "Alex: red comments only (coding as Chicago)")
        self.assertEqual(compare_alex.provenance("mac-gray", ["comments_1"] * 30), "Alex: whole row in red (added or rewritten)")

    def test_cli_writes_only_three_outputs(self):
        out = self.root / "out"
        args = [str(self.workbook), "--deal", "alpha", "--out", str(out), "--alex", str(self.alex),
                "--seed", str(self.seed)]
        with contextlib.redirect_stdout(io.StringIO()) as printed:
            self.assertEqual(compare_alex.main(args), 0)
        self.assertIn("4 of 5 labelled bids aligned", printed.getvalue())
        self.assertEqual(sorted(p.name for p in out.iterdir()), ["alex_bids.csv", "alex_events.csv", "summary.json"])
        self.assertIn("never a target", json.loads((out / "summary.json").read_text())["note"])
        with contextlib.redirect_stderr(io.StringIO()):
            self.assertEqual(compare_alex.main(args), 2)  # output folder must be new or empty
