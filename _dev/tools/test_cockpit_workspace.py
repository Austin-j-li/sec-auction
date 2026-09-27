"""Disposable fixture tests for editable cockpit storage and HTTP boundaries."""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import json
import os
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

import check_lean
from cockpit import data, server
from cockpit.workspace import Missing, WorkspaceError, Conflict
from test_cockpit import build_repo, build_workbook


def fixture(root: Path) -> tuple[data.Cockpit, bytes]:
    build_repo(root)
    path = root / "extraction/alpha-deal.xlsx"
    original = path.read_bytes()
    catalog = {"schema_version": 1, "deals": {"alpha-deal": {
        "name": "Alpha Deal", "default_base": "base-raw",
        "versions": [{"id": "base-raw", "label": "Raw Version 0", "path": "extraction/alpha-deal.xlsx", "sha256": hashlib.sha256(original).hexdigest(), "instruction_version": "Version 0", "kind": "raw", "review_status": "unreviewed"}],
        "findings": [{"id": "F1", "title": "Finding one", "detail": "Review event", "rule": "E2", "source_version": "base-raw", "source_label": "Raw", "source_rows": [2], "evidence": [], "proposed_change": "Consider correction", "needs_recheck": True, "judgment": "unreviewed", "implementation": "not_applied", "verification": "unchecked", "note": "", "actor": None, "at": None}],
        "documents": [{"id": "report", "label": "Report", "source_version": "base-raw", "kind": "report", "path": "_dev/reviews/report.md"}],
    }}}
    (root / "_dev/cockpit").mkdir(parents=True)
    (root / "_dev/reviews").mkdir(parents=True)
    (root / "_dev/reviews/report.md").write_text("A report", encoding="utf-8")
    (root / "_dev/cockpit/catalog.json").write_text(json.dumps(catalog), encoding="utf-8")
    (root / "SEC_Deal_Ledger_Extraction_Instruction.md").write_text("# Synthetic instruction\n\n**Revision of 1 January 2026, Version 0.**\n\nRead the filing.\n", encoding="utf-8")
    return data.Cockpit(root), original


def build_version1_workbook(path: Path) -> Path:
    """A Version 1 fixture with representative synthetic terms: figures, a range, Varies, blanks,
    dates, flags, # references and a two-part deadline outcome."""
    build_workbook(path)
    wb = load_workbook(path)
    ws = wb["Deal ledger"]
    old = [cell.value for cell in ws[1]]
    records = [dict(zip(old, [cell.value for cell in row])) for row in ws.iter_rows(min_row=2, max_col=len(old))]
    for row in ws.iter_rows():
        for cell in row: cell.value = None
    terms = {1: {"Price low": 21.25, "Price high": 21.25, "Stock %": "Part stock", "CVR/earnout": "Y", "CVR/earnout value": 1.25, "Formality": "Informal",
                 "Conditions": "Heavy", "Due diligence": "Incomplete", "Financing": "Contingent", "Regulatory": "Not stated", "Note": "Revised in #2; stated stock range 50\u201375%"},
             2: {"Stock %": 0, "Count": 2, "Formality": "Formal", "Conditions": "Unclear", "Due diligence": "Varies", "Financing": "Varies",
                 "Regulatory": "Concern", "Antitrust": "Y", "Exclusivity": "Requested", "Note": "Cohort; see #1-#3"},
             3: {"Price low": None, "Price high": None, "Stock %": "Not stated"}}
    for column, name in enumerate(check_lean.LEDGER_COLUMNS, 1):
        ws.cell(1, column, name)
    for excel_row, record in enumerate(records, 2):
        if all(value is None for value in record.values()): continue
        record.update(terms.get(record.get("#"), {}))
        for column, name in enumerate(check_lean.LEDGER_COLUMNS, 1):
            ws.cell(excel_row, column, record.get(name))
    wb["Rounds"].cell(2, check_lean.ROUND_COLUMNS.index("Deadline outcome") + 1, "Extended (late bid accepted); Enforced")
    wb.save(path)
    return path


def add_run_version(ws, slug: str, ident: str, workbook: Path, checker: dict | None = None, receipt: dict | None = None, instruction_sha256: str = "d" * 64) -> dict:
    """Register a workbook as an imported run version, as the worker does, optionally with a receipt check.json."""
    folder = ws.root / "_dev/cockpit/state/versions" / slug / ident
    folder.mkdir(parents=True)
    path = folder / f"{slug}.xlsx"
    path.write_bytes(workbook.read_bytes())
    if receipt is not None:
        (folder / "check.json").write_text(json.dumps(receipt), encoding="utf-8")
    conn = ws._connect(write=True)
    conn.execute("INSERT INTO versions (slug, id, label, path, sha256, kind, engine, model, effort, instruction_version, instruction_sha256, filing_sha256, started_by, started_at, finished_at, receipts, checker) VALUES (?,?,?,?,?,'raw','Opus 5.5','claude-opus-5-5','medium',NULL,?,'f','alex','2026-09-24T22:41:00+00:00','2026-09-24T22:52:00+00:00',?,?)",
                 (slug, ident, f"Opus 5.5 · medium · {ident} — Alex", str(path.relative_to(ws.root)), hashlib.sha256(path.read_bytes()).hexdigest(), instruction_sha256, str(folder.relative_to(ws.root)), json.dumps(checker or {})))
    conn.commit(); conn.close()
    return {"id": ident, "path": path, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


class WorkspaceTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit, self.original = fixture(self.root)
        self.ws = self.cockpit.workspace
        self.base = hashlib.sha256(self.original).hexdigest()

    def tearDown(self): self.temp.cleanup()

    def save(self, ops, revision=0, reason="Review edit", actor="austin"):
        return self.ws.edit("alpha-deal", {"revision": revision, "base_sha256": self.base, "reason": reason, "operations": ops}, actor)

    def test_read_only_get_and_immutable_export(self):
        payload = self.cockpit.deal("alpha-deal")
        self.assertEqual(payload["workspace"]["revision"], 0)
        self.assertEqual(payload["ledger"]["rows"][-1]["excel_row"], 8)
        self.assertFalse(self.ws.db_path.exists())
        self.assertEqual(self.ws.export("alpha-deal", "base-raw"), self.original)
        self.assertEqual(self.ws.document("alpha-deal", "report")["text"], "A report")

    def test_working_copy_check_is_kept_per_revision_and_checker_version(self):
        # The deal list rechecks every edited deal on each request; a saved revision is checked once.
        row = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": row["uid"], "values": {"Price low": "12.5"}}])
        self.ws._checks.clear()
        with patch.object(check_lean, "LeanChecker", side_effect=check_lean.LeanChecker) as checker:
            first = self.ws.deal("alpha-deal")
            self.assertEqual(self.ws.deal("alpha-deal")["check"], first["check"])
            self.assertEqual(len(self.cockpit.list_deals()), 1)
            self.assertEqual(checker.call_count, 1)
            self.save([{"type": "update", "sheet": "Deal ledger", "uid": row["uid"], "values": {"Price low": "13"}}], revision=1)
            self.assertEqual(self.ws.deal("alpha-deal")["workspace"]["revision"], 2)
            self.assertEqual(checker.call_count, 2)
            with patch.object(check_lean, "CHECKER_VERSION", "9.9"):
                self.ws.deal("alpha-deal")
            self.assertEqual(checker.call_count, 3)

    def test_deal_without_default_base_uses_its_first_version(self):
        catalog_path = self.root / "_dev/cockpit/catalog.json"
        catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
        del catalog["deals"]["alpha-deal"]["default_base"]
        catalog_path.write_text(json.dumps(catalog), encoding="utf-8")
        self.assertEqual(data.Cockpit(self.root).deal("alpha-deal")["workspace"]["base_sha256"], self.base)

    def test_typed_edit_history_restart_and_restore(self):
        row = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]
        result = self.save([{"type": "update", "sheet": "Deal ledger", "uid": row["uid"], "values": {"Price low": "12.5", "Sort date": "2020-03-04", "Reviewer note": "=source wording"}}])
        self.assertEqual(result["workspace"]["revision"], 1)
        self.assertEqual(result["ledger"]["rows"][0]["cells"]["Price low"], "12.5")
        wb = load_workbook(io.BytesIO(self.ws.export("alpha-deal")))
        self.assertEqual(wb["Deal ledger"]["H2"].value, 12.5)
        columns = [cell.value for cell in wb["Deal ledger"][1]]
        self.assertIsInstance(wb["Deal ledger"].cell(2, columns.index("Sort date") + 1).value, dt.datetime)
        reviewer = wb["Deal ledger"].cell(2, columns.index("Reviewer note") + 1)
        self.assertEqual((reviewer.data_type, reviewer.value), ("s", "=source wording"))
        self.assertEqual(self.root.joinpath("extraction/alpha-deal.xlsx").read_bytes(), self.original)
        reopened = data.Cockpit(self.root)
        self.assertEqual(reopened.deal("alpha-deal")["workspace"]["revision"], 1)
        self.assertIn("12.5", json.dumps(reopened.workspace.history("alpha-deal")))
        self.assertIn("12.5", json.dumps(reopened.workspace.changes("alpha-deal")))
        self.assertTrue(all(change["record_label"] for change in reopened.workspace.history("alpha-deal")["history"][0]["changes"]))
        self.assertTrue(any("Event #1" in change["record_label"] for change in reopened.workspace.history("alpha-deal")["history"][0]["changes"]))
        restored = reopened.workspace.edit("alpha-deal", {"revision": 1, "base_sha256": self.base, "reason": "Restore raw", "operations": [{"type": "restore", "target_revision": 0}]}, "alex")
        self.assertEqual(restored["workspace"]["revision"], 2)
        self.assertEqual(restored["ledger"]["rows"][0]["cells"]["Price low"], "10")
        self.assertEqual(len(reopened.workspace.history("alpha-deal")["history"]), 3)

    def test_post_round_event_can_be_updated_and_cloned_with_typed_cells(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        original = rows[-1]
        updated = self.save([{"type": "update", "sheet": "Deal ledger", "uid": original["uid"], "values": {"Round": "post", "Event": "Merger announced"}}])
        post = updated["ledger"]["rows"][-1]
        self.assertEqual(post["cells"]["Round"], "post")
        clone_values = dict(post["cells"])
        clone_values.pop("#")
        clone_values["Who"] = "Cloned post event"
        clone_values["Count"] = ""
        cloned = self.save([{"type": "insert", "sheet": "Deal ledger", "after_uid": post["uid"], "values": clone_values}], revision=1)
        new_row = cloned["ledger"]["rows"][-1]
        self.assertEqual((new_row["cells"]["Round"], new_row["cells"]["Who"], new_row["cells"]["Count"]), ("post", "Cloned post event", ""))
        wb = load_workbook(io.BytesIO(self.ws.export("alpha-deal")))
        ws = wb["Deal ledger"]
        headers = [cell.value for cell in ws[1]]
        round_cell = ws.cell(ws.max_row, headers.index("Round") + 1)
        price_cell = ws.cell(ws.max_row, headers.index("Price low") + 1)
        date_cell = ws.cell(ws.max_row, headers.index("Sort date") + 1)
        self.assertEqual((round_cell.value, round_cell.data_type), ("post", "s"))
        self.assertEqual((price_cell.value, price_cell.data_type), (10, "n"))
        self.assertIsInstance(date_cell.value, dt.datetime)
        self.assertEqual(wb["Rounds"].cell(2, 2).data_type, "n")
        wb.close()

    def test_version1_bid_term_cells_are_typed_and_offered_as_choices(self):
        coerce = lambda field, value: self.ws._coerce(data.LEDGER_SHEET, field, value)
        self.assertEqual(coerce("Stock %", "40"), 40)
        self.assertEqual(coerce("Stock %", "33.3"), 33.3)
        with self.assertRaisesRegex(WorkspaceError, "Part stock"):
            coerce("Stock %", " 50\u201375 ")
        self.assertEqual(coerce("Stock %", "Part stock"), "Part stock")
        self.assertEqual(coerce("CVR/earnout value", "1.13"), 1.13)
        with self.assertRaises(WorkspaceError): coerce("CVR/earnout value", "up to $2")
        current = self.cockpit.deal("alpha-deal")
        self.assertEqual(current["ledger_schema"], "Version 1")
        self.assertIn("Financing", current["choices"])
        self.assertIn("No deadline stated", current["choices"]["Deadline outcome"])
        version = add_run_version(self.ws, "alpha-deal", "version1-run", build_version1_workbook(self.root / "version1.xlsx"))
        shown = self.cockpit.deal("alpha-deal", version="version1-run")
        self.assertEqual(shown["ledger_schema"], "Version 1")
        self.assertEqual(shown["choices"]["Financing"], ["Committed", "Contingent", "Not needed", "Not stated", "Varies"])
        self.assertEqual(shown["choices"]["Antitrust"], ["Y"])
        self.assertNotIn("All cash", shown["choices"])
        self.assertIn("Extended (late bid accepted)", shown["choices"]["Deadline outcome"])
        self.assertNotIn("Late bids accepted", shown["choices"]["Deadline outcome"])
        self.assertIn("No deadline stated", shown["choices"]["Deadline outcome"])
        self.assertEqual(shown["workspace"]["base_ledger_schema"], "Version 1")
        rebased = self.save([{"type": "rebase", "target_version": version["id"]}])
        self.assertEqual((rebased["ledger_schema"], rebased["workspace"]["base_ledger_schema"]), ("Version 1", "Version 1"))
        self.assertIn("Financing", rebased["choices"])

    def test_atomic_stale_and_deleted_references(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        first, second = rows[0]["uid"], rows[1]["uid"]
        with self.assertRaises(WorkspaceError):
            self.save([{"type": "update", "sheet": "Deal ledger", "uid": first, "values": {"Who": "Changed"}}, {"type": "delete", "sheet": "Deal ledger", "uid": first}])
        self.assertEqual([entry["revision"] for entry in self.ws.history("alpha-deal")["history"]], [0])
        with self.assertRaises(WorkspaceError): self.save([{"type": "delete", "sheet": "Deal ledger", "uid": first, "replacement_uid": None}])
        saved = self.save([{"type": "delete", "sheet": "Deal ledger", "uid": first, "replacement_uid": second}])
        self.assertEqual(saved["ledger"]["rows"][0]["uid"], second)
        self.assertEqual(saved["questions"]["rows"][0]["cells"]["Rows affected"], "#1")
        with self.assertRaises(Conflict): self.save([{"type": "review", "uid": second, "status": "reviewed", "note": ""}])

    def test_insert_move_and_question_flags(self):
        payload = self.cockpit.deal("alpha-deal")
        first = payload["ledger"]["rows"][0]["uid"]
        result = self.save([{"type": "insert", "sheet": "Deal ledger", "after_uid": None, "values": {"Who": "New", "Event": "Contact", "Sort date": "2020-02-29"}}])
        self.assertEqual([row["id"] for row in result["ledger"]["rows"]], list(range(1, 8)))
        self.assertEqual(result["questions"]["rows"][0]["cells"]["Rows affected"], "#2, #3")
        inserted = result["ledger"]["rows"][0]["uid"]
        self.assertEqual(result["ledger"]["rows"][1]["uid"], first)
        moved = self.save([{"type": "move", "sheet": "Deal ledger", "uid": inserted, "after_uid": first}], revision=1)
        self.assertEqual(moved["ledger"]["rows"][0]["uid"], first)
        question = moved["questions"]["rows"][0]["uid"]
        renamed = self.save([{"type": "update", "sheet": "Questions", "uid": question, "values": {"Q": "Q2"}}], revision=2)
        self.assertEqual(renamed["ledger"]["rows"][0]["cells"]["Flag"], "Q2")

    def test_finding_and_review_attribution(self):
        uid = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        result = self.save([{"type": "finding", "id": "F1", "judgment": "supported", "implementation": "not_applied", "verification": "unchecked", "note": "Needs source check"}, {"type": "review", "uid": uid, "status": "needs_decision", "note": "Check date"}], actor="alex")
        self.assertEqual(result["findings"][0]["actor"], "alex")
        self.assertEqual(result["row_review"][uid]["actor"], "alex")
        self.assertEqual(result["findings"][0]["implementation"], "not_applied")

    def test_finding_implementation_can_remain_unassessed(self):
        result = self.save([{"type": "finding", "id": "F1", "judgment": "unreviewed", "implementation": "unassessed", "verification": "unchecked", "note": "Current copy not checked"}], actor="alex")
        finding = result["findings"][0]
        self.assertEqual((finding["judgment"], finding["implementation"], finding["verification"]), ("unreviewed", "unassessed", "unchecked"))
        self.assertEqual(finding["actor"], "alex")
        reopened = data.Cockpit(self.root)
        self.assertEqual(reopened.deal("alpha-deal")["findings"][0]["implementation"], "unassessed")

    def test_huge_range_rejected(self):
        q = self.cockpit.deal("alpha-deal")["questions"]["rows"][0]["uid"]
        with self.assertRaisesRegex(WorkspaceError, "too large"):
            self.save([{"type": "update", "sheet": "Questions", "uid": q, "values": {"Rows affected": "#1-#999999999"}}])

    def test_new_rows_can_anchor_later_batch_operations(self):
        first = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        saved = self.save([
            {"type": "insert", "sheet": "Deal ledger", "client_uid": "new-one", "after_uid": first, "values": {"Who": "First new"}},
            {"type": "insert", "sheet": "Deal ledger", "client_uid": "new-two", "after_uid": "new-one", "values": {"Who": "Second new"}},
            {"type": "move", "sheet": "Deal ledger", "uid": "new-two", "after_uid": first},
            {"type": "review", "uid": "new-one", "status": "reviewed", "note": "Confirmed"},
        ])
        rows = saved["ledger"]["rows"]
        self.assertEqual([row["cells"]["Who"] for row in rows[:3]], ["Alpha", "Second new", "First new"])
        self.assertEqual(saved["row_review"][rows[2]["uid"]]["status"], "reviewed")

    def test_bulk_update_is_recorded_as_the_same_updates_one_row_at_a_time(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        uids = [row["uid"] for row in rows[:3]]
        self.save([{"type": "review", "uid": uids[0], "status": "reviewed", "note": ""}], actor="alex")
        result = self.save([{"type": "bulk_update", "sheet": "Deal ledger", "uids": uids, "values": {"Process": "2", "Round": 3}}], revision=1, reason="Renumber rounds")
        self.assertEqual([(r["cells"]["Process"], r["cells"]["Round"]) for r in result["ledger"]["rows"]], [("2", "3")] * 3 + [("1", "1")] * (len(rows) - 3))
        self.assertEqual(result["row_review"][uids[0]]["actor"], "alex")  # row marks stay, as under an update
        self.assertEqual(result["ledger"]["columns"], self.cockpit.deal("alpha-deal", "base-raw")["ledger"]["columns"])
        bulk = self.ws.history("alpha-deal")["history"][0]
        self.assertEqual((bulk["revision"], bulk["actor"], bulk["reason"], bulk["summary"]), (2, "austin", "Renumber rounds", "6 changes"))
        wb = load_workbook(io.BytesIO(self.ws.export("alpha-deal")))
        columns = [cell.value for cell in wb["Deal ledger"][1]]
        process, round_ = (wb["Deal ledger"].cell(2, columns.index(field) + 1) for field in ("Process", "Round"))
        self.assertEqual((process.value, process.data_type, round_.value), (2, "n", 3))
        # The same edit as three ordinary updates, from the same starting point, records the same changes.
        self.save([{"type": "restore", "target_revision": 1}], revision=2)
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Process": "2", "Round": 3}} for uid in uids], revision=3)
        single = self.ws.history("alpha-deal")["history"][0]
        key = lambda change: (change["uid"], change["field"])
        self.assertEqual(sorted(single["changes"], key=key), sorted(bulk["changes"], key=key))
        self.assertEqual(single["summary"], bulk["summary"])
        authors = self.cockpit.trace.deal_extras("alpha-deal", "austin")["field_authors"]
        self.assertEqual({authors[uid]["Round"]["revision"] for uid in uids}, {4})

    def test_bulk_update_covers_more_rows_than_the_operation_cap(self):
        path = self.root / "extraction/alpha-deal.xlsx"
        wb = load_workbook(path)
        ledger = wb["Deal ledger"]
        template = [cell.value for cell in ledger[2]]
        for number in range(8, 158):
            ledger.append([number if column == 0 else value for column, value in enumerate(template)])
        wb.save(path)
        catalog_path = self.root / "_dev/cockpit/catalog.json"
        catalog = json.loads(catalog_path.read_text(encoding="utf-8"))
        self.base = catalog["deals"]["alpha-deal"]["versions"][0]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        catalog_path.write_text(json.dumps(catalog), encoding="utf-8")
        self.cockpit = data.Cockpit(self.root)
        self.ws = self.cockpit.workspace
        uids = [row["uid"] for row in self.cockpit.deal("alpha-deal")["ledger"]["rows"]]
        self.assertEqual(len(uids), 156)
        with self.assertRaisesRegex(WorkspaceError, "operations required"):
            self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Round": 2}} for uid in uids[:101]])
        result = self.save([{"type": "bulk_update", "sheet": "Deal ledger", "uids": uids, "values": {"Round": "2"}}])
        self.assertEqual({row["cells"]["Round"] for row in result["ledger"]["rows"]}, {"2"})
        self.assertEqual(self.ws.history("alpha-deal")["history"][0]["summary"], "156 changes")

    def test_bulk_update_accepts_only_process_and_round_and_is_atomic(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        uids = [row["uid"] for row in rows[:2]]
        note = {"type": "update", "sheet": "Deal ledger", "uid": uids[0], "values": {"Note": "Would change"}}
        bad = [
            {"sheet": "Deal ledger", "uids": uids, "values": {"Note": "x"}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Round": 2, "Price low": 3}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"#": 9}},
            {"sheet": "Deal ledger", "uids": uids, "values": {}},
            {"sheet": "Rounds", "uids": [self.cockpit.deal("alpha-deal")["rounds"]["rows"][0]["uid"]], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": [], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": uids[0], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": [uids[0], uids[0]], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": [uids[0], "missing"], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": [uids[0], 7], "values": {"Round": 2}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Process": "0"}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Process": ""}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Round": ""}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Round": "-1"}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Round": "1.5"}},
            {"sheet": "Deal ledger", "uids": uids, "values": {"Round": "later"}},
        ]
        for op in bad:
            with self.subTest(op=op), self.assertRaises(WorkspaceError):
                self.save([note, {"type": "bulk_update", **op}])
        self.assertFalse(self.ws.history("alpha-deal")["history"][:-1])
        self.assertEqual(self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["cells"]["Note"], rows[0]["cells"]["Note"])
        post = self.save([{"type": "bulk_update", "sheet": "Deal ledger", "uids": uids, "values": {"Round": "post"}}])
        self.assertEqual([row["cells"]["Round"] for row in post["ledger"]["rows"][:2]], ["post", "post"])

    def test_bulk_update_reaches_rows_inserted_in_the_same_batch(self):
        first = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        saved = self.save([
            {"type": "insert", "sheet": "Deal ledger", "client_uid": "new-one", "after_uid": first, "values": {"Who": "First new", "Round": 1}},
            {"type": "bulk_update", "sheet": "Deal ledger", "uids": [first, "new-one"], "values": {"Process": 2, "Round": 0}},
        ])
        self.assertEqual([(row["cells"]["Who"], row["cells"]["Process"], row["cells"]["Round"]) for row in saved["ledger"]["rows"][:3]],
                         [("Alpha", "2", "0"), ("First new", "2", "0"), ("Alpha", "1", "1")])

    def test_interior_narrative_range_blocks_dangling_delete(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        first, middle, last = rows[0]["uid"], rows[1]["uid"], rows[2]["uid"]
        state = self.save([{"type": "update", "sheet": "Deal ledger", "uid": first, "values": {"Note": "Compare #1-#3"}}])
        with self.assertRaises(WorkspaceError):
            self.save([{"type": "delete", "sheet": "Deal ledger", "uid": middle}], revision=1)
        state = self.save([{"type": "delete", "sheet": "Deal ledger", "uid": middle, "replacement_uid": last}], revision=1)
        self.assertEqual(state["ledger"]["rows"][0]["cells"]["Note"], "Compare #1, #2")

    def test_batch_literals_use_pre_save_event_numbers(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        saved = self.save([
            {"type": "insert", "sheet": "Deal ledger", "client_uid": "new-bid", "after_uid": rows[0]["uid"], "values": {"Who": "New bidder"}},
            {"type": "update", "sheet": "Deal ledger", "uid": rows[2]["uid"], "values": {"Note": "See #2 and unknown #999"}},
        ])
        self.assertEqual(saved["ledger"]["rows"][3]["cells"]["Note"], "See #3 and unknown #999")

    def test_failed_render_rolls_back_and_excel_length_is_bounded(self):
        uid = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        operation = {"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "A change"}}
        with patch.object(self.ws, "_render_xlsx", side_effect=RuntimeError("synthetic render failure")):
            with self.assertRaisesRegex(RuntimeError, "synthetic render failure"):
                self.save([operation])
        self.assertEqual([entry["revision"] for entry in self.ws.history("alpha-deal")["history"]], [0])
        with self.assertRaisesRegex(WorkspaceError, "32767"):
            self.save([{**operation, "values": {"Note": "😀" * 20000}}])
        self.assertEqual([entry["revision"] for entry in self.ws.history("alpha-deal")["history"]], [0])

    def test_new_text_in_blank_cell_gets_required_wrap(self):
        path = self.root / "extraction/alpha-deal.xlsx"
        book = load_workbook(path)
        ledger = book["Deal ledger"]
        column = check_lean.LEDGER_COLUMNS.index("Note") + 1
        self.assertIsNone(ledger.cell(2, column).value)
        self.assertIsNot(ledger.cell(2, column).alignment.wrap_text, True)
        book.save(path)
        catalog_path = self.root / "_dev/cockpit/catalog.json"
        catalog = json.loads(catalog_path.read_text())
        self.base = catalog["deals"]["alpha-deal"]["versions"][0]["sha256"] = hashlib.sha256(path.read_bytes()).hexdigest()
        catalog_path.write_text(json.dumps(catalog))
        self.cockpit = data.Cockpit(self.root)
        self.ws = self.cockpit.workspace
        uid = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Note": "New reviewer text"}}])
        rendered = load_workbook(io.BytesIO(self.ws.export("alpha-deal")))
        self.assertEqual(rendered["Deal ledger"].cell(2, column).value, "New reviewer text")
        self.assertIs(rendered["Deal ledger"].cell(2, column).alignment.wrap_text, True)

    def test_replacement_chain_and_late_deleted_reference_rejected(self):
        rows = self.cockpit.deal("alpha-deal")["ledger"]["rows"]
        first, unreferenced = rows[0]["uid"], rows[-1]["uid"]
        with self.assertRaises(WorkspaceError):
            self.save([
                {"type": "delete", "sheet": "Deal ledger", "uid": first, "replacement_uid": unreferenced},
                {"type": "delete", "sheet": "Deal ledger", "uid": unreferenced},
            ])
        with self.assertRaises(WorkspaceError):
            self.save([
                {"type": "delete", "sheet": "Deal ledger", "uid": unreferenced},
                {"type": "update", "sheet": "Deal ledger", "uid": first, "values": {"Note": "Later see #7"}},
            ])
        self.assertEqual([entry["revision"] for entry in self.ws.history("alpha-deal")["history"]], [0])

    def test_question_rename_updates_narrative_and_late_old_reference_rejected(self):
        payload = self.cockpit.deal("alpha-deal")
        question = payload["questions"]["rows"][0]["uid"]
        round_uid = payload["rounds"]["rows"][0]["uid"]
        ledger_uid = payload["ledger"]["rows"][0]["uid"]
        with self.assertRaises(WorkspaceError):
            self.save([
                {"type": "update", "sheet": "Questions", "uid": question, "values": {"Q": "Q2"}},
                {"type": "update", "sheet": "Rounds", "uid": round_uid, "values": {"How opened": "Old Q1 again"}},
            ])
        saved = self.save([
            {"type": "update", "sheet": "Rounds", "uid": round_uid, "values": {"How opened": "Review Q1"}},
            {"type": "update", "sheet": "Questions", "uid": question, "values": {"Q": "Q2"}},
        ])
        self.assertEqual(saved["rounds"]["rows"][0]["cells"]["How opened"], "Review Q2")
        self.assertEqual(saved["ledger"]["rows"][0]["cells"]["Flag"], "Q2")
        with self.assertRaises(WorkspaceError):
            self.save([{"type": "delete", "sheet": "Questions", "uid": question}], revision=1)

    def test_review_item_r_id_saves_renames_and_deletes_after_refs_are_cleared(self):
        payload = self.cockpit.deal("alpha-deal")
        question = payload["questions"]["rows"][0]["uid"]
        ledger = payload["ledger"]["rows"][0]["uid"]
        rounds = payload["rounds"]["rows"][0]["uid"]
        first = self.save([{"type": "update", "sheet": "Questions", "uid": question, "values": {"Q": "R1"}}])
        self.assertEqual((first["questions"]["rows"][0]["cells"]["Q"], first["ledger"]["rows"][0]["cells"]["Flag"]), ("R1", "R1"))
        second = self.save([{"type": "update", "sheet": "Rounds", "uid": rounds, "values": {"How opened": "Review R1"}},
                            {"type": "update", "sheet": "Questions", "uid": question, "values": {"Q": "R2"}}], revision=1)
        self.assertEqual((second["rounds"]["rows"][0]["cells"]["How opened"], second["ledger"]["rows"][0]["cells"]["Flag"]), ("Review R2", "R2"))
        with self.assertRaisesRegex(WorkspaceError, "referenced"):
            self.save([{"type": "delete", "sheet": "Questions", "uid": question}], revision=2)
        third = self.save([{"type": "update", "sheet": "Rounds", "uid": rounds, "values": {"How opened": "Outreach"}},
                           {"type": "update", "sheet": "Deal ledger", "uid": ledger, "values": {"Flag": ""}},
                           {"type": "delete", "sheet": "Questions", "uid": question}], revision=2)
        self.assertEqual((third["questions"]["rows"], third["ledger"]["rows"][0]["cells"]["Flag"]), ([], ""))


class SchemaAndRebaseTests(unittest.TestCase):
    """Version 1 payloads, recorded checks, comparisons and safe rebasing."""

    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit, self.original = fixture(self.root)
        self.ws = self.cockpit.workspace
        self.base = hashlib.sha256(self.original).hexdigest()
        self.filing = next((self.root / "raw_filing").glob("alpha-deal_*.htm"))
        self.version1 = build_version1_workbook(self.root / "version1.xlsx")

    def tearDown(self): self.temp.cleanup()

    def save(self, ops, reason="Review edit", actor="austin"):
        current = self.ws.deal("alpha-deal")["workspace"]
        return self.ws.edit("alpha-deal", {"revision": current["revision"], "base_sha256": current["base_sha256"], "reason": reason, "operations": ops}, actor)

    def test_schema_label_is_available_when_the_check_is_fatal(self):
        self.assertEqual(data.report_schema({"status": "error"}), "Version 1")
        report = check_lean.LeanChecker(self.version1, self.root / "missing.htm").run()
        self.assertEqual(report["status"], "error")
        payload = data.build_deal_payload("alpha-deal", {}, data.read_workbook(self.version1), self.cockpit.filing(self.filing), report)
        self.assertEqual(payload["ledger_schema"], "Version 1")

    def test_all_versions_use_the_current_checker_and_choices(self):
        for ident, digest in (("run-a", "a" * 64), ("run-b", "b" * 64)):
            add_run_version(self.ws, "alpha-deal", ident, self.version1, instruction_sha256=digest)
            shown = self.ws.deal("alpha-deal", ident)
            self.assertEqual((shown["ledger_schema"], shown["check"]["checker_version"]), ("Version 1", "Version 1"))
            self.assertEqual(shown["choices"], check_lean.choice_lists())
        with self.assertRaisesRegex(WorkspaceError, "Part stock"):
            self.ws._coerce(data.LEDGER_SHEET, "Stock %", "40-60")
        self.assertEqual(self.ws._coerce(data.LEDGER_SHEET, "Stock %", "Part stock"), "Part stock")
        self.assertEqual(self.ws._coerce(data.LEDGER_SHEET, "Stock %", "37.5"), 37.5)

    def test_an_unlisted_value_saves(self):
        payload = self.ws.deal("alpha-deal")
        rounds, ledger = payload["rounds"]["rows"][0]["uid"], payload["ledger"]["rows"][0]["uid"]
        saved = self.save([{"type": "update", "sheet": "Rounds", "uid": rounds, "values": {"Deadline outcome": "Extended; Enforced"}},
                           {"type": "update", "sheet": "Deal ledger", "uid": ledger, "values": {"Formality": "Not on any list"}}])
        self.assertEqual(saved["rounds"]["rows"][0]["cells"]["Deadline outcome"], "Extended; Enforced")
        self.assertEqual(saved["ledger"]["rows"][0]["cells"]["Formality"], "Not on any list")
        self.assertNotIn("Not on any list", saved["choices"]["Formality"])

    def test_compare_versions_shows_changed_terms_in_both_directions(self):
        version = add_run_version(self.ws, "alpha-deal", "opus55-version1", self.version1)
        forward = self.ws.compare("alpha-deal", "base-raw", version["id"])["changes"]
        backward = self.ws.compare("alpha-deal", version["id"], "base-raw")["changes"]
        for changes in (forward, backward):
            fields = {change["field"] for change in changes if change["type"] == "update"}
            self.assertTrue({"Stock %", "CVR/earnout value", "Financing"} <= fields, fields)
        self.assertIn(("Price low", "10", "21.25"), [(c["field"], c["before"], c["after"]) for c in forward])
        self.assertIn(("Stock %", "", "Part stock"), [(c["field"], c["before"], c["after"]) for c in forward])
        self.assertIn(("Stock %", "Part stock", ""), [(c["field"], c["before"], c["after"]) for c in backward])

    def test_run_check_receipt_is_displayed_without_mutation(self):
        receipt = check_lean.LeanChecker(self.version1, self.filing).run()
        add_run_version(self.ws, "alpha-deal", "run-receipt", self.version1, checker={"errors": 9, "warnings": 9}, receipt=receipt)
        add_run_version(self.ws, "alpha-deal", "run-stored", self.version1,
                        checker={"errors": 0, "warnings": 2, "checker_version": "Version 1", "ledger_schema": "Version 1"})
        add_run_version(self.ws, "alpha-deal", "run-bare", self.version1)
        before = {path: path.read_bytes() for path in self.root.rglob("check.json")}
        shown = {version["id"]: version.get("checker") for version in self.ws.deal("alpha-deal")["versions"]}
        self.assertIsNone(shown["base-raw"])
        self.assertEqual(shown["run-receipt"], {"checker_version": "Version 1", "ledger_schema": "Version 1",
                                                "errors": receipt["summary"]["errors"], "warnings": receipt["summary"]["warnings"]})
        self.assertEqual(shown["run-stored"], {"checker_version": "Version 1", "ledger_schema": "Version 1", "errors": 0, "warnings": 2})
        self.assertIsNone(shown["run-bare"])
        self.assertEqual({path: path.read_bytes() for path in self.root.rglob("check.json")}, before)

    def test_version1_round_trip_keeps_every_value_and_the_raw_version(self):
        version = add_run_version(self.ws, "alpha-deal", "opus55-version1", self.version1)
        raw_before = version["path"].read_bytes()
        rebased = self.save([{"type": "rebase", "target_version": version["id"]}], reason="Use the Version 1 run")
        first = rebased["ledger"]["rows"][0]["uid"]
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": first, "values": {"Who": "Alpha Holdings"}}])
        reopened = data.Cockpit(self.root)
        self.assertEqual(reopened.deal("alpha-deal")["ledger"]["rows"][0]["cells"]["Who"], "Alpha Holdings")
        exported = load_workbook(io.BytesIO(reopened.workspace.export("alpha-deal")))
        source = load_workbook(self.version1)
        for sheet in ("Deal ledger", "Rounds", "Questions", "Deal facts"):
            got = [[cell.value for cell in row] for row in exported[sheet].iter_rows()]
            want = [[cell.value for cell in row] for row in source[sheet].iter_rows()]
            want = [row for row in want if any(value is not None for value in row) or row is want[0]]
            if sheet == "Deal ledger":
                want[1][check_lean.LEDGER_COLUMNS.index("Who")] = "Alpha Holdings"
            self.assertEqual(got, want, sheet)
        ledger = exported["Deal ledger"]
        header = [cell.value for cell in ledger[1]]
        cell = lambda row, field: ledger.cell(row, header.index(field) + 1)
        self.assertEqual((cell(2, "Price low").value, cell(2, "Price low").data_type), (21.25, "n"))
        self.assertEqual((cell(2, "CVR/earnout value").value, cell(2, "CVR/earnout value").data_type), (1.25, "n"))
        self.assertEqual((cell(2, "Stock %").value, cell(3, "Stock %").value, cell(3, "Financing").value), ("Part stock", 0, "Varies"))
        self.assertIsNone(cell(4, "Price low").value)
        self.assertIsInstance(cell(2, "Sort date").value, dt.datetime)
        self.assertEqual((cell(2, "Flag").value, cell(3, "Note").value), ("Q1", "Cohort; see #1-#3"))
        self.assertEqual(exported["Questions"].cell(2, check_lean.QUESTION_COLUMNS.index("Rows affected") + 1).value, "#1, #2")
        self.assertEqual(reopened.workspace.deal("alpha-deal")["ledger_schema"], "Version 1")
        self.assertEqual(version["path"].read_bytes(), raw_before)
        self.assertEqual(hashlib.sha256(reopened.workspace.export("alpha-deal", version["id"])).hexdigest(), version["sha256"])

    def test_a_working_copy_keeps_its_base_columns(self):
        uid = self.ws.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        with self.assertRaisesRegex(WorkspaceError, "invalid columns"):
            self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Unknown column": "0"}}])
        version = add_run_version(self.ws, "alpha-deal", "opus55-version1", self.version1)
        rebased = self.save([{"type": "rebase", "target_version": version["id"]}])
        self.assertEqual(rebased["ledger"]["columns"], list(check_lean.LEDGER_COLUMNS))
        with self.assertRaisesRegex(WorkspaceError, "invalid columns"):
            self.save([{"type": "update", "sheet": "Deal ledger", "uid": rebased["ledger"]["rows"][0]["uid"], "values": {"All cash": "Yes"}}])

    def test_a_past_revision_compares_and_downloads_without_a_save(self):
        uid = self.ws.deal("alpha-deal")["ledger"]["rows"][0]["uid"]
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Who": "First edit"}}])
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": uid, "values": {"Who": "Second edit"}}])
        compared = self.ws.compare("alpha-deal", "rev:1", "working")
        self.assertEqual((compared["from_label"], [(c["field"], c["before"], c["after"]) for c in compared["changes"]]), ("Revision 1", [("Who", "First edit", "Second edit")]))
        self.assertEqual([(c["before"], c["after"]) for c in self.ws.compare("alpha-deal", "rev:0", "rev:1")["changes"]], [("Alpha", "First edit")])
        past = load_workbook(io.BytesIO(self.ws.export("alpha-deal", "rev:1")))
        self.assertEqual(past["Deal ledger"]["C2"].value, "First edit")
        self.assertEqual(self.ws.export("alpha-deal", "rev:0"), self.original)
        with self.assertRaises(Missing): self.ws.export("alpha-deal", "rev:9")
        with self.assertRaises(Missing): self.ws.compare("alpha-deal", "rev:01", "working")
        self.assertEqual([entry["revision"] for entry in self.ws.history("alpha-deal")["history"]], [2, 1, 0])

    def test_rebase_preview_counts_and_carried_judgments(self):
        version = add_run_version(self.ws, "alpha-deal", "opus55-version1", self.version1)
        rows = self.ws.deal("alpha-deal")["ledger"]["rows"]
        self.save([{"type": "update", "sheet": "Deal ledger", "uid": rows[0]["uid"], "values": {"Who": "Edited"}},
                   {"type": "review", "uid": rows[0]["uid"], "status": "reviewed"}, {"type": "review", "uid": rows[1]["uid"], "status": "needs_decision"},
                   {"type": "finding", "id": "F1", "judgment": "supported", "implementation": "applied", "verification": "verified", "note": ""}])
        self.save([{"type": "delete", "sheet": "Deal ledger", "uid": rows[-1]["uid"]}])
        trace = self.cockpit.trace
        trace.comment("alpha-deal", {"action": "create", "target": {"kind": "row", "sheet": "Deal ledger", "uid": rows[0]["uid"]}, "body": "On the old row"}, "alex")
        trace.comment("alpha-deal", {"action": "create", "target": {"kind": "deal"}, "body": "Deal level"}, "alex")
        before = self.ws.history("alpha-deal")
        preview = self.ws.rebase_preview("alpha-deal", version["id"])
        self.assertEqual(self.ws.history("alpha-deal"), before)  # nothing saved
        self.assertEqual((preview["current"]["ledger_schema"], preview["target"]["ledger_schema"], preview["current"]["revision"]), ("Version 1", "Version 1", 2))
        stops = preview["stops_applying"]
        self.assertEqual(stops["revisions"], 2)
        self.assertEqual(stops["edits"], 2)  # one cell, one deleted row
        self.assertEqual(stops["row_marks"], {"total": 2, "reviewed": 1, "needs_decision": 1, "unreviewed": 0})
        self.assertEqual(stops["finding_decisions"], {"total": 1, "judgments_kept": 1, "reset": 1})
        self.assertEqual(stops["row_threads"], {"total": 1, "open": 1, "resolved": 0})
        with self.assertRaisesRegex(WorkspaceError, "already the base"): self.ws.rebase_preview("alpha-deal", "base-raw")
        with self.assertRaises(Missing): self.ws.rebase_preview("alpha-deal", "nosuch")

        shown = self.ws.deal("alpha-deal", version["id"])["workspace"]  # the rebase dialog saves from this view
        self.assertEqual((shown["revision"], shown["updated_by"], shown["editable"]), (2, "austin", False))
        rebased = self.save([{"type": "rebase", "target_version": version["id"]}], reason="Use the Version 1 run", actor="alex")
        finding = rebased["findings"][0]
        self.assertEqual((finding["judgment"], finding["implementation"], finding["verification"], finding["actor"]), ("supported", "unassessed", "unchecked", "austin"))
        self.assertEqual(finding["carried_over"], {"revision": 3, "from_base": "base-raw", "actor": "alex"})
        self.assertEqual(rebased["row_review"], {})
        threads = trace.comments("alpha-deal")["threads"]
        self.assertEqual([(t["target"]["kind"], t["target_missing"], t["target_context"]) for t in threads],
                         [("row", True, "on an earlier base (revision 2)"), ("deal", False, None)])
        restored = self.save([{"type": "restore", "target_revision": 2}], reason="Undo the rebase")
        self.assertEqual((restored["findings"][0]["implementation"], len(restored["row_review"])), ("applied", 2))
        self.assertEqual(trace.comments("alpha-deal")["threads"][0]["target_missing"], False)

    def test_a_thread_on_a_row_deleted_by_an_edit_reads_record_removed(self):
        rows = self.ws.deal("alpha-deal")["ledger"]["rows"]
        trace = self.cockpit.trace
        trace.comment("alpha-deal", {"action": "create", "target": {"kind": "row", "sheet": "Deal ledger", "uid": rows[-1]["uid"]}, "body": "Delete this?"}, "alex")
        self.save([{"type": "delete", "sheet": "Deal ledger", "uid": rows[-1]["uid"]}])
        [thread] = trace.comments("alpha-deal")["threads"]
        self.assertEqual((thread["target_missing"], thread["target_context"]), (True, "record removed"))
        version = add_run_version(self.ws, "alpha-deal", "opus55-version1", self.version1)
        trace.comment("alpha-deal", {"action": "create", "target": {"kind": "row", "sheet": "Deal ledger", "uid": rows[0]["uid"]}, "body": "Before any save"}, "alex")
        self.save([{"type": "rebase", "target_version": version["id"]}])
        contexts = [thread["target_context"] for thread in trace.comments("alpha-deal")["threads"]]
        self.assertEqual(contexts, ["record removed", "on an earlier base (revision 1)"])


class MigrationTests(unittest.TestCase):
    def test_add_column_tolerates_a_concurrent_migration(self):
        import sqlite3
        from cockpit.workspace import add_column
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "db.sqlite3"
            first, second = sqlite3.connect(path), sqlite3.connect(path)
            first.execute("CREATE TABLE jobs (id TEXT)"); first.commit()
            second.execute("PRAGMA table_info(jobs)").fetchall()  # the second connection has already looked
            add_column(first, "jobs", "cancelled_by", "TEXT"); first.commit()

            class Stale:  # replays the second connection's stale view, then its ALTER meets the added column
                def execute(self, sql, *args):
                    return iter([(0, "id")]) if sql.startswith("PRAGMA") else second.execute(sql, *args)
            add_column(Stale(), "jobs", "cancelled_by", "TEXT")
            self.assertEqual([row[1] for row in second.execute("PRAGMA table_info(jobs)")], ["id", "cancelled_by"])
            with self.assertRaises(sqlite3.OperationalError): add_column(Stale(), "missing", "x", "TEXT")
            first.close(); second.close()


class HTTPTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.cockpit, _ = fixture(self.root)
        self.httpd = server.make_server(0, self.cockpit, quiet=True)
        self.port = self.httpd.server_address[1]
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def tearDown(self):
        self.httpd.shutdown(); self.httpd.server_close(); self.temp.cleanup()

    def request(self, path, body=None, headers=None):
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=body, headers=headers or {})
        try:
            with urllib.request.urlopen(req, timeout=20) as response: return response.status, json.loads(response.read()) if response.headers.get_content_type() == "application/json" else response.read()
        except urllib.error.HTTPError as exc: return exc.code, json.loads(exc.read())

    def test_session_and_write_guards(self):
        status, session = self.request("/api/session")
        self.assertEqual(status, 200)
        self.assertTrue(session["can_edit"])
        row = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]
        body = json.dumps({"revision": 0, "base_sha256": self.cockpit.deal("alpha-deal")["workspace"]["base_sha256"], "reason": "Test", "operations": [{"type": "review", "uid": row["uid"], "status": "reviewed", "note": "okay"}]}).encode()
        base_headers = {"Content-Type": "application/json", "Origin": f"http://127.0.0.1:{self.port}", "X-Cockpit-CSRF": session["csrf_token"]}
        self.assertEqual(self.request("/api/deal/alpha-deal/edit", body, {**base_headers, "X-Cockpit-CSRF": "bad"})[0], 403)
        self.assertEqual(self.request("/api/deal/alpha-deal/edit", body, {**base_headers, "Origin": "https://evil.invalid"})[0], 403)
        self.assertEqual(self.request("/api/deal/alpha-deal/edit", body, base_headers)[0], 200)
        self.assertEqual(self.request("/api/deal/alpha-deal/edit", body, base_headers)[0], 409)
        self.assertEqual(self.request("/api/document/alpha-deal/..%2F..%2Fsecret")[0], 404)
        self.assertEqual(self.request("/assets/..%2Fserver.py")[0], 404)

    def test_public_identity_requires_configured_host_and_access_email(self):
        with patch.dict(os.environ, {"COCKPIT_PUBLIC_ORIGIN": "https://lines.dealextract.org"}):
            public = {"Host": "lines.dealextract.org", "Cf-Access-Authenticated-User-Email": "junyu.li.24@ucl.ac.uk"}
            status, session = self.request("/api/session", headers=public)
            self.assertEqual(status, 200)
            self.assertEqual((session["user"], session["can_edit"]), ("austin", True))
            status, anonymous = self.request("/api/session", headers={"Host": "lines.dealextract.org"})
            self.assertEqual(status, 200)
            self.assertFalse(anonymous["can_edit"])
            row = self.cockpit.deal("alpha-deal")["ledger"]["rows"][0]
            body = json.dumps({"revision": 0, "base_sha256": self.cockpit.deal("alpha-deal")["workspace"]["base_sha256"], "reason": "Public fixture", "operations": [{"type": "review", "uid": row["uid"], "status": "reviewed", "note": "okay"}]}).encode()
            headers = {**public, "Content-Type": "application/json", "Origin": "https://lines.dealextract.org", "X-Cockpit-CSRF": session["csrf_token"]}
            self.assertEqual(self.request("/api/deal/alpha-deal/edit", body, headers)[0], 200)
            self.assertEqual(self.cockpit.deal("alpha-deal")["row_review"][row["uid"]]["actor"], "austin")


if __name__ == "__main__": unittest.main()
