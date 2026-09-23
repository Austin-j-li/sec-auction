#!/usr/bin/env python3
"""Prepare, launch, and monitor isolated extraction runs (Opus; Sol transport retained)."""

from __future__ import annotations

import argparse
import base64
import contextlib
import datetime as dt
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import shutil
import signal
import subprocess
import sys
import tempfile
import time


BASE = Path(__file__).resolve().parent
PROJECT = BASE.parents[2]
HOME_HOST = Path.home()
SANDBOX_HOME = Path("/home/uctpiaj")
WORK = SANDBOX_HOME / "work"
INSTRUCTION_NAME = "SEC_Deal_Ledger_Extraction_Instruction.md"
CODEX_BIN = Path(os.environ.get("SEC_CODEX_BIN") or shutil.which("codex") or "/missing/codex").resolve()
CODEX_RELEASE = CODEX_BIN.parent.parent
CLAUDE_BIN = Path(os.environ.get("SEC_CLAUDE_BIN") or shutil.which("claude") or "/missing/claude").resolve()
# A long-lived subscription token from `claude setup-token`, readable only by its owner. It reaches
# the sandboxed CLI through an inherited pipe, never argv, environment or disk. The host login's
# credentials file is never shared: a sandboxed refresh would rotate its refresh token, which the
# read-only sandbox cannot write back, and so log the host out.
CLAUDE_TOKEN_FILE = Path(os.environ.get("SEC_CLAUDE_OAUTH_TOKEN_FILE")
                         or HOME_HOST / ".config/sec-extraction/claude-oauth-token")
# Models and effort levels each provider may be prepared with. The first model is the default.
MODELS = {"sol": ("gpt-6-sol", "gpt-5.6-sol"), "opus": ("claude-opus-5-5", "claude-opus-5")}
# Codex's "ultra" level delegates to subagents automatically, so it is not offered.
EFFORTS = {"sol": ("low", "medium", "high", "xhigh", "max"), "opus": ("low", "medium", "high", "xhigh", "max")}
DEFAULT_EFFORT = {"sol": "xhigh", "opus": "medium"}  # Opus: 22 Sep 2026 effort sweep
# Claude Code settings for every Opus run. A classifier refusal fails the run instead of switching
# models; the prompt-cache lifetime stays the subscription default; and the "user hasn't heard
# from you" reminder never fires, since no one reads a sandboxed run while it works.
CLAUDE_ENV = {
    "CLAUDE_CODE_DISABLE_REFUSAL_FALLBACK": "1",
    "CLAUDE_CODE_PROMPT_CACHE_TTL": "1h",
    "CLAUDE_CODE_SILENT_TURN_REMINDER_TURNS": "1000000",
}
# Codex features switched off for every Sol run. By default `codex exec` offers the account's
# ChatGPT app connectors (mail, GitLab, site deploys, a remote shell), web browsing, image
# generation and subagents; a sandboxed extraction gets shell commands and file patches only.
# Code mode stays on because GPT-6-Sol calls every tool through it.
CODEX_DISABLED_FEATURES = (
    "apps", "plugins", "remote_plugin", "plugin_sharing", "browser_use", "browser_use_external",
    "browser_use_full_cdp_access", "in_app_browser", "computer_use", "image_generation", "multi_agent",
    "goals", "sleep_tool", "tool_suggest", "view_image", "skill_search", "skill_mcp_dependency_install",
    "mentions_v2", "workspace_dependencies", "hooks",
)
CODEX_TOKEN_MARGIN_HOURS = 7  # longer than the longest run, so a sandboxed Codex never refreshes its login
MAX_CONTINUATIONS = 2  # Opus 5.5 can end a turn with a progress report before saving the workbook
SHEETS = ["Deal ledger", "Rounds", "Questions", "Deal facts"]
CODEX_SANDBOX = Path("/opt/codex")
CLAUDE_SANDBOX = Path("/opt/claude")
MINIFORGE = HOME_HOST / "miniforge3"  # Ubuntu laptop; absent on the VM, where system Python is used
import openpyxl as _openpyxl
PYLIB_HOST = Path(_openpyxl.__file__).resolve().parents[1]  # site-packages that holds openpyxl
PYLIB = Path("/opt/pylib")
REPORT_NAME = "checker_report.md"
TIMEOUT_SECONDS = 90 * 60
TIMEOUT_BOUNDS = (10 * 60, 6 * 60 * 60)
RUNS = PROJECT / "_dev" / "runs"


def now_iso() -> str:
    return dt.datetime.now(dt.timezone.utc).isoformat()


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    os.replace(tmp, path)


def library_versions() -> dict[str, str | None]:
    """Installed versions of the pinned libraries the extracting agent can import."""
    versions: dict[str, str | None] = {}
    for name in ("openpyxl", "beautifulsoup4", "lxml"):
        try:
            versions[name] = importlib.metadata.version(name)
        except importlib.metadata.PackageNotFoundError:
            versions[name] = None
    return versions


def events(events_path: Path):
    """JSON objects from a provider's event log, skipping lines that are not JSON objects."""
    with events_path.open(encoding="utf-8", errors="replace") as lines:
        for line in lines:
            try:
                event = json.loads(line)
            except ValueError:
                continue
            if isinstance(event, dict):
                yield event


def run_usage(events_path: Path) -> dict[str, object] | None:
    """Tokens and cost from the provider's event log; None if it reports none.

    Claude ends each invocation with one "result" event. Its usage covers that invocation, so
    usage is summed over a run and its continuations; total_cost_usd is cumulative for the
    session, so the last one is the run's cost.
    Codex reports usage per "turn.completed" event and no cost; turns are summed.
    """
    totals: dict[str, float] = {}
    cost = None
    found = False
    for event in events(events_path):
        if event.get("type") not in ("result", "turn.completed"):
            continue
        usage = event.get("usage")
        if not isinstance(usage, dict):
            continue
        found = True
        if event["type"] == "result":
            cost = event.get("total_cost_usd")
        for key, value in usage.items():
            if key.endswith("_tokens") and isinstance(value, (int, float)):
                totals[key] = totals.get(key, 0) + value
    if not found:
        return None
    return {"tokens": totals, "cost_usd": cost}


def claude_results(events_path: Path) -> list[dict]:
    """The Claude CLI's result events, one per invocation."""
    return [event for event in events(events_path) if event.get("type") == "result"]


def claude_summary(events_path: Path) -> dict[str, object] | None:
    """What the Claude CLI reported about a run and its continuations; None if it reported nothing.

    modelUsage and duration_api_ms are cumulative for the session, like total_cost_usd.
    """
    init, results, refusal = None, [], False
    for event in events(events_path):
        if event.get("type") == "system" and event.get("subtype") == "init" and init is None:
            init = event
        elif event.get("type") == "assistant" and (event.get("message") or {}).get("stop_reason") == "refusal":
            refusal = True
        elif event.get("type") == "result":
            results.append(event)
    if init is None and not results and not refusal:
        return None
    init, last = init or {}, results[-1] if results else {}

    def total(field) -> int:
        return sum(field(result.get("usage") or {}) or 0 for result in results)

    return {
        "claude_code_version": init.get("claude_code_version"),
        "session_model": init.get("model"),
        "tools": init.get("tools"),
        "invocations": len(results),
        "served_models": sorted({model for result in results for model in result.get("modelUsage") or {}}),
        "subtype": last.get("subtype"),
        "is_error": last.get("is_error"),
        "stop_reason": last.get("stop_reason"),
        "stop_details": last.get("stop_details"),
        "terminal_reason": last.get("terminal_reason"),
        "api_error_status": last.get("api_error_status"),
        "refusal": refusal or any(result.get("stop_reason") == "refusal" for result in results),
        "num_turns": sum(result.get("num_turns") or 0 for result in results),
        "duration_api_ms": last.get("duration_api_ms"),
        "thinking_tokens": total(lambda usage: (usage.get("output_tokens_details") or {}).get("thinking_tokens")),
        "cache_write_1h_tokens": total(lambda usage: (usage.get("cache_creation") or {}).get("ephemeral_1h_input_tokens")),
        "cache_write_5m_tokens": total(lambda usage: (usage.get("cache_creation") or {}).get("ephemeral_5m_input_tokens")),
        "permission_denials": sum(len(result.get("permission_denials") or []) for result in results),
        "subagents_spawned": sum((result.get("subagent_stats") or {}).get("spawned") or 0 for result in results),
    }


def prepare(args: argparse.Namespace) -> None:
    if bool(args.revise_from) != bool(args.report):
        raise SystemExit("--revise-from and --report must be supplied together")
    if Path(args.filing).name != args.filing or Path(args.deal).name != args.deal or args.deal in ("", ".", ".."):
        raise SystemExit("--filing and --deal must be bare names, not paths")
    if args.revise_from and (not Path(args.revise_from).is_file() or not Path(args.report).is_file()):
        raise SystemExit("revision workbook or findings report is missing")
    model = args.model or MODELS[args.provider][0]
    effort = args.effort or DEFAULT_EFFORT[args.provider]
    if model not in MODELS[args.provider] or effort not in EFFORTS[args.provider]:
        raise SystemExit(f"{args.provider} runs one of {', '.join(MODELS[args.provider])} "
                         f"at effort {', '.join(EFFORTS[args.provider])}")
    timeout = TIMEOUT_SECONDS if args.timeout_minutes is None else args.timeout_minutes * 60
    if not TIMEOUT_BOUNDS[0] <= timeout <= TIMEOUT_BOUNDS[1]:
        raise SystemExit("--timeout-minutes must be between 10 and 360")
    run_dir = Path(args.run_dir).resolve()
    if run_dir.exists() and any(run_dir.iterdir()):
        raise SystemExit(f"refusing to overwrite non-empty run directory: {run_dir}")
    instruction_src = Path(args.instruction).resolve() if args.instruction else PROJECT / INSTRUCTION_NAME
    filing_src = PROJECT / "raw_filing" / args.filing
    if not instruction_src.is_file() or not filing_src.is_file():
        raise SystemExit("instruction or filing source is missing")

    input_dir = run_dir / "input"
    raw_dir = input_dir / "raw_filing"
    output_dir = run_dir / "extraction"
    raw_dir.mkdir(parents=True)
    output_dir.mkdir()
    shutil.copy2(instruction_src, input_dir / INSTRUCTION_NAME)
    shutil.copy2(filing_src, raw_dir / args.filing)

    prompt = (
        f"Read {INSTRUCTION_NAME} in this folder and follow it. "
        f"Extract the deal whose filing is in raw_filing/ and save the finished workbook "
        f"in extraction/ as {args.deal}.xlsx. Work only from this instruction and this filing."
    )
    if args.revise_from:
        # Revision pass: the agent starts from a finished workbook and the checker's findings on it.
        shutil.copy2(args.revise_from, output_dir / f"{args.deal}.xlsx")
        shutil.copy2(args.report, input_dir / REPORT_NAME)
        prompt = (
            f"extraction/{args.deal}.xlsx is a finished ledger for the filing in raw_filing/, made by following "
            f"{INSTRUCTION_NAME} in this folder. {REPORT_NAME} lists what an automatic checker found in it. "
            f"Read the instruction, then go through the findings one by one against the filing. Mechanical errors "
            f"identify failed checks. Mechanical review leads are warnings, not necessarily errors; do not discard "
            f"required evidence just to silence one. Change the workbook only where a "
            f"finding shows a real error under the instruction; do not add rows the instruction would fold into a "
            f"note, and do not change anything no finding points to. Save the revised workbook under the same name, "
            f"and write extraction/revision_notes.md listing each finding, what you decided and why. "
            f"Work only from the instruction, the filing, the workbook and the findings."
        )
    (run_dir / "prompt.txt").write_text(prompt + "\n", encoding="utf-8")
    metadata = {
        "prepared_at": now_iso(),
        "provider": args.provider,
        "deal": args.deal,
        "filing_name": args.filing,
        "instruction_name": INSTRUCTION_NAME,
        "instruction_sha256": sha256(input_dir / INSTRUCTION_NAME),
        "filing_sha256": sha256(raw_dir / args.filing),
        "prompt_sha256": hashlib.sha256((prompt + "\n").encode()).hexdigest(),
        "expected_output": f"{args.deal}.xlsx",
        "mode": "revise" if args.revise_from else "extract",
        "revised_from_sha256": sha256(output_dir / f"{args.deal}.xlsx") if args.revise_from else None,
        "report_sha256": sha256(input_dir / REPORT_NAME) if args.revise_from else None,
        "model": model,
        "effort": effort,
        "timeout_seconds": timeout,
        "runner_sha256": sha256(Path(__file__)),
        "python_version": sys.version,
        "library_versions": library_versions(),
        "sandbox": "bubblewrap fresh home/state/tmp; one instruction; one filing; one output directory",
    }
    write_json(run_dir / "metadata.json", metadata)
    print(json.dumps(metadata, sort_keys=True))


def bwrap_base(run_dir: Path, provider: str, state: Path, scratch: Path) -> list[str]:
    """The sandbox: fresh provider state, one instruction, one filing, one output directory.

    Home and /tmp are scratch directories kept for all invocations of one run, so a continuation
    finds the working files its session made; the worker deletes them with the provider state.
    """
    metadata = prepared_metadata(run_dir, provider)
    instruction = run_dir / "input" / INSTRUCTION_NAME
    filings = list((run_dir / "input" / "raw_filing").glob("*"))
    if len(filings) != 1:
        raise SystemExit(f"expected exactly one filing in {run_dir}")
    filing = filings[0]
    output = run_dir / "extraction"
    for name in ("home", "tmp"):
        (scratch / name).mkdir(parents=True, exist_ok=True)
    h = str(SANDBOX_HOME)
    work = str(WORK)

    command = [
        "bwrap", "--unshare-all", "--share-net", "--die-with-parent", "--new-session", "--clearenv",
        "--ro-bind", "/usr", "/usr",
        "--symlink", "usr/bin", "/bin",
        "--symlink", "usr/sbin", "/sbin",
        "--symlink", "usr/lib", "/lib",
        "--symlink", "usr/lib64", "/lib64",
        "--ro-bind", "/etc/ssl", "/etc/ssl",
        "--ro-bind", "/etc/resolv.conf", "/etc/resolv.conf",
        "--ro-bind", "/etc/hosts", "/etc/hosts",
        "--ro-bind", "/etc/nsswitch.conf", "/etc/nsswitch.conf",
        "--ro-bind", "/etc/alternatives", "/etc/alternatives",
        "--ro-bind", "/etc/passwd", "/etc/passwd",
        "--ro-bind", "/etc/group", "/etc/group",
        "--ro-bind", "/etc/localtime", "/etc/localtime",
        "--proc", "/proc", "--dev", "/dev", "--bind", str(scratch / "tmp"), "/tmp", "--tmpfs", "/run",
        "--dir", "/home", "--bind", str(scratch / "home"), h,
        "--dir", work,
        "--dir", f"{work}/raw_filing",
        "--dir", f"{work}/extraction",
        "--ro-bind", str(instruction), f"{work}/{INSTRUCTION_NAME}",
        "--ro-bind", str(filing), f"{work}/raw_filing/{filing.name}",
        "--bind", str(output), f"{work}/extraction",
        "--setenv", "HOME", h,
        "--setenv", "USER", "uctpiaj",
        "--setenv", "LANG", "C.UTF-8",
        "--setenv", "TMPDIR", "/tmp",
        "--ro-bind", str(PYLIB_HOST), str(PYLIB),
        "--setenv", "PYTHONPATH", str(PYLIB),
    ]
    for certs in ["/etc/ca-certificates", "/etc/pki", "/etc/crypto-policies"]:  # Debian and Red Hat layouts
        if Path(certs).exists():
            command += ["--ro-bind", certs, certs]
    if MINIFORGE.is_dir():
        command += ["--ro-bind", str(MINIFORGE), str(MINIFORGE), "--setenv", "PATH", f"{MINIFORGE}/bin:/usr/bin:/bin"]
    else:
        command += ["--setenv", "PATH", "/usr/bin:/bin"]
    report = run_dir / "input" / REPORT_NAME
    if metadata.get("mode") == "revise":
        if not report.is_file():
            raise SystemExit(f"revision findings are missing: {report}")
        command += ["--ro-bind", str(report), f"{work}/{REPORT_NAME}"]

    if provider == "sol":
        command += [
            "--dir", str(SANDBOX_HOME / ".codex"),
            "--bind", str(state), str(SANDBOX_HOME / ".codex"),
        ]
        command += [
            "--ro-bind", str(CODEX_RELEASE), str(CODEX_SANDBOX),
            "--ro-bind", str(HOME_HOST / ".codex/auth.json"), str(SANDBOX_HOME / ".codex/auth.json"),
            "--setenv", "CODEX_HOME", str(SANDBOX_HOME / ".codex"),
        ]
    else:
        command += [
            "--ro-bind", str(CLAUDE_BIN), str(CLAUDE_SANDBOX),
            "--bind", str(state), str(SANDBOX_HOME / ".claude"),
        ]
        for name, value in CLAUDE_ENV.items():
            command += ["--setenv", name, value]
    command += ["--chdir", work]
    return command


def provider_command(run_dir: Path, provider: str, metadata: dict,
                     prompt: str | None = None, resume: str | None = None) -> list[str]:
    """The provider's argv, built from the verified prepared model and effort.

    A continuation resumes the Claude session with a new prompt; the session lives in the
    temporary provider state, so it is deleted with it.
    """
    prompt = prompt or (run_dir / "prompt.txt").read_text(encoding="utf-8").strip()
    model, effort = metadata["model"], metadata["effort"]
    if provider == "sol":
        codex = CODEX_SANDBOX / "bin/codex"
        disabled = [part for feature in CODEX_DISABLED_FEATURES for part in ("--disable", feature)]
        return [
            str(codex), "exec", "--ephemeral", "--ignore-user-config", "--ignore-rules",
            "--skip-git-repo-check", "--dangerously-bypass-approvals-and-sandbox",
            "--json", "--color", "never", "--model", model,
            "-c", f'model_reasoning_effort="{effort}"', "-c", 'web_search="disabled"', *disabled,
            "-C", str(WORK), prompt,
        ]
    command = [
        str(CLAUDE_SANDBOX), "--print", "--output-format", "stream-json", "--verbose",
        "--model", model, "--effort", effort, "--safe-mode",
        "--disable-slash-commands", "--no-chrome", "--strict-mcp-config",
        "--mcp-config", '{"mcpServers":{}}', "--tools", "Bash,Read,Write", "--allowedTools", "Bash,Read,Write",
        "--disallowedTools", "WebFetch,WebSearch,Agent", "--dangerously-skip-permissions",
    ]
    if resume:
        command += ["--resume", resume]
    return command + [prompt]


def unfinished(run_dir: Path, metadata: dict) -> bool:
    """Whether a run still owes its deliverable: revision notes, or a readable four-sheet workbook."""
    if metadata.get("mode") == "revise":
        return not (run_dir / "extraction" / "revision_notes.md").is_file()
    validation = validate_workbook(run_dir)
    return not validation.get("valid_xlsx") or validation.get("sheet_names") != SHEETS


def continuation_prompt(metadata: dict) -> str:
    if metadata.get("mode") == "revise":
        owed = "extraction/revision_notes.md has not been written. Continue until the revised workbook and its notes are saved."
    else:
        owed = (f"extraction/{metadata['expected_output']} is not yet saved as the finished four-sheet workbook. "
                f"Continue until it is.")
    return f"The task is not finished: {owed} If something blocks you, say what it is."


def validate_workbook(run_dir: Path) -> dict[str, object]:
    metadata = json.loads((run_dir / "metadata.json").read_text(encoding="utf-8"))
    workbook = run_dir / "extraction" / metadata["expected_output"]
    result: dict[str, object] = {"path": str(workbook), "exists": workbook.is_file()}
    if not workbook.is_file():
        result["valid_xlsx"] = False
        result["error"] = "expected workbook missing"
        return result
    try:
        from openpyxl import load_workbook

        result["bytes"] = workbook.stat().st_size
        result["sha256"] = sha256(workbook)
        opened = load_workbook(workbook, read_only=True, data_only=False)
        try:
            # Read-only worksheets are lazy: opening the ZIP alone misses broken sheet XML.
            for sheet in opened:
                sheet.reset_dimensions()
                for _ in sheet.iter_rows(values_only=True):
                    pass
            result["sheet_names"] = opened.sheetnames
            result["valid_xlsx"] = True
        finally:
            opened.close()
    except Exception as exc:  # validation evidence, not repair
        result["valid_xlsx"] = False
        result["error"] = f"{type(exc).__name__}: {exc}"
    return result


@contextlib.contextmanager
def claude_token(provider: str):
    """Sandbox arguments and inherited descriptors that hand the Claude CLI its OAuth token."""
    if provider != "opus":
        yield [], ()
        return
    read, write = os.pipe()
    try:
        os.write(write, CLAUDE_TOKEN_FILE.read_bytes().strip())  # far below the pipe buffer
        os.close(write)
        yield ["--setenv", "CLAUDE_CODE_OAUTH_TOKEN_FILE_DESCRIPTOR", str(read)], (read,)
    finally:
        os.close(read)


def codex_token_hours_left(auth: Path) -> float | None:
    """Hours until the host Codex login's access token expires; None if it has no readable expiry."""
    try:
        payload = json.loads(auth.read_text(encoding="utf-8"))["tokens"]["access_token"].split(".")[1]
        claims = json.loads(base64.urlsafe_b64decode(payload + "=" * (-len(payload) % 4)))
        return (claims["exp"] - time.time()) / 3600
    except (OSError, ValueError, KeyError, IndexError, TypeError, AttributeError):
        return None


def preflight(provider: str) -> Path:
    if not shutil.which("bwrap"):
        raise SystemExit("bubblewrap (bwrap) is required")
    binary = CODEX_BIN if provider == "sol" else CLAUDE_BIN
    if not binary.is_file() or not os.access(binary, os.X_OK):
        raise SystemExit(f"provider executable not found: {binary}; set SEC_CODEX_BIN or SEC_CLAUDE_BIN")
    if provider == "sol" and (binary.parent.name != "bin" or binary.name != "codex"):
        raise SystemExit("SEC_CODEX_BIN must resolve to a standalone release's bin/codex")
    if provider == "opus":
        if not CLAUDE_TOKEN_FILE.is_file():
            raise SystemExit(f"no Claude token for the sandbox: run `claude setup-token` and save the token "
                             f"to {CLAUDE_TOKEN_FILE} (chmod 600), or set SEC_CLAUDE_OAUTH_TOKEN_FILE")
        if CLAUDE_TOKEN_FILE.stat().st_mode & 0o077:
            raise SystemExit(f"{CLAUDE_TOKEN_FILE} must be readable by its owner only (chmod 600)")
        return binary
    credential = HOME_HOST / ".codex/auth.json"
    if not credential.is_file():
        raise SystemExit(f"provider authentication file is missing: {credential}")
    # The Codex login is bound read-only. A sandboxed refresh would rotate its refresh token and
    # log the host out, so a run starts only while the access token outlasts any run.
    hours = codex_token_hours_left(credential)
    if hours is not None and hours < CODEX_TOKEN_MARGIN_HOURS:
        raise SystemExit(f"the host Codex login expires in {hours:.1f} h; renew it on the host (codex login) "
                         f"so a sandboxed run never has to")
    return binary


def prepared_metadata(run_dir: Path, provider: str) -> dict:
    """Verify the prepared bytes before launch and again inside the worker.

    Old extraction metadata already records all necessary hashes. Old revisions
    lacking the report hash must be prepared again; silently trusting them would
    leave a model-visible input outside the integrity check.
    """
    path = run_dir / "metadata.json"
    if not path.is_file():
        raise SystemExit(f"run is not prepared: {run_dir}")
    try:
        metadata = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError) as exc:
        raise SystemExit(f"cannot read prepared metadata: {exc}") from exc
    if not isinstance(metadata, dict):
        raise SystemExit("prepared metadata must be a JSON object")
    if metadata.get("provider") != provider:
        raise SystemExit(f"prepared provider is {metadata.get('provider')}, not {provider}")
    if metadata.get("instruction_name") != INSTRUCTION_NAME:
        raise SystemExit("prepared instruction name does not match the runner")
    if metadata.get("model") not in MODELS[provider] or metadata.get("effort") not in EFFORTS[provider]:
        raise SystemExit("prepared model or effort is not allowed for this provider; prepare a new run directory")
    timeout = metadata.get("timeout_seconds")
    if type(timeout) is not int or not TIMEOUT_BOUNDS[0] <= timeout <= TIMEOUT_BOUNDS[1]:
        raise SystemExit("prepared timeout_seconds is missing or out of range; prepare a new run directory")
    for key in ("filing_name", "expected_output"):
        name = metadata.get(key)
        if not isinstance(name, str) or name in ("", ".", "..") or Path(name).name != name:
            raise SystemExit(f"prepared {key} must be a bare filename")
    mode = metadata.get("mode", "extract")
    if mode not in ("extract", "revise"):
        raise SystemExit(f"unknown prepared mode: {mode}")
    filing = run_dir / "input" / "raw_filing" / metadata["filing_name"]
    if not filing.parent.is_dir() or list(filing.parent.iterdir()) != [filing]:
        raise SystemExit(f"expected exactly the recorded filing in {filing.parent}")
    inputs = [
        ("instruction_sha256", run_dir / "input" / INSTRUCTION_NAME),
        ("filing_sha256", filing),
        ("prompt_sha256", run_dir / "prompt.txt"),
    ]
    if mode == "revise":
        inputs += [
            ("revised_from_sha256", run_dir / "extraction" / metadata["expected_output"]),
            ("report_sha256", run_dir / "input" / REPORT_NAME),
        ]
    for key, source in inputs:
        if not metadata.get(key):
            raise SystemExit(f"prepared metadata lacks {key}; prepare a new run directory")
        try:
            actual = sha256(source)
        except OSError as exc:
            raise SystemExit(f"cannot read prepared input {source}: {exc}") from exc
        if actual != metadata[key]:
            raise SystemExit(f"prepared input hash mismatch: {source}; prepare a new run directory")
    return metadata


def run_worker(args: argparse.Namespace, state: Path, scratch: Path) -> int:
    run_dir = Path(args.run_dir).resolve()
    metadata = prepared_metadata(run_dir, args.provider)
    binary = preflight(args.provider)
    if args.provider == "sol":
        cache = HOME_HOST / ".codex/models_cache.json"
        if cache.is_file():
            shutil.copy2(cache, state / "models_cache.json")

    base = bwrap_base(run_dir, args.provider, state, scratch)
    started = time.monotonic()
    deadline = started + metadata["timeout_seconds"]
    started_at = now_iso()
    record = {
        "provider": args.provider,
        "provider_argv": None,
        "continuation_argv": [],
        **({"provider_environment": CLAUDE_ENV} if args.provider == "opus" else {}),
        "credential_delivery": ("long-lived OAuth token through an inherited pipe (CLAUDE_CODE_OAUTH_TOKEN_FILE_DESCRIPTOR); "
                                "not in argv, environment or run directory" if args.provider == "opus" else
                                "read-only external credential bind; no credential copied to run directory"),
        "runtime_state": "temporary directory outside the run; removed after worker completion",
        "provider_binary": str(binary),
        "provider_binary_sha256": sha256(binary),
        "runner_sha256": sha256(Path(__file__)),
        "started_at": started_at,
    }
    write_json(run_dir / "status.json", {
        "state": "running", "worker_pid": os.getpid(), "started_at": started_at,
    })

    stdout_path = run_dir / "events.jsonl"
    prompt = resume = None
    for attempt in range(MAX_CONTINUATIONS + 1):
        provider_argv = provider_command(run_dir, args.provider, metadata, prompt, resume)
        if attempt == 0:
            record["provider_argv"] = provider_argv
        else:
            record["continuation_argv"].append(provider_argv)
        write_json(run_dir / "command.json", record)
        with claude_token(args.provider) as (token_args, token_fds):
            exit_code, timed_out, client_pid = run_provider(
                base + token_args + provider_argv, run_dir, started_at, max(deadline - time.monotonic(), 1), token_fds)
        if timed_out or exit_code != 0 or args.provider != "opus" or not unfinished(run_dir, metadata):
            break
        # A clean end of turn with the deliverable still owed is a progress report, not a finished
        # task: resume the session and ask for the rest, a bounded number of times.
        results = claude_results(stdout_path)
        last = results[-1] if results else {}
        if (last.get("subtype") != "success" or last.get("is_error")
                or last.get("stop_reason") == "refusal" or not last.get("session_id")):
            break
        prompt, resume = continuation_prompt(metadata), last["session_id"]

    validation = validate_workbook(run_dir)
    write_json(run_dir / "validation.json", validation)
    summary = None
    if args.provider == "opus":
        write_json(run_dir / "provider-results.json", claude_results(stdout_path))
        summary = claude_summary(stdout_path) or {}
    ended_at = now_iso()
    if timed_out:
        outcome, failure_reason = "timed_out", "timeout"
    elif summary and summary["refusal"]:
        outcome, failure_reason = "failed", "provider_refusal"
    elif exit_code != 0:
        outcome, failure_reason = "failed", "provider_exit"
    elif summary is not None and (summary.get("subtype") != "success" or summary.get("is_error")):
        outcome, failure_reason = "failed", "provider_error"
    elif summary is not None and summary["served_models"] != [metadata["model"]]:
        outcome, failure_reason = "failed", "model_mismatch"
    elif not validation.get("exists"):
        outcome, failure_reason = "failed", "workbook_missing"
    elif not validation.get("valid_xlsx"):
        outcome, failure_reason = "failed", "workbook_unreadable"
    elif validation.get("sheet_names") != SHEETS:
        outcome, failure_reason = "failed", "workbook_incomplete"
    elif metadata.get("mode") == "revise" and not (run_dir / "extraction" / "revision_notes.md").is_file():
        outcome, failure_reason = "failed", "revision_notes_missing"
    else:
        outcome, failure_reason = "completed", None
    result = {
        # Completion is execution success only; no checker or substantive review runs here.
        "state": outcome,
        "failure_reason": failure_reason,
        "worker_pid": os.getpid(),
        "client_pid": client_pid,
        "started_at": started_at,
        "ended_at": ended_at,
        "elapsed_seconds": round(time.monotonic() - started, 3),
        "exit_code": exit_code,
        "timed_out": timed_out,
        "continuations": len(record["continuation_argv"]),
        "workbook_exists": validation.get("exists", False),
        "workbook_valid_xlsx": validation.get("valid_xlsx", False),
        "usage": run_usage(stdout_path),
    }
    if summary is not None:
        result["provider"] = summary
    write_json(run_dir / "status.json", result)
    return 0 if outcome == "completed" else 1


def run_provider(command: list[str], run_dir: Path, started_at: str, timeout: float,
                 pass_fds: tuple[int, ...] = ()) -> tuple[int, bool, int]:
    """Run one provider invocation, appending to the run's event and error logs."""
    with (run_dir / "events.jsonl").open("ab") as stdout, (run_dir / "stderr.log").open("ab") as stderr:
        proc = subprocess.Popen(
            command,
            stdin=subprocess.DEVNULL,
            stdout=stdout,
            stderr=stderr,
            start_new_session=True,
            close_fds=True,
            pass_fds=pass_fds,
        )
        write_json(run_dir / "status.json", {
            "state": "running", "worker_pid": os.getpid(), "client_pid": proc.pid,
            "started_at": started_at,
        })
        try:
            return proc.wait(timeout=timeout), False, proc.pid
        except subprocess.TimeoutExpired:
            os.killpg(proc.pid, signal.SIGTERM)
            try:
                return proc.wait(timeout=20), True, proc.pid
            except subprocess.TimeoutExpired:
                os.killpg(proc.pid, signal.SIGKILL)
                return proc.wait(), True, proc.pid


def worker(args: argparse.Namespace) -> int:
    path = Path(args.run_dir).resolve() / "status.json"
    if path.exists():
        raise SystemExit(f"refusing to overwrite prior run status: {path.parent}")
    try:
        with tempfile.TemporaryDirectory(prefix="sec-extraction-state-") as state, \
                tempfile.TemporaryDirectory(prefix="sec-extraction-scratch-") as scratch:
            return run_worker(args, Path(state), Path(scratch))
    except (Exception, SystemExit) as exc:
        # A detached worker must leave an outcome even if input verification or
        # provider startup fails before the normal completion record is written.
        previous = json.loads(path.read_text(encoding="utf-8")) if path.is_file() else {}
        write_json(path, {
            **previous, "state": "failed", "worker_pid": os.getpid(),
            "failure_reason": "worker_error", "error": f"{type(exc).__name__}: {exc}",
            "ended_at": now_iso(),
        })
        raise


def launch(args: argparse.Namespace) -> None:
    run_dir = Path(args.run_dir).resolve()
    prepared_metadata(run_dir, args.provider)
    preflight(args.provider)
    status_path = run_dir / "status.json"
    if status_path.exists():
        old = json.loads(status_path.read_text(encoding="utf-8"))
        if old.get("state") == "running":
            raise SystemExit(f"run already marked running: {run_dir}")
        raise SystemExit(f"refusing to overwrite prior run status: {run_dir}")
    launcher_log = (run_dir / "launcher.log").open("ab")
    argv = [sys.executable, str(Path(__file__).resolve()), "worker", "--provider", args.provider, "--run-dir", str(run_dir)]
    proc = subprocess.Popen(
        argv,
        stdin=subprocess.DEVNULL,
        stdout=launcher_log,
        stderr=launcher_log,
        start_new_session=True,
        close_fds=True,
        cwd=str(BASE),
    )
    launcher_log.close()
    write_json(run_dir / "launch.json", {"launcher_pid": proc.pid, "launched_at": now_iso(), "argv": argv})
    print(json.dumps({"run_dir": str(run_dir), "worker_pid": proc.pid, "provider": args.provider}, sort_keys=True))


def status(args: argparse.Namespace) -> None:
    for path in sorted(Path(args.runs_dir).glob("*")):
        if not path.is_dir():
            continue
        status_path = path / "status.json"
        payload = {"run": path.name, "state": "prepared"}
        if status_path.exists():
            payload.update(json.loads(status_path.read_text(encoding="utf-8")))
        print(json.dumps(payload, sort_keys=True))


def parser() -> argparse.ArgumentParser:
    root = argparse.ArgumentParser()
    sub = root.add_subparsers(dest="command", required=True)
    p = sub.add_parser("prepare")
    p.add_argument("--provider", choices=["sol", "opus"], required=True)
    p.add_argument("--run-dir", required=True)
    p.add_argument("--deal", required=True)
    p.add_argument("--filing", required=True)
    p.add_argument("--instruction", help="candidate instruction file to test (default: the working instruction)")
    p.add_argument("--model", help="model to run (default: the provider's first allowed model, e.g. claude-opus-5-5)")
    p.add_argument("--effort", help="effort level (Opus: low, medium, high, xhigh or max; default: DEFAULT_EFFORT)")
    p.add_argument("--timeout-minutes", type=int, help="wall-clock limit for the run, 10-360 (default: 90)")
    p.add_argument("--revise-from", help="finished workbook to revise (revision pass)")
    p.add_argument("--report", help="checker findings in plain text, required with --revise-from")
    p.set_defaults(func=prepare)
    p = sub.add_parser("launch")
    p.add_argument("--provider", choices=["sol", "opus"], required=True)
    p.add_argument("--run-dir", required=True)
    p.set_defaults(func=launch)
    p = sub.add_parser("worker")
    p.add_argument("--provider", choices=["sol", "opus"], required=True)
    p.add_argument("--run-dir", required=True)
    p.set_defaults(func=worker)
    p = sub.add_parser("status")
    p.add_argument("--runs-dir", type=Path, default=RUNS, help="run root to inspect (default: _dev/runs)")
    p.set_defaults(func=status)
    return root


if __name__ == "__main__":
    parsed = parser().parse_args()
    sys.exit(parsed.func(parsed))
