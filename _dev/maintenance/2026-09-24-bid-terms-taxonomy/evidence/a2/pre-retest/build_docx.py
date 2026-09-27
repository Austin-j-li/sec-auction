"""Build the replacement Questions for Alex (V114_SPEC D19, section 5) as a DOCX.

Reconciled with the v1.14.1 candidate on 26 September 2026 (PIPELINE_UPGRADE_SPEC
WP8.1 and V1141_SPEC section 10 step 8, in ../../../2026-09-26-v1141-streamline/).
The pre-v1.14.1 script, DOCX and renders are kept in pre-v1141/. The 3.3(a)
agreement numbers are still those of the 24 September draft trial: they are to be
recomputed after the v1.14.1 retest (see RETEST_PENDING below).

Usage (in a virtual environment with python-docx installed):
    python build_docx.py --out <path.docx> [--date YYYY-MM-DD] [--check-only]

Every filing excerpt in the case appendix is checked against the filing in
raw_filing/ before the document is written: each quoted fragment (fragments are
separated by " … ") must appear, after whitespace is normalized, in the text of
the printed page or pages the appendix cites. The build stops if one does not.
The filings are only read.
"""

import argparse
import datetime
import re
import sys
from html.parser import HTMLParser
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches

CHECKOUT = Path("/home/uctpiaj/work/Projects/sec-extraction")
FILINGS = {
    "mac-gray": "mac-gray_2013-12-04_DEFM14A.htm",
    "synacor": "synacor_2021-03-03_SCTO-T.htm",
    "penford": "penford_2014-12-29_DEFM14A.htm",
    "stec": "stec_2013-08-08_DEFM14A.htm",
    "meredith": "meredith_2021-11-08_DEFM14A.htm",
    "kraton": "kraton_2021-11-04_DEFM14A.htm",
    "pw": "providence-worcester_2016-09-20_DEFM14A.htm",
    "petsmart": "petsmart_2015-02-02_DEFM14A.htm",
    "datalink": "datalink_2016-11-29_DEFM14A.htm",
}
FILING_NAMES = {
    "mac-gray": "Mac-Gray, 4 December 2013",
    "synacor": "Synacor, Offer to Purchase (SC TO-T), 3 March 2021",
    "penford": "Penford, 29 December 2014",
    "stec": "sTec, 8 August 2013",
    "meredith": "Meredith, 8 November 2021",
    "kraton": "Kraton, 4 November 2021",
    "pw": "Providence and Worcester, 20 September 2016",
    "petsmart": "PetSmart, 2 February 2015",
    "datalink": "Datalink, 29 November 2016",
}

FONT = "Georgia"
INK = RGBColor(0x1F, 0x2A, 0x37)
NAVY = RGBColor(0x1F, 0x4E, 0x79)
GREY = RGBColor(0x5A, 0x64, 0x72)
HEAD_FILL = "E4EBF2"
ANSWER_FILL = "F4F6F8"
ELL = " … "

# ---------------------------------------------------------------------------
# Case appendix. Each part: (pages the fragment is on, page label shown, text).
# Text is quoted exactly from the filing; " … " marks an omission.
# ---------------------------------------------------------------------------

CASES = [
    ("C1", "Mac-Gray, Parties A and B, 2013", "mac-gray", [
        ([35], "p. 35; 9 September", "Party B's revised indication of interest did not include a firm financing commitment."),
        ([35], "p. 35; 10 September", "On September 10, 2013, Party A submitted a revised indication of interest with an all-cash purchase price of $18.00 to $19.00 per share" + ELL
         + "Neither Party A's nor Party C's revised indication of interest included a firm financing commitment."),
        ([36], "p. 36", "Also on September 18, 2013, Party A reiterated in a telephone call to a representative of BofA Merrill Lynch that its previous indication of interest "
         "with an all-cash purchase price of $18.00 to $19.00 per share was its best and final offer."),
        ([36], "p. 36; 18 September", "Party B valued at $21.50 per share, including $19.00 of cash to be paid at closing and the remaining per share price to be paid in the form of options" + ELL
         + "Party B's proposal stated that it valued the Party B options at $2.50 of incremental value per share."),
        ([37], "p. 37; 19 September", "The Special Committee also noted that Party B would likely be financing the transaction with a combination of equity from affiliated funds and third party debt capital"),
    ]),
    ("C2", "Synacor, Company E, 2020", "synacor", [
        ([33], "p. 33", "On September 18, 2020, Company E provided a revised letter of intent to the Company" + ELL + "requesting that the Company enter into exclusivity with it until October 23, 2020."),
        ([34], "p. 34", "On September 23, 2020, the Company and Company E entered into Company E’s letter of intent, and began coordination to commence diligence"),
    ]),
    ("C3", "Penford, Ingredion, 2014", "penford", [
        ([25], "p. 25", "On August 10, 2014, Penford received a letter from Mr. Fortnum" + ELL + "providing an indicative valuation of $18.00 per share in cash"),
        ([27], "p. 27; 14 August", "Mr. Fortnum asked whether Penford would consider whether to enter into exclusive negotiations with Ingredion."),
    ]),
    ("C4", "sTec, WDC, 2013", "stec", [
        ([30], "p. 30", "On May 28, 2013, WDC submitted a written second-round indication of interest at a price per share of $9.15 in cash, with a mark-up of the merger agreement"),
        ([31], "p. 31; 31 May", "WDC discontinued their due diligence efforts with the company at that time."),
        ([32], "p. 32", "On June 10, 2013, WDC submitted a revised written indication of interest to acquire the company at a price range of $6.60 to $7.10 per share in cash." + ELL + "WDC was otherwise prepared to move forward on the transaction terms previously proposed on May 28, 2013."),
    ]),
    ("C5", "Meredith, Party B, 2021", "meredith", [
        ([67], "p. 67", "On April 20, 2021, Party B submitted a revised proposal for $2.33 billion in cash" + ELL + "The proposal was otherwise materially unchanged from its April 14, 2021 proposal"),
    ]),
    ("C6", "Mac-Gray, CSC/Pamplona, 2013", "mac-gray", [
        ([37, 38], "pp. 37–38; 21 to 23 September", "a $15 million \"reverse\" termination fee payable by CSC/Pamplona to Mac-Gray in the event regulatory clearance is not obtained" + ELL
         + "there would be no financing contingency"),
        ([39], "p. 39", "On October 5, 2013, Kirkland delivered to Goodwin Procter a revised draft of the Pamplona commitment letter which, among other things, proposed that, in addition to CSC, "
         "Pamplona would be responsible in the event CSC failed to consummate the merger"),
        ([39], "p. 39; 8 October", "which Kirkland proposed, at the direction of CSC/Pamplona, be capped at $50 million."),
        ([40, 41], "pp. 40–41; 11 October", "agreed to the $11 million Mac-Gray termination fee."),
    ]),
    ("C7", "Kraton, Parent, 2021", "kraton", [
        ([38], "p. 38; 8 September", "while Parent agreed that financing would not be a condition to consummation of the merger," + ELL + "although Parent required debt financing to consummate the merger"),
        ([39], "p. 39; Kraton's counsel, 20 September", "King & Spalding LLP provided a revised draft of the merger agreement" + ELL + "and a revised draft of the debt commitment letter."),
        ([40], "p. 40", "a reverse termination fee of $63 million if the Company terminated the Merger Agreement as a result of the failure of the Debt Financing being available to Parent."),
    ]),
    ("C8", "Kraton, Party K, 2021", "kraton", [
        ([35], "p. 35", "Party K also informed Kraton that it was only interested in acquiring Kraton’s chemical segment"),
        ([35], "p. 35; 6 July", "the proposals submitted by these parties would be further evaluated by the Board in the context of any updated proposals submitted with respect to the potential acquisition of all of Kraton."),
    ]),
    ("C9", "Synacor, reopening and Company E, 2020", "synacor", [
        ([34], "p. 34; Mr. Bhise is Synacor's chief executive", "on October 27, 2020, Mr. Bhise again reached out to Company H to evaluate its interest in a potential transaction with the Company"),
        ([34], "p. 34", "On December 14, 2020, the Company received a revised proposal from Company E" + ELL + "a tender offer at a purchase price of $2.00 per Share for 35% of the outstanding Shares."),
        ([35], "p. 35", "Company E advised the Company that it no longer had interest in its prior proposed transaction structure which included a purchase price of $2.00 per Share for 100% of the Company’s outstanding equity and equity awards."),
    ]),
    ("C10", "Synacor, Company B, 2018–2019", "synacor", [
        ([30], "p. 30", "In October 2018, the Company and a publicly-traded software company (“Company B”) entered into a mutual non-disclosure agreement" + ELL + "“merger of equals”" + ELL + "elected to terminate the process in March 2019."),
    ]),
    ("C11", "Providence and Worcester, G&W, 2016", "pw", [
        ([29], "p. 29", "Party B, G&W and another bidder (“Party D”) also provided mark-ups of the draft merger agreement"),
        ([29], "p. 29", "On July 26, 2016" + ELL + "G&W submitted a revised LOI, which increased its offer to $22.15 per share"),
        ([30], "p. 30; 3 August", "Hinckley Allen provided G&W with a revised merger agreement and voting agreement, reflecting changes from the marked-up documents included with G&W’s LOI."),
    ]),
    ("C12", "sTec, Company D, 2013", "stec", [
        ([30], "p. 30; 28 May", "Company D needed approximately two weeks of additional time to continue its ongoing due diligence review"),
        ([31], "p. 31; after 31 May", "they informed Company D’s representatives that timing for sTec’s process had been delayed, and that Company D had an opportunity to continue in the process."),
        ([32], "p. 32", "On June 5, 2013 representatives of Company D" + ELL + "was disengaging from the process."),
    ]),
    ("C13", "Mac-Gray, due dates, 2013", "mac-gray", [
        ([32], "p. 32", "Over the next two months" + ELL + "18 financial bidders (including Party B and Party C), entered into confidentiality agreements with Mac-Gray" + ELL + "a preliminary indication of interest, which they were instructed to submit by July 23, 2013."),
        ([33], "p. 33", "On July 24, 2013, Party B submitted a preliminary indication of interest" + ELL
         + "Also on July 24, 2013, representatives from Party C called BofA Merrill Lynch and presented an oral preliminary indication of interest"),
        ([33], "p. 33; 25 July", "representatives of BofA Merrill Lynch reviewed the preliminary indications of interest received from Party B, Party C and CSC/Pamplona"),
        ([35], "p. 35", "submit revised written proposals no later than September 9, 2013" + ELL + "On September 10, 2013, Party A submitted a revised indication of interest"),
        ([36], "p. 36; 11 September", "request final indications of interest by September 18, 2013"),
        ([37], "p. 37; 19 September", "The Special Committee then discussed each of the three revised proposals" + ELL
         + "The Special Committee then authorized the Transaction Committee to inform CSC/Pamplona that the Special Committee would be prepared to move into exclusivity at a price of $21.25 per share"),
    ]),
    ("C14", "PetSmart, due date, 2014", "petsmart", [
        ([23], "p. 23", "non-binding preliminary indications of interest would be due on October 30, 2014."),
        ([24], "p. 24", "On October 30, six of the potentially interested parties submitted indications of interest. From October 30 to November 2, 2014, representatives of J.P. Morgan spoke by telephone "
         "with all of the potentially interested parties" + ELL + "to hear the parties’ respective rationales for the price levels suggested in their indications" + ELL
         + "As a result of its discussions with J.P. Morgan, another bidder (which we refer to as “Bidder 2”), which had initially indicated a price of $78.00, increased its indication to a range of $81.00 to $84.00 per share."),
        ([24], "p. 24; 3 November", "The board determined to allow the four bidders that had indicated a price or range at or above $80.00 per share to proceed to the final round"),
    ]),
    ("C15", "Kraton, rounds, 2021", "kraton", [
        ([35], "p. 35; 6 July; the filing prints 2020 for 2021", "the Board would require updated indications of interest to be received by July 19, 2020 that ascribed a greater value to Kraton "
         "before making any determination regarding which potential acquirers, if any, should be invited into the second round of the process."),
        ([36], "p. 36; 20 July", "invite Party A, Party H, and Parent to the second round of the transaction process, and to provide these parties with additional financial and other due diligence materials"),
        ([37], "p. 37", "Following the Board meeting on August 11, 2021" + ELL + "J.P. Morgan provided a final bid procedures letter to Party A, Party H, and Parent "
         "that requested these bidders to submit" + ELL + "final indications of interest by September 15, 2021."),
        ([39], "p. 39", "On September 17, 2021, Party A provided a revised proposal to acquire Kraton for $40.50 per share."),
        ([39], "p. 39; 18 September", "the Board reviewed the updated proposals submitted by Parent, Party A and Party K."),
    ]),
    ("C16", "Datalink, rounds, 2016", "datalink", [
        ([27], "p. 27", "On January 29, 2016, our board convened to consider the proposal from Party A."),
        ([28], "p. 28", "Beginning June 6, Raymond James contacted 13 strategic acquirers (including Party A, Party B and Insight), 10 of which (including Party A, Party B and Insight) signed a non-disclosure agreement" + ELL
         + "Later in June, Raymond James also contacted 14 financial sponsors (including Party C), 13 of which (including Party C) signed a non-disclosure agreement"),
        ([29], "p. 29", "Four strategic parties (in addition to Party A, which submitted the proposal described above on March 29, 2016), including Insight and Party B, and five financial sponsors, including Party C, submitted initial indications of interest."),
        ([29], "p. 29; 27 July", "our board directed our management and the Raymond James team to continue to pursue a transaction with the five parties that submitted the highest-priced initial indications of interest"),
        ([29], "p. 29", "On August 16, 2016, Raymond James provided each of Insight, Party B and Party C with an instruction letter relating to their final proposals." + ELL
         + "submit its proposal by 5:00 p.m. Eastern time on August 30, 2016"),
        ([32], "p. 32", "On October 1, 2016, Raymond James contacted Party B and Party C by telephone, each of whom expressed interest on those calls in reengaging in the process."),
    ]),
    ("C17", "sTec, rounds, 2013", "stec", [
        ([30], "p. 30; 16 May", "After the meeting, at the direction of the board, BofA Merrill Lynch sent final round process letters and a draft merger agreement to WDC and Company D, requesting a response by May 28, 2013."),
        ([30], "p. 30; 29 May", "in response to sTec’s request for non-binding proposals on May 28, 2013" + ELL
         + "our board of directors directed BofA Merrill Lynch to request a “best and final” proposal from WDC and a “best and final” written proposal from Company D by May 30, 2013."),
        ([31], "p. 31; 30 and 31 May", "Our board of directors unanimously determined to move forward with WDC on their proposed terms" + ELL + "WDC was not prepared to move forward with a transaction with sTec at that time."),
    ]),
    ("C18", "sTec, Company H, 2013", "stec", [
        ([29], "p. 29; 15 May", "Company H submitted a written non-binding indication of interest" + ELL + "in the range of $5.00 – $5.75 per share in cash."),
        ([30], "p. 30", "representatives of BofA Merrill Lynch contacted Company H and indicated that the price range Company H had submitted was not sufficient to move them forward in the process. "
         "The Company H representative was told that Company H could submit a revised indication of interest"),
        ([30], "p. 30; 23 May", "Company H remained interested in a potential acquisition of the company but that Company H was not able to increase its indicated value range."),
    ]),
    ("C19", "Penford, Party A, 2014", "penford", [
        ([31], "p. 31", "On October 4, Deutsche Bank followed up with Party A" + ELL + "Party A indicated any offer would be below $17.50 - $18.00 per share in cash. "
         "Deutsche Bank encouraged Party A to submit a letter of interest that could be reviewed by the board of directors."),
        ([32], "p. 32", "On October 13, 2014, Deutsche Bank had a discussion with representatives of Party A, who indicated that its value range for a potential transaction had been reduced from "
         "$17.50 - $18.00 per share to $16.00 - $18.00 per share"),
        ([32], "p. 32", "Later on October 14, 2014, Party A provided a formal letter with its indication of interest" + ELL + "at a price of $16.00 per share"),
    ]),
]


# ---------------------------------------------------------------------------
# Filing text by printed page, for the excerpt check.
# ---------------------------------------------------------------------------

BLOCK = {"p", "div", "br", "tr", "li", "h1", "h2", "h3", "h4", "h5", "h6", "table", "td"}


class _Text(HTMLParser):
    def __init__(self):
        super().__init__()
        self.out = []

    def handle_starttag(self, tag, attrs):
        style = (dict(attrs).get("style") or "").lower().replace(" ", "")
        if tag == "hr" or "page-break-before:always" in style or "page-break-after:always" in style:
            self.out.append("\n<<<BREAK>>>\n")
        elif tag in BLOCK:
            self.out.append("\n")
        if tag == "td":
            self.out.append(" ")

    def handle_endtag(self, tag):
        if tag in BLOCK:
            self.out.append("\n")

    def handle_data(self, data):
        self.out.append(data)


def filing_pages(path):
    """Map each printed page label to its text, dropping running headers and the label."""
    parser = _Text()
    parser.feed(path.read_text(encoding="utf-8", errors="replace"))
    pages = {}
    for chunk in "".join(parser.out).replace("\xa0", " ").split("<<<BREAK>>>"):
        lines = [re.sub(r"\s+", " ", line).strip() for line in chunk.split("\n")]
        lines = [line for line in lines if line]
        if not lines:
            continue
        label = None
        for line in reversed(lines[-3:]):
            if re.fullmatch(r"\d{1,3}", line):
                label = int(line)
                break
        if label is None:
            continue
        body = [line for line in lines if line.lower() != "table of contents" and line != str(label)]
        if label in pages:
            pages[label] = None  # label printed twice (front matter, annexes): ambiguous
        else:
            pages[label] = " ".join(body)
    return pages


def norm(text):
    return re.sub(r"\s+", " ", text).strip()


def check_excerpts():
    cache, problems = {}, []
    for cid, _title, deal, parts in CASES:
        if deal not in cache:
            cache[deal] = filing_pages(CHECKOUT / "raw_filing" / FILINGS[deal])
        pages = cache[deal]
        for page_list, label, text in parts:
            if any(pages.get(p) is None for p in page_list):
                problems.append(f"{cid} {label}: page missing or its label is printed twice")
                continue
            joined = norm(" ".join(pages[p] for p in page_list))
            for fragment in text.split(ELL):
                if norm(fragment) not in joined:
                    problems.append(f"{cid} {label}: not found: {fragment[:80]!r}")
    return problems


# ---------------------------------------------------------------------------
# Document helpers
# ---------------------------------------------------------------------------


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=70, left=110, bottom=70, right=110):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:tcMar")):  # a merged cell is visited once per grid column
        tc_pr.remove(old)
    mar = OxmlElement("w:tcMar")
    for side, value in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tc_pr.append(mar)


def set_table_borders(table, color="A9B4C2", size=4):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:color"), color)
        borders.append(el)
    tbl_pr.append(borders)


TBL_PR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
                "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"]


def order_tbl_pr(table):
    """Put tblPr children in schema order, which Word expects."""
    tbl_pr = table._tbl.tblPr
    children = list(tbl_pr)
    rank = {qn(f"w:{name}"): idx for idx, name in enumerate(TBL_PR_ORDER)}
    for child in children:
        tbl_pr.remove(child)
    for child in sorted(children, key=lambda el: rank.get(el.tag, len(rank))):
        tbl_pr.append(child)


def fixed_layout(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False  # sets tblLayout to fixed
    for row in table.rows:
        for idx, width in enumerate(widths):
            if idx < len(row.cells):
                row.cells[idx].width = Inches(width)
    grid = table._tbl.tblGrid
    for idx, col in enumerate(grid.findall(qn("w:gridCol"))):
        if idx < len(widths):
            col.set(qn("w:w"), str(int(widths[idx] * 1440)))
    order_tbl_pr(table)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    tr_pr.append(el)


def keep_row_together(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    tr_pr.append(el)


def min_row_height(row, inches):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(inches * 1440)))
    el.set(qn("w:hRule"), "atLeast")
    tr_pr.append(el)


NBSP_PAGE = re.compile(r"\b(pp?)\. (?=\d)")


def keep_page_label(text):
    """A non-breaking space after "p." and "pp.", so a page number never starts a line."""
    return NBSP_PAGE.sub("\\1.\u00a0", text)


def add_runs(par, text, size=None, color=None, bold=False, italic=False, nbsp=True):
    """Add text with **bold** spans and {small reference} spans. Quoted filing text passes nbsp=False."""
    if nbsp:
        text = keep_page_label(text)
    for piece in re.split(r"(\*\*[^*]+\*\*|\{[^}]+\})", text):
        if not piece:
            continue
        if piece.startswith("**"):
            run = par.add_run(piece[2:-2])
            run.bold = True
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
        elif piece.startswith("{"):
            par.add_run(piece[1:-1], style="Reference")
        else:
            run = par.add_run(piece)
            run.bold = bold
            run.italic = italic
            if size:
                run.font.size = Pt(size)
            if color is not None:
                run.font.color.rgb = color
    return par


def para(doc_or_cell, text="", style=None, size=None, color=None, bold=False, italic=False,
         after=None, before=None, keep_next=False, align=None):
    par = doc_or_cell.add_paragraph(style=style)
    add_runs(par, text, size=size, color=color, bold=bold, italic=italic)
    fmt = par.paragraph_format
    if after is not None:
        fmt.space_after = Pt(after)
    if before is not None:
        fmt.space_before = Pt(before)
    if keep_next:
        fmt.keep_with_next = True
    if align is not None:
        par.alignment = align
    return par


def cell_text(cell, text, size=9, bold=False, color=None, first=True, after=2):
    par = cell.paragraphs[0] if first else cell.add_paragraph()
    add_runs(par, text, size=size, bold=bold, color=color)
    par.paragraph_format.space_after = Pt(after)
    par.paragraph_format.line_spacing = 1.08
    return par


def header_row(table, labels):
    row = table.rows[0]
    repeat_header(row)
    for cell, label in zip(row.cells, labels):
        set_cell_shading(cell, HEAD_FILL)
        cell_text(cell, label, size=9, bold=True, color=NAVY)


def grid_table(doc, labels, rows, widths, size=9):
    """A short table kept on one page: a shaded header row and text rows (each cell may hold several paragraphs)."""
    table = doc.add_table(rows=1, cols=len(labels))
    set_table_borders(table)
    header_row(table, labels)
    for values in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            chunks = value if isinstance(value, list) else [value]
            for idx, chunk in enumerate(chunks):
                cell_text(cell, chunk, size=size, first=(idx == 0))
    for row in table.rows:
        keep_row_together(row)
        for cell in row.cells:
            set_cell_margins(cell)
    for row in table.rows[:-1]:  # a short table stays on one page
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
    fixed_layout(table, widths)
    spacer(doc)
    return table


def spacer(doc, points=6):
    """A thin empty paragraph: keeps adjacent tables apart without a full blank line."""
    fmt = doc.add_paragraph().paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = Pt(points)


def answer_box(doc, prompt, lines=2):
    """A shaded one-cell box with the answer prompt and room to write."""
    table = doc.add_table(rows=1, cols=1)
    set_table_borders(table)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, ANSWER_FILL)
    set_cell_margins(cell, top=90, bottom=90)
    cell_text(cell, prompt, size=9.5, after=0)
    min_row_height(table.rows[0], 0.35 + 0.22 * lines)
    keep_row_together(table.rows[0])
    fixed_layout(table, [6.7])
    spacer(doc)


def heading(doc, text, level):
    par = doc.add_heading(level=level)
    add_runs(par, text)
    par.paragraph_format.keep_with_next = True
    return par


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, before, after in (("Title", 20, 0, 3), ("Heading 1", 14, 14, 5),
                                      ("Heading 2", 11.5, 10, 3), ("Heading 3", 10.5, 8, 2)):
        style = styles[name]
        style.font.name = FONT
        style.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.italic = False
        style.font.color.rgb = NAVY
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    title_ppr = styles["Title"].element.get_or_add_pPr()
    for border in title_ppr.findall(qn("w:pBdr")):
        title_ppr.remove(border)
    reference = styles.add_style("Reference", 2)  # character style for Austin's reference numbers
    reference.font.size = Pt(7.5)
    reference.font.color.rgb = GREY
    for name in ("Byline", "Excerpt", "Small"):
        style = styles.add_style(name, 1)
        style.base_style = normal
    styles["Byline"].font.color.rgb = GREY
    styles["Byline"].font.size = Pt(11)
    styles["Excerpt"].font.size = Pt(9)
    styles["Excerpt"].paragraph_format.left_indent = Inches(0.25)
    styles["Excerpt"].paragraph_format.space_after = Pt(3)
    styles["Small"].font.size = Pt(9)
    styles["Small"].font.color.rgb = GREY
    bullet = styles["List Bullet"]
    bullet.font.name = FONT
    bullet.font.size = Pt(10.5)
    bullet.paragraph_format.space_after = Pt(3)


def page_number_footer(section):
    par = section.footer.paragraphs[0]
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = par.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)
    run.font.size = Pt(8.5)
    run.font.color.rgb = GREY


# ---------------------------------------------------------------------------
# Content
# ---------------------------------------------------------------------------

DECIDED = [
    ("**Contingent payments.** {D4} A CVR or earnout is part of the price, in its own columns (Price low and high are upfront only), and alone neither makes a bid Heavy nor prevents None.",
     "Mac-Gray, Party B, 18 September 2013: $19.00 cash plus options Party B valued at $2.50, so upfront $19.00 and CVR/earnout $2.50 (p. 36; C1).",
     "**Yes.** The draft Part A listed “a contingent part of the price” among the conditions."),
    ("**Exclusivity.** {D14, R4} Never changes Formality or the Conditions level. Only a period of two weeks or more that the filing ties to remaining diligence alone, and does not call confirmatory, counts toward Heavy; a period that also covers exclusivity or negotiation does not.",
     "Synacor, Company E, 18 September 2020: about five weeks' exclusivity; not Heavy on that, normally Unclear (pp. 33–34; C2).",
     "**Yes.** Question 6(a): A (Company E Heavy) was recommended; B follows your note that exclusivity should not make a formal bid informal."),
    ("**A later exclusivity request.** {D15} Made with no other change, it is its own Exclusivity changed row and leaves the earlier bid unchanged; made with a bid, it is coded on that bid's row only.",
     "Penford, Ingredion: the 14 August 2014 request is its own row; the 10 August bid is unchanged (pp. 25–27; C3).",
     "**Yes.** The rules you confirmed also coded it on the standing bid."),
    ("**Same offer.** {D1, R1} When a bidder says its earlier offer stands (reiterates, confirms, repeats or holds it), the row copies that bid and changes only what the filing says changed: the date and round always, Formality by its own route, and any condition reported anew; the Note reads “Same as #n”. No other row copies an earlier one: a revision is coded from its own communication.",
     "Mac-Gray, Party A, 18 September 2013: reiterates its 10 September offer, which had no firm financing commitment, as best and final; the copy is Heavy on financing and Formal, as it answers the final request (pp. 35–36; C1). "
     "A revision with a new price that is “otherwise” unchanged is not copied: sTec, WDC, 10 June 2013 (p. 32; C4); Meredith, Party B, 20 April 2021 (p. 67; C5).",
     "**Yes.** The rules you confirmed copied all earlier values when the filing said the other terms were unchanged; now only a bid the bidder says stands is copied."),
    ("**Same-price commitment changes.** {D13, H4} A same-price change to the bidder's commitments (reverse fee, sponsor guarantee or damages cap, financing or closing conditions) is a Bid row with the price blank, and not a new price observation; the target's termination fee goes in a dated Note.",
     "Mac-Gray, CSC/Pamplona at $21.25: the 21–23 September 2013 package and the proposals of 5 and 8 October are Bid rows; the target's $11 million fee is a Note (pp. 37–41; C6).",
     "No."),
]

PROVISIONAL = [
    ("**“No financing condition”.** {D5} A bid so stated is Committed, even with an unsigned or draft debt commitment or only highly confident lender support; the Note records the lender documents and any reverse fee.",
     "Kraton, Parent, September 2021: draft debt commitment letter; $63 million reverse fee. Committed (pp. 38–40; C7).",
     "No. Consistent with option B of question 6(b), which was recommended."),
    ("**Partial-company bidders.** {D7} Recorded, but outside whole-company live counts and the auction screen; a bidder switching to a partial offer leaves at the switch, recorded as Withdrew, “continued on a partial basis”.",
     "Kraton, Party K (chemical segment only) (p. 35; C8). Synacor, Company E: Withdrew on 14 December 2020, continued on a partial basis (pp. 34–35; C9).",
     "**Yes.** Question 2: C was recommended; B fits your remark that such bids “do not compete against bids to sell the entire company”."),
    ("**Rounds.** {D8; v1.14.1 D6} After round 1, a round opens only when the target (a) selects who advances and asks them for new offers; (b) first asks for final, binding or best-and-final offers, even from unchanged bidders; (c) with no final round yet, starts definitive negotiation with selected bidders; or (d) asks for offers again after a pause of 30 days or more with no request, or after an exclusivity period with one bidder ended. Asking the same bidders to improve continues a round. Finality is the procedure the target announced.",
     "Reopening under (d): Synacor, 27 October 2020, after Company E's exclusivity ended (pp. 33–34; C2, C9); Datalink, 1 October 2016, after Insight's exclusivity lapsed (p. 32; C16). Maps: 3.1.",
     "**Yes, in part.** Question 3: A, unchanged in substance; the triggers are now a closed list, and “after a suspension” became the 30-day or ended-exclusivity test."),
    ("**Merger-of-equals talks.** {D17} No control or premium test; the counterparty stays outside the counts unless the filing reports that the target is being sold to it, and the talks are dated events.",
     "Synacor, Company B, October 2018 to March 2019 (p. 30; C10).",
     "**Yes.** Question 4: C was recommended and is withdrawn; this is neither A (no rows) nor B (a sale attempt with Company B as a bidder)."),
    ("**Price-only revision: Formality.** {D9; v1.14.1 D1} Formal only if it qualifies on its own (a markup or definitive document with it, or an answer to a final solicitation); a reference back to earlier Formal terms, the target negotiating from an earlier markup, or later document work, is not enough.",
     "Providence and Worcester, G&W, 26 July 2016: no route of its own, so Informal (pp. 29–30; C11). "
     "sTec, WDC, 10 June 2013: refers back to its 28 May terms, but has no route of its own, so Informal (p. 32; C4). You coded G&W Formal, WDC Informal.",
     "**Yes.** Question 5: B was recommended; it also made G&W Formal, as the target negotiated from its markup. WDC's 10 June bid now matches your coding."),
    ("**Price-only revision: conditions.** {question 5, second half; R2} Coded from what the filing says about the revision, from its communication up to that bidder's next bid, exclusivity, exit or signing, forecasts included; silence is Not stated; nothing is copied from an earlier row unless the bidder says its offer stands (2a).",
     "Mac-Gray, Party B, 18 September 2013: the package is silent on financing, but the committee noted on 19 September that Party B “would likely be financing” with third party debt, so Financing Contingent and Heavy (pp. 35–37; C1).",
     "**Yes, in part.** As option B, except that what the filing says after the bid, up to the bidder's next row, now counts."),
    ("**Missed due dates.** {D10} A bidder that misses a due date but carries on (asks for time, is invited to continue, or bids later) gets no exit and no re-entry; the miss is its own dated event. “Did not submit” is used only where participation ends.",
     "sTec, Company D: missed the 28 and 30 May 2013 due dates, asked for time, was invited to continue; live until it withdrew on 5 June (pp. 30–32; C12).",
     "No. As reading 2 of part 2."),
    ("**Deadlines.** {D11} The first that fits: Extended, if a later due date was set for any bidder; Extended (late bid accepted), if the target considered a required response that arrived after the date; Enforced, if after the date the target acted on the bids in hand (evaluated, selected or gave feedback); Passed without action, if bidding continued with no reported step on them; Unclear, if the filing reports nothing after the date.",
     "Mac-Gray: 23 July and 9 September 2013 Extended (late bid accepted), as day-late responses were considered; 18 September Enforced, as the committee acted on the three proposals on 19 September (pp. 32–37; C13). "
     "PetSmart, 30 October 2014: Enforced; the bank heard the parties' rationales and the board selected four bidders on 3 November. Bidder 2's raise revised an on-time bid, so it is not a late required response (pp. 23–24; C14). "
     "sTec, 30 May 2013: Enforced; your coding differs (3.1).",
     "**Yes.** Question 7: A was recommended; now any accepted late required response is an extension (your convention, as Austin reported it), and two outcomes are added. A's prediction (both Enforced) holds."),
]

TWO_B_WIDTHS = [0.3, 2.55, 1.95, 1.35, 0.55]

# 3.3(a): these agreement counts come from the 24 September draft trial. The v1.14.1
# retest has not been run; recompute them from its workbooks (PIPELINE_UPGRADE_SPEC
# WP3 item 8) and then drop this marker from the document.
RETEST_PENDING = "to be recomputed after the v1.14.1 retest"


def provisional_table(doc):
    """2b: one row per item (rule, case, changed, Yes/No to circle), each followed by a shaded comment row.
    Every paragraph of an item row keeps with the next, so an item never parts from its comment row."""
    table = doc.add_table(rows=1, cols=5)
    set_table_borders(table)
    header_row(table, ["#", "Rule", "Case", "Changed?", "Agree?"])
    comment_rows = []
    for idx, (rule, case, changed) in enumerate(PROVISIONAL, start=1):
        cells = table.add_row().cells
        cell_text(cells[0], str(idx))
        cell_text(cells[1], rule)
        cell_text(cells[2], case)
        cell_text(cells[3], changed)
        for pos, answer in enumerate(("**Yes**", "", "**No**")):
            cell_text(cells[4], answer, first=(pos == 0)).alignment = WD_ALIGN_PARAGRAPH.CENTER
        for cell in cells:
            for par in cell.paragraphs:
                par.paragraph_format.keep_with_next = True
        comment = table.add_row()
        merged = comment.cells[0].merge(comment.cells[4])
        set_cell_shading(merged, ANSWER_FILL)
        cell_text(merged, "Comment:", size=9, after=0)
        min_row_height(comment, 0.5)
        comment_rows.append(comment)
    for row in table.rows:
        keep_row_together(row)
        for cell in row.cells:
            set_cell_margins(cell)
    fixed_layout(table, TWO_B_WIDTHS)
    for row in comment_rows:  # a merged cell spans the full width
        row.cells[0].width = Inches(sum(TWO_B_WIDTHS))
    spacer(doc)


def build(out_path, date):
    doc = Document()
    configure_styles(doc)
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    for side in ("left_margin", "right_margin"):
        setattr(section, side, Inches(0.9))
    section.top_margin, section.bottom_margin = Inches(0.8), Inches(0.8)
    page_number_footer(section)
    core = doc.core_properties
    core.title = "Questions on coding conventions, v1.14.1"
    core.author = "Austin Li"
    core.comments = "Replaces Open_Questions_for_Alex_2026-09-24.docx"
    core.last_modified_by = "Austin Li"
    core.revision = 1
    stamp = datetime.datetime.combine(date, datetime.time(12, 0))
    core.created = core.modified = stamp

    day = f"{date.day} {date:%B %Y}"
    title = doc.add_paragraph(style="Title")
    add_runs(title, "Questions on coding conventions, v1.14.1")
    para(doc, f"For Alex Gorbenko · {day}", style="Byline", after=12)

    heading(doc, "1. What changed", 1)
    para(doc, "The bid-terms columns you confirmed on 24 September are in v1.14.1 of the extraction instruction; this document replaces that day's questions. "
              "Austin has since decided some points and adopted provisional answers to others in v1.14.1, a shorter version that gives each recurring case a default; a later version changes any you reject. "
              "Some answers, decided and provisional, depart from a recommendation you saw or from a rule or wording you confirmed; each is marked **Yes** under “Changed”.")
    para(doc, "**How to answer.** Circle Yes, No or an option, or write a comment. "
              "Question numbers are those of 24 September; pages are the filings' printed pages; C1 to C19 are the excerpts in part 4; "
              "small grey labels are Austin's references.", after=4)

    heading(doc, "2a. Decided, for information (comments welcome)", 1)
    grid_table(doc, ["Rule", "Case", "Changed?"], [list(row) for row in DECIDED], [2.75, 2.45, 1.5])
    answer_box(doc, "Comments on 2a:", lines=2)

    heading(doc, "2b. Provisional: yes or no", 1)
    provisional_table(doc)
    para(doc, "As confirmed on 24 September, a regulatory concern rules out None without by itself setting the level; it needs no question.", style="Small")

    heading(doc, "3. Open decisions", 1)

    heading(doc, "3.1 Round maps under the listed triggers {Decision 1}", 2)
    para(doc, "A round opens only at one of the listed target acts (2b, item 3). Where the target selects who advances, asks for nothing, and later sends its first request for final offers, "
              "the maps differ (Kraton, Datalink): under A the selection and the letter are one round, since the selection asked for no offers and a round is a stage in which bidders are asked for offers; "
              "under B the selection is a round of its own and the letter opens another. "
              "For sTec, the question is whether 16 May began the final round. Maps come from a re-check of the current ledgers, updated for v1.14.1.")
    heading(doc, "Kraton (pp. 33–39; C15)", 3)
    grid_table(doc, ["Round", "A. Selection and final letter are one round: two rounds (recommended)", "B. Selection and final letter each open a round: three rounds (current ledger)"], [
        ["1", "24 May to 20 July 2021: outreach; due 29 June, then 19 July after the 6 July request to improve; Not final", "Same"],
        ["2", "20 July admission of Parties A, H and Parent; diligence; final bid procedures letter, due 15 September; Announced as final; "
              "Extended (late bid accepted), as Party A's 17 September bid was considered", "20 July admission and diligence; no request for bids; Not final"],
        ["3", "—", "Final bid procedures letter after 11 August; due 15 September; Announced as final"],
    ], [0.6, 3.05, 3.05])
    para(doc, "A is recommended because the 20 July selection asked for no offers, so under B round 2 would hold no request, due date or bid. "
              "Your voice note (item 3) starts round 2 on 6 July, making the formal stage round 3; v1.14.1 keeps 6 July in round 1, as asking the same bidders to improve continues a round.", style="Small")
    heading(doc, "Datalink (pp. 27–32; C16)", 3)
    grid_table(doc, ["Round", "A. Selection and final letter are one round: four rounds", "B. Selection and final letter each open a round: five rounds (standing map)"], [
        ["1", "January to May 2016: bilateral talks with Party A, from 29 January", "Same"],
        ["2", "June outreach, after more than 30 days with no request; indications due 18 and 21 July", "Same"],
        ["3", "27 July selection of five bidders, diligence, then the 16 August letter; due 30 August; Announced as final", "27 July selection and diligence only; no request; Not final"],
        ["4", "1 October re-approach to Parties B and C after Insight's exclusivity lapsed", "16 August letter; due 30 August; Announced as final"],
        ["5", "—", "1 October re-approach"],
    ], [0.6, 3.05, 3.05])
    para(doc, "Map A is what the listed triggers give, on the same reading as Kraton's A. Austin's ruling of 21 September (five rounds, January's talks with Party A as round 1) stands until he decides otherwise. "
              "Your notes start round 1 at the bank's first contacts (June); both maps keep Austin's round 1.", style="Small")
    heading(doc, "sTec (pp. 27–32; C4, C17)", 3)
    grid_table(doc, ["Round", "Map A: 16 May began the final round (recommended)", "Map C: 16 May not final (your voice note)"], [
        ["1", "April 2013 outreach; indications due 3 May; Not final", "Same"],
        ["2", "From 15 or 16 May: letters to WDC and Company D; due 28 May (Extended), then best and final by 30 May (Enforced: no later due date was set); Announced as final; "
              "through WDC's withdrawal (31 May), return (10 June) and $6.85 bid (14 June) to signing", "16 May letters, due 28 May; Enforced; Not final"],
        ["3", "—", "From the 29 May request for best and final proposals, due 30 May; Enforced; Announced as final; to signing"],
    ], [0.6, 3.05, 3.05])
    para(doc, "Map A follows the target's “final round process letters”; map C, the minutes' “request for non-binding proposals” (both p. 30) and your voice note "
              "(item 6: “this is not the final round”). "
              "Your spring coding agrees with map A on the round count but extends the round to 10 June, where map A has 30 May Enforced (pp. 31–32). "
              "Map B, the current ledger, follows map A to 30 May, then opens round 3 after WDC's withdrawal on 31 May; "
              "no listed trigger supports it, as it came days, not 30 days, after the 30 May request, and no exclusivity period had ended.", style="Small", keep_next=True)
    answer_box(doc, "Kraton and Datalink, circle one:        **A**        **B**        Comment:\n"
                    "sTec, circle one:        **Map A**        **Map B**        **Map C**        Comment:", lines=3)

    heading(doc, "3.2 How counts are used in estimation (question 1) {Decision 3, R3}", 2)
    para(doc, "**For information: the ledger no longer records ranges.** The unnamed members of a reported total with no reported offer are one row, "
              "recorded as not submitting by the first due date after they appear and marked inferred; its count is the total less the members recorded by name, "
              "with the arithmetic in one Note sentence. A named party that may belong to the total stays inside it. A qualified figure (“more than ten”) keeps its qualifier and leaves the count blank. "
              "Datalink: 13 signers did not bid, that is 23 signers less Parties A, B and C, Insight and the six other bidders (pp. 28–29; C16). "
              "Mac-Gray: the 16 unnamed financial signers, by 23 July (pp. 32–33; C13). How should estimation use these counts?")
    para(doc, "**A.** As recorded, the same as reported counts.\n"
              "**B.** As recorded, with a robustness check that sets aside the counts marked inferred; closest to the option recommended on 24 September.\n"
              "**C.** Another way (say which).", after=4, keep_next=True)
    answer_box(doc, "Circle one:        **A**        **B**        **C**        Comment:")

    heading(doc, "3.3 How analysis reads the ledger {Decision 3b}", 2)
    para(doc, "Analysis computes every variant below and picks none.", after=4)
    para(doc, "**(a) Which Formality reading is primary?** Agreement with your labels on trial extractions made with the 24 September draft "
              f"(a check, not a tuning target; {RETEST_PENDING}):", after=4, keep_next=True)
    grid_table(doc, ["Reading", "Bid counts as formal when", "Mac-Gray: agree, of 13", "Providence and Worcester: agree, of 14"], [
        ["T0", "Formality as recorded", "13", "11"],
        ["T1", "Formal and not Heavy (Unclear counts as not Heavy)", "11", "14"],
        ["T1u", "As T1, but Unclear counts as Heavy", "11", "13"],
        ["T2", "Formal, with a single price, not a range", "12", "11"],
        ["T3", "Formal, in a round announced as final or inferred final", "13", "11"],
    ], [0.7, 3.2, 1.3, 1.5])
    para(doc, "Your spring coding matches T0 (and T3) fully on Mac-Gray and T1 fully on Providence and Worcester; no reading matches both. "
              "The trial ledgers have Formal with Heavy conditions where you coded Formal (Mac-Gray, Parties A and B, 18 September 2013; Party A's copies its earlier uncommitted financing (2a), "
              "and Party B's Financing Contingent comes from the committee's 19 September remark, which 2b item 6 now uses) and where you coded Informal (Providence and Worcester, Party D on 20 July and 1 August 2016 "
              "and G&W on 21 July; diligence periods of three weeks or more). "
              "Under v1.14.1 a bid with committed financing and nothing said about diligence is Unclear rather than Light, which can move bids between T1 and T1u.", style="Small", keep_next=True)
    answer_box(doc, "(a) Circle one:        **T0**        **T1**        **T1u**        **T2**        **T3**        Comment:")
    para(doc, "**(b) Settled, for information: same-price commitment changes are not new price observations.** "
              "They stay Bid rows with the price blank (2a; C6), so the price is observed once. Austin decided this; it was option B.", after=6)
    para(doc, "**(c) Do inferred exits enter as dropouts or as censoring?** The ledger infers an unreported exit at the first point that applies and marks it as inferred. "
              "For example, Mac-Gray's 16 unnamed financial signers with no reported offer are one row, recorded as not submitting by 23 July, count 16 (3.2; C13). "
              "\n**A.** Dropout, as for a reported exit.\n"
              "**B.** Censoring: nothing is inferred about the bidder's value.", after=4, keep_next=True)
    answer_box(doc, "(c) Circle one:        **A** dropout        **B** censoring        Comment:")

    heading(doc, "3.4 Readings from part 2", 2)
    para(doc, "**sTec, Company H** (pp. 29–30; C18). **Settled, for information.** Austin decided that the target dropped H: it was told its range was not enough to advance, "
              "and the final-round letters of 16 May went to WDC and Company D only, so H is recorded as dropped by the target by 16 May, "
              "with the reason “would not improve earlier offer” from its 23 May reply. Your coding has H's own drop.", after=6)
    para(doc, "**Penford, Party A** (pp. 31–32; C19). Proposed: 4 October (any offer would be below $17.50–$18.00) and 13 October (a reduced value range told to the bank) are valuation statements, not bids; "
              "the 14 October letter at $16.00 is its bid. The ledger raises no Question on it. This is what was proposed on 24 September.", keep_next=True)
    answer_box(doc, "Agree?        **Yes**        **No**        Comment:")

    heading(doc, "3.5 Which source governs", 2)
    para(doc, "When your spring hand coding and August voice notes disagree, which governs? They differ on whether sTec's 16 May began the final round (3.1); "
              "on Providence and Worcester the voice note says Party A “never made it out of round one”, while the hand coding has it drop out on 22 July 2016.")
    para(doc, "**A.** The voice notes.\n**B.** The hand coding.\n**C.** As so far: a general rule stated in the notes governs; a remark about one deal's facts is a reading to confirm with you.", after=4, keep_next=True)
    answer_box(doc, "Circle one:        **A**        **B**        **C**        Comment:")

    heading(doc, "4. Case appendix", 1)
    para(doc, "Excerpts are quoted exactly; “…” marks an omission; a date after a page number is the event date where the excerpt lacks it. "
              "Filings (DEFM14A unless stated): " + "; ".join(FILING_NAMES[deal] for deal in FILINGS) + ".", style="Small")
    for cid, title_text, _deal, parts in CASES:
        heading(doc, f"{cid}. {title_text}", 3)
        for idx, (_pages, label, text) in enumerate(parts):
            par = doc.add_paragraph(style="Excerpt")
            add_runs(par, f"“{text}”", nbsp=False)
            run = par.add_run(f"  ({keep_page_label(label)})")
            run.font.color.rgb = GREY
            par.paragraph_format.keep_with_next = idx < len(parts) - 1  # keep a case's excerpts on one page

    doc.save(out_path)


def paragraph_text(par):
    """A paragraph's text, with a line break or tab read as a space."""
    parts = []
    for el in par.iter(qn("w:t"), qn("w:br"), qn("w:tab")):
        parts.append(el.text or "" if el.tag == qn("w:t") else " ")
    return "".join(parts)


def word_count(path):
    """Words in every paragraph of the body, tables included (a merged cell is one w:tc, so it counts once).
    Returns (whole document, parts 1-3, case appendix); parts 1-3 include the title and byline."""
    body = Document(path).element.body
    total = before = 0
    in_appendix = False
    for par in body.iter(qn("w:p")):
        text = paragraph_text(par)
        if text.strip() == "4. Case appendix":
            in_appendix = True
        words = len(text.split())
        total += words
        if not in_appendix:
            before += words
    return total, before, total - before


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--out", type=Path)
    parser.add_argument("--date", help="YYYY-MM-DD for the byline (default: today, UTC)")
    parser.add_argument("--check-only", action="store_true")
    args = parser.parse_args()
    problems = check_excerpts()
    if problems:
        print("\n".join(problems))
        sys.exit(1)
    print(f"excerpts checked: {sum(len(parts) for *_x, parts in CASES)} parts in {len(CASES)} cases, all found on the cited pages")
    if args.check_only:
        return
    if args.out is None:
        parser.error("--out is required unless --check-only")
    date = (datetime.date.fromisoformat(args.date) if args.date
            else datetime.datetime.now(datetime.timezone.utc).date())
    build(args.out, date)
    total, parts_1_3, appendix = word_count(args.out)
    print(f"wrote {args.out}; words: {total} ({parts_1_3} in parts 1-3, {appendix} in the case appendix)")


if __name__ == "__main__":
    main()
