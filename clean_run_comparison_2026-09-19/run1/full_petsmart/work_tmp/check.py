# -*- coding: utf-8 -*-
"""Mechanical checks on the built workbook + source/quote verification."""
import json, re, sys
from datetime import datetime
from openpyxl import load_workbook

WB = '/home/uctpiaj/work/extraction/petsmart.xlsx'
DATA = json.load(open('/home/uctpiaj/work/work_tmp/pages.json'))

def norm(t):
    if t is None: return ''
    t = str(t).replace('\xa0', ' ').replace('\u200b', '').replace('\u2009', ' ').replace('\u202f', ' ')
    return re.sub(r'\s+', ' ', t).strip()

# full normalized filing text for quote checking
FULL = []
for d in DATA:
    for t in d['paras']:
        nt = norm(t)
        if nt: FULL.append(nt)
FULLTEXT = ' '.join(FULL)
FULLTEXT_NQ = FULLTEXT.replace('“','').replace('”','').replace('"','')

wb = load_workbook(WB)
ws = wb['Deal ledger']
HEADERS = [c.value for c in ws[1]]
H = {name: i + 1 for i, name in enumerate(HEADERS)}
rows = []
for r in ws.iter_rows(min_row=2, values_only=False):
    rows.append(r)

def val(row, name):
    return row[H[name] - 1].value

problems = []
results = {}

# 1. ids / order
ids = [val(r, 'Row id') for r in rows]
if len(ids) != len(set(ids)): problems.append(('ids', 'duplicate row ids'))
nums = [val(r, '#') for r in rows]
if nums != sorted(nums): problems.append(('order', '# not ascending'))
results['unique_row_ids'] = len(ids) == len(set(ids))

# working-date order and bounds
order_ok = True
prev = None
for r in rows:
    wd = val(r, 'Working date'); num = val(r, '#')
    if wd is None:
        problems.append(('working', f"{val(r,'Row id')} missing working date")); order_ok = False; continue
    if prev and wd < prev:
        problems.append(('order', f"{val(r,'Row id')} working {wd} < previous {prev}")); order_ok = False
    prev = wd
    df, dt = val(r, 'Date from'), val(r, 'Date to')
    if df and wd < df: problems.append(('bounds', f"{val(r,'Row id')} working before from"))
    if dt and wd > dt: problems.append(('bounds', f"{val(r,'Row id')} working after to"))
    if df and dt and df > dt: problems.append(('bounds', f"{val(r,'Row id')} from > to"))
results['working_order_and_bounds'] = order_ok and not [p for p in problems if p[0] in ('order','bounds','working')]

# basis/method pairs
pair_ok = True
BASES = {"Reported day","Inferred day","Reported interval","Approximate window","Relative only","Undated"}
METHODS = {"Reported","Inferred","Assigned: deadline","Assigned: decision day","Assigned: midpoint","Assigned: bound","Assigned: sequence"}
for r in rows:
    b, m = val(r, 'Date basis'), val(r, 'Date method')
    if b is None or m is None:
        problems.append(('pair', f"{val(r,'Row id')} missing basis/method")); pair_ok = False; continue
    if b not in BASES or m not in METHODS:
        problems.append(('pair', f"{val(r,'Row id')} bad basis/method {b}/{m}")); pair_ok = False
    if (b == 'Reported day') != (m == 'Reported'):
        problems.append(('pair', f"{val(r,'Row id')} Reported mismatch {b}/{m}")); pair_ok = False
    if (b == 'Inferred day') != (m == 'Inferred'):
        problems.append(('pair', f"{val(r,'Row id')} Inferred mismatch {b}/{m}")); pair_ok = False
results['date_pairs'] = pair_ok

# controlled values vs Lists
wl = wb['Lists']
lists = {}
for col in range(1, wl.max_column + 1):
    key = wl.cell(row=1, column=col).value
    vals = [wl.cell(row=i, column=col).value for i in range(2, wl.max_row + 1)]
    vals = [v for v in vals if v is not None]
    if key: lists[key] = set(vals)
ctrl_ok = True
for name, allowed in lists.items():
    if name not in H: continue
    for r in rows:
        v = val(r, name)
        if v is not None and str(v) not in allowed:
            problems.append(('ctrl', f"{val(r,'Row id')} {name}={v!r} not in list")); ctrl_ok = False
results['controlled_values'] = ctrl_ok

# conditions detail format on Bid rows
cd_re = re.compile(r"^Fin: (committed|no contingency|not needed|represented|not committed|not stated); DD required: (none|confirmatory|substantive \([^)]*\)|not stated); DD open: (yes|no|not disclosed); Excl: (requested \([^)]*\)|granted|none|not stated)$")
cd_ok = True
for r in rows:
    if val(r, 'What happened') in ('Bid', 'Bid reaffirmed'):
        cd = val(r, 'Conditions detail')
        if not cd or not cd_re.match(str(cd)):
            problems.append(('cond', f"{val(r,'Row id')} bad/absent conditions detail: {cd!r}")); cd_ok = False
results['conditions_detail'] = cd_ok

# round opened / finality / anchor / deadline treatment
rnd_ok = True
for r in rows:
    if val(r, 'What happened') == 'Round opened':
        if not val(r, 'Round finality'):
            problems.append(('round', f"{val(r,'Row id')} no finality")); rnd_ok = False
    if val(r, 'What happened') in ('Deadline', 'Deadline revised'):
        if not val(r, 'Deadline treatment'):
            problems.append(('deadline', f"{val(r,'Row id')} no treatment")); rnd_ok = False
    if val(r, 'Row id') == 'R011' and 'R1 anchor:' not in str(val(r, 'Terms or outcome')):
        problems.append(('round', 'R011 missing R1 anchor tag')); rnd_ok = False
results['round_deadline_fields'] = rnd_ok

# inferred outcome rows start with basis
inf_ok = True
for r in rows:
    ob = val(r, 'Outcome basis')
    if ob and str(ob).startswith('Inferred'):
        if not str(val(r, 'Terms or outcome')).startswith(ob):
            problems.append(('inferred', f"{val(r,'Row id')} terms not starting with basis")); inf_ok = False
results['inferred_rows'] = inf_ok

# outcome rows have outcome basis filled
oc_ok = True
OUTCOMES = {'Dropped by target','Withdrew','Did not submit','Participation paused','Re-entered','Not selected at signing','Joined group'}
for r in rows:
    if val(r, 'What happened') in OUTCOMES and not val(r, 'Outcome basis'):
        problems.append(('outcome', f"{val(r,'Row id')} no outcome basis")); oc_ok = False
results['outcome_basis'] = oc_ok

# counts (SUMIFS equivalents, Include=Yes)
def sif(what=None, rnd=None):
    tot = 0
    for r in rows:
        if val(r, 'Include') != 'Yes': continue
        if what and val(r, 'What happened') != what: continue
        if rnd is not None and val(r, 'Round') != rnd: continue
        c = val(r, 'Count')
        if isinstance(c, (int, float)): tot += c
    return tot
counts = {
    'nda_units': sif('NDA signed'),
    'r1_bid_units': sif('Bid', 1),
    'r1_non_submitters': sif('Did not submit', 1),
    'r2_bid_units': sif('Bid', 2),
    'dropped_units': sif('Dropped by target'),
    'joined_units': sif('Joined group'),
    'not_selected': sif('Not selected at signing'),
    'did_not_submit_all': sif('Did not submit'),
}
expect = {'nda_units': 15, 'r1_bid_units': 6, 'r1_non_submitters': 9, 'r2_bid_units': 2,
          'dropped_units': 2, 'joined_units': 2, 'not_selected': 1, 'did_not_submit_all': 11}
counts_ok = counts == expect
if not counts_ok: problems.append(('counts', f"{counts} != {expect}"))
results['counts'] = counts_ok

# source ids exist; page of first source matches Page column
wst = wb['Source text']
src_ids = {}
for r in wst.iter_rows(min_row=3, values_only=True):
    if r[0]: src_ids[r[0]] = r
src_ok = True
src_first = {}
for r in rows:
    src = str(val(r, 'Source') or '')
    page = val(r, 'Page')
    first = None
    for sid in [s.strip() for s in src.split(';') if s.strip()]:
        found = sid if sid in src_ids else None
        if found is None:
            for k in src_ids:
                if k.startswith(sid) and len(k) == len(sid) + 1 and k[-1] in 'ab':
                    found = k; break
        if found is None:
            problems.append(('source', f"{val(r,'Row id')} unresolved source id {sid}")); src_ok = False
        elif first is None:
            first = found
    if first:
        src_first[val(r, 'Row id')] = first
        sp = src_ids[first][2]
        if page is not None and sp is not None:
            if str(page) != str(sp):
                problems.append(('page', f"{val(r,'Row id')} Page={page} vs first source {first} page={sp}"))
results['source_links_resolve'] = src_ok

# quote verification (depth-tracked)
def quoted_runs(why):
    out, buf, depth = [], [], 0
    for ch in str(why):
        if ch == '“':
            depth += 1
            if depth == 1: buf = []  # start new outer
            continue
        if ch == '”':
            depth = max(0, depth - 1)
            if depth == 0 and buf:
                out.append(''.join(buf)); buf = []
            continue
        if depth >= 1:
            buf.append(ch)
        elif depth == 0 and buf:
            out.append(''.join(buf)); buf = []
    if buf: out.append(''.join(buf))
    return out

def runs_between_quotes(why):
    # returns raw depth>=1 segments split at nested-quote boundaries
    segs = []
    depth = 0; buf = []
    for ch in str(why):
        if ch == '“':
            if depth >= 1 and buf:
                segs.append(''.join(buf)); buf = []
            depth += 1
            continue
        if ch == '”':
            depth = max(0, depth - 1)
            if buf:
                segs.append(''.join(buf)); buf = []
            continue
        if depth >= 1:
            buf.append(ch)
    if buf: segs.append(''.join(buf))
    return [norm(s) for s in segs if len(norm(s)) >= 25]
q_fail = []
for r in rows:
    why = val(r, 'Why and evidence')
    if not why: continue
    for seg in runs_between_quotes(why):
        if seg not in FULLTEXT:
            q_fail.append((val(r, 'Row id'), seg[:120]))
results['quotes_verified'] = len(q_fail) == 0

# formulas present in Summary; workbook features
wsm = wb['Summary']
formula_cells = []
for row in wsm.iter_rows():
    for c in row:
        if isinstance(c.value, str) and c.value.startswith('='):
            formula_cells.append((c.coordinate, c.value))
results['summary_formulas'] = len(formula_cells) >= 5

# data validations
dvs = []
for dv in ws.data_validations.dataValidation:
    dvs.append((str(dv.sqref), dv.formula1))
results['data_validations'] = len(dvs)

# column outline groups
outlines = [(k, v.outline_level, v.hidden) for k, v in ws.column_dimensions.items() if v.outline_level]
results['outline_columns'] = len(outlines)

# conditional formatting
cf = ws.conditional_formatting
cf_ranges = [str(x.sqref) for x in cf]
cf_rules = sum(len(x.rules) for x in cf)
cf_ok = cf_rules >= 2
results['conditional_formatting'] = f"{cf_rules} rules on {cf_ranges}"

# hyperlinks in Source column
hlinks = 0
for r in rows:
    c = r[H['Source'] - 1]
    if c.hyperlink is not None: hlinks += 1
results['source_hyperlinks'] = hlinks

# hidden sheets & order
results['sheet_order'] = wb.sheetnames
results['hidden'] = {s: wb[s].sheet_state for s in wb.sheetnames}

# merged cells
results['merged_cells'] = {s: len(wb[s].merged_cells.ranges) for s in wb.sheetnames}

out = {"problems": problems, "results": results, "counts": counts, "formulas": formula_cells,
       "quote_failures": q_fail}
print(json.dumps(out, indent=1, default=str))
